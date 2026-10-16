from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

path = Path(r"E:\Working\Bieu Mau\Biên bản sữa chữa.xlsx")
wb = load_workbook(path, data_only=False)
report = {"workbook": str(path), "sheets": []}

for ws in wb.worksheets:
    images = []
    for img in ws._images:
        anchor = img.anchor
        if hasattr(anchor, "_from"):
            images.append(
                {
                    "from": [anchor._from.row, anchor._from.col],
                    "to": [anchor.to.row, anchor.to.col] if hasattr(anchor, "to") else None,
                    "width": img.width,
                    "height": img.height,
                }
            )
        else:
            images.append({"anchor": str(anchor), "width": img.width, "height": img.height})

    report["sheets"].append(
        {
            "title": ws.title,
            "dimensions": ws.calculate_dimension(),
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "column_widths": {k: v.width for k, v in ws.column_dimensions.items()},
            "row_heights": {str(k): v.height for k, v in ws.row_dimensions.items() if v.height},
            "view": ws.views.sheetView[0].view,
            "zoom": ws.views.sheetView[0].zoomScale,
            "show_gridlines": ws.views.sheetView[0].showGridLines,
            "print_area": str(ws.print_area),
            "orientation": ws.page_setup.orientation,
            "paper_size": ws.page_setup.paperSize,
            "scale": ws.page_setup.scale,
            "fit_to_width": ws.page_setup.fitToWidth,
            "fit_to_height": ws.page_setup.fitToHeight,
            "margins": {
                "left": ws.page_margins.left,
                "right": ws.page_margins.right,
                "top": ws.page_margins.top,
                "bottom": ws.page_margins.bottom,
            },
            "images": images,
            "formulas": {
                cell.coordinate: cell.value
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            },
            "rows_10_26": [
                [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
                for row in range(10, ws.max_row + 1)
            ],
            "style_ids_11_20": {
                str(row): [ws.cell(row, col).style_id for col in range(1, ws.max_column + 1)]
                for row in range(11, min(20, ws.max_row) + 1)
            },
        }
    )

print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
