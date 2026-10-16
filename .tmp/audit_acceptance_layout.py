from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

before_path = Path(r"E:\Working\Bieu Mau\Biên Bản Nghiệm Thu.xlsx")
after_path = Path(r"E:\Working\Bieu Mau\Biên Bản Nghiệm Thu_v2.xlsx")
before = load_workbook(before_path, data_only=False)
after = load_workbook(after_path, data_only=False)

issues = []
value_changes = []

if before.sheetnames != after.sheetnames:
    issues.append({"sheetnames": [before.sheetnames, after.sheetnames]})

for sheet_name in before.sheetnames:
    ws1 = before[sheet_name]
    ws2 = after[sheet_name]
    properties = {
        "freeze_panes": (ws1.freeze_panes, ws2.freeze_panes),
        "print_area": (str(ws1.print_area), str(ws2.print_area)),
        "merged_cells": (
            sorted(str(x) for x in ws1.merged_cells.ranges),
            sorted(str(x) for x in ws2.merged_cells.ranges),
        ),
        "view": (
            ws1.views.sheetView[0].view,
            ws2.views.sheetView[0].view,
        ),
        "zoom": (
            ws1.views.sheetView[0].zoomScale,
            ws2.views.sheetView[0].zoomScale,
        ),
        "gridlines": (
            ws1.views.sheetView[0].showGridLines,
            ws2.views.sheetView[0].showGridLines,
        ),
        "orientation": (ws1.page_setup.orientation, ws2.page_setup.orientation),
        "paper_size": (ws1.page_setup.paperSize, ws2.page_setup.paperSize),
        "scale": (ws1.page_setup.scale, ws2.page_setup.scale),
        "fit_width": (ws1.page_setup.fitToWidth, ws2.page_setup.fitToWidth),
        "fit_height": (ws1.page_setup.fitToHeight, ws2.page_setup.fitToHeight),
        "images": (len(ws1._images), len(ws2._images)),
    }
    for key, (left, right) in properties.items():
        if left != right:
            issues.append({"sheet": sheet_name, "property": key, "before": left, "after": right})

    cols = set(ws1.column_dimensions) | set(ws2.column_dimensions)
    for col in cols:
        left = ws1.column_dimensions[col].width
        right = ws2.column_dimensions[col].width
        if left != right:
            issues.append({"sheet": sheet_name, "column": col, "before": left, "after": right})

    rows = set(ws1.row_dimensions) | set(ws2.row_dimensions)
    for row in rows:
        left = ws1.row_dimensions[row].height
        right = ws2.row_dimensions[row].height
        if left != right:
            issues.append({"sheet": sheet_name, "row": row, "before": left, "after": right})

    max_row = max(ws1.max_row, ws2.max_row)
    max_col = max(ws1.max_column, ws2.max_column)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            c1 = ws1.cell(row, col)
            c2 = ws2.cell(row, col)
            if c1.value != c2.value:
                value_changes.append(
                    {
                        "sheet": sheet_name,
                        "cell": c1.coordinate,
                        "before": c1.value,
                        "after": c2.value,
                    }
                )
            if c1.style_id != c2.style_id:
                issues.append(
                    {
                        "sheet": sheet_name,
                        "cell": c1.coordinate,
                        "property": "style_id",
                        "before": c1.style_id,
                        "after": c2.style_id,
                    }
                )
            h1 = c1.hyperlink.target if c1.hyperlink else None
            h2 = c2.hyperlink.target if c2.hyperlink else None
            if h1 != h2:
                issues.append(
                    {
                        "sheet": sheet_name,
                        "cell": c1.coordinate,
                        "property": "hyperlink",
                        "before": h1,
                        "after": h2,
                    }
                )

print(json.dumps({"layout_issues": issues, "value_changes": value_changes}, ensure_ascii=False, indent=2))
if issues:
    raise SystemExit(1)
