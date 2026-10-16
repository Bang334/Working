from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\Working\Bieu Mau")
SOURCE_TEMPLATE = BASE / "Biên bản sữa chữa.xlsx"
SOURCE_TEST = BASE / "Biên bản sữa chữa_TEST_14_BAN_GHI.xlsx"
TARGET_TEMPLATE = BASE / "Biên Bản Nghiệm Thu.xlsx"
TARGET_TEST = BASE / "Biên Bản Nghiệm Thu_TEST_14_BAN_GHI.xlsx"


def image_signature(ws):
    result = []
    for image in ws._images:
        anchor = image.anchor
        result.append(
            {
                "from": [anchor._from.row, anchor._from.col],
                "to": [anchor.to.row, anchor.to.col],
                "width": image.width,
                "height": image.height,
            }
        )
    return result


def sheet_layout(ws):
    return {
        "freeze": str(ws.freeze_panes) if ws.freeze_panes else None,
        "columns": {key: dim.width for key, dim in ws.column_dimensions.items()},
        "view": ws.views.sheetView[0].view,
        "zoom": ws.views.sheetView[0].zoomScale,
        "gridlines": ws.views.sheetView[0].showGridLines,
        "orientation": ws.page_setup.orientation,
        "paper": ws.page_setup.paperSize,
        "scale": ws.page_setup.scale,
        "margins": [
            ws.page_margins.left,
            ws.page_margins.right,
            ws.page_margins.top,
            ws.page_margins.bottom,
        ],
        "images": image_signature(ws),
    }


issues = []
checks = {}

source_template = load_workbook(SOURCE_TEMPLATE, data_only=False)
source_test = load_workbook(SOURCE_TEST, data_only=False)
checks["source_sheetnames"] = source_test.sheetnames
if source_template.sheetnames != source_test.sheetnames:
    issues.append("Tên/thứ tự sheet nguồn thử không khớp mẫu.")

for name in source_template.sheetnames:
    left = sheet_layout(source_template[name])
    right = sheet_layout(source_test[name])
    if left != right:
        issues.append({"source_layout_mismatch": name, "before": left, "after": right})
    for row in range(1, 12):
        for col in range(1, source_template[name].max_column + 1):
            c1 = source_template[name].cell(row, col)
            c2 = source_test[name].cell(row, col)
            if c1.value != c2.value or c1._style != c2._style:
                issues.append({"top_block_changed": f"{name}!{c1.coordinate}"})

primary = source_test["CX Technology (VN)"]
quote = source_test["Sheet1"]
checks["source_record_count_primary"] = sum(
    1 for row in range(12, 26) if primary.cell(row, 2).value
)
checks["source_record_count_quote"] = sum(
    1 for row in range(12, 26) if quote.cell(row, 2).value
)
checks["primary_total_formula"] = primary["H27"].value
checks["quote_total_formula"] = quote["I27"].value
checks["primary_print_area"] = str(primary.print_area)
checks["primary_merges"] = sorted(str(x) for x in primary.merged_cells.ranges)
checks["quote_merges"] = sorted(str(x) for x in quote.merged_cells.ranges)

if checks["source_record_count_primary"] != 14 or checks["source_record_count_quote"] != 14:
    issues.append("Số bản ghi nguồn thử không bằng 14 trên cả hai sheet.")
if primary["H27"].value != "=SUM(H12:H25)":
    issues.append("Công thức tổng sheet CX không đúng.")
if quote["I27"].value != "=SUM(I12:I25)":
    issues.append("Công thức tổng sheet báo giá không đúng.")
if str(primary.print_area) != "'CX Technology (VN)'!$A$1:$I$43":
    issues.append("Vùng in sheet CX chưa mở rộng đúng.")
if len(primary._images) != 1 or len(quote._images) != 1:
    issues.append("Logo không được bảo toàn trên cả hai sheet.")

target_template = load_workbook(TARGET_TEMPLATE, data_only=False)
target_test = load_workbook(TARGET_TEST, data_only=False)
if target_template.sheetnames != target_test.sheetnames:
    issues.append("Tên/thứ tự sheet nghiệm thu thử không khớp mẫu.")

expected_notes = {
    "RFR900NA55RSVHBA036": "Không có pin trước khi nhận sửa chữa / 送修前無電池",
    "EF550R4ALAWBH109": "Vỏ máy trầy xước nhẹ trước khi nhận / 送修前機殼有輕微刮痕",
    "EF550R4ALAWBI661": "Thiếu dây đeo tay khi bàn giao / 交付時缺少手腕帶",
    "EF550R4ALAWBI677": "Tem serial bị mờ trước khi sửa chữa / 維修前序號標籤模糊",
    "EF550RA4LAWBI666": "Không có pin trước khi nhận sửa chữa / 送修前無電池",
    "EF550RA4LAWBI678": "Không có pin và nắp che cổng sạc / 無電池且缺少充電埠護蓋",
}
actual_notes = {}

for name in target_template.sheetnames:
    ws1 = target_template[name]
    ws2 = target_test[name]
    if sheet_layout(ws1) != sheet_layout(ws2):
        issues.append({"target_layout_mismatch": name})
    if sorted(str(x) for x in ws1.merged_cells.ranges) != sorted(
        str(x) for x in ws2.merged_cells.ranges
    ):
        issues.append({"target_merge_mismatch": name})
    max_row = max(ws1.max_row, ws2.max_row)
    max_col = max(ws1.max_column, ws2.max_column)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            c1 = ws1.cell(row, col)
            c2 = ws2.cell(row, col)
            if c1.style_id != c2.style_id:
                issues.append({"target_style_changed": f"{name}!{c1.coordinate}"})
            h1 = c1.hyperlink.target if c1.hyperlink else None
            h2 = c2.hyperlink.target if c2.hyperlink else None
            if h1 != h2:
                issues.append({"target_hyperlink_changed": f"{name}!{c1.coordinate}"})

    if name == "Nghiệm Thu":
        continue
    note_cell = None
    for row in ws2.iter_rows():
        for cell in row:
            text = str(cell.value or "")
            if "Ghi chú" in text or "備註" in text:
                note_cell = ws2.cell(cell.row + 1, cell.column)
                break
        if note_cell:
            break
    actual_notes[name] = note_cell.value if note_cell else None

for name, expected in expected_notes.items():
    if actual_notes.get(name) != expected:
        issues.append({"wrong_note": name, "actual": actual_notes.get(name)})

for name, value in actual_notes.items():
    if name not in expected_notes and value not in (None, ""):
        issues.append({"unexpected_note": name, "actual": value})

checks["assigned_notes"] = {k: v for k, v in actual_notes.items() if v}
checks["blank_note_sheets"] = sorted(k for k, v in actual_notes.items() if not v)
checks["unknown_serial_sheet_created"] = "RFR900NA55RSVTEST901" in target_test.sheetnames

print(json.dumps({"issues": issues, "checks": checks}, ensure_ascii=False, indent=2))
if issues:
    raise SystemExit(1)
