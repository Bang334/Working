from pathlib import Path
import sys

from openpyxl import load_workbook


BASE = Path(r"E:\Working\Bieu Mau")

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

for path in sorted(BASE.glob("*.xlsx")):
    if path.name.startswith("~$"):
        continue
    print(f"\nFILE {path.name}")
    workbook = load_workbook(path, data_only=False, read_only=False)
    print("sheets=", workbook.sheetnames)
    for worksheet in workbook.worksheets:
        matches = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value in (None, ""):
                    continue
                text = str(cell.value)
                is_note_label = "ghi chú" in text.lower() or "備註" in text
                is_detail_note = worksheet.title != "Nghiệm Thu" and cell.coordinate in ("D6", "F6")
                is_source_footer = worksheet.title in ("CX Technology (VN)", "Sheet1") and cell.row >= 18
                if is_note_label or is_detail_note or is_source_footer:
                    matches.append((cell.coordinate, text))
        if matches:
            print("SHEET", worksheet.title)
            for match in matches:
                print(repr(match))
