# -*- coding: utf-8 -*-
"""
=============================================================================
HONG MING GLOBAL - CX TECHNOLOGY PAYMENT REQUEST GENERATOR (請款單 / ĐƠN YÊU CẦU THANH TOÁN)
=============================================================================
Dựa trên mẫu chuẩn 1:1 từ phôi "đơn yêu cầu NCC thanh toán (1).xlsx" (Sheet CX)
Bóc tách dữ liệu từ:
1. Don_Xin_Thanh_Toan_HVM007-20261021-2_CÔNG TY TNHH TOÀN CẦU HONG MING.xlsx (Đợt 3 - 20% = 46.800.000 VND)
2. 02_BIEN_BAN_NGHIEM_THU_BAN_GIAO_CX_.docx (Hợp đồng HMV151020250001 ngày 09/12/2025, tổng 234.000.000 VND)
=============================================================================
"""

import os
import sys
import shutil
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def generate_cx_payment_request(
    template_path: str = r"C:\Users\BANG\Downloads\đơn yêu cầu NCC thanh toán (1).xlsx",
    output_dir: str = r"C:\Users\BANG\Downloads",
    workspace_dir: str = r"E:\Working\Ho So Du An\Don Xin Thanh Toan"
):
    print("=" * 70)
    print("Khởi tạo Phiếu yêu cầu thanh toán (請款單) cho CX TECHNOLOGY...")
    print("=" * 70)

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Không tìm thấy file mẫu tại: {template_path}")

    # 1. Mở workbook mẫu
    wb = openpyxl.load_workbook(template_path, data_only=False)
    if "CX" not in wb.sheetnames:
        raise ValueError("Không tìm thấy sheet 'CX' trong file mẫu!")

    # 2. Tạo workbook chuyên biệt chỉ chứa sheet CX để gửi khách hàng
    # Tạo bản sao sạch: Xóa các sheet khác ngoài 'CX'
    sheets_to_remove = [s for s in wb.sheetnames if s != "CX"]
    for s in sheets_to_remove:
        del wb[s]

    ws = wb["CX"]
    ws.title = "CX_請款單"

    # 3. Cập nhật các trường thông tin chuẩn xác theo Hợp đồng HMV151020250001
    form_no = "HM21102026-001"
    apply_date = "21/10/2026"
    due_date = "31/10/2026"

    # G9: Form NO. | H9: Số Form
    ws["G9"].value = "Form NO."
    ws["H9"].value = form_no

    # B10: 客 戶 名 稱：/ Tên khách hàng： | C10: CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)
    ws["B10"].value = "客 戶 名 稱：\nTên khách hàng："
    ws["C10"].value = "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)"

    # H10: 申 請 日 期：/ Ngày nộp đơn： | I10: 21/10/2026
    ws["H10"].value = "申 請 日 期：\nNgày nộp đơn："
    ws["I10"].value = apply_date
    ws["I10"].number_format = "@"

    # B11: 專 案 名 稱：/ Tên dự án：
    project_title = (
        "Căn cứ hợp đồng số : HMV151020250001 ký ngày 09/12/2025\n"
        "RFID 硬體 (RFID Hardware) – Phần mềm tùy chỉnh chức năng hệ thống RFID\n"
        "根据合同编号：HMV151020250001 簽約日期：2025年12月09日\n"
        "RFID硬體（RFID Hardware）－RFID系統客製化軟體"
    )
    ws["B11"].value = "專 案 名 稱：\nTên dự án："
    ws["C11"].value = project_title
    ws.row_dimensions[11].height = 98.4

    # B12: 應 付 款 日 期： | C12: 31/10/2026 | F12: 付 款 方 式： | H12: 匯款\nChuyển khoản
    ws["B12"].value = "應 付 款 日 期：\nNgày đến hạn thanh toán:"
    ws["C12"].value = due_date
    ws["C12"].number_format = "@"
    ws["F12"].value = "付 款 方 式：\nPhương thức thanh toán："
    ws["H12"].value = "匯款\nChuyển khoản"
    ws.row_dimensions[12].height = 41.4

    # Bảng chi tiết thanh toán (Row 14 đến Row 21)
    # Hàng 15: Đợt 3 (20% Nghiệm thu đạt tiêu chuẩn)
    item_desc = (
        "驗收合格尾款：20%\n"
        "合約：HMV151020250001（RFID 系統功能客製化軟體）\n"
        "Thanh toán đợt Giai đoạn 3 (20% Nghiệm thu đạt tiêu chuẩn theo Biên bản nghiệm thu)"
    )
    ws["A15"].value = 1
    ws["B15"].value = item_desc
    ws["C15"].value = "VND"
    ws["D15"].value = 234000000  # Tổng giá trị hợp đồng
    ws["E15"].value = 1
    ws["F15"].value = "=D15*E15"
    ws["G15"].value = None       # Thuế suất (để trống theo chuẩn phần mềm không VAT)
    ws["H15"].value = "=F15*G15"
    ws["I15"].value = "=(F15+H15)*20%"  # 234.000.000 * 20% = 46.800.000 VND
    ws.row_dimensions[15].height = 76.0

    # Các hàng 16 đến 20: Giữ nguyên cấu trúc
    ws["A16"].value = 2
    ws["F16"].value = "=D16*E16"
    ws["H16"].value = "=F16*G16"
    ws["I16"].value = "=F16+H16"

    ws["A17"].value = 3
    ws["F17"].value = "=D17*E17"
    ws["H17"].value = "=F17*G17"
    ws["I17"].value = "=F17+H17"

    ws["A18"].value = 4
    ws["F18"].value = "=D18*E18"
    ws["H18"].value = "=F18*G18"
    ws["I18"].value = "=F18+H18"

    ws["A19"].value = 5
    ws["F19"].value = "=D19*E19"
    ws["H19"].value = "=F19*G19"
    ws["I19"].value = "=F19+H19"

    ws["A20"].value = 6
    ws["F20"].value = "=D20*E20"
    ws["H20"].value = "=F20*G20"
    ws["I20"].value = "=F20+H20"

    # Row 21: Tổng cộng
    ws["A21"].value = "合計 Tổng cộng"
    ws["F21"].value = "=SUM(I15:I20)"
    ws.row_dimensions[21].height = 25.0

    # Row 23: Ghi chú trích dẫn đầy đủ 3 đợt thanh toán từ Hợp đồng & Biên bản nghiệm thu
    remark_text = (
        "備註：Căn cứ Mục A – Phương thức thanh toán của Hợp đồng số HMV151020250001:\n"
        "1. Sau khi ký kết hợp đồng: thanh toán 40% giá trị hợp đồng = 93.600.000 VNĐ (Chín mươi ba triệu sáu trăm nghìn đồng).\n"
        "1. 簽約完成時：支付合約金額40%＝93,600,000越盾（玖仟參佰陸拾萬越盾整）。\n"
        "2. Sau khi triển khai sản phẩm đưa vào sử dụng: thanh toán 40% giá trị hợp đồng = 93.600.000 VNĐ (Chín mươi ba triệu sáu trăm nghìn đồng).\n"
        "2. 產品安裝後：支付合約金額40%＝93,600,000越盾（玖仟參佰陸拾萬越盾整）。\n"
        "3. Khi nghiệm thu đạt tiêu chuẩn: thanh toán 20% còn lại = 46.800.000 VNĐ (Bốn mươi sáu triệu tám trăm nghìn đồng).\n"
        "3. 驗收合格時：支付尾款20%＝46,800,000越盾（肆仟陸佰捌拾萬越盾整）。\n\n"
        "**以上價格不包含VAT增值稅，稅率依發票開立時的稅率為準。Giá nêu trên chưa bao gồm thuế GTGT (VAT). Thuế suất áp dụng theo mức thuế suất tại thời điểm xuất hóa đơn.\n"
    )
    ws["B23"].value = remark_text
    ws.row_dimensions[23].height = 158.0  # Chiều cao an toàn chống cắt chữ

    # Row 25 & 26: Tài khoản ngân hàng chuẩn của Hong Ming Global
    ws["B25"].value = "銀行帳戶名稱"
    ws["C25"].value = "CÔNG TY TNHH TOÀN CẦU HONG MING"
    ws["G25"].value = "銀行帳號"
    ws["H25"].value = "604080006546(VND)"
    ws["B26"].value = "銀行名稱"
    ws["C26"].value = "CATHAY UNITED BANK – HO CHI MINH CITY BRANCH"

    # Row 28: Khối chữ ký
    ws["B28"].value = "買方主管"
    ws["F28"].value = "申請單位"

    # Cấu hình in ấn chuẩn xác
    ws.print_area = "$A$1:$I$29"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.scale = 59     # Scale 59% vừa khít A4 chuẩn mẫu
    if ws.views.sheetView:
        ws.views.sheetView[0].zoomScale = 85

    # 4. Lưu file an toàn
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(workspace_dir, exist_ok=True)

    filename = f"{form_no}_Đơn_Yêu_Cầu_Thanh_Toán_CX_Technology_HMV151020250001.xlsx"
    target_downloads = os.path.join(output_dir, filename)
    target_workspace = os.path.join(workspace_dir, filename)

    try:
        wb.save(target_downloads)
        print(f" -> Đã lưu thành công vào Downloads: {target_downloads}")
    except PermissionError:
        target_downloads = os.path.join(output_dir, f"{form_no}_Đơn_Yêu_Cầu_Thanh_Toán_CX_Technology_v2.xlsx")
        wb.save(target_downloads)
        print(f" -> File đang mở, đã lưu bản dự phòng Downloads: {target_downloads}")

    try:
        shutil.copy2(target_downloads, target_workspace)
        print(f" -> Đã đồng bộ vào Hồ Sơ Dự Án: {target_workspace}")
    except Exception as exc:
        print(f" -> Cảnh báo copy sang Workspace: {exc}")

    # 5. Cập nhật luôn sheet CX vào chính file mẫu gốc để đồng bộ tiện lợi
    try:
        wb_full = openpyxl.load_workbook(template_path, data_only=False)
        if "CX" in wb_full.sheetnames:
            ws_f = wb_full["CX"]
            ws_f["H9"].value = form_no
            ws_f["C10"].value = "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)"
            ws_f["I10"].value = apply_date
            ws_f["I10"].number_format = "@"
            ws_f["C11"].value = project_title
            ws_f.row_dimensions[11].height = 98.4
            ws_f["C12"].value = due_date
            ws_f["B15"].value = item_desc
            ws_f["C15"].value = "VND"
            ws_f["D15"].value = 234000000
            ws_f["E15"].value = 1
            ws_f["F15"].value = "=D15*E15"
            ws_f["G15"].value = None
            ws_f["H15"].value = "=F15*G15"
            ws_f["I15"].value = "=(F15+H15)*20%"
            ws_f["B23"].value = remark_text
            ws_f.row_dimensions[23].height = 158.0
            wb_full.save(template_path)
            print(f" -> Đã đồng bộ trực tiếp vào sheet 'CX' của file mẫu: {template_path}")
    except PermissionError:
        print(f" -> File mẫu gốc đang mở trong Excel, bỏ qua bước ghi đè mẫu gốc.")
    except Exception as exc:
        print(f" -> Lưu file mẫu gốc: {exc}")

    print("=" * 70)
    print("HOÀN TẤT TẠO PHIẾU YÊU CẦU THANH TOÁN CX TECHNOLOGY!")
    print(f"File chính: {target_downloads}")
    print("=" * 70)
    return target_downloads


if __name__ == "__main__":
    generate_cx_payment_request()
