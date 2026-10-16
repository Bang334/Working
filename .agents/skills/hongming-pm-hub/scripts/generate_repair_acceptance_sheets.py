# -*- coding: utf-8 -*-
r"""
Script tự động hoàn thiện Biên bản nghiệm thu thiết bị sửa chữa CX Technology
File: C:\Users\BANG\Downloads\Biên Bản Nghiệm Thu_維修設備驗收記錄_HDSC-HMV06032026001.xlsx
Dựa trên mẫu 3 sheet đã làm của người dùng, tạo tiếp 9 sheet còn lại cho đủ 12 thiết bị:
- Giữ nguyên cấu trúc, định dạng ô vàng chờ chèn ảnh
- Thiết lập tiêu đề, thông tin STT, Bộ phận, Lỗi ban đầu, Hạng mục kiểm tra, Kết quả đạt
- Ghi chú pin cho máy không có pin (STT 11, STT 12)
- Tạo liên kết Hyperlink 2 chiều (từ sheet Nghiệm Thu sang từng máy và nút quay về)
"""

import openpyxl
import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# 9 thiết bị cần tạo sheet mới (từ STT 4 đến 12)
NEW_DEVICES = [
    {
        "stt": "4",
        "sheet_name": "EF550R4ALAWBH109",
        "serial": "EF550R4ALAWBH109",
        "model_extra": "(BF606-1677)",
        "bo_phan": "LH",
        "loi": "1. Barcode lỗi / 條碼掃描故障  ;  2. Không nhận pin / 無法辨識電池",
        "hang_muc": "1. Kiểm tra chức năng quét barcode / 檢查條碼掃描功能  ;  2. Kiểm tra nhận dạng và sạc pin / 檢查電池辨識與充電功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 9
    },
    {
        "stt": "5",
        "sheet_name": "EF550R4ALAWBI654",
        "serial": "EF550R4ALAWBI654",
        "model_extra": "(BF606-1677)",
        "bo_phan": "CNC",
        "loi": "Màn hình lỗi / 螢幕顯示異常",
        "hang_muc": "Kiểm tra hiển thị màn hình / 檢查螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 11
    },
    {
        "stt": "6",
        "sheet_name": "EF550R4ALAWBI661",
        "serial": "EF550R4ALAWBI661",
        "model_extra": "(BF606-1677)",
        "bo_phan": "XI MẠ",
        "loi": "Màn hình lỗi / 螢幕顯示異常",
        "hang_muc": "Kiểm tra hiển thị màn hình / 檢查螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 12
    },
    {
        "stt": "7",
        "sheet_name": "EF550R4ALAWBI677",
        "serial": "EF550R4ALAWBI677",
        "model_extra": "(BF606-1677)",
        "bo_phan": "NC",
        "loi": "Màn hình lỗi / 螢幕顯示異常",
        "hang_muc": "Kiểm tra hiển thị màn hình / 檢查螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 13
    },
    {
        "stt": "8",
        "sheet_name": "EF550RA4LAWBI657",
        "serial": "EF550RA4LAWBI657",
        "model_extra": "(BF606-1677)",
        "bo_phan": "X-TR",
        "loi": "Màn hình lỗi sau khi bật nguồn / 開機後螢幕顯示錯誤",
        "hang_muc": "Kiểm tra khởi động & hiển thị màn hình / 檢查開機及螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 14
    },
    {
        "stt": "9",
        "sheet_name": "EF550R4ALAWBI641",
        "serial": "EF550R4ALAWBI641",
        "model_extra": "(BF606-1677)",
        "bo_phan": "X-TR",
        "loi": "Không sạc được pin / 無法充電",
        "hang_muc": "Kiểm tra chức năng sạc / 檢查充電功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 15
    },
    {
        "stt": "10",
        "sheet_name": "EF550RA4LAWBI686",
        "serial": "EF550RA4LAWBI686",
        "model_extra": "(BF606-1677)",
        "bo_phan": "Alume\n鋁製裝組",
        "loi": "Màn hình lỗi / Alume 螢幕顯示異常",
        "hang_muc": "Kiểm tra hiển thị màn hình / 檢查螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": None,
        "row_in_summary": 16
    },
    {
        "stt": "11",
        "sheet_name": "EF550RA4LAWBI666",
        "serial": "EF550RA4LAWBI666",
        "model_extra": "(BF606-1677)",
        "bo_phan": "DG\n包裝組",
        "loi": "Màn hình lỗi / 螢幕顯示異常",
        "hang_muc": "Kiểm tra hiển thị màn hình / 檢查螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": "KHÔNG CÓ PIN \n無電池",
        "row_in_summary": 17
    },
    {
        "stt": "12",
        "sheet_name": "EF550RA4LAWBI678",
        "serial": "EF550RA4LAWBI678",
        "model_extra": "(BF606-1677)",
        "bo_phan": "DG\n包裝組",
        "loi": "Màn hình lỗi / 螢幕顯示異常",
        "hang_muc": "Kiểm tra hiển thị màn hình / 檢查螢幕顯示功能",
        "ket_qua": "Đạt / 合格",
        "ghi_chu": "KHÔNG CÓ PIN \n無電池",
        "row_in_summary": 18
    }
]

def complete_repair_acceptance_workbook(file_path):
    print(f"Loading workbook: {file_path}")
    wb = openpyxl.load_workbook(file_path)

    # Sheet mẫu để nhân bản: EF550R4ALAWBI670
    template_ws = wb['EF550R4ALAWBI670']
    summary_ws = wb['Nghiệm Thu']

    # 1. Thêm Hyperlink quay về cho 3 sheet mẫu sẵn có
    for existing_name in ['RFR900NA55RSVHBA036', 'EF550R4ALAWBI670', 'EF550R4ALAWBH111']:
        if existing_name in wb.sheetnames:
            ws = wb[existing_name]
            ws['A8'].value = '← Quay về bảng tổng hợp "Nghiệm Thu" / 返回總表'
            ws['A8'].hyperlink = "#'Nghiệm Thu'!A1"

    # Thêm hyperlink từ sheet Nghiệm Thu cho 3 thiết bị đầu
    summary_ws['G6'].hyperlink = "#'RFR900NA55RSVHBA036'!A1"
    summary_ws['G7'].hyperlink = "#'EF550R4ALAWBI670'!A1"
    summary_ws['G8'].hyperlink = "#'EF550R4ALAWBH111'!A1"

    # 2. Tạo 9 sheet mới cho STT 4 đến 12 theo quy tắc RF (3 ảnh) và EF (5 ảnh)
    rf_template = wb['RFR900NA55RSVHBA036']
    ef_template = wb['EF550R4ALAWBI670']

    for dev in NEW_DEVICES:
        sname = dev["sheet_name"]
        serial = dev["serial"].upper()
        
        # Áp dụng quy tắc phân loại mã máy:
        # - Mã bắt đầu bằng 'RF' (RFR900 Gun): 3 ảnh (Cột A, B, C) + Ghi chú (Cột D)
        # - Mã bắt đầu bằng 'EF' (EF550 Handheld): 5 ảnh (Cột A, B, C, D, E) + Ghi chú (Cột F)
        if serial.startswith('RF'):
            active_template = rf_template
            img_cols = ['A', 'B', 'C']
            note_col = 'D'
            print(f"Creating sheet: {sname} (Mã RF: Áp dụng mẫu 3 ảnh)...")
        else:
            active_template = ef_template
            img_cols = ['A', 'B', 'C', 'D', 'E']
            note_col = 'F'
            print(f"Creating sheet: {sname} (Mã EF: Áp dụng mẫu 5 ảnh)...")

        # Xóa sheet cũ nếu đã tồn tại để tạo mới chuẩn xác
        if sname in wb.sheetnames:
            del wb[sname]

        # Copy từ template tương ứng
        new_ws = wb.copy_worksheet(active_template)
        new_ws.title = sname

        # Cập nhật thông tin chi tiết
        new_ws['A1'].value = f"HÌNH ẢNH KIỂM TRA / 檢查圖片 — {dev['serial']} {dev['model_extra']}"
        new_ws['A2'].value = f"STT: {dev['stt']}    |    Bộ phận / 部門: {dev['bo_phan']}    |    Kết quả / 結果: {dev['ket_qua']}"
        new_ws['A3'].value = f"Lỗi ban đầu / 故障說明: {dev['loi']}"
        new_ws['A4'].value = f"Hạng mục KT / 檢查項目: {dev['hang_muc']}"

        # Đảm bảo các ô chèn ảnh trống (màu vàng nhạt) để người dùng chèn ảnh sau
        for col_letter in img_cols:
            cell = new_ws[f"{col_letter}6"]
            cell.value = None

        # Cột Ghi chú
        cell_note = new_ws[f"{note_col}6"]
        if dev['ghi_chu']:
            cell_note.value = dev['ghi_chu']
        else:
            cell_note.value = None

        # Dòng 7 hướng dẫn
        new_ws['A7'].value = "⬆  Chèn hình ảnh vào ô màu vàng: Insert → Pictures → Place in Cell  /  請將圖片插入黃色儲存格"

        # Dòng 8 nút quay về kèm Hyperlink
        new_ws['A8'].value = '← Quay về bảng tổng hợp "Nghiệm Thu" / 返回總表'
        new_ws['A8'].hyperlink = "#'Nghiệm Thu'!A1"

        # Cập nhật Hyperlink trong bảng tổng hợp "Nghiệm Thu"
        r_sum = dev['row_in_summary']
        summary_cell = summary_ws[f"G{r_sum}"]
        summary_cell.value = f"→ {sname}"
        summary_cell.hyperlink = f"#'{sname}'!A1"

    print(f"\nAll 12 sheets ready. Total sheets in workbook: {len(wb.sheetnames)}")
    print(f"Sheet list: {wb.sheetnames}")

    # Lưu lại file
    completed_name = "Biên Bản Nghiệm Thu_維修設備驗收記錄_HDSC-HMV06032026001_HOAN_TAT.xlsx"
    downloads_completed = os.path.join(os.path.dirname(file_path), completed_name)
    working_copy = os.path.join(r"e:\Working", completed_name)

    # 1. Lưu file chính đã hoàn tất
    wb.save(downloads_completed)
    print(f"Saved completed file to Downloads: {downloads_completed}")

    # 2. Dùng shutil.copy2 để nhân bản sang e:\Working (tránh lỗi openpyxl closed image buffer)
    import shutil
    shutil.copy2(downloads_completed, working_copy)
    print(f"Saved completed file to Working: {working_copy}")

    # 3. Cố gắng ghi đè vào file gốc nếu người dùng không khóa file
    try:
        shutil.copy2(downloads_completed, file_path)
        print(f"Saved directly to original file: {file_path}")
    except PermissionError:
        print("Note: File gốc đang mở trong Excel/WPS. Bản hoàn tất được lưu an toàn tại '_HOAN_TAT.xlsx'.")

if __name__ == "__main__":
    target = r"C:\Users\BANG\Downloads\Biên Bản Nghiệm Thu_維修設備驗收記錄_HDSC-HMV06032026001.xlsx"
    complete_repair_acceptance_workbook(target)
