# -*- coding: utf-8 -*-
"""Kiểm tra file Phiếu yêu cầu thanh toán có giữ nguyên bố cục sheet 712TR."""

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List

import openpyxl


def _widths(ws) -> Dict[str, Any]:
    return {column: ws.column_dimensions[column].width for column in "ABCDEFGHI"}


def _rows(ws) -> Dict[int, Any]:
    return {
        row: (ws.row_dimensions[row].height, bool(ws.row_dimensions[row].hidden))
        for row in range(1, 30)
    }


def verify(template_path: Path, output_path: Path) -> Dict[str, Any]:
    source_wb = openpyxl.load_workbook(template_path, data_only=False)
    output_wb = openpyxl.load_workbook(output_path, data_only=False)
    source = source_wb["712TR"]
    output = output_wb["712TR"]
    checks: Dict[str, bool] = {
        "single_sheet_712TR": output_wb.sheetnames == ["712TR"],
        "freeze_panes": output.freeze_panes == source.freeze_panes,
        "sheet_view": output.sheet_view.view == source.sheet_view.view,
        "zoom_scale": output.sheet_view.zoomScale == source.sheet_view.zoomScale,
        "gridlines": output.sheet_view.showGridLines == source.sheet_view.showGridLines,
        "print_area": output.print_area == source.print_area,
        "page_setup": (
            output.page_setup.orientation,
            output.page_setup.paperSize,
            output.page_setup.scale,
            output.page_setup.fitToWidth,
            output.page_setup.fitToHeight,
        )
        == (
            source.page_setup.orientation,
            source.page_setup.paperSize,
            source.page_setup.scale,
            source.page_setup.fitToWidth,
            source.page_setup.fitToHeight,
        ),
        "page_margins": output.page_margins == source.page_margins,
        "merged_cells": sorted(str(item) for item in output.merged_cells.ranges)
        == sorted(str(item) for item in source.merged_cells.ranges),
        "column_widths": _widths(output) == _widths(source),
        "row_heights_and_hidden": _rows(output) == _rows(source),
        "styles_A1_I29": all(
            output.cell(row, column)._style == source.cell(row, column)._style
            for row in range(1, 30)
            for column in range(1, 10)
        ),
        "formula_F15": output["F15"].value == "=D15*E15",
        "formula_H15": output["H15"].value == "=F15*G15",
        "formula_total": output["F21"].value == "=SUM(I15:I20)",
        "no_images": len(output._images) == len(source._images) == 0,
    }
    failures: List[str] = [name for name, passed in checks.items() if not passed]
    return {
        "status": "pass" if not failures else "fail",
        "template_sheet": "712TR",
        "checks": checks,
        "failures": failures,
    }


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(r"E:\Working\Bieu Mau\đơn yêu cầu NCC thanh toán-Gốc.xlsx"),
    )
    args = parser.parse_args()
    result = verify(args.template, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
