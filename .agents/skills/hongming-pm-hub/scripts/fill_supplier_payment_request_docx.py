# -*- coding: utf-8 -*-
"""Fill the approved supplier payment request DOCX without changing layout.

The file below is the live golden template and is never overwritten:
E:\\Working\\Bieu Mau\\Mau_Goc_De_Nghi_Thanh_Toan_NCC.docx

Only existing text nodes are edited. Paragraph properties, run properties,
table grids, cell widths, borders, row heights, margins and section settings
remain untouched.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


# Locate the project independently of the drive letter on this computer.
_WORKSPACE_ROOT = next((p for p in __import__('pathlib').Path(__file__).resolve().parents if (p / 'pm_portal' / 'app').is_dir()), __import__('pathlib').Path.cwd())
def _workspace_path(relative):
    return str(_WORKSPACE_ROOT / relative.replace(chr(92), '/'))

DEFAULT_TEMPLATE = Path(_workspace_path('Bieu Mau\\Mau_Goc_De_Nghi_Thanh_Toan_NCC.docx'))
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


@dataclass(frozen=True)
class Replacement:
    find: str
    replace: str
    expected: int | None = None
    required: bool = True


def _unique_cells(table) -> Iterable[Any]:
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            yield cell


def _iter_paragraphs(container) -> Iterable[Any]:
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for cell in _unique_cells(table):
            yield from _iter_paragraphs(cell)


def _iter_all_paragraphs(doc) -> Iterable[Any]:
    seen: set[int] = set()
    containers = [doc]
    for section in doc.sections:
        containers.extend(
            [
                section.header,
                section.footer,
                section.first_page_header,
                section.first_page_footer,
                section.even_page_header,
                section.even_page_footer,
            ]
        )
    for container in containers:
        for paragraph in _iter_paragraphs(container):
            key = id(paragraph._p)
            if key in seen:
                continue
            seen.add(key)
            yield paragraph


def _set_text(node, value: str) -> None:
    node.text = value
    if value[:1].isspace() or value[-1:].isspace():
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def _find_matches(text: str, needle: str) -> list[tuple[int, int]]:
    matches: list[tuple[int, int]] = []
    cursor = 0
    while True:
        start = text.find(needle, cursor)
        if start < 0:
            return matches
        end = start + len(needle)
        matches.append((start, end))
        cursor = end


def _locate_start(bounds: list[tuple[int, int]], position: int) -> tuple[int, int]:
    for idx, (start, end) in enumerate(bounds):
        if start <= position < end:
            return idx, position - start
    raise ValueError(f"Không xác định được vị trí bắt đầu {position}.")


def _locate_end(bounds: list[tuple[int, int]], position: int) -> tuple[int, int]:
    for idx, (start, end) in enumerate(bounds):
        if start < position <= end:
            return idx, position - start
    raise ValueError(f"Không xác định được vị trí kết thúc {position}.")


def replace_in_paragraph(paragraph, needle: str, replacement: str) -> int:
    """Replace text across split Word runs while preserving all run properties."""
    text_nodes = list(paragraph._p.xpath(".//w:t"))
    if not text_nodes:
        return 0

    original_parts = [node.text or "" for node in text_nodes]
    full_text = "".join(original_parts)
    matches = _find_matches(full_text, needle)
    if not matches:
        return 0

    bounds: list[tuple[int, int]] = []
    cursor = 0
    for part in original_parts:
        bounds.append((cursor, cursor + len(part)))
        cursor += len(part)

    located: list[tuple[int, int, int, int]] = []
    for start, end in matches:
        start_idx, start_offset = _locate_start(bounds, start)
        end_idx, end_offset = _locate_end(bounds, end)
        located.append((start_idx, start_offset, end_idx, end_offset))

    # Process from the end so earlier character offsets stay valid.
    for start_idx, start_offset, end_idx, end_offset in reversed(located):
        if start_idx == end_idx:
            current = text_nodes[start_idx].text or ""
            _set_text(
                text_nodes[start_idx],
                current[:start_offset] + replacement + current[end_offset:],
            )
            continue

        first = text_nodes[start_idx].text or ""
        last = text_nodes[end_idx].text or ""
        _set_text(text_nodes[start_idx], first[:start_offset] + replacement)
        for idx in range(start_idx + 1, end_idx):
            _set_text(text_nodes[idx], "")
        _set_text(text_nodes[end_idx], last[end_offset:])

    return len(matches)


def _layout_fingerprint(doc) -> str:
    """Hash layout/style XML only; visible text nodes are intentionally excluded."""
    tags = {
        qn("w:sectPr"),
        qn("w:tblPr"),
        qn("w:tblGrid"),
        qn("w:trPr"),
        qn("w:tcPr"),
        qn("w:pPr"),
        qn("w:rPr"),
    }
    digest = hashlib.sha256()
    parts = [doc.part]
    for section in doc.sections:
        parts.extend(
            [
                section.header.part,
                section.footer.part,
                section.first_page_header.part,
                section.first_page_footer.part,
                section.even_page_header.part,
                section.even_page_footer.part,
            ]
        )
    seen_parts: set[str] = set()
    for part in parts:
        part_name = str(part.partname)
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        root = part.element
        for element in root.iter():
            if element.tag in tags:
                digest.update(etree.tostring(element, with_tail=False))
    return digest.hexdigest()


def _parse_replacements(payload: dict[str, Any]) -> list[Replacement]:
    raw = payload.get("replacements")
    if raw is None:
        raise ValueError("JSON phải có mảng 'replacements'.")
    if not isinstance(raw, list):
        raise TypeError("'replacements' phải là một mảng.")

    replacements: list[Replacement] = []
    seen: set[str] = set()
    for idx, item in enumerate(raw):
        if not isinstance(item, dict):
            raise TypeError(f"replacements[{idx}] phải là object.")
        needle = str(item.get("find", ""))
        value = str(item.get("replace", ""))
        if not needle:
            raise ValueError(f"replacements[{idx}].find không được rỗng.")
        if "\n" in value or "\r" in value:
            raise ValueError(
                f"replacements[{idx}].replace có xuống dòng. "
                "Hãy thay riêng từng dòng để bảo toàn run/style của mẫu."
            )
        if needle in seen:
            raise ValueError(f"Chuỗi tìm bị lặp: {needle!r}")
        seen.add(needle)
        expected = item.get("expected")
        if expected is not None:
            expected = int(expected)
        replacements.append(
            Replacement(
                find=needle,
                replace=value,
                expected=expected,
                required=bool(item.get("required", True)),
            )
        )
    return replacements


def _next_fallback(path: Path) -> Path:
    candidate = path.with_name(f"{path.stem}_v2{path.suffix}")
    counter = 3
    while candidate.exists():
        candidate = path.with_name(f"{path.stem}_v{counter}{path.suffix}")
        counter += 1
    return candidate


def fill_payment_request(
    template_path: Path,
    output_path: Path,
    replacements: list[Replacement],
) -> dict[str, Any]:
    template = template_path.resolve()
    output = output_path.resolve()
    if not template.exists():
        raise FileNotFoundError(f"Không tìm thấy mẫu chuẩn: {template}")
    if template == output:
        raise ValueError("Không được ghi đè mẫu gốc. Hãy chọn tên file đầu ra khác.")

    doc = Document(str(template))
    before_fingerprint = _layout_fingerprint(doc)
    counts: dict[str, int] = {}
    paragraphs = list(_iter_all_paragraphs(doc))

    for spec in replacements:
        count = sum(
            replace_in_paragraph(paragraph, spec.find, spec.replace)
            for paragraph in paragraphs
        )
        counts[spec.find] = count
        if spec.expected is not None and count != spec.expected:
            raise ValueError(
                f"Chuỗi {spec.find!r}: tìm thấy {count}, yêu cầu đúng {spec.expected}."
            )
        if spec.required and count == 0:
            raise ValueError(f"Không tìm thấy chuỗi bắt buộc: {spec.find!r}")

    after_fingerprint = _layout_fingerprint(doc)
    if before_fingerprint != after_fingerprint:
        raise RuntimeError("Phát hiện thay đổi layout/style trong bộ nhớ; đã hủy lưu file.")

    output.parent.mkdir(parents=True, exist_ok=True)
    temp_handle = tempfile.NamedTemporaryFile(
        prefix="payment_request_",
        suffix=".docx",
        dir=output.parent,
        delete=False,
    )
    temp_path = Path(temp_handle.name)
    temp_handle.close()
    actual_output = output
    try:
        doc.save(str(temp_path))
        saved_doc = Document(str(temp_path))
        saved_fingerprint = _layout_fingerprint(saved_doc)
        if saved_fingerprint != before_fingerprint:
            raise RuntimeError("File sau khi lưu không còn khớp dấu vân tay layout/style.")
        try:
            os.replace(temp_path, output)
        except PermissionError:
            actual_output = _next_fallback(output)
            os.replace(temp_path, actual_output)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    return {
        "status": "success",
        "template": str(template),
        "output": str(actual_output),
        "layout_fingerprint": before_fingerprint,
        "replacement_counts": counts,
    }


def _load_payload(value: str) -> dict[str, Any]:
    candidate = Path(value)
    if candidate.is_file():
        return json.loads(candidate.read_text(encoding="utf-8"))
    return json.loads(value)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Điền Phiếu yêu cầu NCC thanh toán từ mẫu chuẩn, không đổi layout."
    )
    parser.add_argument("--data", "-d", required=True, help="File JSON hoặc chuỗi JSON.")
    parser.add_argument("--output", "-o", required=True, type=Path)
    parser.add_argument("--template", "-t", type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args()

    payload = _load_payload(args.data)
    replacements = _parse_replacements(payload)
    result = fill_payment_request(args.template, args.output, replacements)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
