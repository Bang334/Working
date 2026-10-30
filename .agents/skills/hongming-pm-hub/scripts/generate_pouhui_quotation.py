# -*- coding: utf-8 -*-
"""
Script tạo báo giá hệ thống đo đếm điện năng 4 trạm biến áp cho Nhà máy Pou Hui Việt Nam
Chuẩn 1:1 theo mẫu ECO PRIME_BHT Báo Giá Đồng Hồ Điện_24-07-2026.xlsx
Bảo toàn 100% logo, font chữ, màu sắc, border, công thức tính tiền sống và bố cục in ấn.
"""

import openpyxl
from copy import copy
import math
import os
import sys
# Locate the project independently of the drive letter on this computer.
_WORKSPACE_ROOT = next((p for p in __import__('pathlib').Path(__file__).resolve().parents if (p / 'pm_portal' / 'app').is_dir()), __import__('pathlib').Path.cwd())
def _workspace_path(relative):
    return str(_WORKSPACE_ROOT / relative.replace(chr(92), '/'))

sys.stdout.reconfigure(encoding='utf-8')

def compute_cell_lines(text, col_width_chars):
    if not text:
        return 1
    paragraphs = str(text).split('\n')
    total_rendered = 0
    usable_width = max(col_width_chars - 2.5, 5.0)
    for p in paragraphs:
        if not p.strip():
            total_rendered += 1
            continue
        weight = 0.0
        for ch in p:
            code = ord(ch)
            if (0x4E00 <= code <= 0x9FFF) or (0x3400 <= code <= 0x4DBF) or (0xF900 <= code <= 0xFAFF) or code > 1279:
                weight += 1.95
            elif ch.isupper():
                weight += 1.2
            else:
                weight += 0.95
        lines_for_p = max(1, math.ceil(weight / usable_width))
        total_rendered += lines_for_p
    return max(1, total_rendered)

def copy_cell_style(src_cell, dst_cell):
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = copy(src_cell.number_format)
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)

def create_pouhui_quotation(template_path, output_path):
    wb = openpyxl.load_workbook(template_path, data_only=False)
    
    # Sử dụng sheet BHT làm sheet chính cho Pou Hui
    ws = wb['BHT']
    ws.title = 'PouHui'
    
    # Xóa sheet phụ BHT2 nếu có để file báo giá gửi khách hàng chỉ có 1 sheet duy nhất, chuyên nghiệp
    if 'BHT2' in wb.sheetnames:
        del wb['BHT2']
        
    # Độ rộng cột chuẩn của mẫu ECO PRIME
    col_widths = {
        'A': 5.33,
        'B': 33.33,
        'C': 30.44,
        'D': 52.78,
        'E': 27.33,
        'F': 24.55,
        'G': 25.33,
        'H': 39.22
    }
    for col_l, w in col_widths.items():
        ws.column_dimensions[col_l].width = w

    # --- 1. CẬP NHẬT THÔNG TIN NGƯỜI MUA (BÊN MUA - POU HUI VIETNAM) ---
    ws['F2'].value = "CÔNG TY TNHH VIỆT NAM PHỔ HUỆ Y TẾ CÔNG NGHỆ (POU HUI)"
    ws['F3'].value = "Ban Quản lý Dự án & Cơ điện 機電與專案部"
    ws['F4'].value = "(Liên hệ khách hàng)"
    ws['F5'].value = ""
    ws['F6'].value = "Lô D7/I, Đường số 1B, KCN Vĩnh Lộc, Phường Bình Hưng Hòa B, Q. Bình Tân, TP. Hồ Chí Minh"
    ws['F7'].value = "(Email khách hàng)"
    
    # --- 2. CẬP NHẬT TIÊU ĐỀ BÁO GIÁ & NGÀY THÁNG ---
    ws['A9'].value = "Báo giá hệ thống đo đếm điện năng 4 trạm biến áp 高壓變電站電能計量系統硬體報價訂購確認單"
    ws['A10'].value = "報價日期 Ngày báo giá: 21/10/2026"
    
    # --- 3. TIÊU ĐỀ GIẢI PHÁP & TỶ GIÁ ---
    ws['A12'].value = "1) Giải pháp nâng cấp hệ thống đo đếm điện năng và giám sát 4 trạm biến áp cao thế 普惠醫工4座高壓變電站電能計量與能耗監控系統方案"
    ws['H12'].value = 26130  # Tỷ giá USD/VND hiện hành

    # --- 4. DANH MỤC 8 HẠNG MỤC THIẾT BỊ THEO SƠ ĐỒ THIẾT KẾ & KHẢO SÁT ---
    items = [
        {
            "stt": 1,
            "name": "Đồng hồ đa năng PM210-4-V\nPM210-4-V 三相單迴路多功能需量總表",
            "category": "Đồng hồ đa năng\n多功能電表",
            "desc": "Đồng hồ đo đếm điện năng 3 pha đa chức năng PM210-4-V, đo công suất, điện áp, dòng điện, điện năng tiêu thụ (KWh/KVAh/KVARh), hỗ trợ kết nối trực tiếp biến dòng mềm Rogowski, giao tiếp Modbus RTU/RS485. 1 bộ / trạm biến áp (Trạm 1~Trạm 4).\nPM210-4-V三相單迴路多功能需量總表, 具備電力參數測量、電能計量功能, 直連羅氏線圈比流器, Modbus RTU/RS485通訊。每座變電站配置1台 (1號~4號高壓變電站)。",
            "qty": 4,
            "unit": "BỘ",
            "price": 320.0
        },
        {
            "stt": 2,
            "name": "Biến dòng mềm Rogowski CT-1000RV (1000A)\n三相羅氏比流器 CT-1000RV (1000A)",
            "category": "Biến dòng Rogowski\n羅氏比流器",
            "desc": "Biến dòng kiểu vòng mềm Rogowski coil CT-1000RV, dòng định mức 1000A, độ chính xác cao, lắp đặt kẹp mềm không cần cắt điện dừng máy nhà máy. Dùng cho Trạm 1, Trạm 2, Trạm 3 (máy biến áp 630KVA, cáp XLPE 325mm² x 2 sợi/pha).\nCT-1000RV三相羅氏比流器 (1000A), 柔性開口式線圈, 免停電快速施工安裝。適用於1號、2號、3號高壓變電站 (630KVA變壓器, XLPE 325mm²x2線徑)。",
            "qty": 3,
            "unit": "BỘ",
            "price": 145.0
        },
        {
            "stt": 3,
            "name": "Biến dòng mềm Rogowski CT-2000RV (2000A)\n三相羅氏比流器 CT-2000RV (2000A)",
            "category": "Biến dòng Rogowski\n羅氏比流器",
            "desc": "Biến dòng kiểu vòng mềm Rogowski coil CT-2000RV, dòng định mức 2000A (dải đo dòng sơ cấp 1200A), đường kính vòng lớn luồn qua bó 3 sợi cáp/pha, không cần cắt điện. Dùng cho Trạm 4 (máy biến áp 750KVA, cáp XLPE 250mm² x 3 sợi/pha).\nCT-2000RV三相羅氏比流器 (2000A), 大口徑柔性開口式線圈, 適用於單相3條電纜束線, 免停電快速安裝。適用於4號高壓變電站 (750KVA變壓器, XLPE 250mm²x3線徑)。",
            "qty": 1,
            "unit": "BỘ",
            "price": 185.0
        },
        {
            "stt": 4,
            "name": "Gateway thu thập dữ liệu CC2051\nCC2051 工業級資料收集閘道器",
            "category": "Gateway thu thập dữ liệu\n數據採集網關",
            "desc": "Bộ điều khiển/thu thập dữ liệu biên công nghiệp CC2051 (Edge Gateway), cổng RS-485 Modbus RTU thu thập tín hiệu từ các đồng hồ đo điện PM210, chuyển đổi sang Modbus TCP / MQTT qua cổng mạng RJ45 Ethernet, truyền dữ liệu liên tục về hệ thống EMS trung tâm.\nCC2051工業級邊緣計算與資料收集閘道器, 具備RS-485 Modbus RTU接口採集PM210電表數據, 經由RJ45以太網轉換為Modbus TCP/MQTT協議, 即時上傳至中央能源管理系統 (EMS)。",
            "qty": 2,
            "unit": "BỘ",
            "price": 275.0
        },
        {
            "stt": 5,
            "name": "Bộ nguồn công nghiệp 24VDC (PW)\n工業級電源供應器 (PW)",
            "category": "Nguồn cấp công nghiệp\n電源供應器",
            "desc": "Bộ nguồn công nghiệp gắn Din-rail Meanwell/tương đương, điện áp vào dải rộng AC 85~265V, đầu ra DC 24V ổn định, chống sét lan truyền và quá tải, cấp nguồn liên tục cho Gateway CC2051 và đồng hồ PM210. 1 bộ / trạm biến áp.\n工業級導軌式開關電源供應器, 寬電壓輸入 AC 85~265V, 輸出 DC 24V 穩定電源, 具備過載與防雷保護, 為 PM210 電表及 CC2051 閘道器供電。每座變電站1台。",
            "qty": 4,
            "unit": "BỘ",
            "price": 38.0
        },
        {
            "stt": 6,
            "name": "Bộ chia mạng công nghiệp 5-Port\n工業級5埠乙太網交換機 (LAN SWITCH)",
            "category": "Thiết bị mạng\n網絡設備",
            "desc": "Switch mạng công nghiệp 5 cổng RJ45 10/100Mbps gắn thanh ray, vỏ kim loại chuẩn công nghiệp chống nhiễu điện từ, kết nối truyền dữ liệu từ các Gateway CC2051 về mạng nội bộ LAN nhà máy.\n工業級5埠 10/100Mbps 導軌式交換機, 金屬抗干擾外殼, 適用於工廠高壓變電站環境, 將 CC2051 網關連接至廠區局域網 (LAN)。",
            "qty": 1,
            "unit": "BỘ",
            "price": 75.0
        },
        {
            "stt": 7,
            "name": "Cáp truyền thông RS485 & Vật tư phụ trọn gói\nRS485隔離通訊線及全套安裝輔材",
            "category": "Vật tư phụ & Cáp mạng\n工程輔材",
            "desc": "Cáp truyền thông chống nhiễu chuyên dụng UL2464 24AWG 2P (vỏ bọc chống nhiễu kép), cáp mạng Cat.5e/Cat.6 chống nhiễu UTP/FTP, ống luồn bảo vệ chống cháy, đầu cosse bọc cách điện, dây tiếp địa, cầu chì bảo vệ và phụ kiện đấu nối trọn gói.\nUL2464 24AWG 2P 工業隔離屏蔽通訊線、Cat.5e/Cat.6 屏蔽網線、阻燃保護套管、冷壓絕緣端子、接地線、保護保險絲及全套配線輔材。",
            "qty": 1,
            "unit": "GÓI",
            "price": 220.0
        },
        {
            "stt": 8,
            "name": "Chi phí thi công, đấu nối & cấu hình nghiệm thu\n現場安裝、接線、系統調試與技術移交",
            "category": "Dịch vụ kỹ thuật\n工程服務",
            "desc": "Nhân công kỹ thuật thi công lắp đặt tại 4 trạm biến áp nhà máy Pou Hui: định vị lắp đặt đồng hồ, luồn cuộn biến dòng Rogowski, kéo rải cáp truyền thông, đấu nối nguồn, cài đặt địa chỉ ID Modbus, kiểm tra truyền nhận dữ liệu qua Gateway CC2051 lên hệ thống, nghiệm thu và hướng dẫn vận hành.\n普惠工廠4座高壓變電站現場施工: 電表固定安裝、羅氏線圈套接、通訊線布放、電源接線、Modbus通信參數設定、CC2051網關聯調、系統驗收測試及操作移交。",
            "qty": 1,
            "unit": "GÓI",
            "price": 500.0
        }
    ]

    # --- 5. HỦY TẤT CẢ CÁC VÙNG MERGE TỪ HÀNG 20 TRỞ ĐI ĐỂ TÁI LẬP CHUẨN XÁC ---
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 20:
            ws.unmerge_cells(str(rng))

    # --- 6. GHI DỮ LIỆU 8 HẠNG MỤC (HÀNG 13 ĐẾN 20) ---
    start_row = 13
    for idx, item in enumerate(items):
        r = start_row + idx
        # Tính toán chiều cao hàng chống cắt chữ
        lines_b = compute_cell_lines(item["name"], col_widths['B'])
        lines_c = compute_cell_lines(item["category"], col_widths['C'])
        lines_d = compute_cell_lines(item["desc"], col_widths['D'])
        max_lines = max(lines_b, lines_c, lines_d)
        calc_height = max(75.0, max_lines * 19.5 + 16.0)
        ws.row_dimensions[r].height = calc_height
        
        # Format từng ô dựa trên mẫu hàng 13
        ref_row = 13
        for col_idx in range(1, 9):
            cell = ws.cell(r, col_idx)
            ref_cell = ws.cell(ref_row, col_idx)
            copy_cell_style(ref_cell, cell)
            
        # Điền giá trị
        ws.cell(r, 1).value = item["stt"]
        ws.cell(r, 1).alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
        
        ws.cell(r, 2).value = item["name"]
        ws.cell(r, 2).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        ws.cell(r, 3).value = item["category"]
        ws.cell(r, 3).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        ws.cell(r, 4).value = item["desc"]
        ws.cell(r, 4).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        ws.cell(r, 5).value = item["qty"]
        ws.cell(r, 5).alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
        ws.cell(r, 5).number_format = f'0" {item["unit"]} "'
        
        ws.cell(r, 6).value = item["price"]
        ws.cell(r, 6).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
        ws.cell(r, 6).number_format = '"US$"#,##0.0_);\\("US$"#,##0.0\\)'
        
        ws.cell(r, 7).value = f"=F{r}*E{r}"
        ws.cell(r, 7).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
        ws.cell(r, 7).number_format = '"US$"#,##0.0_);\\("US$"#,##0.0\\)'
        
        ws.cell(r, 8).value = f"=G{r}*H12"
        ws.cell(r, 8).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
        ws.cell(r, 8).number_format = '[$VND]\\ #,##0_);\\([$VND]\\ #,##0\\)'

    last_item_row = start_row + len(items) - 1  # 20

    # --- 7. HÀNG 以下空白 (HÀNG 21) ---
    blank_row = last_item_row + 1  # 21
    ws.row_dimensions[blank_row].height = 18.0
    for col_idx in range(1, 9):
        c = ws.cell(blank_row, col_idx)
        c.font = openpyxl.styles.Font(name="微軟正黑體", size=10.0, bold=False)
        c.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
        # Viền: trên thin, dưới medium
        c.border = openpyxl.styles.Border(
            top=openpyxl.styles.Side(style='thin', color='000000'),
            bottom=openpyxl.styles.Side(style='medium', color='000000'),
            left=openpyxl.styles.Side(style='medium', color='000000') if col_idx == 1 else None,
            right=openpyxl.styles.Side(style='medium', color='000000') if col_idx == 8 else None
        )
    ws.cell(blank_row, 1).value = "以下空白"
    ws.merge_cells(start_row=blank_row, start_column=1, end_row=blank_row, end_column=8)

    # --- 8. KHỐI TỔNG TIỀN (HÀNG 22, 23, 24) ---
    peach_fill = openpyxl.styles.PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    
    # Hàng 22: Tạm tính 小計
    r_sub = blank_row + 1  # 22
    ws.row_dimensions[r_sub].height = 38.0
    ws.cell(r_sub, 6).value = "Tạm tính 小計"
    ws.cell(r_sub, 6).fill = peach_fill
    ws.cell(r_sub, 6).font = openpyxl.styles.Font(name="Microsoft YaHei", size=12.0, bold=True)
    ws.cell(r_sub, 6).alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
    ws.cell(r_sub, 6).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='thin', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='thin', color='000000')
    )
    
    ws.cell(r_sub, 7).value = f"=SUM(G{start_row}:G{last_item_row})"
    ws.cell(r_sub, 7).font = openpyxl.styles.Font(name="Arial", size=13.0, bold=True)
    ws.cell(r_sub, 7).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    ws.cell(r_sub, 7).number_format = '"US$"#,##0.00_);\\("US$"#,##0.00\\)'
    ws.cell(r_sub, 7).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='thin', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='thin', color='000000')
    )
    
    ws.cell(r_sub, 8).value = f"=G{r_sub}*H12"
    ws.cell(r_sub, 8).font = openpyxl.styles.Font(name="Arial", size=13.0, bold=True)
    ws.cell(r_sub, 8).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    ws.cell(r_sub, 8).number_format = '[$VND]\\ #,##0_);\\([$VND]\\ #,##0\\)'
    ws.cell(r_sub, 8).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='medium', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='thin', color='000000')
    )

    # Hàng 23: Thuế VAT 8% 增值稅
    r_vat = r_sub + 1  # 23
    ws.row_dimensions[r_vat].height = 38.0
    ws.cell(r_vat, 6).value = "Thuế VAT 8% 增值稅"
    ws.cell(r_vat, 6).fill = peach_fill
    ws.cell(r_vat, 6).font = openpyxl.styles.Font(name="Microsoft YaHei", size=12.0, bold=True)
    ws.cell(r_vat, 6).alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
    ws.cell(r_vat, 6).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='thin', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='thin', color='000000')
    )
    
    ws.cell(r_vat, 7).value = f"=G{r_sub}*0.08"
    ws.cell(r_vat, 7).font = openpyxl.styles.Font(name="Arial", size=13.0, bold=True)
    ws.cell(r_vat, 7).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    ws.cell(r_vat, 7).number_format = '"US$"#,##0.00_);\\("US$"#,##0.00\\)'
    ws.cell(r_vat, 7).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='thin', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='thin', color='000000')
    )
    
    ws.cell(r_vat, 8).value = f"=H{r_sub}*0.08"
    ws.cell(r_vat, 8).font = openpyxl.styles.Font(name="Arial", size=13.0, bold=True)
    ws.cell(r_vat, 8).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    ws.cell(r_vat, 8).number_format = '[$VND]\\ #,##0_);\\([$VND]\\ #,##0\\)'
    ws.cell(r_vat, 8).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='medium', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='thin', color='000000')
    )

    # Hàng 24: Tổng cộng 總金額
    r_tot = r_vat + 1  # 24
    ws.row_dimensions[r_tot].height = 38.0
    ws.cell(r_tot, 6).value = "Tổng cộng 總金額"
    ws.cell(r_tot, 6).fill = peach_fill
    ws.cell(r_tot, 6).font = openpyxl.styles.Font(name="Microsoft YaHei", size=12.0, bold=True)
    ws.cell(r_tot, 6).alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
    ws.cell(r_tot, 6).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='thin', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='medium', color='000000')
    )
    
    ws.cell(r_tot, 7).value = f"=G{r_sub}*1.08"
    ws.cell(r_tot, 7).font = openpyxl.styles.Font(name="Arial", size=13.0, bold=True)
    ws.cell(r_tot, 7).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    ws.cell(r_tot, 7).number_format = '"US$"#,##0.00_);\\("US$"#,##0.00\\)'
    ws.cell(r_tot, 7).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='thin', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='medium', color='000000')
    )
    
    ws.cell(r_tot, 8).value = f"=H{r_sub}*1.08"
    ws.cell(r_tot, 8).font = openpyxl.styles.Font(name="Arial", size=13.0, bold=True)
    ws.cell(r_tot, 8).alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")
    ws.cell(r_tot, 8).number_format = '[$VND]\\ #,##0_);\\([$VND]\\ #,##0\\)'
    ws.cell(r_tot, 8).border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color='000000'),
        right=openpyxl.styles.Side(style='medium', color='000000'),
        top=openpyxl.styles.Side(style='thin', color='000000'),
        bottom=openpyxl.styles.Side(style='medium', color='000000')
    )

    # Merge vùng trống A22:E24 bên trái khối tổng cộng
    for r_m in range(r_sub, r_tot + 1):
        for c_m in range(1, 6):
            cell = ws.cell(r_m, c_m)
            cell.border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style='medium', color='000000') if c_m == 1 else None,
                bottom=openpyxl.styles.Side(style='medium', color='000000') if r_m == r_tot else None
            )
    ws.merge_cells(start_row=r_sub, start_column=1, end_row=r_tot, end_column=5)

    # --- 9. TIÊU ĐỀ KHÁCH HÀNG & GHI CHÚ (HÀNG 25) ---
    r_hdr = r_tot + 1  # 25
    ws.row_dimensions[r_hdr].height = 24.0
    
    # Bên trái: ※客戶確認 Xác nhận khách hàng (A25:C25)
    for c_idx in range(1, 4):
        c = ws.cell(r_hdr, c_idx)
        c.fill = peach_fill
        c.border = openpyxl.styles.Border(
            left=openpyxl.styles.Side(style='medium', color='000000') if c_idx == 1 else None,
            right=openpyxl.styles.Side(style='thin', color='000000') if c_idx == 3 else None,
            top=openpyxl.styles.Side(style='medium', color='000000'),
            bottom=None
        )
    ws.cell(r_hdr, 1).value = "※客戶確認 Xác nhận khách hàng"
    ws.cell(r_hdr, 1).font = openpyxl.styles.Font(name="Arial", size=12.0, bold=True)
    ws.cell(r_hdr, 1).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=r_hdr, start_column=1, end_row=r_hdr, end_column=3)

    # Bên phải: ◆ 備註聲明 Ghi chú và tuyên bố (D25:H25)
    for c_idx in range(4, 9):
        c = ws.cell(r_hdr, c_idx)
        c.fill = peach_fill
        c.border = openpyxl.styles.Border(
            left=openpyxl.styles.Side(style='thin', color='000000') if c_idx == 4 else None,
            right=openpyxl.styles.Side(style='medium', color='000000') if c_idx == 8 else None,
            top=openpyxl.styles.Side(style='medium', color='000000'),
            bottom=None
        )
    ws.cell(r_hdr, 4).value = "◆ 備註聲明"
    ws.cell(r_hdr, 4).font = openpyxl.styles.Font(name="Microsoft YaHei", size=11.0, bold=True)
    ws.cell(r_hdr, 4).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=r_hdr, start_column=4, end_row=r_hdr, end_column=8)

    # --- 10. CHỮ KÝ VÀ 3 ĐIỀU KHOẢN (HÀNG 26, 27, 28) ---
    r_n1 = r_hdr + 1  # 26
    r_n2 = r_n1 + 1   # 27
    r_n3 = r_n2 + 1   # 28
    
    ws.row_dimensions[r_n1].height = 54.0
    ws.row_dimensions[r_n2].height = 54.0
    ws.row_dimensions[r_n3].height = 58.0

    notes_text = [
        "1. 以上報價交期：報價單簽回後，另日程安排後回覆；\n   Thời gian giao hàng của báo giá trên: sau khi báo giá được ký và gửi lại, lịch trình cụ thể sẽ được thông báo sau;",
        "2. 報價有效期限：2026/11/21 止，報價簽回視同專案確認，不可任意終止、取消本專案；交貨後保固一年\n    Thời hạn hiệu lực của báo giá: đến hết ngày 21/11/2026. Việc ký và gửi lại báo giá được xem là xác nhận dự án, không được phép đơn phương chấm dứt hoặc hủy bỏ dự án này; Bảo hành 1 năm kể từ ngày giao hàng.",
        "3. 付款方式 : 訂金 35%、 交貨 55%、驗收 10% TT30Day。\n     Điều khoản thanh toán: [Đặt cọc: 35%, Khi giao hàng: 55%, Sau khi nghiệm thu: Thanh toán 10% còn lại theo điều khoản TT 30 ngày (TT30 Days).]"
    ]

    for n_idx, (r_curr, n_text) in enumerate(zip([r_n1, r_n2, r_n3], notes_text)):
        is_last = (n_idx == 2)
        # Bên trái: Signature col A-C
        for c_idx in range(1, 4):
            c = ws.cell(r_curr, c_idx)
            c.fill = peach_fill
            c.border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style='medium', color='000000') if c_idx == 1 else None,
                right=openpyxl.styles.Side(style='thin', color='000000') if c_idx == 3 else None,
                bottom=openpyxl.styles.Side(style='medium', color='000000') if is_last else None
            )
        # Bên phải: Notes col D-H
        for c_idx in range(4, 9):
            c = ws.cell(r_curr, c_idx)
            c.fill = peach_fill
            c.border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style='thin', color='000000') if c_idx == 4 else None,
                right=openpyxl.styles.Side(style='medium', color='000000') if c_idx == 8 else None,
                bottom=openpyxl.styles.Side(style='medium', color='000000') if is_last else None
            )
        ws.cell(r_curr, 4).value = n_text
        ws.cell(r_curr, 4).font = openpyxl.styles.Font(name="Microsoft YaHei", size=10.5, bold=True)
        ws.cell(r_curr, 4).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r_curr, start_column=4, end_row=r_curr, end_column=8)

    # Merge ô chữ ký A26:C28
    ws.cell(r_n1, 1).value = "（請簽名用印） Chữ ký và con dấu"
    ws.cell(r_n1, 1).font = openpyxl.styles.Font(name="Arial", size=11.0, bold=True)
    ws.cell(r_n1, 1).alignment = openpyxl.styles.Alignment(horizontal="left", vertical="top")
    ws.merge_cells(start_row=r_n1, start_column=1, end_row=r_n3, end_column=3)

    # --- 11. THIẾT LẬP VÙNG IN & CHẾ ĐỘ XEM CHUẨN 1:1 ---
    ws.print_area = f"$A$1:$H${r_n3}"
    ws.views.sheetView[0].view = "pageBreakPreview"
    ws.views.sheetView[0].zoomScale = 60
    ws.views.sheetView[0].showGridLines = None  # Không làm loang rìa ngoài

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.scale = 38
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Lưu an toàn bọc try-except
    try:
        wb.save(output_path)
        print(f"Đã lưu thành công file báo giá Pou Hui tại: {output_path}")
    except PermissionError:
        alt_path = output_path.replace(".xlsx", "_v2.xlsx")
        wb.save(alt_path)
        print(f"File gốc đang mở, đã lưu dự phòng tại: {alt_path}")

if __name__ == "__main__":
    tmpl = _workspace_path('ChuaHieu\\ECO PRIME_BHT Báo Giá Đồng Hồ Điện_24-07-2026.xlsx')
    out1 = _workspace_path('ChuaHieu\\ECO PRIME_Báo Giá Đồng Hồ Điện & CC2051_Pou Hui_20261021.xlsx')
    create_pouhui_quotation(tmpl, out1)
    
    # Đồng bộ 1 bản ra thư mục gốc E:\Working
    out2 = _workspace_path('ECO PRIME_Báo Giá Đồng Hồ Điện & CC2051_Pou Hui_20261021.xlsx')
    create_pouhui_quotation(tmpl, out2)
