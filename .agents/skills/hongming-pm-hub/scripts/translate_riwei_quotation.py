# -*- coding: utf-8 -*-
"""
Script dịch toàn bộ Báo giá Sunstige (日威) sang song ngữ Việt - Trung
File: ●SQ2026日威報價單(SMB-MES-QMS-WMS)-海外.xlsx
Dịch toàn bộ 6 sheet:
1. 報價金額與經銷價參考
2. SMB機聯網
3. MES
4. QMS
5. WMS
6. IQA關鍵製程(MES Lite)
Giữ nguyên dòng tiếng Trung, đặt dòng tiếng Việt chuẩn thuật ngữ ngay bên dưới.
Tự động căn chỉnh chiều cao dòng (Anti-Clipping) và bật Wrap Text.
"""

import os
import sys
import openpyxl
from openpyxl.styles import Alignment, Font
from copy import copy

# Locate the project independently of the drive letter on this computer.
_WORKSPACE_ROOT = next((p for p in __import__('pathlib').Path(__file__).resolve().parents if (p / 'pm_portal' / 'app').is_dir()), __import__('pathlib').Path.cwd())
def _workspace_path(relative):
    return str(_WORKSPACE_ROOT / relative.replace(chr(92), '/'))

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Từ điển ánh xạ toàn bộ 356 chuỗi tiếng Trung sang tiếng Việt chuẩn chuyên ngành
TRANSLATIONS = {
    # 1. Thông tin chung, Tiêu đề & Pháp nhân
    "日威科技股份有限公司": "Công ty Cổ phần Công nghệ Sunstige (Sunstige Technology Co., Ltd.)",
    "報價訂購確認單": "BẢNG BÁO GIÁ KIÊM ĐƠN ĐẶT HÀNG (QUOTATION & ORDER CONFIRMATION)",
    "報價訂購確認單(MES)": "BẢNG BÁO GIÁ KIÊM ĐƠN ĐẶT HÀNG (HỆ THỐNG MES)",
    "買方名稱：": "Tên Bên Mua:",
    "賣方名稱：": "Tên Bên Bán:",
    "聯 絡 人：": "Người liên hệ:",
    "聯 絡 人： ": "Người liên hệ:",
    "曾家彥": "Tăng Gia Ngạn (Ray Tseng)",
    "聯絡電話：": "Số điện thoại:",
    "傳真電話：": "Số Fax:",
    "聯絡地址：": "Địa chỉ liên lạc:",
    "330桃園市桃園區永安路1015號4樓": "Tầng 4, Số 1015 đường Vĩnh An, Quận Đào Viên, TP. Đào Viên 330, Đài Loan",
    "電子郵件：": "Email:",
    "報價日期：2025年10月8日": "Ngày báo giá: Ngày 08 tháng 10 năm 2025",
    "報價單號：": "Số báo giá:",
    "大項": "Hạng mục lớn (Phân hệ)",
    "報價項目": "Hạng mục báo giá",
    "單價(Taiwan終端)": "Đơn giá (Khách hàng cuối Đài Loan)",
    "海外代理參考價格": "Giá tham khảo đại lý nước ngoài",
    "備註": "Ghi chú",
    "產品名稱": "Tên sản phẩm",
    "項目": "Hạng mục / Module",
    "產 品 內 容 說 明": "MÔ TẢ CHI TIẾT NỘI DUNG SẢN PHẨM",
    "數量": "Số lượng",
    "客戶單價": "Đơn giá khách hàng",
    "單價": "Đơn giá",
    "經銷價": "Giá đại lý",
    "小計": "Thành tiền",
    "總價": "Tổng giá (Chưa VAT)",
    "稅額": "Tiền thuế",
    "加總": "Tổng cộng thanh toán",
    "一式": "1 Gói",
    "1式": "1 Gói",
    "1套": "1 Bộ",
    "個案": "Theo từng dự án cụ thể",
    "個案\n(選配)": "Theo dự án cụ thể\n(Tùy chọn thêm)",
    "個案客製(依小時計價)": "Tùy biến theo dự án (Tính phí theo giờ công)",
    "個案Excel 工單匯入": "Import lệnh sản xuất từ Excel theo yêu cầu dự án",
    "※客戶確認（請簽名用印）": "XÁC NHẬN CỦA KHÁCH HÀNG (Vui lòng ký tên và đóng dấu)",
    "◆ 備註聲明": "◆ GHI CHÚ VÀ TUYÊN BỐ ĐIỀU KHOẢN",

    # 2. Bảng giá & Đại lý nước ngoài (Sheet 1)
    "由比例50%、30%、20%等分原則做拆分\n50 (產品-日威)\n30 (顧問)\n20 (業務)": 
        "Phân chia theo nguyên tắc tỷ lệ 50%, 30%, 20%:\n50% (Sản phẩm - Sunstige)\n30% (Chuyên gia tư vấn giải pháp)\n20% (Kinh doanh / Sales)",
    "SMB 射出機(IO連線)-Web系統": "Hệ thống Web kết nối SMB cho máy ép nhựa (Kết nối I/O)",
    "1台15000元\n一案30萬(20台)，多加機台新增費用": "15.000 TWD/máy\nGói dự án 300.000 TWD (20 máy), thêm máy tính thêm chi phí",
    "7500/台": "7.500 TWD/máy",
    "未滿10台，另收安裝費2萬\n安裝擷取設備，取得設備狀態、生產次數、計算週期時間": 
        "Dưới 10 máy, thu thêm phí lắp đặt 20.000 TWD\nLắp đặt thiết bị thu thập dữ liệu, lấy trạng thái máy, số lần sản xuất, tính thời gian chu kỳ (Cycle Time)",
    "SMB 射出機(完整連線)-Web系統": "Hệ thống Web kết nối SMB cho máy ép nhựa (Kết nối hoàn chỉnh qua Controller)",
    "1台20000元\n一案50萬(25台)，多加機台新增費用": "20.000 TWD/máy\nGói dự án 500.000 TWD (25 máy), thêm máy tính thêm chi phí",
    "10000/台": "10.000 TWD/máy",
    "未滿10台，另收安裝費2萬\n與控制器通訊，取得設備設定參數、實際涮": 
        "Dưới 10 máy, thu thêm phí lắp đặt 20.000 TWD\nGiao tiếp với bộ điều khiển, lấy thông số cài đặt máy, thông số thực tế",
    "IQA 關鍵製程管理系統": "Hệ thống quản lý công đoạn trọng yếu IQA (MES Lite)",
    "750000 (25台)，多加機台新增費用": "750.000 TWD (cho 25 máy), thêm máy tính thêm chi phí",
    "15000/台": "15.000 TWD/máy",
    "未滿10台，另收安裝費5萬": "Dưới 10 máy, thu thêm phí lắp đặt 50.000 TWD",
    "MES 生產製造管理系統": "Hệ thống quản lý sản xuất chế tạo MES",
    "1000000/套": "1.000.000 TWD/bộ",
    "50台設備為基準，多一台加 5000": "Tiêu chuẩn định mức 50 thiết bị, thêm 1 máy cộng thêm 5.000 TWD",
    "模具管理(MES選配)": "Quản lý khuôn mẫu (Module tùy chọn cho MES)",
    "180000/套": "180.000 TWD/bộ",
    "設備管理(MES選配)": "Quản lý thiết bị máy móc (Module tùy chọn cho MES)",
    "生產前動作(MES選配)": "Quản lý thao tác chuẩn bị trước sản xuất (Module tùy chọn cho MES)",
    "90000/套": "90.000 TWD/bộ",
    "QMS品質管理系統\n(首檢/巡檢/進料檢/成品檢/出貨檢)": "Hệ thống quản lý chất lượng QMS\n(Kiểm tra đầu ca / Tuần tra / Đầu vào / Thành phẩm / Xuất xưởng)",
    "400000/套": "400.000 TWD/bộ",
    "QMS品質管理系統(首檢/巡檢)": "Hệ thống quản lý chất lượng QMS (Kiểm tra đầu ca / Tuần tra)",
    "250000/套": "250.000 TWD/bộ",
    "WMS倉儲管理系統": "Hệ thống quản lý kho vận WMS",
    "600000/套": "600.000 TWD/bộ",
    "其他": "Hạng mục khác",
    "串接ERP": "Tích hợp hệ thống ERP",
    "150000~300000/式": "150.000 ~ 300.000 TWD/gói",
    "串接內容依照與ERP廠商、客戶溝通後評估難易度報價，此價格為基本價格，會調整": 
        "Nội dung tích hợp được báo giá dựa trên đánh giá độ phức tạp sau khi trao đổi với NCC ERP và khách hàng; đây là mức giá cơ sở và có thể điều chỉnh",
    "系統無維護狀態工時": "Chi phí giờ công khi hệ thống hết bảo hành và không ký hợp đồng bảo trì",
    "系統已過保固並且不簽維護，工時評估，最低工時4hr，2200元/時\n流程:確認-->報價-->回簽-->進行": 
        "Hệ thống đã hết bảo hành và không ký bảo trì, đánh giá giờ công tối thiểu 4 giờ, đơn giá 2.200 TWD/giờ\nQuy trình: Xác nhận --> Báo giá --> Ký xác nhận --> Thực hiện",
    "IO硬體模組 ": "Module phần cứng I/O",
    "可提供BOM表由SI自行採購": "Có thể cung cấp bảng BOM vật tư để SI (đơn vị tích hợp) tự mua",
    "RS232/485轉RJ45硬體": "Phần cứng chuyển đổi RS232/485 sang RJ45 (Ethernet)",
    "5000/台": "5.000 TWD/chiếc",
    "80000/式": "80.000 TWD/gói",
    "可提供廠牌由SI自行採購": "Có thể cung cấp danh sách nhãn hiệu để SI tự mua",
    "設備通訊費用": "Chi phí giao tiếp / truyền thông thiết bị",
    "可通訊設備分為三類型，\n本報價為射出機、CNC參考\n若為PLC則依實際處理內容報價": 
        "Thiết bị có thể giao tiếp chia làm 3 loại:\nBáo giá này là giá tham khảo cho máy ép nhựa và máy CNC;\nNếu là PLC thì báo giá theo nội dung xử lý thực tế",
    "工業電腦 (Server)": "Máy tính công nghiệp (Máy chủ Server)",
    "160000/台": "160.000 TWD/bộ",
    "MES、QMS、WMS主機適用\n1. 工業伺服器，研華ACP-4340\n     CPU: i7\n     RAM: 64G\n     Storage: 500 SSD+ 2TB*4 HDD (Raid 6)\n2. 作業系統: Win Server IoT 2022 Standard + 5 Cal\n3. 資料庫: SQL Svr Std RUNTIME 2022 IoT ESD OEI + 5 Clt Std (SERVER綁定版)": 
        "Áp dụng cho máy chủ MES, QMS, WMS:\n1. Server công nghiệp Advantech ACP-4340 (CPU i7, RAM 64GB, Ổ cứng 500GB SSD + 2TB*4 HDD - Raid 6)\n2. Hệ điều hành: Win Server IoT 2022 Standard + 5 CAL\n3. Cơ sở dữ liệu: SQL Server Standard RUNTIME 2022 IoT ESD OEI + 5 Client Standard (Bản quyền gắn theo Server)",
    "工業電腦 (設備資料收集)": "Máy tính công nghiệp IPC (Thu thập dữ liệu thiết bị)",
    "52000/台": "52.000 TWD/bộ",
    "SMB機聯網主機適用\n工業電腦 IPC(一廠區一台或25台設備左右一台)\n1. 研華 MIC 770\n    CPU: i5-13500,\n    RAM: DDR5 SO-DIMM 32GBRAM\n    Storage: 500GB SSD+ 500G HDD, \n2. 作業系統: Windows 10 IOT\n3. 資料庫: open source (MariaDB+Clickhouse)": 
        "Áp dụng cho máy chủ SMB IoT kết nối máy:\nMáy tính công nghiệp IPC (1 chiếc cho 1 phân xưởng hoặc khoảng 25 thiết bị)\n1. Advantech MIC 770 (CPU i5-13500, RAM DDR5 SO-DIMM 32GB, Ổ cứng 500GB SSD + 500GB HDD)\n2. Hệ điều hành: Windows 10 IoT\n3. Cơ sở dữ liệu: Mã nguồn mở (MariaDB + Clickhouse)",

    # 3. Điều khoản & Ghi chú (Các Sheet 2, 3, 4, 5, 6)
    "1.報價簽回視同專案確認，不可任意終止、取消本專案；報價有效日期15日.": 
        "1. Việc ký xác nhận gửi lại báo giá được xem như xác nhận dự án, không được tùy ý chấm dứt hoặc hủy bỏ dự án; Báo giá có hiệu lực trong vòng 15 ngày.",
    "1.報價簽回視同專案確認，不可任意終止、取消本專案；報價有效日期30日.": 
        "1. Việc ký xác nhận gửi lại báo giá được xem như xác nhận dự án, không được tùy ý chấm dứt hoặc hủy bỏ dự án; Báo giá có hiệu lực trong vòng 30 ngày.",
    "2.付款方式: (1)訂金30%，於簽約後30日付款。\n                     (2)安裝完成50%。\n                     (3)驗收20%，驗收後30日付款。": 
        "2. Phương thức thanh toán:\n   (1) Đặt cọc 30%, thanh toán trong vòng 30 ngày sau khi ký hợp đồng.\n   (2) Hoàn thành lắp đặt thanh toán 50%.\n   (3) Nghiệm thu 20%, thanh toán trong vòng 30 ngày sau khi nghiệm thu.",
    "2.付款方式: (1)訂金30%，於簽約後30日付款。\n                   (2)安裝完成50%。\n                   (3)驗收20%(驗收後30日付款)。": 
        "2. Phương thức thanh toán:\n   (1) Đặt cọc 30%, thanh toán trong vòng 30 ngày sau khi ký hợp đồng.\n   (2) Hoàn thành lắp đặt thanh toán 50%.\n   (3) Nghiệm thu 20%, thanh toán trong vòng 30 ngày sau khi nghiệm thu.",
    "2.付款方式:\n                     (1)訂金30%，於簽約後付款(簽約後30日付款)。\n                     (2)系統安裝完成付款50%。\n                     (3)驗收後付款20%(驗收後30日付款)。\n      註：付款方式:轉帳、匯款、支票擇一。 ": 
        "2. Phương thức thanh toán:\n   (1) Đặt cọc 30%, thanh toán trong vòng 30 ngày sau khi ký hợp đồng.\n   (2) Hoàn thành lắp đặt hệ thống thanh toán 50%.\n   (3) Nghiệm thu thanh toán 20% (thanh toán trong vòng 30 ngày sau khi nghiệm thu).\n   *Ghi chú: Phương thức thanh toán có thể chọn chuyển khoản, nộp tiền hoặc séc.",
    " 日威帳戶：台灣中小企銀  金融銀行代號：0503006 帳號：30012192688 ": 
        "Tài khoản Sunstige: Ngân hàng Taiwan Business Bank - Mã ngân hàng: 0503006 - Số tài khoản: 30012192688",
    "3.系統於驗收後，保固一年，如需延長保固，一年維護費為報價單費用10%，另訂維護合約。": 
        "3. Hệ thống được bảo hành 01 năm sau khi nghiệm thu; nếu cần gia hạn bảo hành, phí bảo trì hàng năm bằng 10% giá trị báo giá và sẽ ký hợp đồng bảo trì riêng.",
    "3.系統於驗收後，保固一年，如需延長保固，一年維護費用為報價單費用10%。": 
        "3. Hệ thống được bảo hành 01 năm sau khi nghiệm thu; nếu cần gia hạn bảo hành, phí bảo trì hàng năm bằng 10% giá trị báo giá.",
    " 4、相關硬體另外評估與購買(例:手持裝置、工作站或平板)": 
        "4. Phần cứng liên quan sẽ được đánh giá và mua riêng (ví dụ: máy quét cầm tay PDA, máy trạm hoặc máy tính bảng)",
    "\nMS備註  : \n           1、相關硬體另外評估與購買(例:手持裝置、工作站或平板)": 
        "Ghi chú MES:\n1. Phần cứng liên quan sẽ được đánh giá và mua riêng (ví dụ: máy cầm tay PDA, máy trạm hoặc máy tính bảng)",
    "備註：\n1.  \n建議可提供電視螢幕供現場使用，自行接硬體顯卡呈現畫面(瀏覽器：Chrome)\n2. 提供遠端連線：能快速進行遠端處理，若無對外連線，到場系統維護會有相關的維護費用(車馬費與工時費)15000/日(工作日)會依照北中南調整出差價格\n3.如使用IO、外掛SENSER，部分參數欄位會無資料": 
        "Ghi chú:\n1. Khuyến nghị chuẩn bị màn hình TV tại xưởng, tự kết nối xuất hình để hiển thị (Trình duyệt: Chrome)\n2. Cung cấp kết nối từ xa: Giúp xử lý sự cố nhanh chóng; nếu không có kết nối ra ngoài, việc kỹ sư đến hiện trường bảo trì sẽ tính chi phí (công tác phí và giờ công) 15.000 TWD/ngày làm việc (điều chỉnh theo khu vực)\n3. Nếu dùng I/O hoặc cảm biến ngoài, một số trường thông số nội tại của máy sẽ không có dữ liệu",
    "備註 ：\n1.網路架構：各機台網路線→Switch/Hub→工業電腦 IPC→SERVER\n2. 建議可提供電視螢幕供現場使用，自行接硬體顯卡呈現畫面(瀏覽器：Chrome)\n    現場自備平板(建議10吋以上，支援Android系統)或電腦供現場報工使用\n3. 提供遠端連線：能快速進行遠端處理，若無對外連線，到場系統維護會有相關的維護費用(車馬費與工時費)": 
        "Ghi chú:\n1. Kiến trúc mạng: Dây mạng máy móc -> Switch/Hub -> Máy tính công nghiệp IPC -> SERVER\n2. Khuyến nghị chuẩn bị TV tại xưởng để trình chiếu (Trình duyệt: Chrome); Xưởng tự trang bị máy tính bảng (khuyến nghị 10 inch trở lên, Android) hoặc máy tính để báo công\n3. Cung cấp kết nối từ xa: Giúp xử lý sự cố nhanh; nếu không có kết nối từ xa, bảo trì tại chỗ sẽ tính thêm chi phí đi lại và giờ công",
    "MES備註 ： \n            1、工廠須提供現場網路環境、遠端連線：Switch\\hub的裝設完成，包含SERVER或現場硬體使用的\n                   網路環境)， 網域資安安全須自行評估與負責\n            2、提供遠端連線：當系統遇到問題，能快速進行遠端處理，若無對外連線，後續維護可能會有相關\n                  費用，車馬費與工時費)\n            3、需提供在公司內部網域使用的固定IP\n                  (機台網路環境設定為192.168.20.x的網域，機台IP設定範圍由192.168.20.11~192.168.20.199)\n                  #網路架構：各機台網路線→Switch/Hub→工業電腦 IPC→SERVER\n            4、 建議可提供電視螢幕供現場使用，自行接硬體顯卡呈現畫面(瀏覽器：Chrome)\n                  現場自備平板/支援Android或電腦供現場報工使用(建議10吋以上)，如使用IOS，請提供硬體給\n                 日威做測試使用。": 
        "Ghi chú MES:\n1. Nhà máy cần cung cấp môi trường mạng tại xưởng và kết nối từ xa: Hoàn thành lắp đặt Switch/Hub cho Server và phần cứng hiện trường; an toàn an ninh mạng do nhà máy tự đánh giá và chịu trách nhiệm\n2. Cung cấp kết nối từ xa: Giúp xử lý nhanh khi hệ thống gặp sự cố; nếu không có kết nối ra ngoài, bảo trì sau này sẽ phát sinh chi phí đi lại và giờ công\n3. Cần cấp IP tĩnh trong mạng nội bộ công ty (Dải IP mạng máy móc là 192.168.20.x, IP máy từ 192.168.20.11 ~ 192.168.20.199; Kiến trúc: Cáp mạng máy -> Switch -> IPC -> SERVER)\n4. Khuyến nghị TV tại xưởng hiển thị màn hình (Chrome); tự trang bị máy tính bảng Android (khuyến nghị >=10 inch) hoặc máy tính để báo công; nếu dùng iOS vui lòng cấp thiết bị để Sunstige thử nghiệm",
    "軟體報價為標準內容，如有客製化需求另外評估開發工時。硬體為建議內容，會依工廠管理需求評估所需硬體與費用": 
        "Báo giá phần mềm là nội dung tiêu chuẩn; nếu có yêu cầu tùy biến sẽ đánh giá riêng giờ công phát triển. Phần cứng là đề xuất; sẽ đánh giá cấu hình và chi phí theo nhu cầu thực tế của nhà máy.",

    # 4. Sheet SMB機聯網
    "WEB設備可視與資訊管理系統(SMB機聯網)": "Hệ thống quản lý thông tin và trực quan hóa thiết bị trên nền tảng WEB (SMB Machine IoT)",
    "通訊連線": "Kết nối giao tiếp truyền thông",
    "跟控制器通訊，另外報價": "Giao tiếp với bộ điều khiển (Controller), báo giá riêng",
    "設備資訊": "Thông tin thiết bị",
    "設備運行現況：全廠設備稼動率比較": "Hiện trạng vận hành thiết bị: So sánh tỷ lệ hoạt động (OEE / 稼動率) toàn nhà máy",
    "設備狀況：稼動歷程、生產次數累計(狀態持續時間、生產模次、生產週期、設備稼動率)": 
        "Tình trạng thiết bị: Lịch sử vận hành, tích lũy số lần sản xuất (Thời lượng trạng thái, số khuôn dập/ép, chu kỳ sản xuất, tỷ lệ hoạt động)",
    "設備相關資訊：實際值、設定值、設備歷程記錄、異常記錄、設備狀況": 
        "Thông tin liên quan thiết bị: Giá trị thực tế, giá trị cài đặt, nhật ký vận hành thiết bị, nhật ký bất thường, tình trạng thiết bị",
    "統計圖表": "Biểu đồ thống kê",
    "狀況統計：全廠設備運行狀態統計查詢": "Thống kê trạng thái: Tra cứu thống kê trạng thái vận hành của toàn bộ thiết bị nhà máy",
    "日稼動率比較圖：選擇多機台及班別進行當日稼動率比較": "Biểu đồ so sánh tỷ lệ hoạt động hàng ngày: Chọn nhiều máy và ca làm việc để so sánh tỷ lệ hoạt động trong ngày",
    "月稼動率比較圖：選擇多機台進行當月稼動率比較": "Biểu đồ so sánh tỷ lệ hoạt động hàng tháng: Chọn nhiều máy để so sánh tỷ lệ hoạt động trong tháng",
    "稼動率趨勢圖：單一設備時間區間內稼動趨勢變化": "Biểu đồ xu hướng tỷ lệ hoạt động: Biến động xu hướng hoạt động của từng máy theo khoảng thời gian",
    "稼動歷程：單一設備稼動歷程及生產次數累計查詢\n#日期、班別、設備": 
        "Lịch sử hoạt động: Tra cứu lịch sử hoạt động và tích lũy số lần sản xuất của từng máy\n#Theo ngày, ca làm việc, thiết bị",
    "參數運行監視\n(通訊機台)": "Giám sát vận hành thông số\n(Áp dụng cho máy kết nối truyền thông)",
    "參數圖表：繪製參數運行趨勢圖": "Biểu đồ thông số: Vẽ đồ thị xu hướng biến thiên thông số vận hành",
    "相關性: 參數間關連性圖表": "Tính tương quan: Biểu đồ phân tích mối liên hệ tương quan giữa các thông số",
    "設定值歷程記錄：查詢參數設定值變更紀錄": "Nhật ký giá trị cài đặt: Tra cứu lịch sử thay đổi thông số cài đặt",
    "成型條件\n(通訊機台)": "Điều kiện tạo hình / đúc ép\n(Áp dụng cho máy kết nối truyền thông)",
    "設定值檢視：檢視參數設定值": "Xem giá trị cài đặt: Xem lại các giá trị thông số đã thiết lập",
    "設定值管理：管理已設定的模板": "Quản lý giá trị cài đặt: Quản lý các mẫu (template) thông số đã lưu",
    "異常紀錄\n(通訊機台)": "Nhật ký sự cố bất thường\n(Áp dụng cho máy kết nối truyền thông)",
    "異常紀錄觸發時間及其值": "Thời gian kích hoạt và giá trị ghi nhận của sự cố bất thường",
    "警報中心": "Trung tâm cảnh báo",
    "警報履歷": "Lịch sử cảnh báo",
    "廠區運行總覽：可自訂設備layout位置": "Tổng quan vận hành xưởng: Có thể tùy chỉnh sơ đồ bố trí (Layout) thiết bị trong xưởng",
    "電視看板": "Màn hình TV hiển thị trực quan (Andon/Kanban TV)",
    "今日生產訊息": "Thông tin sản xuất trong ngày",
    "平板操作系統-設備管理": "Hệ điều hành trên máy tính bảng - Quản lý thiết bị",
    "開始保養操作紀錄": "Ghi nhận thao tác bắt đầu bảo trì máy",
    "設備異常單: 單據列表、新增單據、異常分級、異常內容、上傳資料": 
        "Phiếu sự cố thiết bị: Danh sách phiếu, tạo mới phiếu, phân cấp sự cố, nội dung bất thường, tải lên tài liệu/hình ảnh",
    "設備保養/維修單: 單據列表、保養/維修狀態、保養工時": 
        "Phiếu bảo dưỡng/sửa chữa thiết bị: Danh sách phiếu, trạng thái bảo dưỡng/sửa chữa, số giờ công bảo dưỡng",
    "設備檢驗項目建立": "Thiết lập các hạng mục kiểm tra định kỳ thiết bị",
    "設定": "Cài đặt hệ thống",
    "廠區管理、機台管理、參數管理、狀態管理、班別管理(固定班別，可告知修改)、設備總覽欄位管理、選單管理、系統設定": 
        "Quản lý khu xưởng, quản lý máy móc, quản lý thông số, quản lý trạng thái, quản lý ca kíp (ca cố định, có thể yêu cầu sửa đổi), quản lý trường dữ liệu tổng quan thiết bị, quản lý menu chức năng, cài đặt hệ thống",
    "帳號與權限管理": "Quản lý tài khoản và phân quyền",
    "1、角色管理(系統管理員、觀察員、作業員)\n2、用戶管理(編號、名稱、角色、密碼)": 
        "1. Quản lý vai trò (Quản trị hệ thống, Người giám sát/quan sát, Công nhân vận hành)\n2. Quản lý người dùng (Mã số, họ tên, vai trò, mật khẩu)",
    "異質系統串接\n(選配)": "Tích hợp hệ thống bên thứ ba\n(Tùy chọn thêm)",
    "異質系統串接、功能開發、報表、憑證、標籤機…等": 
        "Tích hợp hệ thống bên thứ ba, phát triển tính năng tùy chỉnh, báo cáo biểu mẫu, chứng từ, máy in tem nhãn...vv",

    # 5. Sheet MES
    "SUNSTIGE MES- 生產製造管理系統\n(50台)": "HỆ THỐNG ĐIỀU HÀNH SẢN XUẤT SUNSTIGE MES\n(Quy mô 50 máy)",
    "基本資料管理": "Quản lý dữ liệu gốc (Master Data)",
    "工廠: 建立廠區": "Nhà máy: Thiết lập các khu vực nhà xưởng",
    "產線: 建立廠區產線": "Chuyền sản xuất: Thiết lập các dây chuyền trong xưởng",
    "製程: 建立廠內製程": "Công đoạn: Thiết lập các công đoạn sản xuất nội bộ",
    "設備/工作台: 設備/工作台名稱、所屬產線": "Thiết bị / Bàn thao tác: Tên thiết bị/bàn làm việc, thuộc chuyền sản xuất nào",
    "材料: 材料編號、材料名稱、材料種類": "Nguyên vật liệu: Mã vật liệu, tên vật liệu, chủng loại vật liệu",
    "產品: 產品編號、名稱、規格、可用模具、加工流程、標準工時、標準批量": 
        "Sản phẩm: Mã sản phẩm, tên, quy cách, khuôn tương thích, quy trình gia công, định mức giờ công tiêu chuẩn, cỡ lô tiêu chuẩn",
    "加工流程: 建立產品生產流程順序": "Quy trình gia công: Thiết lập trình tự công đoạn sản xuất sản phẩm (Routing)",
    "異常原因管理: 建立不同製程對應的生產異常原因": "Quản lý nguyên nhân bất thường: Thiết lập danh mục lý do sự cố tương ứng từng công đoạn",
    "工作日管理: 設定工作日、休假日": "Quản lý ngày làm việc: Cài đặt ngày làm việc, ngày nghỉ lễ/ca",
    "線別休息區段: 建立計畫休息時間": "Thời gian nghỉ giải lao theo chuyền: Thiết lập thời gian nghỉ có kế hoạch",
    "工單系統管理": "Quản lý hệ thống lệnh sản xuất (Work Order Management)",
    "工單管理: 新增工單內容(產品、預計產量、客戶、預計起訖日)、建立生產批、完工、撤單": 
        "Quản lý lệnh sản xuất: Tạo mới nội dung lệnh (sản phẩm, sản lượng dự kiến, khách hàng, ngày bắt đầu/kết thúc dự kiến), tạo lô sản xuất, hoàn thành, hủy lệnh",
    "工單排程可視化\n1、檢視選擇：(1)、小時、日 (2)、客戶   \n2、可直接移動、調整時間、機台\n3、直接連接工單資料做修改\n4、顯示目前工單生產現況\n      (1)以顏色代表狀態。\n      (2)標示實際生產數量/預計生產數量\n5、工單變更生產機台時,可限制工單變更的機台": 
        "Trực quan hóa tiến độ xếp lịch lệnh sản xuất:\n1. Tùy chọn xem: (1) Theo giờ, ngày (2) Theo khách hàng\n2. Có thể trực tiếp kéo thả, điều chỉnh thời gian, chuyển máy\n3. Liên kết trực tiếp dữ liệu lệnh sản xuất để chỉnh sửa\n4. Hiển thị hiện trạng sản xuất: (1) Màu sắc thể hiện trạng thái (2) Số lượng thực tế / dự kiến\n5. Có thể giới hạn danh sách máy được phép chuyển đổi",
    "生產批管理: 投產、生產批設定 (派工設備、計算預定生產時間)、結批、中止": 
        "Quản lý lô sản xuất (Lot/Batch): Đưa vào sản xuất, thiết lập lô (giao máy, tính thời gian sản xuất dự kiến), đóng lô, tạm dừng lô",
    "工作站報工: 入/出站、暫停/復工、參數、工程圖、SOP": 
        "Báo công tại trạm làm việc: Vào/ra trạm (Check-in/Check-out), tạm dừng/tiếp tục, thông số kỹ thuật, bản vẽ kỹ thuật, SOP",
    "作業指導書: 開啟圖檔 or PDF": "Hướng dẫn thao tác (SOP): Mở file hình ảnh hoặc file PDF",
    "工程圖: 開啟圖檔 or PDF": "Bản vẽ kỹ thuật: Mở file hình ảnh hoặc file PDF",
    "參數管理: 建立機種標準生產參數": "Quản lý thông số: Thiết lập thông số sản xuất tiêu chuẩn cho từng model",
    "入/出站: 紀錄工單進出站時間": "Vào/ra trạm: Ghi nhận thời gian lệnh sản xuất vào và ra khỏi từng công đoạn",
    "人員上工/下工: 可綁定進站人員，進行人員工時收集": "Công nhân vào/hết ca: Có thể liên kết nhân viên vào trạm để thu thập giờ công lao động",
    "數量/缺陷回報: 支援設備與手動回報、收集缺陷項目": "Báo cáo số lượng / phế phẩm: Hỗ trợ báo cáo tự động từ thiết bị và nhập thủ công, thu thập danh mục lỗi khuyết tật",
    "暫停/復工: 暫停原因收集、生產暫停起訖時間收集": "Tạm dừng / Tiếp tục sản xuất: Thu thập nguyên nhân tạm dừng, thu thập thời gian bắt đầu và kết thúc gián đoạn",
    "入庫作業: 生產完成後，指定入庫位置": "Nhập kho thành phẩm: Chỉ định vị trí ô kệ nhập kho sau khi sản xuất hoàn thành",
    "生產監控": "Giám sát sản xuất thời gian thực",
    "生產總覽": "Tổng quan tình hình sản xuất",
    "設備總覽：設備狀態與生產績效": "Tổng quan thiết bị: Trạng thái máy và hiệu suất sản xuất",
    "生產批總覽: 每一製程站狀態、數量、製程進度位置": "Tổng quan lô sản xuất: Trạng thái từng công đoạn, số lượng, vị trí tiến độ quy trình",
    "工單生產進度: 今日完成數、累積完成數、生產進度、工單狀態": "Tiến độ lệnh sản xuất: Số lượng hoàn thành hôm nay, lũy kế hoàn thành, % tiến độ, trạng thái lệnh",
    "可視化工單生產排程: 已工單角度查詢排產時間": "Lịch sản xuất trực quan theo lệnh: Tra cứu thời gian lập lịch theo góc độ lệnh sản xuất",
    "可視化設備生產排程: 已設備角度查詢排產時間，與現場開工狀況": "Lịch sản xuất trực quan theo thiết bị: Tra cứu thời gian xếp lịch và tình trạng vận hành theo từng máy",
    "生產紀錄WebApp\n(現場人員操作)\n平板或工業電腦操作": 
        "Ứng dụng WebApp ghi nhận sản xuất\n(Dành cho thao tác tại xưởng)\nSử dụng trên máy tính bảng hoặc máy tính công nghiệp IPC",
    "生產任務列表: 顯示當日可執行任務 (以設備區分)": "Danh sách nhiệm vụ sản xuất: Hiển thị nhiệm vụ khả dụng trong ngày (phân theo từng thiết bị)",
    "前置作業報工": "Báo công thao tác chuẩn bị trước sản xuất (Setup/Changeover)",
    "暫停/停機原因回報": "Báo cáo lý do tạm dừng / dừng máy",
    "不良品原因回報": "Báo cáo nguyên nhân hàng lỗi / phế phẩm",
    "不良發生狀況: 統計查詢選擇工單的不良發生時間與數量比例": "Tình trạng phát sinh phế phẩm: Thống kê thời gian và tỷ lệ số lượng lỗi theo lệnh sản xuất được chọn",
    "整體效率": "Hiệu suất tổng thể (OEE)",
    "在製整體效率: 各製程下工單OEE": "OEE bán thành phẩm (WIP): Chỉ số OEE của các lệnh sản xuất theo từng công đoạn",
    "在製產能與不良: 各製程下工單生產產能、不良發生狀況": "Năng suất và phế phẩm WIP: Năng suất sản xuất và tình trạng lỗi phát sinh theo từng công đoạn",
    "在製稼動生產狀況: 各製程下工單生產稼動時間累計、百分比": "Tình trạng hoạt động WIP: Thời gian chạy máy lũy kế và % tỷ lệ hoạt động của lệnh sản xuất theo công đoạn",
    "在製生產進度: 各製程下生產的工單進度、良率": "Tiến độ sản xuất WIP: Tiến độ hoàn thành và tỷ lệ đạt (Yield rate) theo công đoạn",
    "在製生產週期: 各製程下工單平均週期、生產效率、完成仍需時間、預計生產總時、生產落後時間": 
        "Chu kỳ sản xuất WIP: Chu kỳ bình quân (Cycle Time), hiệu suất sản xuất, thời gian dự kiến còn lại, tổng thời gian dự kiến, thời gian chậm trễ",
    "在製稼動生產狀況: 統計查詢選擇工單的生產時間與稼動率": "Tình trạng hoạt động WIP: Tra cứu thống kê thời gian sản xuất và tỷ lệ hoạt động của lệnh được chọn",
    "在製生產統計清單: 統計查詢選擇工單的關鍵指標": "Danh sách thống kê sản xuất WIP: Tra cứu thống kê các chỉ số KPI then chốt của lệnh được chọn",
    "設備生產效率: 統計查詢選擇工單於不同設備生產的效率": "Hiệu suất sản xuất theo máy: Tra cứu thống kê hiệu suất lệnh sản xuất khi chạy trên các máy khác nhau",
    "設備生產良率狀況: 統計查詢選擇工單於不同設備生產的良率": "Tỷ lệ đạt theo máy: Tra cứu tỷ lệ đạt của lệnh khi gia công trên các máy khác nhau",
    "報表管理": "Quản lý báo cáo thống kê",
    "工單生產狀況": "Tình trạng sản xuất của lệnh",
    "工單數量統計: 工單執行進度報表": "Thống kê số lượng lệnh: Báo cáo tiến độ thực hiện lệnh sản xuất",
    "工單生產分析表: 工單、數量、人時、機時、工時差異": "Bảng phân tích sản xuất lệnh: Lệnh sản xuất, số lượng, giờ công người, giờ máy, độ lệch thời gian",
    "工單製程暫停紀錄: 工單暫停/解除時間、人員、原因": "Nhật ký tạm dừng công đoạn: Thời gian dừng/tiếp tục, nhân sự, lý do gián đoạn",
    "產能效率": "Năng suất và hiệu suất",
    "生產日報表: 數量、工時、效率報表": "Báo cáo sản xuất hàng ngày: Báo cáo số lượng, giờ công, hiệu suất",
    "人員工時紀錄表: 查詢人員每天工作總時與操作的工單": "Bảng ghi nhận giờ công nhân sự: Tra cứu tổng thời gian làm việc hàng ngày và các lệnh thao tác",
    "生產履歷查詢: 歷史使用紀錄、生產績效、使用次數": "Tra cứu lý lịch sản xuất: Nhật ký sử dụng lịch sử, hiệu suất sản xuất, số lần sử dụng",
    "生產批入出站紀錄: 工單入/出站紀錄": "Nhật ký vào/ra trạm theo lô: Lịch sử vào/ra trạm của lệnh sản xuất",
    "個案ERP串接(工單匯入匯出)": "Tích hợp hệ thống ERP theo yêu cầu (Import/Export lệnh sản xuất)",
    "下行製令單轉工單(一個中繼表單或API)\n#ERP工單轉匯MES": "Tải lệnh sản xuất từ ERP sang MES (qua bảng trung gian hoặc API)\n#Đồng bộ Work Order từ ERP sang MES",
    "上行報工ERP(一個中繼表單或API)\n#報工資料抛轉": "Đẩy dữ liệu báo công lên ERP (qua bảng trung gian hoặc API)\n#Đồng bộ kết quả sản lượng về ERP",
    "模具: 模具編號、模具名稱、主/共模、客戶、廠商、模穴數、標準週期、可用設備、位置": 
        "Khuôn mẫu: Mã khuôn, tên khuôn, khuôn chính/khuôn chung, khách hàng, nhà cung cấp, số lòng khuôn (Cavity), chu kỳ tiêu chuẩn, máy tương thích, vị trí lưu kho",
    "倉庫及儲位: 設定生產入庫位置紀錄 (可與WMS串接)": "Kho và vị trí ô kệ: Thiết lập ghi nhận vị trí nhập kho sản xuất (có thể liên kết WMS)",
    "位置列表: 建立模具移動位置點": "Danh sách vị trí: Thiết lập các điểm luân chuyển khuôn",
    "新增模具移動紀錄": "Thêm mới nhật ký di chuyển khuôn",
    "移動記錄查詢: 模具移動歷程、移入時間 & 位置": "Tra cứu nhật ký di chuyển: Lịch sử dịch chuyển khuôn, thời gian chuyển vào & vị trí",
    "保養/維修單: 單據列表、保養/維修狀態、保養工時": "Phiếu bảo dưỡng/sửa chữa: Danh sách phiếu, trạng thái bảo dưỡng/sửa chữa, số giờ công",
    "基本資料：建立生產前動作代碼、主要名稱(例：上模、清料管…)/中文，次要名稱/外文。": 
        "Dữ liệu gốc: Thiết lập mã thao tác chuẩn bị trước sản xuất, tên chính (ví dụ: lên khuôn, làm sạch nòng xilanh...)/tiếng Trung, tên phụ/tiếng nước ngoài.",
    "現場：登入帳密、選設備編號，進入生產前作業(依照基本資料)\n1、上模、下模(技術人員)#人員進出、時間累計\n2、清料管(技術人員)\n3、調機(技術人員)#人員進出、時間累計\n4、首件確認時間點選\n5、生產明細#報工良品、不良品數量/原因(與量產的不良品原因同列表)\n6、作業結束#換量產的作業人員登入": 
        "Tại hiện trường: Đăng nhập tài khoản, chọn mã máy, vào tác vụ chuẩn bị sản xuất (theo dữ liệu gốc):\n1. Lên khuôn, hạ khuôn (Kỹ thuật viên) #Ghi nhận ra/vào, lũy kế thời gian\n2. Vệ sinh nòng xilanh / làm sạch vật liệu (Kỹ thuật viên)\n3. Căn chỉnh máy / lên thông số (Kỹ thuật viên) #Ghi nhận ra/vào, lũy kế thời gian\n4. Xác nhận mốc thời gian làm mẫu thử đầu ca (First Article)\n5. Chi tiết sản xuất: Báo công số lượng hàng đạt, phế phẩm / nguyên nhân (chung danh mục lý do lỗi của sản xuất hàng loạt)\n6. Kết thúc chuẩn bị: Bàn giao để công nhân sản xuất hàng loạt đăng nhập",
    "統計列表：前置作業明細表\n搜尋條件：工單編號、機台、人員、日期區間": 
        "Bảng kê thống kê: Bảng chi tiết thao tác chuẩn bị trước sản xuất\nĐiều kiện tìm kiếm: Mã lệnh sản xuất, số hiệu máy, nhân sự, khoảng ngày",

    # 6. Sheet QMS
    "SUNSTIGE QMS-\n品質管理系統": "HỆ THỐNG QUẢN LÝ CHẤT LƯỢNG SUNSTIGE QMS",
    "業務夥伴: 客戶/廠商基本資料": "Đối tác kinh doanh: Dữ liệu gốc khách hàng / nhà cung cấp",
    "品號主檔: \n建立所有要檢驗的品號資料，為日常檢驗單據或其他相關作業的［品號］輸入來源": 
        "Danh mục mã hàng (Part Number Master):\nThiết lập toàn bộ dữ liệu mã hàng cần kiểm tra; là nguồn nhập [Mã hàng] cho các phiếu kiểm tra hàng ngày và các quy trình liên quan",
    "品管類別檢驗項目建立:\n定義品管類別及品號須經過那些檢查項目，日後在該料件檢驗時，依據所建之品管類別＋品號＋製程自動帶出檢驗項目，簡化人員需查標準書之時間": 
        "Thiết lập hạng mục kiểm tra theo phân loại quản lý chất lượng:\nĐịnh nghĩa nhóm chất lượng và các tiêu chí kiểm tra tương ứng; khi kiểm tra mã hàng, hệ thống tự động tải các hạng mục dựa trên [Nhóm chất lượng + Mã hàng + Công đoạn], tiết kiệm thời gian tra cứu tiêu chuẩn",
    "檢驗項目建立:\n建立品質檢驗中所有之檢驗項目、儀器、方式": 
        "Thiết lập hạng mục kiểm tra:\nXây dựng toàn bộ các tiêu chí kiểm tra chất lượng, dụng cụ/thiết bị đo và phương pháp kiểm tra",
    "簽核設定 (異常單使用): \n可依發現時機發現單位設定不同的簽核流程：首檢／巡檢／出貨檢/進料檢": 
        "Thiết lập quy trình duyệt (Dùng cho phiếu sự cố):\nCó thể thiết lập luồng phê duyệt khác nhau tùy thời điểm phát hiện: Kiểm tra đầu ca / Tuần tra / Xuất xưởng / Nhập liệu",
    "日常檢驗作業": "Nghiệp vụ kiểm tra chất lượng hàng ngày",
    "進料檢驗單: 新增單據、檢驗列表、檢驗輸入、判定": 
        "Phiếu kiểm tra nguyên liệu đầu vào (IQC): Tạo phiếu mới, danh sách kiểm tra, nhập số đo kiểm tra, đánh giá kết quả (Pass/Fail)",
    "首件檢驗單: 新增單據、檢驗列表、檢驗輸入、判定": 
        "Phiếu kiểm tra sản phẩm đầu ca (FAI - First Article Inspection): Tạo phiếu mới, danh sách kiểm tra, nhập kết quả, đánh giá kết luận",
    "製程巡檢單: 新增單據、檢驗列表、檢驗輸入、判定": 
        "Phiếu kiểm tra tuần tra công đoạn (IPQC): Tạo phiếu mới, danh sách kiểm tra, nhập kết quả, đánh giá kết luận",
    "成品檢驗單: 新增單據、檢驗列表、檢驗輸入、判定": 
        "Phiếu kiểm tra thành phẩm (FQC - Final QC): Tạo phiếu mới, danh sách kiểm tra, nhập kết quả, đánh giá kết luận",
    "出貨檢驗單: 新增單據、檢驗列表、檢驗輸入、判定": 
        "Phiếu kiểm tra trước khi xuất xưởng (OQC - Outgoing QC): Tạo phiếu mới, danh sách kiểm tra, nhập kết quả, đánh giá kết luận",
    "異常單:\n製品判定不良時，提出改善方案，於首檢/巡檢/出貨檢/進料檢的過程中，判定為料件不良，可自動發起本單或出廠後因客訴，填寫本單。並依簽核流程由相關單位進行表單內容填寫及審核": 
        "Phiếu xử lý sự cố chất lượng (NCR - Non-Conformance Report):\nKhi phát hiện sản phẩm lỗi, đề xuất phương án cải tiến; trong quá trình kiểm tra IQC/FAI/IPQC/FQC/OQC hoặc khi nhận khiếu nại khách hàng, hệ thống tự động khởi tạo phiếu này. Sau đó luân chuyển theo quy trình phê duyệt để các bộ phận liên quan giải trình và thẩm xét",
    "新增臨時檢驗項目": "Thêm hạng mục kiểm tra đột xuất / tạm thời",
    "不良原因管理:\n後續檢驗時可將只要輸入不良原因代號即可，可減少使用者輸入資料的時間": 
        "Quản lý mã nguyên nhân lỗi:\nKhi kiểm tra chỉ cần nhập mã lỗi đại diện, giúp rút ngắn thời gian nhập liệu của nhân viên",
    "圖示化檢驗模組": "Module kiểm tra trực quan trên bản vẽ / Layout",
    "LAYOUT上傳\n※僅限上傳JPG格式的檔案\n#一個產品可上傳6張圖\n檢驗LAYOUT設定/拖曳元件設定檢驗項目": 
        "Tải lên bản vẽ Layout:\n*Chỉ chấp nhận file định dạng JPG\n#Mỗi sản phẩm có thể tải tối đa 6 hình ảnh\nCài đặt Layout kiểm tra / Kéo thả vị trí linh kiện để gắn tiêu chí kiểm tra",
    "新增LAYOUT圖示頁籤，點選檢驗項目位置後輸入檢驗值\n頁面上可區分”已量測”、”未量測”項目，方便QC人員第一時間確認\n※原先檢驗畫面會新增功能畫面[LAYOUT]的頁籤，可切換檢驗畫面成LAYOUT圖示\n※LAYOUT上輸入的檢驗值，同樣回寫到表單畫面作對應檢視\n※USER在檢驗單裡面，如有手動臨時新增檢驗項目，新增的檢驗項目需在列表中輸入檢驗結果，圖面上僅會出現原先LAYOUT有設定的項目\n※LAYOUT的檢視僅出現在QC表單狀態為處理中，單據已結案後，僅保留檢驗數據，不提供歷史圖檔的檢視功能": 
        "Bổ sung tab trực quan Layout: Nhấp vào vị trí điểm kiểm tra trên ảnh để nhập giá trị đo\nGiao diện phân biệt rõ điểm \"Đã đo\" và \"Chưa đo\" giúp QC xác nhận tức thì\n*Màn hình kiểm tra bổ sung tab [LAYOUT] để chuyển đổi giữa dạng bảng và sơ đồ trực quan\n*Giá trị nhập trên Layout tự động đồng bộ về bảng dữ liệu\n*Nếu QC thêm tiêu chí tạm thời thì nhập tại bảng, Layout chỉ thể hiện các điểm đã định vị sẵn\n*Xem Layout khả dụng khi phiếu ở trạng thái Đang xử lý; sau khi đóng phiếu chỉ lưu trữ dữ liệu số đo",
    "管理報表": "Báo cáo quản trị chất lượng",
    "進料檢驗明細表": "Bảng chi tiết kiểm tra nguyên liệu đầu vào (IQC Details)",
    "首件檢驗明細表": "Bảng chi tiết kiểm tra mẫu đầu ca (FAI Details)",
    "製程巡檢明細表": "Bảng chi tiết kiểm tra tuần tra công đoạn (IPQC Details)",
    "成品檢驗明細表": "Bảng chi tiết kiểm tra thành phẩm (FQC Details)",
    "出貨檢驗明細表": "Bảng chi tiết kiểm tra trước xuất hàng (OQC Details)",
    "異常單明細表": "Bảng chi tiết các phiếu sự cố chất lượng (NCR Details)",
    "檢驗單良率統計表:\n依品號的檢驗記錄，統計良率": 
        "Bảng thống kê tỷ lệ đạt (Yield rate) theo phiếu kiểm tra:\nThống kê tỷ lệ đạt dựa trên lịch sử kiểm tra của từng mã hàng",
    "製程能力統計圖:\n對特定品號下的檢驗項目進行X-R chart 圖表": 
        "Biểu đồ năng lực quy trình (SPC / Cp - Cpk):\nVẽ biểu đồ kiểm soát X-bar R chart cho tiêu chí kiểm tra của mã hàng cụ thể",
    "廠商(IQC)走勢圖:\n彙總進貨廠商進貨與檢驗狀況": 
        "Biểu đồ xu hướng nhà cung cấp (IQC):\nTổng hợp tình hình nhập hàng và kết quả kiểm tra chất lượng theo từng NCC",
    "廠商(IQC)零件不良評比總表:\n比較不同供應商間的交貨不良狀況分析": 
        "Bảng đánh giá so sánh tỷ lệ lỗi linh kiện NCC:\nPhân tích so sánh tình hình lỗi giao hàng giữa các nhà cung cấp khác nhau",
    "廠商(IQC)零件不良類型分析表:\n彙總各廠商於查詢區間的交貨不良狀況分析": 
        "Bảng phân tích chủng loại lỗi linh kiện NCC:\nTổng hợp phân tích các dạng lỗi phát sinh theo từng NCC trong khoảng thời gian tra cứu",
    "儀錶板": "Bảng thông tin điều hành (Dashboard)",
    "圓餅圖：統計 待檢驗/檢驗中/檢驗完成，的資料筆數佔比及總筆數。": 
        "Biểu đồ tròn: Thống kê tỷ lệ % và tổng số lượng phiếu Chờ kiểm tra / Đang kiểm tra / Đã hoàn thành",
    "圓餅圖：統計 最終判定為ACC/AOD/REJ，的資料筆數佔比及總筆數": 
        "Biểu đồ tròn: Thống kê tỷ lệ % và tổng số lượng phán định Chấp nhận (Pass/ACC) / Chấp nhận có điều kiện (Concession/AOD) / Từ chối (Reject/REJ)",
    "長條圖：統計最終判定為拒收的筆數，列出前15名筆數最多的品號": 
        "Biểu đồ cột: Thống kê số lần bị từ chối nhập (REJ), liệt kê Top 15 mã hàng có số lần bị từ chối nhiều nhất",
    "長條圖：統計最終判定為特採及退貨的筆數，列出前15名筆數最多的品號": 
        "Biểu đồ cột: Thống kê số lần mua đặc cách (Concession) và trả hàng (Return), liệt kê Top 15 mã hàng phát sinh nhiều nhất",
    "長條圖：統計每月異常單筆數，下展：長條圖點到該月時，Show異常明細": 
        "Biểu đồ cột: Thống kê số lượng phiếu sự cố hàng tháng; Drill-down: Nhấp vào tháng sẽ hiển thị chi tiết các phiếu sự cố",
    "長條圖：統計年度不良原因，列出各原因的筆數及累計百分比": 
        "Biểu đồ Pareto cột & đường: Thống kê nguyên nhân lỗi trong năm, liệt kê số lần xuất hiện và tỷ lệ % lũy kế của từng nguyên nhân",
    "TOP15 異常關注產品": "Top 15 sản phẩm cần theo dõi đặc biệt về sự cố chất lượng",
    "TOP15 異常關注零件": "Top 15 linh kiện cần theo dõi đặc biệt về sự cố chất lượng",
    "IQC、OQC若與ERP串接收費\nIQC、OQC若與日威WMS做串接不另外收費": 
        "Tích hợp IQC, OQC với ERP sẽ tính phí;\nTích hợp IQC, OQC với hệ thống Sunstige WMS được miễn phí (không thu thêm phí)",
    "FAI、IPQC、FQC與日威MES做串接不另外收費": 
        "Tích hợp FAI, IPQC, FQC với hệ thống Sunstige MES được miễn phí (không thu thêm phí)",

    # 7. Sheet WMS
    "Sunstige WMS- 倉儲管理系統": "HỆ THỐNG QUẢN LÝ KHO VẬN SUNSTIGE WMS",
    "基本資料設定": "Thiết lập dữ liệu gốc (Master Data)",
    "倉庫設定: 工廠列表、倉庫列表、儲存區列表": "Thiết lập kho bãi: Danh sách nhà máy, danh sách kho, danh sách khu vực lưu trữ (Storage Area)",
    "條碼管理: 條碼設置、條碼列印": "Quản lý mã vạch: Cài đặt quy tắc mã vạch, in mã vạch / tem nhãn",
    "入庫作業": "Tác vụ nhập kho (Inbound Operations)",
    "收貨單、上架單、銷退入庫單、領用退庫單、生產入庫單、\n托外入庫單": 
        "Phiếu nhận hàng, phiếu lên kệ, phiếu nhập trả hàng bán, phiếu trả kho vật tư đã lĩnh, phiếu nhập kho sản xuất, phiếu nhập hàng gia công ngoài",
    "出庫作業": "Tác vụ xuất kho (Outbound Operations)",
    "銷貨出庫單、批揀單、退貨出庫單、領料出庫單、報廢出庫單": 
        "Phiếu xuất hàng bán, phiếu nhặt hàng theo lô (Wave picking), phiếu xuất trả nhà cung cấp, phiếu xuất lĩnh vật tư, phiếu xuất hủy phế phẩm",
    "庫存管理": "Quản lý tồn kho (Inventory Management)",
    "庫存查詢、庫存分析、庫存調整單、庫存凍結/解凍、庫存明細表\n、庫存明細帳、庫存異動明細查詢、批號異動明細查詢\n、產品序號異動查詢、調撥單": 
        "Tra cứu tồn kho, phân tích tồn kho, phiếu điều chỉnh kho, đóng băng/giải phóng tồn kho, bảng chi tiết tồn kho, sổ chi tiết kho, tra cứu biến động tồn kho, tra cứu lịch sử theo số lô (Lot Tracking), tra cứu lịch sử số Serial Number, phiếu điều chuyển kho",
    "盤點作業": "Tác vụ kiểm kê kho (Stocktake Operations)",
    "盤點底稿、盤盈虧明細表": "Bảng kê kiểm kê kho, bảng chi tiết chênh lệch thừa / thiếu kiểm kê",
    "電子看板": "Màn hình hiển thị điện tử (Warehouse Digital Signage / Kanban)",
    "進貨待驗看板": "Bảng hiển thị hàng nhập chờ kiểm tra chất lượng (Waiting for IQC)",
    "進貨入庫看板": "Bảng hiển thị hàng đã kiểm tra chờ nhập kho lên kệ",
    "已完工待入庫看板": "Bảng hiển thị thành phẩm đã xong chờ nhập kho",
    "出貨看版": "Bảng hiển thị đơn hàng chờ xuất kho",
    "\nAPP模組功能\n(需使用PDA)": "TÍNH NĂNG TRÊN THIẾT BỊ CẦM TAY PDA (MODULE MOBILE APP)",
    "1. 採購收貨、採購上架、調撥入庫、儲位調整\n2. 銷退入庫、領用退庫、生產入庫、托外入庫": 
        "1. Nhận hàng mua, lên kệ mua hàng, điều chuyển nhập kho, điều chỉnh vị trí ô kệ\n2. Nhập kho trả hàng bán, nhập trả vật tư lĩnh, nhập kho thành phẩm sản xuất, nhập kho hàng gia công ngoài",
    "1. 銷貨庫存分配、揀貨下架、調撥出庫\n2. 退貨出庫、領料出庫、報廢出庫": 
        "1. Phân bổ tồn kho xuất bán, nhặt hàng hạ kệ, điều chuyển xuất kho\n2. Xuất trả hàng nhà cung cấp, xuất lĩnh vật tư, xuất hủy phế phẩm",
    "1. 庫存盤點": "1. Kiểm kê tồn kho trên thiết bị PDA",
    "串接ERP派發單據: \n(進貨單、生產入庫單、盤點單、銷貨單、庫存異動單、領料單、調撥單、退料單、退貨單)": 
        "Đồng bộ các chứng từ phát hành từ ERP:\n(Phiếu mua hàng, phiếu nhập kho sản xuất, phiếu kiểm kê, phiếu bán hàng, phiếu biến động kho, phiếu lĩnh vật liệu, phiếu điều chuyển, phiếu trả vật liệu, phiếu trả hàng NCC)",

    # 8. Sheet IQA關鍵製程(MES Lite)
    "SUNSTIGE日威\n現場集成應用系統\n(IQA關鍵製程)": "HỆ THỐNG ỨNG DỤNG TÍCH HỢP HIỆN TRƯỜNG SUNSTIGE\n(QUY TRÌNH TRỌNG YẾU IQA / MES LITE)",
    "基本資料\n建立、管理": "Thiết lập và quản lý dữ liệu gốc (Master Data)",
    "角色管理": "Quản lý vai trò người dùng",
    "用戶角色": "Vai trò người dùng (Phân quyền chức năng)",
    "客戶主檔": "Danh mục khách hàng",
    "客戶(編號、名稱)": "Khách hàng (Mã khách hàng, tên khách hàng)",
    "廠商主檔": "Danh mục nhà cung cấp",
    "廠商(編號、名稱)": "Nhà cung cấp (Mã NCC, tên NCC)",
    "材料主檔": "Danh mục nguyên vật liệu",
    "材料(編號、名稱、種類、色號、品牌、供應商)": "Nguyên vật liệu (Mã vật liệu, tên, loại, mã màu, thương hiệu, nhà cung cấp)",
    "模具主檔(編號、客戶、廠商名稱…)": "Danh mục khuôn mẫu (Mã khuôn, khách hàng, tên nhà cung cấp...)",
    "機台(編號、名稱)": "Máy sản xuất (Mã máy, tên máy)",
    "產品(編號、模具編號、名稱、規格)": "Sản phẩm (Mã sản phẩm, mã khuôn, tên, quy cách)",
    "工單基本資料": "Dữ liệu cơ bản của lệnh sản xuất",
    "工單資料建立、派工": "Tạo mới dữ liệu lệnh sản xuất, phân công giao việc (Dispatching)",
    "點選工單編號": "Nhấp chọn mã lệnh sản xuất",
    "上料管控\n#條碼由WMS產生(同時導入)或自行產生。": 
        "Kiểm soát nạp liệu vào máy (Feeder/Loading Control):\n#Mã vạch do WMS tạo (nếu triển khai đồng thời) hoặc do khách hàng tự tạo",
    "上料紀錄": "Nhật ký nạp nguyên vật liệu vào máy",
    "關鍵料": "Nguyên vật liệu trọng yếu (Critical Components)",
    "ESOP: 上傳產品在各製程對應的SOP或工程圖": "ESOP (Electronic SOP): Tải lên quy trình thao tác chuẩn hoặc bản vẽ kỹ thuật tương ứng từng công đoạn",
    "品質檢結果輸入/判斷 (FAI、IPQC、OQC、FQC)": "Nhập và phán định kết quả kiểm tra chất lượng (FAI - Đầu ca, IPQC - Tuần tra, OQC - Xuất xưởng, FQC - Thành phẩm)",
    "人員自行輸入": "Nhân viên tự nhập tay dữ liệu",
    "QC模板": "Biểu mẫu tiêu chuẩn kiểm tra chất lượng (QC Template)",
    "其他資料收集": "Thu thập dữ liệu khác",
    "收集項目設定: 建立其他手動收集項目": "Cài đặt tiêu chí thu thập: Thiết lập các nội dung cần thu thập thủ công khác",
    "其他手動資料收集明細紀錄": "Nhật ký chi tiết thu thập dữ liệu thủ công khác",
    "工單人員記錄表": "Bảng ghi nhận công nhân theo lệnh sản xuất",
    "工單數量變更記錄": "Nhật ký thay đổi số lượng của lệnh sản xuất",
    "發起異常作業": "Khởi tạo quy trình báo cáo sự cố bất thường",
    "再發防止作業": "Tác vụ phòng ngừa tái diễn sự cố (CAPA - Corrective & Preventive Action)",
    "上傳檔案": "Tải lên tệp đính kèm / hình ảnh chứng minh",
    "產線生產日報表": "Báo cáo sản xuất hàng ngày theo dây chuyền",
    "生產統計\n(工單)": "Thống kê sản xuất (Theo lệnh sản xuất)",
    "缺陷發生狀況": "Tình trạng phát sinh lỗi khuyết tật",
    "生產良率狀況": "Tình hình tỷ lệ đạt (Yield rate) trong sản xuất",
    "年度生產不良細項 TOP10": "Top 10 chi tiết dạng lỗi sản xuất nhiều nhất trong năm",

    # 9. Các thuật ngữ khác xuất hiện trong các sheet
    "Dashboard\n儀錶板": "Bảng điều khiển trực quan Dashboard",
    "ERP串接\n(只限工單資料)": "Tích hợp hệ thống ERP\n(Chỉ áp dụng cho dữ liệu lệnh sản xuất Work Order)",
    "MES- 平板操作系統 (WebAPP模式)": "Hệ thống thao tác trên máy tính bảng MES (Chế độ WebAPP)",
    "X-R管制圖:\n對特定品號下的檢驗項目進行X-R chart 圖表": "Biểu đồ kiểm soát X-bar R chart:\nVẽ biểu đồ kiểm soát X-R chart cho tiêu chí kiểm tra của mã hàng cụ thể",
    "今日生產訊息\n工單排程": "Thông tin sản xuất trong ngày\nTiến độ xếp lịch lệnh sản xuất",
    "代理人設定": "Cài đặt người ủy quyền / người thay thế",
    "使用者登入": "Đăng nhập người dùng",
    "使用者管理": "Quản lý người dùng",
    "使用者管理、角色管理": "Quản lý người dùng, quản lý vai trò",
    "保養/維修單列表": "Danh sách phiếu bảo dưỡng / sửa chữa",
    "保養紀錄": "Nhật ký bảo dưỡng thiết bị",
    "參數履歷監視：關鍵資料異動紀錄": "Giám sát lịch sử thông số: Nhật ký biến động các dữ liệu then chốt",
    "在製稼動生產狀況: 各製程下工單生產稼動時間累計、百分比": "Tình trạng vận hành WIP: Thời gian chạy máy lũy kế và % tỷ lệ hoạt động của lệnh theo công đoạn",
    "平板操作系統 (WebAPP模式)": "Hệ điều hành trên máy tính bảng (Chế độ WebAPP)",
    "日常檢驗資料": "Dữ liệu kiểm tra chất lượng hàng ngày",
    "暫停代碼(調機、保養…)\n#建立中文與其他語言": "Mã tạm dừng (Chỉnh máy, bảo dưỡng...)\n#Thiết lập tiếng Trung và các ngôn ngữ khác",
    "異常代碼(缺料…)\n#建立中文與其他語言": "Mã sự cố bất thường (Thiếu nguyên liệu...)\n#Thiết lập tiếng Trung và các ngôn ngữ khác",
    "異常單列表": "Danh sách các phiếu sự cố chất lượng (NCR List)",
    "異常單明細": "Chi tiết phiếu sự cố chất lượng",
    "當月生產檢驗明細": "Chi tiết kiểm tra sản xuất trong tháng",
    "當月生產檢驗狀況": "Tình hình kiểm tra sản xuất trong tháng",
    "當月進料檢驗明細": "Chi tiết kiểm tra nguyên liệu đầu vào trong tháng (IQC)",
    "當月進料檢驗狀況": "Tình hình kiểm tra nguyên liệu đầu vào trong tháng (IQC)",
    "產線生產清單": "Danh sách sản xuất theo dây chuyền",
    "產品名稱、稼動率、數量": "Tên sản phẩm, tỷ lệ hoạt động, số lượng",
    "異常單: 新增單據、單據列表、異常內容、上傳資料、異常分級": "Phiếu sự cố: Thêm phiếu mới, danh sách phiếu, nội dung sự cố, tải lên tài liệu, phân cấp sự cố",
    "設備保養\n(選配)": "Bảo dưỡng thiết bị\n(Tùy chọn thêm)",
    "設備檢驗項目維護": "Bảo trì và cập nhật các hạng mục kiểm tra thiết bị",
    "設備異常單: 單據列表、新增單據、異常分級、異常內容、上傳資料": "Phiếu sự cố thiết bị: Danh sách phiếu, thêm mới phiếu, phân cấp sự cố, nội dung bất thường, tải lên tài liệu",
    "資料收集範圍設定: 建立需收集的製程與項目": "Cài đặt phạm vi thu thập dữ liệu: Thiết lập các công đoạn và tiêu chí cần thu thập",
    "資料收集設定": "Cài đặt thu thập dữ liệu",
    "製程巡檢記錄表": "Bảng ghi nhận kiểm tra tuần tra công đoạn (IPQC Record)",
    "製程資料設定": "Cài đặt dữ liệu công đoạn sản xuất",
    "模具管理\n(現場操作畫面)": "Quản lý khuôn mẫu\n(Giao diện thao tác tại xưởng)",
    "模具管理\n(選配)": "Quản lý khuôn mẫu\n(Tùy chọn thêm)",
    "模具變更記錄": "Nhật ký thay đổi khuôn mẫu",
    "生產前動作管理\n(選配)": "Quản lý thao tác chuẩn bị trước sản xuất\n(Tùy chọn thêm)",
    "生產清單": "Danh sách sản xuất",
    "生產線平台": "Nền tảng dây chuyền sản xuất",
    "選單管理": "Quản lý menu chức năng",
    "部門主檔": "Danh mục phòng ban",
    "上行報工ERP(一個中繼表單或API)": "Đẩy dữ liệu báo công lên ERP (qua bảng trung gian hoặc API)",
    "下行製令單轉工單(一個中繼表單或API)": "Tải lệnh sản xuất từ ERP sang MES (qua bảng trung gian hoặc API)",
    "串接MES": "Tích hợp / Liên kết hệ thống MES",
    "基本資料: 產品列表、單位列表、選單管理、業務夥伴、編碼原則": "Dữ liệu gốc: Danh sách sản phẩm, danh sách đơn vị tính, quản lý menu, đối tác kinh doanh, quy tắc đặt mã",
    "報工": "Báo công (Báo cáo sản lượng hoàn thành)",
    "工單生產進度": "Tiến độ sản xuất của lệnh sản xuất",
    "模具管理": "Quản lý khuôn mẫu",
    "生產履歷查詢: 模具歷史使用紀錄、生產績效、使用次數": "Tra cứu lý lịch sản xuất: Nhật ký lịch sử sử dụng khuôn, hiệu suất sản xuất, số lần sử dụng",
    "統計圖": "Biểu đồ thống kê",
    "設備管理": "Quản lý thiết bị máy móc",
    "設備點檢作業(作業區工作平台)": "Nghiệp vụ kiểm tra định kỳ thiết bị (Nền tảng thao tác tại khu vực làm việc)"
}

def translate_workbook(input_path, output_path):
    print(f"Loading workbook: {input_path}")
    wb = openpyxl.load_workbook(input_path)

    # Chuẩn hóa từ điển bỏ khoảng trắng đầu/cuối của key
    norm_dict = {k.strip(): v for k, v in TRANSLATIONS.items()}

    total_translated_cells = 0

    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\nProcessing Sheet: {sname}...")
        sheet_translated = 0

        for r in range(1, ws.max_row + 1):
            max_lines_in_row = 1
            has_translated_in_row = False

            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                val = cell.value

                if val is not None and isinstance(val, str) and val.strip():
                    cleaned_val = val.strip()

                    # Khớp từ điển chuẩn
                    if cleaned_val in norm_dict:
                        vi_trans = norm_dict[cleaned_val]
                        bilingual_text = f"{cleaned_val}\n{vi_trans}"
                        cell.value = bilingual_text
                        
                        current_align = cell.alignment
                        cell.alignment = Alignment(
                            horizontal=current_align.horizontal or 'left',
                            vertical=current_align.vertical or 'center',
                            wrap_text=True
                        )
                        
                        sheet_translated += 1
                        total_translated_cells += 1
                        has_translated_in_row = True

                        lines = bilingual_text.count('\n') + 1
                        if lines > max_lines_in_row:
                            max_lines_in_row = lines

            # Căn chỉnh chiều cao dòng (Anti-Clipping)
            if has_translated_in_row:
                current_h = ws.row_dimensions[r].height
                computed_h = max_lines_in_row * 18.0 + 8.0
                if current_h is not None:
                    ws.row_dimensions[r].height = max(current_h * 1.5, computed_h)
                else:
                    ws.row_dimensions[r].height = computed_h

        print(f"Sheet '{sname}': Translated {sheet_translated} cells.")

    print(f"\nTotal translated cells across all sheets: {total_translated_cells}")
    
    wb.save(output_path)
    print(f"Successfully saved to: {output_path}")

if __name__ == "__main__":
    src_file = _workspace_path('●SQ2026日威報價單(SMB-MES-QMS-WMS)-海外.xlsx')
    dst_file = _workspace_path('●SQ2026日威報價單(SMB-MES-QMS-WMS)-海外_SongNgu_Viet_Trung.xlsx')
    
    translate_workbook(src_file, dst_file)

    # Thử ghi đè file gốc nếu không bị khóa
    try:
        translate_workbook(src_file, src_file)
        print("Successfully updated the original file as well!")
    except PermissionError:
        print("Notice: Original file is currently open in Excel/WPS. Translated copy saved to SongNgu file.")
