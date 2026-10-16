# -*- coding: utf-8 -*-
"""Tạo bộ dữ liệu thử phức tạp từ mẫu Biên bản sửa chữa của CX Technology.

Script nhân bản workbook mẫu, mở rộng cả hai bảng lên 14 bản ghi và giữ nguyên
logo, font, màu, viền, merge, chế độ xem, công thức, cấu hình in và chữ ký.
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import os
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries


DEFAULT_TEMPLATE = Path(r"E:\Working\Bieu Mau\Biên bản sữa chữa.xlsx")
DEFAULT_OUTPUT = Path(r"E:\Working\Bieu Mau\Biên bản sữa chữa_TEST_14_BAN_GHI.xlsx")

ITEMS: List[Dict[str, Any]] = [
    {
        "serial": "RFR900NA55RSVHBA036",
        "issue": "Cò bắn phản hồi chậm, quét chập chờn\n扳機反應遲緩，掃描不穩定",
        "department": "CT2",
        "result": "Thay cụm cò bắn / 更換扳機模組",
        "qty": 1,
        "unit_price": 1650000,
        "lead_time": "3–4週",
    },
    {
        "serial": "EF550R4ALAWBI670",
        "issue": "Không đọc được barcode 2D sau khi máy ngủ\n休眠喚醒後無法讀取二維條碼",
        "department": "CT2",
        "result": "Cập nhật firmware & hiệu chuẩn / 更新韌體並校準",
        "qty": 1,
        "unit_price": 950000,
        "lead_time": "2–3週",
    },
    {
        "serial": "EF550R4ALAWBH111",
        "issue": "Màn hình nhấp nháy khi mức pin dưới 30%\n電量低於30%時螢幕閃爍",
        "department": "洗料",
        "result": "Thay cáp màn hình / 更換螢幕排線",
        "qty": 1,
        "unit_price": 1850000,
        "lead_time": "4–5週",
    },
    {
        "serial": "EF550R4ALAWBH109",
        "issue": "Không nhận pin, chân tiếp xúc bị oxy hóa\n無法辨識電池，接點氧化",
        "department": "LH",
        "result": "Vệ sinh & thay đầu nối pin / 清潔並更換電池接頭",
        "qty": 1,
        "unit_price": 1250000,
        "lead_time": "3–4週",
    },
    {
        "serial": "EF550R4ALAWBI654",
        "issue": "Cảm ứng lệch tại góc phải màn hình\n螢幕右側觸控偏移",
        "department": "CNC",
        "result": "Thay màn hình cảm ứng / 更換觸控螢幕",
        "qty": 1,
        "unit_price": 4200000,
        "lead_time": "5–6週",
    },
    {
        "serial": "EF550R4ALAWBI661",
        "issue": "Cổng sạc lỏng, sạc lúc được lúc không\n充電埠鬆動，充電不穩定",
        "department": "XI MẠ",
        "result": "Thay cụm cổng sạc / 更換充電埠模組",
        "qty": 1,
        "unit_price": 1450000,
        "lead_time": "3–4週",
    },
    {
        "serial": "EF550R4ALAWBI677",
        "issue": "Máy tự khởi động lại khi quét liên tục\n連續掃描時設備自動重啟",
        "department": "NC",
        "result": "Sửa bo nguồn & kiểm tra tải / 維修電源板並進行負載測試",
        "qty": 1,
        "unit_price": 2750000,
        "lead_time": "5–6週",
    },
    {
        "serial": "EF550RA4LAWBI657",
        "issue": "Màn hình trắng sau khi bật nguồn\n開機後螢幕白屏",
        "department": "X-TR",
        "result": "Thay IC hiển thị / 更換顯示IC",
        "qty": 1,
        "unit_price": 3250000,
        "lead_time": "5–6週",
    },
    {
        "serial": "EF550R4ALAWBI641",
        "issue": "Pin không sạc đầy, dừng ở mức 62%\n電池僅充至62%即停止",
        "department": "X-TR",
        "result": "Hiệu chuẩn mạch sạc / 校準充電電路",
        "qty": 1,
        "unit_price": 850000,
        "lead_time": "2–3週",
    },
    {
        "serial": "EF550RA4LAWBI686",
        "issue": "Loa cảnh báo không phát âm thanh\n警示喇叭無聲",
        "department": "Alume\n鋁製裝組",
        "result": "Thay loa & kiểm tra âm lượng / 更換喇叭並測試音量",
        "qty": 1,
        "unit_price": 780000,
        "lead_time": "2–3週",
    },
    {
        "serial": "EF550RA4LAWBI666",
        "issue": "Màn hình có sọc dọc, nút nguồn kém nhạy\n螢幕出現直線，電源鍵反應不良",
        "department": "DG\n包裝組",
        "result": "Thay màn hình & nút nguồn / 更換螢幕與電源鍵",
        "qty": 1,
        "unit_price": 4650000,
        "lead_time": "5–6週",
    },
    {
        "serial": "EF550RA4LAWBI678",
        "issue": "Máy không kết nối Wi-Fi 5 GHz\n設備無法連接5GHz Wi-Fi",
        "department": "DG\n包裝組",
        "result": "Thay module Wi-Fi / 更換Wi-Fi模組",
        "qty": 1,
        "unit_price": 2950000,
        "lead_time": "4–5週",
    },
    {
        "serial": "RFR900NA55RSVTEST901",
        "issue": "Cò bắn nứt vỏ, tín hiệu USB gián đoạn\n扳機外殼破裂，USB訊號中斷",
        "department": "CNC",
        "result": "Thay vỏ & cáp USB / 更換外殼與USB線",
        "qty": 1,
        "unit_price": 2150000,
        "lead_time": "4–5週",
    },
    {
        "serial": "EF550RA4LAWTEST902",
        "issue": "Camera không lấy nét và đèn flash yếu\n相機無法對焦且閃光燈亮度不足",
        "department": "QA",
        "result": "Thay camera & đèn flash / 更換相機與閃光燈",
        "qty": 1,
        "unit_price": 3850000,
        "lead_time": "5–6週",
    },
]

NOTES = "\n".join(
    [
        "1. RFR900NA55RSVHBA036 (BF606-1677) – Không có pin trước khi nhận sửa chữa / 送修前無電池",
        "2. EF550R4ALAWBH109 (BF606-1677) – Vỏ máy trầy xước nhẹ trước khi nhận / 送修前機殼有輕微刮痕",
        "3. EF550R4ALAWBI661 (BF606-1677) – Thiếu dây đeo tay khi bàn giao / 交付時缺少手腕帶",
        "4. EF550R4ALAWBI677 (BF606-1677) – Tem serial bị mờ trước khi sửa chữa / 維修前序號標籤模糊",
        "5. EF550RA4LAWBI666 (BF606-1677) – Không có pin trước khi nhận sửa chữa / 送修前無電池",
        "6. EF550RA4LAWBI678 (BF606-1677) – Không có pin và nắp che cổng sạc / 無電池且缺少充電埠護蓋",
        "7. RFR900NA55RSVTEST901 (RF-TEST) – Vỏ cò bắn đã nứt trước khi tiếp nhận / 接收前扳機外殼已破裂",
        "8. Thiết bị được đóng gói chung trong 03 thùng carton; đề nghị đối chiếu số lượng khi nhận / 設備共裝於三個紙箱，收貨時請核對數量",
    ]
)


def _configure_stdout() -> None:
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")


def _copy_cell(source, target, row_delta: int = 0) -> None:
    if source.data_type == "f" and isinstance(source.value, str):
        try:
            target.value = Translator(source.value, origin=source.coordinate).translate_formula(
                target.coordinate
            )
        except Exception:
            target.value = source.value
    else:
        target.value = source.value
    if source.has_style:
        target._style = copy.copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)
    target.hyperlink = copy.copy(source.hyperlink)
    target.comment = copy.copy(source.comment)


def _shift_range(cell_range: str, row_delta: int) -> str:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    return (
        f"{get_column_letter(min_col)}{min_row + row_delta}:"
        f"{get_column_letter(max_col)}{max_row + row_delta}"
    )


def _move_footer_down(ws, old_footer_start: int, row_delta: int) -> None:
    old_max_row = ws.max_row
    old_max_col = ws.max_column
    footer_merges = [
        str(merged)
        for merged in ws.merged_cells.ranges
        if range_boundaries(str(merged))[1] >= old_footer_start
    ]

    for merged in footer_merges:
        ws.unmerge_cells(merged)

    # Đi từ dưới lên để vùng đích không ghi đè dữ liệu nguồn khi hai vùng chồng nhau.
    for row in range(old_max_row, old_footer_start - 1, -1):
        new_row = row + row_delta
        ws.row_dimensions[new_row].height = ws.row_dimensions[row].height
        ws.row_dimensions[new_row].hidden = ws.row_dimensions[row].hidden
        for col in range(1, old_max_col + 1):
            _copy_cell(ws.cell(row, col), ws.cell(new_row, col), row_delta)

    for merged in footer_merges:
        ws.merge_cells(_shift_range(merged, row_delta))


def _style_data_rows(ws, new_data_end: int) -> None:
    template_row = 12
    for row in range(12, new_data_end + 1):
        for col in range(1, ws.max_column + 1):
            source = ws.cell(template_row, col)
            target = ws.cell(row, col)
            if source.has_style:
                target._style = copy.copy(source._style)
            target.font = copy.copy(source.font)
            target.fill = copy.copy(source.fill)
            target.border = copy.copy(source.border)
            target.alignment = copy.copy(source.alignment)
            target.protection = copy.copy(source.protection)


def _weighted_text_length(text: str) -> float:
    total = 0.0
    for char in text:
        code = ord(char)
        if (
            0x3400 <= code <= 0x4DBF
            or 0x4E00 <= code <= 0x9FFF
            or 0xF900 <= code <= 0xFAFF
        ):
            total += 1.95
        elif char.isupper():
            total += 1.2
        else:
            total += 0.95
    return total


def _estimated_lines(text: Any, column_width: float) -> int:
    if text in (None, ""):
        return 1
    usable_width = max(8.0, float(column_width or 13.0) * 1.35)
    lines = 0
    for paragraph in str(text).splitlines() or [""]:
        lines += max(1, math.ceil(_weighted_text_length(paragraph) / usable_width))
    return lines


def _apply_anti_clipping(ws, row: int, max_col: int) -> None:
    max_lines = 1
    for col in range(2, max_col + 1):
        cell = ws.cell(row, col)
        alignment = copy.copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "center"
        cell.alignment = alignment
        column_dimension = ws.column_dimensions.get(get_column_letter(col))
        width = column_dimension.width if column_dimension and column_dimension.width else 13.0
        max_lines = max(max_lines, _estimated_lines(cell.value, width))
    ws.row_dimensions[row].height = max(45.0, max_lines * 19.5 + 16.0)


def _set_common_item_values(ws, item: Dict[str, Any], row: int, quote_sheet: bool) -> None:
    ws.cell(row, 1).value = row - 11
    ws.cell(row, 2).value = item["serial"]
    ws.cell(row, 3).value = item["issue"]
    ws.cell(row, 4).value = item["department"]
    ws.cell(row, 5).value = item["result"]
    ws.cell(row, 6).value = item["qty"]
    if quote_sheet:
        ws.cell(row, 7).value = item["lead_time"]
        ws.cell(row, 8).value = item["unit_price"]
        ws.cell(row, 9).value = f"=F{row}*H{row}"
    else:
        ws.cell(row, 7).value = item["unit_price"]
        ws.cell(row, 8).value = f"=F{row}*G{row}"
    _apply_anti_clipping(ws, row, 9 if quote_sheet else 8)


def _expand_sheet(ws, quote_sheet: bool) -> Dict[str, Any]:
    old_footer_start = 18
    old_count = 6
    new_count = len(ITEMS)
    row_delta = new_count - old_count
    new_data_end = 11 + new_count
    new_footer_start = old_footer_start + row_delta

    _move_footer_down(ws, old_footer_start, row_delta)
    _style_data_rows(ws, new_data_end)

    for row, item in enumerate(ITEMS, start=12):
        _set_common_item_values(ws, item, row, quote_sheet)

    if quote_sheet:
        total_row = 19 + row_delta
        ws.cell(total_row, 9).value = f"=SUM(I12:I{new_data_end})"
    else:
        blank_row = 18 + row_delta
        total_row = 19 + row_delta
        note_row = 20 + row_delta
        ws.cell(total_row, 8).value = f"=SUM(H12:H{new_data_end})"
        ws.cell(note_row, 1).value = NOTES
        note_lines = len(NOTES.splitlines())
        ws.row_dimensions[note_row].height = max(68.55, note_lines * 19.5 + 16.0)
        ws.row_dimensions[blank_row].height = 24.0

        # Giữ nguyên vùng in của mẫu nhưng mở rộng đúng bằng số dòng bổ sung.
        ws.print_area = f"$A$1:$I${35 + row_delta}"

    return {
        "sheet": ws.title,
        "record_count": new_count,
        "data_range": f"A12:{get_column_letter(9 if quote_sheet else 8)}{new_data_end}",
        "footer_start": new_footer_start,
        "total_row": 19 + row_delta,
    }


def _safe_save(workbook, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=output_path.suffix,
        dir=str(output_path.parent),
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        try:
            os.replace(temp_path, output_path)
            return output_path
        except PermissionError:
            fallback = output_path.with_name(f"{output_path.stem}_v2{output_path.suffix}")
            os.replace(temp_path, fallback)
            return fallback
    finally:
        if temp_path.exists():
            temp_path.unlink()


def generate(template_path: Path, output_path: Path) -> Dict[str, Any]:
    workbook = load_workbook(template_path, data_only=False)
    if "CX Technology (VN)" not in workbook.sheetnames or "Sheet1" not in workbook.sheetnames:
        raise ValueError("Workbook mẫu không có đủ hai sheet CX Technology (VN) và Sheet1.")

    primary = _expand_sheet(workbook["CX Technology (VN)"], quote_sheet=False)
    quotation = _expand_sheet(workbook["Sheet1"], quote_sheet=True)
    saved_path = _safe_save(workbook, output_path)

    return {
        "status": "success",
        "template": str(template_path),
        "saved_path": str(saved_path),
        "record_count": len(ITEMS),
        "specific_note_count": 7,
        "generic_note_count": 1,
        "sheets": [primary, quotation],
    }


def main() -> int:
    _configure_stdout()
    parser = argparse.ArgumentParser(description="Tạo file sửa chữa test 14 bản ghi từ mẫu CX.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(generate(args.template, args.output), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
