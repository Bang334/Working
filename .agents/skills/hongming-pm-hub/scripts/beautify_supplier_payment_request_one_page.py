# -*- coding: utf-8 -*-
"""Dàn lại phiếu yêu cầu thanh toán NCC thành một trang A4 dọc.

Script đọc dữ liệu từ file Word gốc, giữ nguyên câu chữ và tạo một
bản trình bày gọn, cân đối hơn. Nếu file đích đang bị khóa, script
tự động lưu sang tên ``_v2``.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor


FONT_LATIN = "Times New Roman"
FONT_SANS = "Arial"
FONT_CJK = "Microsoft JhengHei"
BLUE = "1F4E79"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F8FC"
GRID = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_margins(cell, top=45, bottom=45, left=65, right=65):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcMar"))
    if old is not None:
        tc_pr.remove(old)
    tc_mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>'))


def set_table_borders(table, color=GRID, size="5", outer_color=None):
    outer = outer_color or color
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{size}" w:space="0" w:color="{outer}"/>'
        f'<w:left w:val="single" w:sz="{size}" w:space="0" w:color="{outer}"/>'
        f'<w:bottom w:val="single" w:sz="{size}" w:space="0" w:color="{outer}"/>'
        f'<w:right w:val="single" w:sz="{size}" w:space="0" w:color="{outer}"/>'
        f'<w:insideH w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(parse_xml(xml))


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    tbl_pr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            '<w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
            '<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/>'
            '</w:tblBorders>'
        )
    )


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_widths(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(total * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        keep_row_together(row)
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                break
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_run(run, font=FONT_LATIN, size=8.0, bold=False, italic=False, color=BLACK):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_fonts.set(qn("w:cs"), font)


def format_paragraph(
    paragraph,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    before=0,
    after=0,
    line=1.0,
    keep_next=False,
):
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.widow_control = True
    return paragraph


def write_paragraph(
    paragraph,
    text,
    font=FONT_LATIN,
    size=8.0,
    bold=False,
    italic=False,
    color=BLACK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    before=0,
    after=0,
    line=1.0,
    keep_next=False,
):
    paragraph.clear()
    format_paragraph(paragraph, align, before, after, line, keep_next)
    run = paragraph.add_run(text)
    style_run(run, font, size, bold, italic, color)
    return paragraph


def write_cell(
    cell,
    text,
    font=FONT_LATIN,
    size=8.0,
    bold=False,
    italic=False,
    color=BLACK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    line=1.0,
    fill=None,
    top=45,
    bottom=45,
    left=65,
    right=65,
):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        set_cell_shading(cell, fill)
    set_cell_margins(cell, top, bottom, left, right)
    write_paragraph(cell.paragraphs[0], text, font, size, bold, italic, color, align, line=line)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    write_paragraph(
        p,
        text,
        font=FONT_SANS,
        size=9.8,
        bold=True,
        color=BLUE,
        before=4,
        after=3,
        keep_next=True,
    )
    return p


def note_lines(cell_text):
    return [line.strip() for line in cell_text.splitlines()]


def parse_note_blocks(cell_text):
    lines = note_lines(cell_text)
    while lines and not lines[0]:
        lines.pop(0)
    if lines and lines[0].rstrip("：:").strip() in {"備註", "Ghi chú"}:
        lines.pop(0)
    tax_note = ""
    content = []
    for line in lines:
        if not line:
            continue
        if line.startswith("**") or "VAT" in line and "以上價格" in line:
            tax_note = line.lstrip("*").strip()
        else:
            content.append(line)
    blocks = []
    for idx in range(0, len(content), 2):
        pair = content[idx : idx + 2]
        if pair:
            blocks.append(pair)
    return blocks, tax_note


def add_note_block(cell, number, lines):
    set_cell_margins(cell, top=65, bottom=65, left=80, right=80)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(cell, WHITE if number % 2 else PALE_BLUE)
    first = cell.paragraphs[0]
    first.clear()
    format_paragraph(first, before=0, after=1, line=1.0)
    number_run = first.add_run(f"{number:02d}  ")
    style_run(number_run, FONT_SANS, 9.5, True, False, MID_BLUE)
    if lines:
        zh_run = first.add_run(lines[0])
        style_run(zh_run, FONT_CJK, 8.4, False, False, BLACK)
    if len(lines) > 1:
        p_vi = cell.add_paragraph()
        format_paragraph(p_vi, before=0, after=0, line=1.0)
        p_vi.paragraph_format.left_indent = Pt(18)
        vi_run = p_vi.add_run(lines[1])
        style_run(vi_run, FONT_LATIN, 8.4, False, False, BLACK)


def build_one_page(source: Path, target: Path) -> Path:
    src = Document(str(source))
    if len(src.tables) < 5 or len(src.paragraphs) < 6:
        raise ValueError("File nguồn không đúng cấu trúc phiếu thanh toán mong đợi.")

    src_meta, src_pay, src_notes, src_bank, src_sig = src.tables[:5]
    blocks, tax_note = parse_note_blocks(src_notes.cell(0, 0).text)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.85)
    section.bottom_margin = Cm(0.85)
    section.left_margin = Cm(0.75)
    section.right_margin = Cm(0.75)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    total_width = 7.68

    # Compact company header and form number.
    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(header, [5.25, 2.43])
    remove_table_borders(header)
    left = header.cell(0, 0)
    set_cell_margins(left, 0, 0, 0, 40)
    p = left.paragraphs[0]
    write_paragraph(p, src.paragraphs[0].text.replace(" ", ""), FONT_CJK, 14.5, True, align=WD_ALIGN_PARAGRAPH.LEFT, line=0.95)
    p2 = left.add_paragraph()
    write_paragraph(p2, src.paragraphs[1].text, FONT_SANS, 9.2, True, align=WD_ALIGN_PARAGRAPH.LEFT, line=0.95)
    right = header.cell(0, 1)
    set_cell_margins(right, 0, 0, 40, 0)
    form_text = src.paragraphs[5].text.replace("     ", " ")
    write_paragraph(right.paragraphs[0], form_text, FONT_SANS, 8.8, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.RIGHT, line=0.95)

    title = doc.add_paragraph()
    write_paragraph(
        title,
        f"{src.paragraphs[2].text.replace(' ', '')}   {src.paragraphs[3].text}",
        FONT_CJK,
        17.0,
        True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=3,
        after=4,
        line=0.9,
        keep_next=True,
    )

    salute = doc.add_table(rows=1, cols=2)
    salute.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(salute, [5.4, 2.28])
    remove_table_borders(salute)
    write_cell(salute.cell(0, 0), src.paragraphs[4].text, FONT_CJK, 8.5, False, align=WD_ALIGN_PARAGRAPH.LEFT, top=28, bottom=28, left=0, right=30)
    write_cell(salute.cell(0, 1), "", FONT_LATIN, 7.8, align=WD_ALIGN_PARAGRAPH.RIGHT, top=20, bottom=20, left=30, right=0)

    # Applicant/project metadata.
    meta = doc.add_table(rows=3, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(meta, [1.20, 4.03, 1.15, 1.30])
    set_table_borders(meta, GRID, "4")
    labels = (
        (0, 0, src_meta.cell(0, 0).text),
        (0, 2, src_meta.cell(0, 2).text),
        (1, 0, src_meta.cell(1, 0).text),
        (2, 0, src_meta.cell(2, 0).text),
        (2, 2, src_meta.cell(2, 2).text),
    )
    for r, c, text in labels:
        write_cell(meta.cell(r, c), text, FONT_CJK, 8.25, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.RIGHT, fill=LIGHT_BLUE, line=1.0, top=62, bottom=62, left=55, right=55)
    write_cell(meta.cell(0, 1), src_meta.cell(0, 1).text, FONT_LATIN, 9.2, True, line=1.0, top=62, bottom=62)
    write_cell(meta.cell(0, 3), src_meta.cell(0, 3).text, FONT_LATIN, 9.2, False, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, top=62, bottom=62)
    project = meta.cell(1, 1).merge(meta.cell(1, 3))
    write_cell(project, src_meta.cell(1, 1).text, FONT_LATIN, 8.7, False, line=1.0, top=65, bottom=65, left=75, right=75)
    write_cell(meta.cell(2, 1), src_meta.cell(2, 1).text, FONT_LATIN, 9.0, False, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, top=58, bottom=58)
    write_cell(meta.cell(2, 3), src_meta.cell(2, 3).text, FONT_CJK, 8.7, False, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, top=58, bottom=58)

    add_section_heading(doc, "NỘI DUNG THANH TOÁN   請款內容")

    pay = doc.add_table(rows=3, cols=9)
    pay.alignment = WD_TABLE_ALIGNMENT.CENTER
    pay_widths = [0.32, 1.75, 0.44, 0.85, 0.42, 0.96, 0.50, 0.52, 1.92]
    set_table_widths(pay, pay_widths)
    set_table_borders(pay, GRID, "4", outer_color=MID_BLUE)
    set_repeat_header(pay.rows[0])
    for idx in range(9):
        write_cell(
            pay.cell(0, idx),
            src_pay.cell(0, idx).text,
            FONT_CJK,
            7.8,
            True,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            line=0.95,
            fill=BLUE,
            top=62,
            bottom=62,
            left=28,
            right=28,
        )
    for idx in range(9):
        align = WD_ALIGN_PARAGRAPH.CENTER
        if idx == 1:
            align = WD_ALIGN_PARAGRAPH.LEFT
        elif idx in (3, 5, 8):
            align = WD_ALIGN_PARAGRAPH.RIGHT
        write_cell(
            pay.cell(1, idx),
            src_pay.cell(1, idx).text,
            FONT_CJK if idx == 1 else FONT_LATIN,
            8.55,
            idx == 8,
            color=BLACK,
            align=align,
            line=1.0,
            fill=WHITE,
            top=72,
            bottom=72,
            left=35,
            right=35,
        )
    total_label = pay.cell(2, 0).merge(pay.cell(2, 4))
    total_value = pay.cell(2, 5).merge(pay.cell(2, 8))
    write_cell(total_label, src_pay.cell(2, 0).text, FONT_CJK, 9.2, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.RIGHT, fill=LIGHT_BLUE, line=1.0, top=60, bottom=60)
    write_cell(total_value, src_pay.cell(2, 5).text, FONT_LATIN, 10.6, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.RIGHT, fill=LIGHT_BLUE, line=1.0, top=60, bottom=60)

    add_section_heading(doc, "ĐIỀU KHOẢN THANH TOÁN   付款條款")

    # Hiển thị số hàng đúng bằng số lượng điều khoản thực tế: "Bao nhiêu điều khoản bấy nhiêu hàng"
    notes_rows = max(1, len(blocks))
    notes = doc.add_table(rows=notes_rows, cols=1)
    notes.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(notes, [total_width])
    set_table_borders(notes, GRID, "4", outer_color=GRID)
    for idx in range(1, notes_rows + 1):
        cell = notes.cell(idx - 1, 0)
        lines = blocks[idx - 1] if idx - 1 < len(blocks) else []
        add_note_block(cell, idx, lines)

    if tax_note:
        tax = doc.add_paragraph()
        write_paragraph(tax, tax_note, FONT_CJK, 8.4, True, color=BLACK, before=4, after=7, line=1.0, keep_next=True)

    add_section_heading(doc, "THÔNG TIN CHUYỂN KHOẢN   匯款資訊")
    bank = doc.add_table(rows=2, cols=4)
    bank.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(bank, [1.28, 3.05, 1.05, 2.30])
    set_table_borders(bank, GRID, "4", outer_color=GRID)
    write_cell(bank.cell(0, 0), src_bank.cell(0, 0).text, FONT_CJK, 8.6, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, fill=LIGHT_BLUE, line=1.0, top=65, bottom=65)
    write_cell(bank.cell(0, 1), src_bank.cell(0, 1).text, FONT_SANS, 8.9, True, line=1.0, top=65, bottom=65)
    write_cell(bank.cell(0, 2), src_bank.cell(0, 2).text, FONT_CJK, 8.6, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, fill=LIGHT_BLUE, line=1.0, top=65, bottom=65)
    write_cell(bank.cell(0, 3), src_bank.cell(0, 3).text, FONT_SANS, 9.4, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, top=65, bottom=65)
    write_cell(bank.cell(1, 0), src_bank.cell(1, 0).text, FONT_CJK, 8.6, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, fill=LIGHT_BLUE, line=1.0, top=65, bottom=65)
    bank_name = bank.cell(1, 1).merge(bank.cell(1, 3))
    write_cell(bank_name, src_bank.cell(1, 1).text, FONT_SANS, 8.9, True, line=1.0, top=65, bottom=65)

    signatures = doc.add_table(rows=2, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(signatures, [total_width / 2, total_width / 2])
    remove_table_borders(signatures)
    signatures.rows[0].height = Pt(20)
    signatures.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    signatures.rows[1].height = Pt(72)
    signatures.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for col in range(2):
        write_cell(signatures.cell(0, col), src_sig.cell(0, col).text, FONT_CJK, 9.8, True, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER, top=55, bottom=0)
        write_cell(
            signatures.cell(1, col),
            src_sig.cell(1, col).text,
            FONT_LATIN,
            8.7,
            False,
            italic=True,
            color=BLACK,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            line=0.9,
            top=0,
            bottom=0,
        )
        signatures.cell(1, col).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(target))
        return target
    except PermissionError:
        fallback = target.with_name(f"{target.stem}_v2{target.suffix}")
        doc.save(str(fallback))
        return fallback


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    result = build_one_page(args.source, args.target)
    print(result)


if __name__ == "__main__":
    main()
