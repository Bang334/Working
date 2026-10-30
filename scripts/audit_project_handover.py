"""Read-only inventory and content extraction for the 30 Oct PM handover."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import hashlib
import json
import re
import sqlite3
import warnings
from pathlib import Path
from collections import Counter
from datetime import datetime
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import openpyxl
from pypdf import PdfReader

ROOT=Path('D:/Working')
SOURCE=Path(r'C:\Users\HM-Sales01\Downloads\drive-download-20261030T014758Z-1-001')
OUT=ROOT/'.qa/project_handover_20261030'
warnings.filterwarnings('ignore',category=UserWarning)

def dump(path,data):
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf-8')

def extract(path):
    ext=path.suffix.lower()
    if ext=='.pdf':
        reader=PdfReader(path)
        pages=[p.extract_text() or '' for p in reader.pages]
        return '\n\n'.join(f'[Trang {i+1}]\n{t}' for i,t in enumerate(pages)),{'pages':len(pages),'scan_pages':[i+1 for i,t in enumerate(pages) if len(re.sub(r'\s','',t))<40]}
    if ext=='.xlsx':
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        chunks=[];sheets=[]
        for ws in wb:
            sheets.append({'name':ws.title,'rows':ws.max_row,'cols':ws.max_column})
            for row in ws.iter_rows():
                cells=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                if cells:chunks.append(f'[{ws.title}] '+' | '.join(cells))
        wb.close()
        return '\n'.join(chunks),{'sheets':sheets}
    if ext in ('.docx','.pptx'):
        with ZipFile(path) as z:
            names=[n for n in z.namelist() if (ext=='.docx' and re.match(r'word/(document|header\d+|footer\d+)\.xml$',n)) or (ext=='.pptx' and re.match(r'ppt/slides/slide\d+\.xml$',n))]
            chunks=[]
            for name in names:
                node=ET.fromstring(z.read(name))
                lines=[]
                for para in node.iter():
                    if para.tag.endswith('}p'):
                        text=''.join(n.text or '' for n in para.iter() if n.tag.endswith('}t'))
                        if text.strip():lines.append(text)
                chunks.append(f'[{name}]\n'+'\n'.join(lines))
            return '\n\n'.join(chunks),{}
    return '',{'unread_reason':'Ảnh cần kiểm tra trực quan' if ext in ('.jpg','.png') else 'Word .doc cũ cần chuyển đổi'}

def main():
    OUT.mkdir(parents=True,exist_ok=True);(OUT/'text').mkdir(exist_ok=True)
    conn=sqlite3.connect(f'file:{ROOT / "database/pm_portal.db"}?mode=ro',uri=True);conn.row_factory=sqlite3.Row
    tables={r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    snapshot={table:[dict(r) for r in conn.execute(f'SELECT * FROM {table}')] for table in ['contracts','project_receivables','payment_milestones','documents','purchase_orders'] if table in tables}
    conn.close();dump(OUT/'database_snapshot.json',snapshot)
    files=[];hashes={}
    for i,path in enumerate(sorted(p for p in SOURCE.rglob('*') if p.is_file()),1):
        relative=path.relative_to(SOURCE)
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        item={'id':i,'customer':relative.parts[0],'relative':str(relative),'path':str(path),'extension':path.suffix.lower(),'bytes':path.stat().st_size,'sha256':digest}
        if digest in hashes:
            item['duplicate_of']=hashes[digest]
            item['text_file']=files[hashes[digest]-1].get('text_file')
        else:
            hashes[digest]=i
            try:
                text,metadata=extract(path)
                item.update(metadata);item['text_chars']=len(text)
                if text:
                    dest=OUT/'text'/f'{i:03d}.txt';dest.write_text(text,encoding='utf-8');item['text_file']=str(dest)
            except Exception as exc:
                item['extraction_error']=f'{type(exc).__name__}: {exc}'
        files.append(item)
        if i%40==0:print(f'Processed {i}',flush=True)
    dump(OUT/'inventory.json',{'source':str(SOURCE),'files':files})
    groups=[]
    for folder in sorted(p for p in SOURCE.iterdir() if p.is_dir()):
        entries=[f for f in files if f['customer']==folder.name]
        groups.append({'customer':folder.name,'files':len(entries),'unique':sum('duplicate_of' not in f for f in entries),'types':dict(Counter(f['extension'] for f in entries))})
    dump(OUT/'summary.json',{'customers':groups,'files':len(files),'unique_hashes':len(hashes),'duplicates':sum('duplicate_of' in f for f in files),'errors':[f for f in files if 'extraction_error' in f],'scans':[{'id':f['id'],'relative':f['relative'],'pages':f.get('pages'),'scan_pages':f.get('scan_pages')} for f in files if f.get('scan_pages')]})
    print(json.dumps({'files':len(files),'unique':len(hashes),'duplicates':len(files)-len(hashes),'database':{t:len(r) for t,r in snapshot.items()}},ensure_ascii=False))

if __name__=='__main__':main()
