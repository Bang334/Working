"""Bound print ranges and verify compact item tables without touching cell data."""
from pathlib import Path
import zipfile,xml.etree.ElementTree as E,os
import openpyxl
P=Path('D:/Working/outputs/seiyo_20261030')
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
q=lambda x:'{'+N+'}'+x
for kind,name,last,area in [('invoice','INVOICE_HONG_MING_SEIYO_20261030_GON.xlsx',55,"'INV'!$A$16:$L$55"),('packing','PKL_HONG_MING_SEIYO_20261030_GON.xlsx',37,"'Packing List'!$A$1:$H$37")]:
 path=P/name;temp=P/(name+'.tmp')
 with zipfile.ZipFile(path) as src,zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as dst:
  for f in src.infolist():
   data=src.read(f.filename)
   if f.filename=='xl/workbook.xml':
    root=E.fromstring(data);dn=root.find(q('definedNames'))
    if dn is None:dn=E.SubElement(root,q('definedNames'))
    pa=next((d for d in dn if d.get('name')=='_xlnm.Print_Area'),None)
    if pa is None:pa=E.SubElement(dn,q('definedName'),{'name':'_xlnm.Print_Area','localSheetId':'0'})
    pa.text=area;data=E.tostring(root,encoding='utf-8',xml_declaration=True)
   dst.writestr(f,data)
 try:os.replace(temp,path)
 except PermissionError:
  path=path.with_stem(path.stem+'_v2');os.replace(temp,path)
 w=openpyxl.load_workbook(path,data_only=False).active
 v=openpyxl.load_workbook(path,data_only=True).active
 items=[c.value for row in w for c in row if c.column==1 and isinstance(c.value,(int,float))]
 assert items==list(range(1,22 if kind=='invoice' else 21)),items
 if kind=='invoice':
  assert v['L55'].value==712020000
  assert w['L55'].value=='=SUM(L33:L42,L44:L54)'
 else:
  assert all(v[f'E{r}'].value is None for r in range(13,33))
  assert v['G33'].value in [None,'']
 for row in w.iter_rows(min_row=last+1,max_row=max(last+1,w.max_row)):
  assert all(c.value is None for c in row)
 assert not any(c.data_type=='e' for row in v for c in row)
 print('Verified',kind,'items',len(items),'print',w.print_area)
