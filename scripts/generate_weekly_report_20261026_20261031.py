import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor

def set_run_props(run, font_name='微軟正黑體', size_pt=11.0, bold=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color_rgb is not None:
        run.font.color.rgb = color_rgb

def create_weekly_report():
    src_path = r'd:\Working\Ho So Du An\Bao Cao\Tuan\Bao_Cao_Tuan_BANG_20261019-20261024.pptx'
    dst_path = r'd:\Working\Ho So Du An\Bao Cao\Tuan\Bao_Cao_Tuan_BANG_20261026-20261031.pptx'

    prs = pptx.Presentation(src_path)

    # --- Slide 1: Cover ---
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.name == '副標題 4' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "2026/10/26– 2026/10/31\nNGUYEN SY KIM BANG"
            set_run_props(run, font_name='微軟正黑體', size_pt=20.0, bold=False, color_rgb=RGBColor(255, 255, 255))

    # --- Slide 2: Schedule Table ---
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.name == '標題 1' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "10/26–10/31 This Week’s Schedule（10/31 调整）"
            set_run_props(run, font_name='微軟正黑體', size_pt=18.0, bold=True, color_rgb=RGBColor(255, 255, 255))
        elif shape.has_table:
            table = shape.table
            headers = ['日期', '10/26 (一)', '10/27 (二)', '10/28 (三)', '10/29 (四)', '10/30 (五)', '10/31 (六)', '11/01 (日)']
            for c_idx, h in enumerate(headers):
                cell = table.cell(0, c_idx)
                cell.text_frame.clear()
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = h
                set_run_props(run, font_name='Arial', size_pt=10.0, bold=True, color_rgb=RGBColor(255, 255, 255))

            morning_tasks = [
                '上午時段 (8:00~12:00)',
                'Theo dõi công nợ; nhận hàng Darfon kiểm tra, lên danh sách. 追蹤應收帳款；接收達方（Darfon）設備點收檢驗並建立清冊。',
                'Khảo sát hiện trường dự án (buổi sáng). 專案現場實地勘查（上午時段）。',
                'Đi gặp gỡ và làm việc với khách hàng. 外出拜訪客戶洽談專案業務。',
                'Rà soát các công việc được bàn giao. 研讀並梳理各項交接之專案工作。',
                '（Hôm nay: công việc Thứ 6 sẽ cập nhật sau） （今日週五工作事項後續補充更新）',
                '（Công việc Thứ 7 sẽ cập nhật sau） （週六工作事項後續補充更新）',
                ''
            ]
            for c_idx, task in enumerate(morning_tasks):
                cell = table.cell(1, c_idx)
                cell.text_frame.clear()
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = task
                size = 9.5 if c_idx > 0 else 10.0
                bold = (c_idx == 0)
                set_run_props(run, font_name='Arial', size_pt=size, bold=bold, color_rgb=RGBColor(255, 255, 255))

            afternoon_tasks = [
                '下午時段 (13:00~17:00)',
                'Đọc mail, xem cách viết mail và tham khảo mail mẫu. 研讀專案郵件，學習郵件撰寫規範並參考過往信件範例。',
                'Tiếp tục khảo sát hiện trường dự án (buổi chiều). 持續進行專案現場實地勘查（下午時段）。',
                'Nghiên cứu thông tin cuộc họp sáng; nhận bàn giao từ anh Sung. 研讀上午會議資訊；接收 Sung 主管交接專案工作。',
                'Theo dõi công nợ dự án, đôn đốc thu hồi nợ Vonmax. 追蹤專案應收帳款，催收弘聲（Vonmax）到期款項。',
                '（Hôm nay: công việc Thứ 6 sẽ cập nhật sau） （今日週五工作事項後續補充更新）',
                '（Công việc Thứ 7 sẽ cập nhật sau） （週六工作事項後續補充更新）',
                ''
            ]
            for c_idx, task in enumerate(afternoon_tasks):
                cell = table.cell(3, c_idx)
                cell.text_frame.clear()
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = task
                size = 9.5 if c_idx > 0 else 10.0
                bold = (c_idx == 0)
                set_run_props(run, font_name='Arial', size_pt=size, bold=bold, color_rgb=RGBColor(255, 255, 255))

    # Helper function to fill daily slides with proper run styling
    def fill_daily_slide(slide, title_text, items):
        # Update title
        for shape in slide.shapes:
            if shape.name == 'TextBox 11' and shape.has_text_frame:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = title_text
                set_run_props(run, font_name='微軟正黑體', size_pt=16.0, bold=True, color_rgb=RGBColor(255, 255, 255))

        # Find content text box
        content_shape = None
        for shape in slide.shapes:
            if shape.name in ['TextBox 2', 'TextBox 3'] and shape.has_text_frame:
                content_shape = shape
                break
        
        if content_shape:
            tf = content_shape.text_frame
            tf.clear()
            for i, item in enumerate(items):
                # Paragraph 1: Title (Bold, White, 13pt)
                p1 = tf.add_paragraph() if (i > 0 or len(tf.paragraphs[0].text) > 0) else tf.paragraphs[0]
                p1.space_before = Pt(8.0) if i > 0 else Pt(0.0)
                p1.space_after = Pt(2.0)
                run1 = p1.add_run()
                run1.text = item['title']
                set_run_props(run1, font_name='微軟正黑體', size_pt=13.0, bold=True, color_rgb=RGBColor(255, 255, 255))

                # Paragraph 2: Vietnamese detail (F0F0F0, 11pt)
                p2 = tf.add_paragraph()
                p2.space_before = Pt(0.0)
                p2.space_after = Pt(2.0)
                run2 = p2.add_run()
                run2.text = item['vn']
                set_run_props(run2, font_name='微軟正黑體', size_pt=11.0, bold=False, color_rgb=RGBColor(240, 240, 240))

                # Paragraph 3: Chinese detail (BEBEBE, 11pt)
                p3 = tf.add_paragraph()
                p3.space_before = Pt(0.0)
                p3.space_after = Pt(6.0)
                run3 = p3.add_run()
                run3.text = item['cn']
                set_run_props(run3, font_name='微軟正黑體', size_pt=11.0, bold=False, color_rgb=RGBColor(190, 190, 190))

    # --- Slide 3: 10/26 ---
    items_1026 = [
        {
            'title': '1. Theo dõi công nợ & Nhận hàng Darfon kiểm tra / 專案應收帳款追蹤與接收達方（Darfon）設備檢驗',
            'vn': 'Sáng theo dõi tình hình công nợ các dự án; tiếp nhận lô hàng thiết bị từ đối tác Darfon, tiến hành mở thùng kiểm tra số lượng, linh kiện, đối chiếu model và lập danh sách kiểm kê chi tiết.',
            'cn': '上午追蹤各專案應收帳款狀況；接收達方（Darfon）送達之專案設備貨物，開箱清點數量、核對規格型號與配件完整性並建立設備驗收清冊。'
        },
        {
            'title': '2. Nghiên cứu quy chuẩn và văn hóa viết email công vụ / 研讀商務郵件撰寫規範與過往信件範本',
            'vn': 'Buổi chiều đọc hệ thống email dự án, tìm hiểu quy chuẩn trao đổi thông tin, văn hóa gửi email công vụ và tham khảo các email mẫu trước đó để nâng cao nghiệp vụ giao tiếp văn bản.',
            'cn': '下午研讀專案歷史郵件系統，深入學習商務郵件溝通規範與對外信件禮儀，參考標準郵件範本以提升書面業務對接能力。'
        }
    ]
    fill_daily_slide(prs.slides[2], "10/26 当周排程工作结果记录", items_1026)

    # --- Slide 4: 10/27 ---
    items_1027 = [
        {
            'title': '1. Khảo sát hiện trường dự án (Buổi sáng) / 專案現場實地勘查（上午時段）',
            'vn': 'Sáng đi khảo sát thực tế tại hiện trường nhà máy dự án; kiểm tra vị trí tủ điện phân phối, khảo sát các đường cáp và đánh giá phương án thi công đi dây.',
            'cn': '上午前往專案工廠現場進行實地勘查；檢視配電盤安裝位置、通訊線路佈設路徑並評估現場施工可行性。'
        },
        {
            'title': '2. Khảo sát hiện trường dự án (Buổi chiều) / 專案現場實地勘查（下午時段）',
            'vn': 'Buổi chiều tiếp tục khảo sát khu vực máy móc sản xuất, vị trí dự kiến lắp đặt Gateway thu thập dữ liệu và hạ tầng mạng kết nối; ghi nhận các yêu cầu kỹ thuật đặc thù tại hiện trường.',
            'cn': '下午持續勘查現場生產機台區域、評估資料收集閘道器（Gateway）與通訊網路節點佈設環境；記錄現場各項具體技術需求。'
        },
        {
            'title': '3. Tổng hợp số liệu và lập biên bản khảo sát hiện trường / 彙整現場勘查數據與記錄',
            'vn': 'Tổng hợp các số liệu đo đạc kỹ thuật, đối chiếu hình ảnh chụp hiện trường để làm cơ sở xây dựng phương án kỹ thuật và báo giá chi tiết.',
            'cn': '彙總整理現場測量數據與勘查實拍照片，作為後續編制工程技術方案及精確報價之依據。'
        }
    ]
    fill_daily_slide(prs.slides[3], "10/27 当周排程工作结果记录", items_1027)

    # --- Slide 5: 10/28 ---
    items_1028 = [
        {
            'title': '1. Gặp gỡ và làm việc trực tiếp với khách hàng / 外出拜訪客戶洽談專案業務',
            'vn': 'Sáng đi gặp gỡ trực tiếp khách hàng; trao đổi về tiến độ triển khai dự án, thảo luận về nhu cầu kỹ thuật và các nội dung phối hợp thực hiện giữa hai bên.',
            'cn': '上午隨同主管外出拜訪客戶；洽談專案執行進度、深入探討現場技術需求及雙方業務配合事項。'
        },
        {
            'title': '2. Nghiên cứu tài liệu và nội dung cuộc họp sáng / 研讀與消化上午客戶會議相關資訊',
            'vn': 'Chiều nghiên cứu, rà soát lại các thông tin kỹ thuật, yêu cầu của khách hàng được ghi nhận trong cuộc họp sáng để chuẩn bị các phương án xử lý tiếp theo.',
            'cn': '下午研讀並整理上午客戶會議相關技術資訊與客戶反饋要點，研擬後續各項應對與對接措施。'
        },
        {
            'title': '3. Tiếp nhận bàn giao công việc từ anh Sung / 接收 Sung 主管交接專案工作',
            'vn': 'Họp trao đổi và tiếp nhận bàn giao toàn diện các đầu việc dự án, hồ sơ tài liệu và các nội dung công việc quan trọng cần tiếp tục theo dõi từ anh Sung.',
            'cn': '與 Sung 主管對接，全面接收各項專案工作任務、業務檔案與後續需持續推進之重點工作。'
        }
    ]
    fill_daily_slide(prs.slides[4], "10/28 当周排程工作结果记录", items_1028)

    # --- Slide 6: 10/29 ---
    items_1029 = [
        {
            'title': '1. Nghiên cứu và phân loại các công việc được bàn giao / 梳理與研讀交接之各項專案工作',
            'vn': 'Sáng rà soát chi tiết toàn bộ hồ sơ, hợp đồng và các danh mục công việc đã được anh Sung bàn giao; phân loại mức độ ưu tiên và lập kế hoạch thực hiện.',
            'cn': '上午詳細研讀並分類 Sung 主管交接之各項合約文件與工作清單，釐清優先順序並排定推進時程。'
        },
        {
            'title': '2. Theo dõi công nợ dự án & Đòi nợ Vonmax / 追蹤專案應收帳款，催收弘聲（Vonmax）款項',
            'vn': 'Chiều theo dõi, cập nhật bảng công nợ các dự án đang triển khai; liên hệ khách hàng Vonmax để đôn đốc việc thanh toán các khoản công nợ đến hạn theo hợp đồng.',
            'cn': '下午檢視並更新各在建專案之應收帳款追蹤表；積極聯繫客戶弘聲（Vonmax）財務與窗口，嚴格催收各筆到期款項。'
        }
    ]
    fill_daily_slide(prs.slides[5], "10/29 当周排程工作结果记录", items_1029)

    # --- Slide 7: 10/30 ---
    items_1030 = [
        {
            'title': '1. Công việc Thứ 6 (hôm nay sẽ cập nhật sau) / 今日週五工作事項（後續補充）',
            'vn': 'Rà soát hồ sơ tiếp nhận dự án, đối chiếu công nợ và các nội dung công việc phát sinh trong ngày Thứ 6 sẽ được cập nhật bổ sung sau khi kết thúc ngày làm việc.',
            'cn': '核對專案交接檔案、核算各專案款項，週五當日其餘具體執行事項將於下班前陸續補充更新。'
        }
    ]
    fill_daily_slide(prs.slides[6], "10/30 当周排程工作结果记录", items_1030)

    # --- Slide 8: 10/31 ---
    items_1031 = [
        {
            'title': '1. Công việc Thứ 7 (sẽ cập nhật sau) / 週六工作事項（後續補充）',
            'vn': 'Các nội dung công việc dự kiến phát sinh trong ngày Thứ 7 sẽ được cập nhật và hoàn thiện sau.',
            'cn': '週六預定執行事項與進度將於後續補充更新。'
        }
    ]
    fill_daily_slide(prs.slides[7], "10/31 当周排程工作结果记录", items_1031)

    # --- Slide 9: Next Week Cover ---
    slide9 = prs.slides[8]
    for shape in slide9.shapes:
        if shape.name == '副標題 4' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "2026/11/02– 2026/11/07\nNGUYEN SY KIM BANG"
            set_run_props(run, font_name='微軟正黑體', size_pt=20.0, bold=False, color_rgb=RGBColor(255, 255, 255))

    # --- Slide 10: Next Week Schedule Table ---
    slide10 = prs.slides[9]
    for shape in slide10.shapes:
        if shape.name == '標題 1' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "11/02–11/07 Next Week’s Tasks – Scheduled（11/01 調整）"
            set_run_props(run, font_name='微軟正黑體', size_pt=18.0, bold=True, color_rgb=RGBColor(255, 255, 255))
        elif shape.has_table:
            table = shape.table
            headers_next = ['日期', '11/02 (一)', '11/03 (二)', '11/04 (三)', '11/05 (四)', '11/06 (五)', '11/07 (六)', '11/08 (日)']
            for c_idx, h in enumerate(headers_next):
                cell = table.cell(0, c_idx)
                cell.text_frame.clear()
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = h
                set_run_props(run, font_name='Arial', size_pt=10.0, bold=True, color_rgb=RGBColor(255, 255, 255))

            morning_tasks_next = [
                '上午時段 (8:00~12:00)',
                'Theo dõi công nợ, kiểm tra danh mục Darfon. 追蹤應收帳款，核對達方設備清冊。',
                'Đôn đốc công nợ Vonmax và các dự án đến hạn. 催收 Vonmax 及到期款項。',
                'Theo dõi công nợ, đối chiếu các khoản thu dự án. 追蹤各專案應收帳款。',
                'Rà soát tiến độ các đầu việc anh Sung bàn giao. 盤點 Sung 主管交接各專案工作。',
                'Theo dõi công nợ các dự án đang triển khai. 追蹤各在建專案應收帳款。',
                'Sắp xếp, lưu trữ hồ sơ dự án và công nợ tuần. 整理歸檔專案文件與資料。',
                ''
            ]
            for c_idx, task in enumerate(morning_tasks_next):
                cell = table.cell(1, c_idx)
                cell.text_frame.clear()
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = task
                size = 9.5 if c_idx > 0 else 10.0
                bold = (c_idx == 0)
                set_run_props(run, font_name='Arial', size_pt=size, bold=bold, color_rgb=RGBColor(255, 255, 255))

            afternoon_tasks_next = [
                '下午時段 (13:00~17:00)',
                'Xử lý các đầu việc được anh Sung bàn giao. 落實執行交接之各項專案工作。',
                'Hoàn thiện phương án kỹ thuật sau khảo sát. 完善現場實地勘查後技術方案。',
                'Chuẩn bị chứng từ giao hàng thiết bị Darfon. 準備達方設備出貨交接憑證。',
                'Theo dõi công nợ và cập nhật bảng theo dõi. 持續追蹤應收帳款並更新表單。',
                'Tổng kết công việc tuần, chuẩn bị báo cáo tuần. 總結本週各項工作，編製週報。',
                'Rà soát lại toàn bộ công việc tuần mới. 盤點各項待辦業務，準備新週排程。',
                ''
            ]
            for c_idx, task in enumerate(afternoon_tasks_next):
                cell = table.cell(3, c_idx)
                cell.text_frame.clear()
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = task
                size = 9.5 if c_idx > 0 else 10.0
                bold = (c_idx == 0)
                set_run_props(run, font_name='Arial', size_pt=size, bold=bold, color_rgb=RGBColor(255, 255, 255))

    # --- Slide 11: To Do List ---
    slide11 = prs.slides[10]
    for shape in slide11.shapes:
        if shape.name == 'TextBox 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            
            p_vn_title = tf.paragraphs[0]
            run_vn_t = p_vn_title.add_run()
            run_vn_t.text = "Kế hoạch tuần tới"
            set_run_props(run_vn_t, font_name='微軟正黑體', size_pt=13.0, bold=True, color_rgb=RGBColor(255, 255, 255))

            p_cn_title = tf.add_paragraph()
            p_cn_title.space_after = Pt(6.0)
            run_cn_t = p_cn_title.add_run()
            run_cn_t.text = "下週工作計畫"
            set_run_props(run_cn_t, font_name='微軟正黑體', size_pt=13.0, bold=True, color_rgb=RGBColor(255, 255, 255))

            todo_vn = [
                "1. Tiếp tục theo dõi sát sao công nợ các dự án và đôn đốc các khoản thanh toán của khách hàng (đặc biệt Vonmax, CX).",
                "2. Hoàn thiện phương án kỹ thuật và các hồ sơ liên quan sau đợt khảo sát thực tế hiện trường dự án.",
                "3. Quản lý, kiểm soát danh mục thiết bị Darfon đã nhận và triển khai hiệu quả các nội dung công việc anh Sung bàn giao."
            ]
            for item in todo_vn:
                p = tf.add_paragraph()
                p.space_after = Pt(3.0)
                run = p.add_run()
                run.text = item
                set_run_props(run, font_name='微軟正黑體', size_pt=11.0, bold=False, color_rgb=RGBColor(240, 240, 240))

            todo_cn = [
                "1. 持續緊盯專案應收帳款，嚴格催收客戶到期款項（特別是 Vonmax、CX 等）。",
                "2. 完善專案現場實地勘查後之工程技術方案與相關對接文件。",
                "3. 妥善管理達方（Darfon）進貨設備清冊，穩步推進落實 Sung 主管交接之各項專案任務。"
            ]
            for item in todo_cn:
                p = tf.add_paragraph()
                p.space_after = Pt(3.0)
                run = p.add_run()
                run.text = item
                set_run_props(run, font_name='微軟正黑體', size_pt=11.0, bold=False, color_rgb=RGBColor(190, 190, 190))

    # Save presentation safely
    try:
        prs.save(dst_path)
        print(f"Successfully generated: {dst_path}")
    except PermissionError:
        dst_v2 = dst_path.replace(".pptx", "_v2.pptx")
        prs.save(dst_v2)
        print(f"File was locked. Successfully saved to fallback: {dst_v2}")

if __name__ == '__main__':
    create_weekly_report()
