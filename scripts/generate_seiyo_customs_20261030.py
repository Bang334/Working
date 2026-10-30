"""Read approved order; preserve original XLSX packages while updating authored cells."""
from pathlib import Path
import json, sys, zipfile, copy
import xml.etree.ElementTree as ET
import openpyxl

ROOT=Path('D:/Working/outputs/seiyo_20261030')
SRC=Path('C:/Users/HM-Sales01/Downloads')
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def q(t): return '{'+NS+'}'+t

def prepare():
 ROOT.mkdir(parents=True,exist_ok=True)
 w=openpyxl.load_workbook(SRC/'mau.xlsx',data_only=True).active
 rows=list(range(16,26))+list(range(27,38))
 units=['CÁI','CÁI','CÁI','CÁI','CÁI','CÁI','CÁI','CÁI','CÁI','CÁI','LÔ (350 m)','LÔ (960 m)','CÁI','CÁI','CÁI','CÁI','HỘP','LÔ (2.100 m)','LÔ (30 m)','LÔ','GÓI']
 descriptions=[
 'Tủ mạng 6U, màu đen, D500\n6U 黑色壁掛機櫃，深度500mm',
 'Ổ cắm điện PDU 6 cổng\n機架式電源排插 6孔',
 'Phích cắm CN 1P-16A\n插頭 CN 1P-16A',
 'Thanh quản lý cáp mạng ngang\n水平理線架',
 'Switch 24 cổng Gigabit + 4 SFP TL-SG3428\n24埠 Gigabit + 4埠 SFP 交換器',
 'Switch 16 cổng Gigabit + 2 SFP TL-SG2218\n16埠 Gigabit + 2埠 SFP 交換器',
 'Bộ nguồn 24V DC HDR-150-24\n24V DC 電源轉換器 HDR-150-24',
 'Hộp đấu nối sắt 500×400×150mm\n鐵製接線盒 500×400×150mm',
 'Bộ chuyển đổi quang điện 10/100/1000\n光電轉換器 1G',
 'Module quang SFP GBIC-1.25G\n光纖模組 SFP GBIC-1.25G',
 'Dây điện Cu/PVC/PVC 3C×1,5mm², 350m\n電線 Cu/PVC/PVC 3C×1.5mm²，350米',
 'Cáp quang SM 4FO, 960m\n單模4芯光纖纜線，960米',
 'Hộp ODF 24 khớp SC đôi, cố định, 1U\n光纖熔接盤 24埠，雙SC，1U',
 'Hộp ODF 8 khớp SC đôi\n光纖熔接盤 8埠，雙SC',
 'Hộp ODF 4 khớp SC đôi\n光纖熔接盤 4埠，雙SC',
 'Dây nhảy quang SC/LC 3m\n光纖跳線 SC/LC，3米',
 'Cáp mạng CAT6 UTP, 4 đôi\nCAT6 UTP 網路線，4對',
 'Ống PVC D25, D20, tổng 2.100m\nPVC 管 D25、D20，共2,100米',
 'Máng cáp có nắp 100×100×1,2mm, 30m\n附蓋電纜槽 100×100×1.2mm，30米',
 'Vật tư phụ: đầu cút, ốc vít, băng dính, ruột gà\n配件：接頭、螺絲、膠帶、軟管等',
 'Chi phí thi công\n施工費'
 ]
 items=[]
 for i,r in enumerate(rows):
  items.append(dict(no=i+1,row=r,name=w[f'B{r}'].value,description=descriptions[i],detail=w[f'D{r}'].value,qty=w[f'E{r}'].value,unit=units[i],price=w[f'F{r}'].value,total=w[f'G{r}'].value))
 assert sum(x['qty']*x['price'] for x in items)==712020000
 inv={}; pkl={}
 # Clear variable item data only; retain styles, merged cells and all page settings.
 old=openpyxl.load_workbook(SRC/'Copy of HongMinh - Seiyo ct Bắc Ninh RV5 25.8.2026 CHECK.xlsx').active
 for row in old.iter_rows(min_row=33,max_row=62,max_col=12):
  for c in row:
   if c.value is not None and c.row!=39:inv[c.coordinate]=None
 inv.update({'B6':'INVOICE / 發票 — DỰ THẢO / 草稿','C20':'INVOICE / 發票 — DỰ THẢO / 草稿','L8':'DATE: 30/10/2026','J21':'DATE: 30/10/2026','D21':None,'J27':'INCOTERM : FOB','A32':'DANH MỤC HÀNG HÓA VÀ DỊCH VỤ / 貨品與服務明細','A39':'DANH MỤC HÀNG HÓA VÀ DỊCH VỤ (tiếp) / 貨品與服務明細（續）'})
 slots=list(range(33,39))+list(range(40,63))
 short_names=['Tủ mạng 6U\n6U 壁掛機櫃','Ổ cắm PDU\n電源排插','Phích cắm\n插頭','Thanh cáp ngang\n水平理線架','Switch quản lý\n管理型交換器','Switch quản lý\n管理型交換器','Bộ nguồn 24V\n24V 電源','Hộp đấu nối\n接線盒','Bộ đổi quang điện\n光電轉換器','Module quang\n光纖模組','Dây điện\n電線','Cáp quang\n光纖纜線','ODF 24埠','ODF 8埠','ODF 4 cổng\n4埠光纖熔接盤']
 for it,r in zip(items,slots):
  desc=it['description'].replace('\n',' / ') if r>=46 else it['description']
  inv.update({f'A{r}':it['no'],f'C{r}':desc,f'H{r}':it['detail'],f'J{r}':str(int(it['qty']))+' '+it['unit'],f'K{r}':it['price'],f'L{r}':{'formula':f'=K{r}*{it["qty"]}'}})
  if r not in [47] and not 50<=r<=61: inv[f'B{r}']=it['name']
 for i,r in enumerate(slots[:15]):
  if r!=47:inv[f'B{r}']=short_names[i]
 inv['B46']='Hộp ODF\n光纖熔接盤'
 inv['B49']='Vật tư và dịch vụ\n材料與服務'
 inv['L63']={'formula':'=SUM(L33:L38,L40:L62)'}
 inv['K63']='TỔNG / 總價'
 old=openpyxl.load_workbook(SRC/'PKL_HONG_MING_SEIYO_20260818 FN.xlsx').active
 for row in old.iter_rows(min_row=13,max_row=41,max_col=8):
  for c in row:
   if c.value is not None:pkl[c.coordinate]=None
 pkl.update({'B4':'PACKING LIST / 裝箱單 — DỰ THẢO / 草稿','G7':None,'G8':'30/10/2026','A41':'TOTAL / 合計','D41':None,'G41':{'formula':'=IF(COUNT(G13:G32)=20,SUM(G13:G32),"")'}})
 for it,r in zip(items[:-1],range(13,33)):
  pkl.update({f'A{r}':it['no'],f'B{r}':it['description'],f'D{r}':str(int(it['qty']))+' '+it['unit'],f'E{r}':None,f'G{r}':{'formula':f'=IF(E{r}="","",E{r}*{it["qty"]})'}})
 data={'items':items,'invoice':inv,'packing':pkl,'total':712020000}
 (ROOT/'payload.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
 print('Prepared 21 invoice lines / 20 packing lines; total VND 712020000')

def transplant(kind,template,output):
 data=json.loads((ROOT/'payload.json').read_text(encoding='utf-8'))[kind]
 with zipfile.ZipFile(ROOT/(kind+'_authored.xlsx')) as a:
  ar=ET.fromstring(a.read('xl/worksheets/sheet1.xml'))
  authored={c.attrib['r']:c for c in ar.iter(q('c'))}
  ss=[]
  if 'xl/sharedStrings.xml' in a.namelist():
   ss=[''.join(x.itertext()) for x in ET.fromstring(a.read('xl/sharedStrings.xml'))]
 with zipfile.ZipFile(SRC/template) as z:
  root=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
  sd=root.find(q('sheetData')); cells={c.attrib['r']:c for c in root.iter(q('c'))}
  for addr in data:
   cell=cells.get(addr)
   if cell is None:
    rn=''.join(t for t in addr if t.isdigit()); row=next((r for r in sd if r.attrib['r']==rn),None)
    if row is None:row=ET.SubElement(sd,q('row'),{'r':rn})
    cell=ET.SubElement(row,q('c'),{'r':addr})
   for child in list(cell):
    if child.tag in [q('f'),q('v'),q('is')]:cell.remove(child)
   cell.attrib.pop('t',None)
   if data[addr] is None:continue
   ac=authored.get(addr)
   assert ac is not None,addr
   if ac.attrib.get('t')=='s':
    cell.attrib['t']='inlineStr'; ins=ET.SubElement(cell,q('is')); t=ET.SubElement(ins,q('t'));t.text=ss[int(ac.find(q('v')).text)]
   else:
    if ac.attrib.get('t'):cell.attrib['t']=ac.attrib['t']
    for child in ac:
     if child.tag in [q('f'),q('v'),q('is')]:cell.append(copy.deepcopy(child))
  # XML cells must stay in column order; original dimensions/settings remain untouched.
  from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
  for row in sd:
   row[:]=sorted(row,key=lambda c:column_index_from_string(coordinate_from_string(c.attrib['r'])[0]))
  target=ROOT/output
  try:
   with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as out:
    for entry in z.infolist():out.writestr(entry,ET.tostring(root,encoding='utf-8',xml_declaration=True) if entry.filename=='xl/worksheets/sheet1.xml' else z.read(entry.filename))
  except PermissionError:
   target=target.with_stem(target.stem+'_v2')
   with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as out:
    for entry in z.infolist():out.writestr(entry,ET.tostring(root,encoding='utf-8',xml_declaration=True) if entry.filename=='xl/worksheets/sheet1.xml' else z.read(entry.filename))
  v=openpyxl.load_workbook(target,data_only=True); f=openpyxl.load_workbook(target,data_only=False)
  original=openpyxl.load_workbook(SRC/template)
  assert str(f.active.merged_cells)==str(original.active.merged_cells)
  assert f.active.freeze_panes==original.active.freeze_panes
  for k,d in original.active.column_dimensions.items():assert d.width==f.active.column_dimensions[k].width
  for k,d in original.active.row_dimensions.items():assert d.height==f.active.row_dimensions[k].height
  if kind=='invoice':assert v.active['L63'].value==712020000,v.active['L63'].value
  else:
   assert all(v.active[f'E{r}'].value is None for r in range(13,41))
   assert v.active['G41'].value in [None,'']
  print('Verified',target)

if __name__=='__main__':
 if len(sys.argv)==1:prepare()
 else:
  transplant('invoice','Copy of HongMinh - Seiyo ct Bắc Ninh RV5 25.8.2026 CHECK.xlsx','INVOICE_HONG_MING_SEIYO_20261030_DU_THAO.xlsx')
  transplant('packing','PKL_HONG_MING_SEIYO_20260818 FN.xlsx','PKL_HONG_MING_SEIYO_20261030_DU_THAO.xlsx')
