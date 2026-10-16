# -*- coding: utf-8 -*-
"""Sinh và kiểm tra hợp đồng thử tải với 10 sản phẩm song ngữ."""

from __future__ import annotations

import argparse
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR))

from dien_hop_dong_tu_dong import fill_contract
from generate_jinxun_mes_contract import DEFAULT_TEMPLATE, JINXUN_CONTRACT_DATA


PRODUCTS = (
    ("1", "Hệ thống quản lý kho WMS Standard\nWMS Standard 倉庫管理系統", "185.000.000 VND", "1 SET", "185.000.000 VND"),
    ("2", "Phần mềm quản lý sản xuất MES Pro (bản quyền 30 máy)\nMES Pro 生產管理軟體（30台設備授權版）", "420.000.000 VND", "1 SET", "420.000.000 VND"),
    ("3", "Gói tích hợp API kết nối ERP\nERP 系統 API 整合服務", "95.000.000 VND", "1 SET", "95.000.000 VND"),
    ("4", "Thiết bị thu thập dữ liệu IoT công nghiệp\n工業物聯網數據採集設備", "28.500.000 VND", "5 SET", "142.500.000 VND"),
    ("5", "Máy quét mã vạch công nghiệp\n工業條碼掃描器", "12.800.000 VND", "10 SET", "128.000.000 VND"),
    ("6", "Máy tính bảng công nghiệp 10 inch\n10吋工業平板電腦", "18.600.000 VND", "6 SET", "111.600.000 VND"),
    ("7", "Bộ chuyển mạch mạng công nghiệp\n工業網路交換器", "9.500.000 VND", "8 SET", "76.000.000 VND"),
    ("8", "Dịch vụ lắp đặt và cấu hình hệ thống\n系統安裝與設定服務", "75.000.000 VND", "1 SET", "75.000.000 VND"),
    ("9", "Dịch vụ đào tạo nhân viên vận hành\n操作人員教育訓練服務", "32.000.000 VND", "2 SET", "64.000.000 VND"),
    ("10", "Gói bảo trì hệ thống trong 12 tháng\n12個月系統維護服務", "85.000.000 VND", "1 SET", "85.000.000 VND"),
)


def build_data() -> dict:
    data = deepcopy(JINXUN_CONTRACT_DATA)
    data["hang_muc"] = [
        {
            "stt": stt,
            "ten_san_pham": product_name,
            "don_gia": unit_price,
            "so_luong": quantity,
            "tong_gia": total_price,
        }
        for stt, product_name, unit_price, quantity, total_price in PRODUCTS
    ]
    data["tong_tien_truoc_giam"] = "1.382.100.000 VND"
    data.pop("tong_tien_sau_giam", None)
    data["tong_tien_hop_dong_so"] = "1.382.100.000"
    data["tong_so_luong"] = "36 SET"
    return data


def verify_output(output_path: Path) -> None:
    document = Document(output_path)
    table = document.tables[0]
    assert len(table.rows) == 12, f"Cần 12 hàng, hiện có {len(table.rows)}"

    for row_index, expected in enumerate(PRODUCTS, start=1):
        cells = [cell.text.strip() for cell in table.rows[row_index].cells]
        stt, product_name, unit_price, quantity, total_price = expected
        assert cells[0] == stt, (row_index, cells[0], stt)
        assert cells[1] == product_name, (row_index, cells[1], product_name)
        assert cells[2] == unit_price, (row_index, cells[2], unit_price)
        assert cells[3] == quantity, (row_index, cells[3], quantity)
        assert cells[4] == total_price, (row_index, cells[4], total_price)

        tr_pr = table.rows[row_index]._tr.trPr
        tr_height = tr_pr.find(qn("w:trHeight")) if tr_pr is not None else None
        if tr_height is not None:
            assert tr_height.get(qn("w:hRule")) != "exact", (
                f"Hàng {row_index} đang dùng chiều cao exact, có nguy cơ cắt chữ"
            )

    summary = [cell.text.strip() for cell in table.rows[-1].cells]
    assert "總共 Tổng Cộng" in summary[0]
    assert summary[2] == ""
    assert summary[3] == "36 SET"
    assert summary[4] == "1.382.100.000 VND"

    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    full_text += "\n" + "\n".join(
        cell.text for table_item in document.tables for row in table_item.rows for cell in row.cells
    )
    assert "{{" not in full_text and "}}" not in full_text
    assert "1.382.100.000 VND" in full_text
    assert len(re.findall(r"^Đợt\s+\d+", full_text, flags=re.MULTILINE)) == 3


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    result = fill_contract(str(args.template), str(args.output), build_data())
    actual_output = Path(result["output_path"])
    verify_output(actual_output)
    print(f"ĐẠT: {actual_output}")
    print("- 10 sản phẩm, 36 SET, tổng 1.382.100.000 VND.")
    print("- Mọi dòng cho phép tự giãn chiều cao; không có hRule=exact.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
