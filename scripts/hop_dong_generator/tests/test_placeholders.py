# -*- coding: utf-8 -*-
"""Kiểm thử placeholder bị Word chia qua nhiều run."""

from pathlib import Path
import sys

from docx import Document


PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR))

from dien_hop_dong_tu_dong import (
    get_paragraph_all_text,
    replace_next_placeholder_preserve_style,
    strip_placeholder_braces,
    update_party_a_tax_line,
)


def main() -> int:
    doc = Document()
    doc.add_paragraph("Bên A (買方): MẪU")
    paragraph = doc.add_paragraph()
    for fragment in (
        "代表人：{{HUANG ",
        "CU}}                    MST 稅號: {{3700",
        "660198}}",
    ):
        paragraph.add_run(fragment)
    doc.add_paragraph("Bên B (賣方): HONG MING")

    assert update_party_a_tax_line(doc, "0801382054", "HUANG YU")
    result = get_paragraph_all_text(paragraph)
    assert result == "代表人：HUANG YU                    MST 稅號: 0801382054", result
    assert "{{" not in result and "}}" not in result

    split_braces = doc.add_paragraph()
    for fragment in ("Giữ nội dung {", "{mẫu chia run}", "}"):
        split_braces.add_run(fragment)
    strip_placeholder_braces(doc)
    assert get_paragraph_all_text(split_braces) == "Giữ nội dung mẫu chia run"

    styled = doc.add_paragraph()
    styled.add_run("Bên A (Bên Mua買方): ")
    styled.add_run("{{")
    content_run = styled.add_run("TÊN CŨ")
    content_run.bold = True
    styled.add_run("}}")
    assert replace_next_placeholder_preserve_style(styled, "TÊN MỚI")
    assert styled.text == "Bên A (Bên Mua買方): TÊN MỚI"
    assert any(run.text == "TÊN MỚI" and run.bold for run in styled.runs)
    print("ĐẠT: placeholder MST/đại diện chia qua nhiều run được thay đúng.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
