# -*- coding: utf-8 -*-
"""Kiểm tra hồi quy nội dung hợp đồng JINXUN trong toàn bộ OOXML.

Khác với ``python-docx``, bộ kiểm tra này đọc mọi đoạn ``w:p`` trong các
thành phần XML của DOCX, kể cả chữ nằm trong textbox/shape.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

EXPECTED_PARTY_TEXT = (
    "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM",
    "Thôn Trụ, Phường Việt Hòa, Thành phố Hải Phòng, Việt Nam",
    "0801382054",
    "HUANG YU",
    "Tổng giám đốc",
)

EXPECTED_TAX_IDS = {"0801382054", "0318280158"}
PAYMENT_PERCENTAGES = (35, 55, 10)
EXPECTED_PRODUCTS = (
    ("1", "Nghiệp vụ PM", "110.000.000 VND"),
    ("2", "Phiên bản cấp phép cho 25 máy", "588.000.000 VND"),
    ("3", "API cấp liệu tùy chỉnh", "109.583.184 VND"),
)
FORBIDDEN_TEXT = (
    "2400976255",
    "CHENG FENG",
    "晟峰",
    "Báo giá triển khai Báo giá triển khai",
    "T/T 60 ngày",
    "T/T 60 天",
    "越南金迅科技有限公司",
    "越南海防市越和坊柱村",
    "宏明環球有限公司",
    "總經理",
    "{{",
    "}}",
)

EXPECTED_CHINESE_PAYMENT_TITLES = (
    "簽約定金",
    "設備安裝款",
    "驗收尾款",
)


def extract_ooxml_paragraphs(docx_path: Path) -> list[tuple[str, str]]:
    paragraphs: list[tuple[str, str]] = []
    with ZipFile(docx_path) as package:
        for part_name in package.namelist():
            if not part_name.endswith(".xml"):
                continue
            try:
                root = ET.fromstring(package.read(part_name))
            except ET.ParseError:
                continue
            for paragraph in root.iter(W_NS + "p"):
                text = "".join(
                    node.text or "" for node in paragraph.iter(W_NS + "t")
                ).strip()
                if text:
                    paragraphs.append((part_name, text))
    return paragraphs


def check_contract(docx_path: Path) -> int:
    paragraphs = extract_ooxml_paragraphs(docx_path)
    full_text = "\n".join(text for _, text in paragraphs)
    failures: list[str] = []

    document = Document(docx_path)
    body_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text]

    def normalized(text: str) -> str:
        return re.sub(r"\s+", " ", text).strip()

    party_a_heading = next(
        (paragraph for paragraph in body_paragraphs if paragraph.text.startswith("Bên A")),
        None,
    )
    expected_party_a_heading = normalized(
        "Bên A (Bên Mua買方): CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM"
    )
    if party_a_heading is None or normalized(party_a_heading.text) != expected_party_a_heading:
        failures.append("Dòng tên Bên A đã bị đổi chữ ngoài placeholder {{...}}")
    elif len(party_a_heading.runs) < 2 or not all(
        run.underline for run in party_a_heading.runs[:2]
    ):
        failures.append("Định dạng gạch chân 'Bên A' của mẫu không được bảo toàn")

    party_a_tax_line = next(
        (
            paragraph
            for paragraph in body_paragraphs
            if "MST" in paragraph.text and "0801382054" in paragraph.text
        ),
        None,
    )
    if party_a_tax_line is None or "HUANG YU" in party_a_tax_line.text:
        failures.append("Dòng 代表人/MST đã bị chèn thêm tên ngoài placeholder")

    party_a_address_line = next(
        (
            paragraph
            for paragraph in body_paragraphs
            if paragraph.text.startswith("Địa chỉ")
            and "Thôn Trụ" in paragraph.text
        ),
        None,
    )
    expected_address_line = normalized(
        "Địa chỉ 地址 : Thôn Trụ, Phường Việt Hòa, Thành phố Hải Phòng, Việt Nam"
    )
    if party_a_address_line is None or normalized(party_a_address_line.text) != expected_address_line:
        failures.append("Dòng địa chỉ Bên A đã bị thêm chữ ngoài placeholder {{...}}")

    party_a_role_line = next(
        (
            paragraph
            for paragraph in body_paragraphs
            if paragraph.text.startswith("Người đại diện") and "HUANG YU" in paragraph.text
        ),
        None,
    )
    expected_role_line = normalized(
        "Người đại diện: HUANG YU Chức vụ: Tổng giám đốc"
    )
    if party_a_role_line is None or normalized(party_a_role_line.text) != expected_role_line:
        failures.append(
            "Dòng đại diện/chức vụ không giữ nguyên nhãn của mẫu hoặc bị lặp 代表人/職務"
        )

    party_a_start = next(
        (i for i, paragraph in enumerate(body_paragraphs) if paragraph.text.startswith("Bên A")),
        -1,
    )
    party_b_start = next(
        (i for i, paragraph in enumerate(body_paragraphs) if paragraph.text.startswith("Bên B")),
        len(body_paragraphs),
    )
    if party_a_start >= 0:
        party_a_text = "\n".join(
            paragraph.text for paragraph in body_paragraphs[party_a_start:party_b_start]
        )
        if party_a_text.count("代表人") != 1:
            failures.append(
                f"Khối Bên A có {party_a_text.count('代表人')} lần 代表人; yêu cầu đúng 1 lần"
            )

    if not document.tables:
        failures.append("Không còn bảng sản phẩm trong hợp đồng")
    else:
        product_rows = document.tables[0].rows
        if len(product_rows) != 5:
            failures.append(
                f"Bảng sản phẩm có {len(product_rows)} hàng; cần 5 hàng "
                "(tiêu đề + 3 sản phẩm + tổng)"
            )
        for row_index, (stt, product_text, price) in enumerate(
            EXPECTED_PRODUCTS, start=1
        ):
            if row_index >= len(product_rows):
                failures.append(f"Thiếu hàng sản phẩm số {stt}")
                continue
            cells_text = [cell.text for cell in product_rows[row_index].cells]
            row_text = "\n".join(cells_text)
            if cells_text[0].strip() != stt:
                failures.append(f"STT hàng sản phẩm {row_index} không phải {stt}")
            if product_text not in row_text or price not in row_text:
                failures.append(
                    f"Hàng sản phẩm {stt} thiếu nội dung/đơn giá đã xác nhận"
                )
        if product_rows:
            summary_text = "\n".join(
                cell.text for cell in product_rows[-1].cells
            )
            for expected_summary in (
                "807.583.184 VND",
                "Sau Giảm Giá",
                "781.530.000 VND",
            ):
                if expected_summary not in summary_text:
                    failures.append(
                        f"Dòng tổng bảng sản phẩm thiếu: {expected_summary}"
                    )

    if len(document.tables) < 2:
        failures.append("Không còn bảng chữ ký trong hợp đồng")
    else:
        signature_table = document.tables[-1]
        signature_a = normalized(signature_table.rows[0].cells[0].text)
        signature_b = normalized(signature_table.rows[0].cells[1].text)
        if signature_a != "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM":
            failures.append("Tên công ty Bên A ở bảng chữ ký không đúng hoặc bị thêm chữ Hán")
        if signature_b != "CÔNG TY TNHH TOÀN CẦU HONG MING":
            failures.append("Tên công ty Bên B ở bảng chữ ký không còn nguyên mẫu tiếng Việt")

    for expected in EXPECTED_PARTY_TEXT:
        if expected not in full_text:
            failures.append(f"Thiếu dữ liệu do người dùng cung cấp: {expected}")

    for forbidden in FORBIDDEN_TEXT:
        if forbidden in full_text:
            failures.append(f"Còn nội dung cũ hoặc tự thêm ngoài yêu cầu: {forbidden}")

    tax_ids = set(re.findall(r"(?<!\d)\d{10}(?!\d)", full_text))
    unexpected_tax_ids = tax_ids - EXPECTED_TAX_IDS
    missing_tax_ids = EXPECTED_TAX_IDS - tax_ids
    if unexpected_tax_ids:
        failures.append(
            "Có MST 10 chữ số ngoài dữ liệu cho phép: "
            + ", ".join(sorted(unexpected_tax_ids))
        )
    if missing_tax_ids:
        failures.append("Thiếu MST: " + ", ".join(sorted(missing_tax_ids)))
    for tax_id in EXPECTED_TAX_IDS:
        occurrences = len(re.findall(rf"(?<!\d){re.escape(tax_id)}(?!\d)", full_text))
        if occurrences != 1:
            failures.append(
                f"MST {tax_id} xuất hiện {occurrences} lần; yêu cầu đúng 1 lần"
            )

    installment_paragraphs = [
        text for _, text in paragraphs if re.match(r"^Đợt\s+\d+", text)
    ]
    chinese_installment_paragraphs = [
        text for _, text in paragraphs if "支付合約總金額" in text
    ]
    if len(installment_paragraphs) != 3:
        failures.append(
            f"Số đợt thanh toán không đúng: thấy {len(installment_paragraphs)}, cần 3"
        )
    for index, percentage in enumerate(PAYMENT_PERCENTAGES, start=1):
        if not any(
            re.match(rf"^Đợt\s+{index}\b", text)
            and f"{percentage}%" in text
            for text in installment_paragraphs
        ):
            failures.append(f"Đợt {index} không có đúng tỷ lệ {percentage}%")
        expected_title = EXPECTED_CHINESE_PAYMENT_TITLES[index - 1]
        if not any(
            text.lstrip().startswith(f"{expected_title}：")
            and f"{percentage}%" in text
            for text in chinese_installment_paragraphs
        ):
            failures.append(
                f"Đợt tiếng Trung {index} phải giữ tiêu đề gốc {expected_title} "
                f"và đúng tỷ lệ {percentage}%"
            )

    if len(chinese_installment_paragraphs) != 3:
        failures.append(
            "Số đợt thanh toán tiếng Trung không đúng: "
            f"thấy {len(chinese_installment_paragraphs)}, cần 3"
        )
    if any(re.match(r"^第\s*\d+\s*期", text.strip()) for text in chinese_installment_paragraphs):
        failures.append("Generator đã tự thêm cách ghi 第…期 không có trong mẫu")

    payment_start = next(
        (i for i, paragraph in enumerate(document.paragraphs) if "簽約定金" in paragraph.text),
        -1,
    )
    payment_end = next(
        (
            i
            for i, paragraph in enumerate(document.paragraphs)
            if i > payment_start and "Bàn giao交付" in paragraph.text
        ),
        len(document.paragraphs),
    )
    if payment_start >= 0:
        for paragraph in document.paragraphs[payment_start:payment_end]:
            num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
            if num_pr is not None and not paragraph.text.strip():
                failures.append(
                    "Còn đoạn đánh số tự động rỗng trong phần thanh toán (hiển thị số 4 thừa)"
                )
                break

    relevant_terms = (
        "JINXUN",
        "金迅",
        "MST",
        "稅號",
        "税号",
        "Địa chỉ",
        "地址",
        "Đợt ",
    )
    print(f"Tệp: {docx_path}")
    print("Các đoạn liên quan trong toàn bộ OOXML:")
    for part_name, text in paragraphs:
        if any(term in text for term in relevant_terms):
            print(f"- [{part_name}] {text}")

    if failures:
        print("\nKẾT QUẢ: KHÔNG ĐẠT")
        for failure in failures:
            print(f"- {failure}")
        return 1

    print("\nKẾT QUẢ: ĐẠT")
    print("- Không có MST 10 chữ số ngoài dữ liệu đã cung cấp.")
    print("- Có đủ dữ liệu Việt - Trung do người dùng cung cấp.")
    print("- Tỷ lệ thanh toán đúng 35% - 55% - 10%.")
    print("- Không còn T/T 60 ngày hoặc câu dẫn báo giá bị lặp.")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Kiểm tra dữ liệu pháp nhân và thanh toán trong hợp đồng JINXUN."
    )
    parser.add_argument("docx", type=Path, help="Tệp DOCX cần kiểm tra")
    args = parser.parse_args()
    if not args.docx.is_file():
        parser.error(f"Không tìm thấy tệp: {args.docx}")
    return check_contract(args.docx)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
