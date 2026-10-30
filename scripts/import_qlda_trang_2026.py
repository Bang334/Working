"""Import reviewed QLDA TRANG AR data; backup, atomic writes, preserve acceptance."""
import argparse
import hashlib
import json
import re
import shutil
import sqlite3
import sys
import warnings
from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 29)


def text(v):
    return str(v).strip() if v is not None else ''


def number(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def iso(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    try:
        return date.fromisoformat(text(v)).isoformat()
    except ValueError:
        return None


def status(gross, received, due):
    if gross > 0 and received >= gross:
        return 'PAID'
    if due:
        days = (date.fromisoformat(due) - TODAY).days
        return 'DELAYED' if days < 0 else ('URGENT' if days <= 7 else 'PENDING')
    return 'PENDING'


def parse(source):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb['3. Theo Dõi Công Nợ Dự Án']
    report = {'source': str(source), 'sheet': ws.title, 'warnings': [str(w.message) for w in caught],
              'removed_duplicates': [], 'review': [], 'source_totals': {c:ws[f'{c}55'].value for c in ['I','J','K']}}
    report['source_total_limitation']='Dòng tổng nguồn chỉ SUM(I2:I45), SUM(J2:J45), SUM(K2:K45), bỏ sót dòng 46–49 của CHENGFENG; tổng giá trị HĐ E55 chỉ SUM(E2:E18) và cộng lặp HĐ. Không dùng dòng tổng nguồn làm tổng kiểm soát.'
    # Same contract + installment amount; rows 27/28 are the later detailed schedule.
    for old, new in [(8,27),(9,28)]:
        assert ws[f'D{old}'].value == ws[f'D{new}'].value == 'DFV20260505'
        assert ws[f'E{old}'].value == ws[f'E{new}'].value
        assert ws[f'F{old}'].value == ws[f'F{new}'].value
        report['removed_duplicates'].append({'removed_row':old,'kept_row':new,'reason':'Darfon: cùng HĐ, tổng giá trị và khoản thanh toán; giữ bộ lịch 30/60/10 ở dòng 27–29.'})
    items = []
    for r in range(2,55):
        v = [ws.cell(r,c).value for c in range(1,18)]
        if r in [8,9] or not any(v[1:4]):
            continue
        issues = []
        project, customer, ref = map(text, v[1:4])
        if not project:
            pct = round(number(v[5])/number(v[4])*100) if number(v[4]) else 0
            project = f'{pct}% Mở rộng MES Lite giai đoạn 2 / MES Lite 第二階段擴充' if ref == 'DFV20260505' else f'Khoản thanh toán dòng {r} / 第 {r} 列款項'
            issues.append('Tên hạng mục trống: dùng mô tả theo tỷ lệ và dòng nguồn; cần đối chiếu điều kiện đợt.')
        net, vat = number(v[5]), number(v[6])
        gross = number(v[8])
        if v[8] is None and net:
            gross = round(net*(1+vat),2)
            issues.append('Ô phải thu trống: tính từ khoản chưa thuế × (1 + VAT); không phải số đã thu.')
        tax = number(v[7]) if v[7] is not None else round(net*vat,2)
        if v[5] is None and gross:
            net = gross-tax
        received = number(v[9])
        if 0 < received-gross < 1:
            issues.append(f'Sửa sai số công thức đã thu {received} thành {gross}; nguồn cộng nhầm tỷ lệ VAT.')
            received = gross
        debt = round(gross-received,2)
        if abs(number(v[10])-debt) > .01:
            issues.append(f'Nợ nguồn {v[10]} khác phải thu − đã thu {debt}; dùng phép trừ đối soát.')
        for col, raw in zip(['L','M','N'],v[11:14]):
            if raw and not iso(raw):
                issues.append(f'{col}{r} không phải ngày hợp lệ: {raw}; giữ trong ghi chú.')
        raw_status = text(v[15])
        if ('已付款' in raw_status or 'Đã thanh toán' in raw_status) and received < gross:
            issues.append('Nguồn ghi đã thanh toán nhưng chưa có số tiền thực nhận tương ứng: chưa tự xác nhận thu tiền.')
        due, paid = iso(v[12]), iso(v[13])
        if received and not paid:
            issues.append('Có số đã thu nhưng thiếu ngày thu hợp lệ; không tự tạo ngày.')
        code_match = re.search(r'\b(?:HM[A-Z]*\d[A-Z0-9-]*|DFV\d+|SC-\d{4}-\d+)\b',ref)
        code = code_match.group(0) if code_match else None
        notes = f'Nguồn: {source.name} | {ws.title}!A{r}:Q{r}\nTham chiếu gốc: {ref}\nTrạng thái nguồn: {raw_status}\nGhi chú nguồn: {text(v[16])}'
        if issues:
            notes += '\nCẦN ĐỐI CHIẾU: ' + ' | '.join(issues)
            report['review'].append({'row':r,'issues':issues})
        items.append(dict(row=r,order_no=int(number(v[0])) or r,project_name=project,customer_name=customer,
            contract_code=code or ref or None,total_contract_amount=number(v[4]),net_amount=net,vat_rate=vat,
            vat_amount=tax,gross_amount=gross,received_amount=received,remaining_debt=debt,
            customer_promised_date=iso(v[11]),due_date=due,paid_date=paid,doc_status=text(v[14]) or None,
            collection_status=status(gross,received,due),notes=notes,code=code))
    groups = defaultdict(list)
    for it in items:
        if it['code']:
            groups[it['code']].append(it)
    for code, rows in groups.items():
        totals = sorted({i['total_contract_amount'] for i in rows})
        if len(totals)>1:
            msg = f'{code}: nhiều tổng giá trị {totals}; chưa xác nhận giá trị hợp đồng và không gộp các khoản khác hạng mục.'
            report['review'].append({'contract':code,'issues':[msg]})
            for i in rows:
                i['notes'] += '\nCẦN ĐỐI CHẾU: '+msg
    return items,groups,report


def rows(c, query, params=()):
    return [dict(r) for r in c.execute(query,params)]


def protected(c):
    return {'documents':rows(c,"SELECT * FROM documents WHERE doc_type NOT IN ('HOP_DONG','YEU_CAU_THANH_TOAN','DON_XIN_THANH_TOAN') ORDER BY id"),
            'acceptance_cases':rows(c,"SELECT * FROM workflow_cases WHERE kind='acceptance' ORDER BY id"),
            'acceptance_records':{t:rows(c,f"SELECT * FROM {t} WHERE case_id IN (SELECT id FROM workflow_cases WHERE kind='acceptance') ORDER BY rowid") for t in ['workflow_files','workflow_events','workflow_mail','workflow_receipts','project_case_links','project_facts']},
            'other_tables':{t:rows(c,f'SELECT * FROM {t} ORDER BY rowid') for t in ['procurement_orders','quotations','custom_reminders','beneficiaries','payment_voucher_requests','emails']},
            'files':{str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'Ho So Du An'/'Nghiem Thu').rglob('*') if p.is_file()}}


def apply(db, source, items, groups, report):
    c=sqlite3.connect(db);c.row_factory=sqlite3.Row
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    folder=ROOT/'database'/'backups'/('qlda_trang_'+stamp)
    folder.mkdir(parents=True)
    backup=sqlite3.connect(folder/'pm_portal_before.db');c.backup(backup);backup.close()
    shutil.copy2(source,folder/source.name)
    before=protected(c)
    try:
        c.execute('BEGIN IMMEDIATE')
        old_ids=[r[0] for r in c.execute("SELECT id FROM workflow_cases WHERE kind IN ('payment','contract')")]
        for cid in old_ids:
            for table in ['workflow_files','workflow_events','workflow_mail','workflow_receipts','project_case_links','project_facts']:
                c.execute(f'DELETE FROM {table} WHERE case_id=?',(cid,))
        c.execute("DELETE FROM workflow_cases WHERE kind IN ('payment','contract')")
        c.execute('DELETE FROM project_receivables')
        c.execute('DELETE FROM payment_milestones')
        # Keep exact ID of a retained contract so existing document associations remain valid.
        existing={r['contract_code']:r['id'] for r in c.execute('SELECT * FROM contracts')}
        obsolete=set(existing)-set(groups)
        for code in obsolete:
            c.execute("DELETE FROM documents WHERE contract_code=? AND doc_type IN ('HOP_DONG','YEU_CAU_THANH_TOAN','DON_XIN_THANH_TOAN')",(code,))
        c.execute('DELETE FROM contracts')
        linked={}
        for code, group in groups.items():
            totals={i['total_contract_amount'] for i in group}
            ambiguous=len(totals)!=1
            total=next(iter(totals)) if not ambiguous else 0
            paid=sum(i['received_amount'] for i in group) if not ambiguous else 0
            note=f'Nhập từ {source.name}; dòng '+','.join(str(i['row']) for i in group)+'. Chỉ tổng hợp các khoản có trong file, không suy diễn phần tiền các đợt không được liệt kê.'
            if ambiguous:
                note+=' CẦN ĐỐI CHIẾU tổng giá trị: '+str(sorted(totals))+'. Giá trị 0 là chưa xác định.'
            st='Cần đối chiếu dữ liệu nguồn' if ambiguous else ('Đã thu đủ theo file' if paid>=total and total else 'Theo dõi công nợ theo file mới')
            customer=group[0]['customer_name'];project=group[0]['project_name']
            cur=c.execute('INSERT INTO contracts (id,contract_code,customer_name_vi,project_name_vi,total_amount,paid_amount,currency,status,notes) VALUES (?,?,?,?,?,?,?,?,?)',
                (existing.get(code),code,customer,project,total,paid,'VND',st,note))
            contract_id=cur.lastrowid
            if ambiguous:
                continue
            for i in group:
                pct=round(i['net_amount']/total*100,4) if total else None
                st='PAID' if i['collection_status']=='PAID' else ('WAITING' if i['due_date'] else 'PENDING_TRIGGER')
                cur=c.execute('INSERT INTO payment_milestones (contract_id,contract_code,milestone_name,percentage,amount,trigger_condition,due_date,status,status_label,paid_date,triggered,action_required) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                    (contract_id,code,i['project_name'],pct,i['gross_amount'],i['notes'],i['due_date'],st,
                     'Đã thu theo file' if st=='PAID' else 'Theo dõi / đối chiếu nguồn',i['paid_date'],int(bool(i['due_date'])),'Đối chiếu ghi chú nguồn trước khi xử lý'))
                linked[i['row']]=cur.lastrowid
        fields=['order_no','project_name','customer_name','contract_code','total_contract_amount','net_amount','vat_rate','vat_amount','gross_amount','received_amount','remaining_debt','customer_promised_date','due_date','paid_date','doc_status','collection_status','notes','linked_milestone_id']
        for i in items:
            i['linked_milestone_id']=linked.get(i['row'])
            c.execute('INSERT INTO project_receivables ('+','.join(fields)+') VALUES ('+','.join('?' for _ in fields)+')',[i.get(k) for k in fields])
        assert protected(c)==before,'Protected data changed'
        assert c.execute('SELECT COUNT(*) FROM project_receivables').fetchone()[0]==len(items)
        assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert not c.execute('SELECT id FROM project_receivables WHERE ABS(gross_amount-received_amount-remaining_debt)>.01').fetchall()
        c.commit()
    except Exception:
        c.rollback();raise
    report.update(contracts=len(groups),receivables=len(items),milestones=len(linked),backup=str(folder/'pm_portal_before.db'),
                  preserved_acceptance=len(before['acceptance_cases']),preserved_acceptance_files=len(before['files']),
                  totals={k:round(sum(i[k] for i in items),2) for k in ['gross_amount','received_amount','remaining_debt']})
    (folder/'import_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    lines=['# Kết quả cập nhật hợp đồng và công nợ — 29/10/2026',
           f'Nguồn: {source.name}, sheet {report["sheet"]}.',
           f'Đã nhập {len(groups)} mã hợp đồng và {len(items)} dòng công nợ; {len(linked)} mốc liên kết. Các khoản không có mã hợp đồng vẫn được giữ trong công nợ.',
           f'Giữ nguyên {len(before["files"])} tệp nghiệm thu, dữ liệu nghiệm thu và các phân hệ khác. Tệp Word/Excel hợp đồng cũ được giữ trên ổ đĩa làm hồ sơ lịch sử; dữ liệu hợp đồng cũ ngoài nguồn mới được loại khỏi danh sách hoạt động.',
           'Đã loại 2 dòng Darfon trùng (8, 9), giữ dòng 27, 28 và 29. Không loại các dòng CHENGFENG cùng số tiền vì có thể là những đợt khác nhau.',
           report['source_total_limitation'],
           'Tổng sau đối soát (còn bao gồm các khoản cần xác nhận, chưa phải xác nhận công nợ cuối cùng):']
    lines += [f'- {k}: {v:,.0f} VND' for k,v in report['totals'].items()]
    lines += ['','## Các điểm cần đối chiếu',
              'Vonmax HMV060620250001 có tổng giá trị 545.454.545 và 1.100.000.000 VND. Bản ghi hợp đồng tạm ghi giá trị 0 với trạng thái cần đối chiếu (0 không phải giá trị đã xác nhận); giữ nguyên từng khoản công nợ và chưa tạo liên kết mốc để tránh cộng nhầm.',
              'SEIYO HMV22092026SY0001 lấy trạng thái chưa thu từ file mới, thay dữ liệu cũ ghi thu 90%.',
              'Số đã thu chỉ tổng hợp khoản có trong file; không tự bổ sung khoản cọc các đợt không được liệt kê.']
    for issue in report['review']:
        lines.append('- '+str(issue.get('row',issue.get('contract')))+': '+'; '.join(issue['issues']))
    lines += ['','Bản sao lưu trước cập nhật: '+report['backup']]
    (folder/'BAO_CAO_CAP_NHAT.md').write_text('\n\n'.join(lines),encoding='utf-8')
    report['report_path']=str(folder/'BAO_CAO_CAP_NHAT.md')
    c.close()
    return report


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--db',type=Path,default=ROOT/'database'/'pm_portal.db');p.add_argument('--apply',action='store_true');a=p.parse_args()
    items,groups,report=parse(a.source)
    if a.apply:
        report=apply(a.db,a.source,items,groups,report)
    else:
        report.update(receivables=len(items),contracts=len(groups),contract_codes=list(groups))
    print(json.dumps(report,ensure_ascii=False,indent=2,default=str))
