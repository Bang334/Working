import docx
from pathlib import Path

doc = docx.Document("E:/Bien_ban_nghiem_thu_HMV100120260001.docx")
with open("template_full_details.txt", "w", encoding="utf-8") as f:
    f.write("=== SECTIONS ===\n")
    for s_idx, sec in enumerate(doc.sections):
        f.write(f"Section {s_idx}: top={sec.top_margin.pt}, bot={sec.bottom_margin.pt}, left={sec.left_margin.pt}, right={sec.right_margin.pt}\n")
        f.write(f"  Header paragraphs: {len(sec.header.paragraphs)}\n")
        for hp in sec.header.paragraphs:
            f.write(f"    H: {hp.text}\n")
        f.write(f"  Footer paragraphs: {len(sec.footer.paragraphs)}\n")
        for fp in sec.footer.paragraphs:
            f.write(f"    F: {fp.text}\n")

    f.write("\n=== ALL PARAGRAPHS ===\n")
    for p_idx, p in enumerate(doc.paragraphs):
        f.write(f"\n--- Paragraph {p_idx} (style={p.style.name}) ---\n")
        f.write(f"Full Text: {p.text}\n")
        for r_idx, r in enumerate(p.runs):
            f.write(f"  Run {r_idx}: text='{r.text}' | font={r.font.name} | size={r.font.size.pt if r.font.size else None} | bold={r.bold} | italic={r.italic} | color={r.font.color.rgb if r.font.color else None}\n")

    f.write("\n=== ALL TABLES ===\n")
    for t_idx, t in enumerate(doc.tables):
        f.write(f"\n=================== TABLE {t_idx} (rows={len(t.rows)}, cols={len(t.columns)}) ===================\n")
        for r_idx, row in enumerate(t.rows):
            f.write(f"-- Row {r_idx} --\n")
            for c_idx, cell in enumerate(row.cells):
                p_texts = []
                for p in cell.paragraphs:
                    runs_str = " + ".join([f"'{r.text}'(fn={r.font.name},sz={r.font.size.pt if r.font.size else None},b={r.bold},c={r.font.color.rgb if r.font.color else None})" for r in p.runs])
                    p_texts.append(f"P: {p.text} [{runs_str}]")
                f.write(f"  Cell [{r_idx}, {c_idx}]: {' // '.join(p_texts)}\n")

print("DUMPED FULL DETAILS TO template_full_details.txt")
