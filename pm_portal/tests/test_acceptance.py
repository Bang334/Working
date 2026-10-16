"""Integration and unit tests for Acceptance Module."""

import asyncio
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

# Add app to path
BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

import openpyxl
import httpx
from app.main import app


class ASGITestClient:
    """Small sync facade for the installed httpx/Starlette version combination."""

    @staticmethod
    async def _request(method, url, **kwargs):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport, base_url="http://testserver"
        ) as async_client:
            return await async_client.request(method, url, **kwargs)

    def get(self, url, **kwargs):
        return asyncio.run(self._request("GET", url, **kwargs))

    def post(self, url, **kwargs):
        return asyncio.run(self._request("POST", url, **kwargs))


client = ASGITestClient()


class TestAcceptanceModule(unittest.TestCase):
    """Test suite for Acceptance Module (Biên Bản Nghiệm Thu & Kiểm Tra Thiết Bị)."""

    def setUp(self):
        self.patcher_index = patch("app.database.index_single_document")
        self.mock_index_doc = self.patcher_index.start()
        self.sample_payload = {
            "tieu_de_viet": "BIÊN BẢN KIỂM TRA & NGHIỆM THU THIẾT BỊ SAU SỬA CHỮA",
            "tieu_de_trung": "維修設備檢查及驗收記錄",
            "so_hop_dong": "HM-FINEPRO-20261009.001",
            "so_bien_ban": "001/2026/BBNT-SC",
            "ngay_nghiem_thu": "09/10/2026",
            "ben_a": {
                "ten_cong_ty": "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)",
                "ten_tieng_trung": "CX Technology (VN) 股份有限公司",
                "dia_chi": "KCX Tân Thuận, Phường Tân Thuận Đông, Quận 7, TP.HCM",
                "mst": "0301445778"
            },
            "hang_muc": [
                {
                    "stt": 1,
                    "serial_number": "RFR900NA55RSWLBA574",
                    "loi_ban_dau": "Cò bắn không nhả\n扳機扣住不鬆",
                    "bo_phan": "CNC",
                    "hang_muc_kiem_tra": "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常",
                    "ket_qua": "Đạt / 合格",
                    "tong_gia": "550.000 VND"
                },
                {
                    "stt": 2,
                    "serial_number": "EF550RA4LAWBHI677",
                    "loi_ban_dau": "Màn hình lỗi\n螢幕故障",
                    "bo_phan": "CNC",
                    "hang_muc_kiem_tra": "Kiểm tra hiển thị màn hình & đầu quét barcode\n檢查螢幕顯示及條碼掃描功能",
                    "ket_qua": "Đạt / 合格",
                    "tong_gia": "8.000.000 VND"
                },
                {
                    "stt": 3,
                    "serial_number": "RFR900NA55RSWLBA565",
                    "loi_ban_dau": "Cò bắn không nhả\n扳機扣住不鬆",
                    "bo_phan": "CNC",
                    "hang_muc_kiem_tra": "Kiểm tra cò bắn & gá chuyển đổi Mounting Adapter\n檢查扳機及轉接支架功能",
                    "ket_qua": "Đạt / 合格",
                    "tong_gia": "3.500.000 VND"
                },
                {
                    "stt": 4,
                    "serial_number": "EF550RA4LAWBHI641",
                    "loi_ban_dau": "Cò bắn & Sạc không vô điện\n扳機故障及無法充電",
                    "bo_phan": "Xi treo\n吊鍍",
                    "hang_muc_kiem_tra": "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常",
                    "ket_qua": "Đạt / 合格",
                    "tong_gia": "550.000 VND"
                },
                {
                    "stt": 5,
                    "serial_number": "EF550RA4LAWBHI657",
                    "loi_ban_dau": "Cò bắn & Sạc không vô điện\n扳機故障及無法充電",
                    "bo_phan": "Xi treo\n吊鍍",
                    "hang_muc_kiem_tra": "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常",
                    "ket_qua": "Đạt / 合格",
                    "tong_gia": "550.000 VND"
                },
                {
                    "stt": 6,
                    "serial_number": "EF550RA4LAWBHI650",
                    "loi_ban_dau": "Cò bắn & Sạc không vô điện\n扳機故障及無法充電",
                    "bo_phan": "Xi treo\n吊鍍",
                    "hang_muc_kiem_tra": "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常",
                    "ket_qua": "Đạt / 合格",
                    "tong_gia": "550.000 VND"
                }
            ],
            "tong_so_luong": "6 Thiết bị",
            "tong_tien_so": "13.700.000",
            "ghi_chu_1": "Thiết bị đã được kiểm tra chức năng, bảo trì/sửa chữa và chạy thử đạt tiêu chuẩn kỹ thuật trước khi bàn giao / 設備已完成功能檢查、維修更換零件並測試正常運作後交付。",
            "ghi_chu_3": "Chi tiết kết quả kiểm tra & linh kiện thay thế / 檢驗結果與更換零件明細:\n   - STT 1, 4, 5, 6: Không phát hiện lỗi phát sinh thêm (No error founds) / 檢查未發現異常。\n   - STT 2 (EF550RA4LAWBHI677): Xử lý cụm Barcode / 檢修條碼掃描模組。\n   - STT 3 (RFR900NA55RSWLBA565): Thay thế Mounting Adapter / 更換轉接支架。",
            "tao_cac_sheet_anh": False
        }

    def tearDown(self):
        self.patcher_index.stop()

    @classmethod
    def tearDownClass(cls):
        try:
            from app.database import get_db_connection
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("DELETE FROM documents WHERE file_path LIKE '%Temp%' OR file_path LIKE '%tmp%'")
            conn.commit()
            conn.close()
        except Exception:
            pass

    def test_01_acceptance_page_render(self):
        """Verify the acceptance UI page renders with HTTP 200."""
        response = client.get("/acceptance")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Biên Bản Nghiệm Thu", response.text)
        self.assertIn("CX Technology", response.text)
        self.assertIn('id="repairFileInput"', response.text)
        self.assertIn('id="repairProcessButton"', response.text)
        self.assertIn("Tạo Nghiệm Thu Từ File Biên Bản Sửa Chữa", response.text)

    def test_02_acceptance_preview(self):
        """Verify live preview rendering."""
        response = client.post("/api/v1/acceptance/preview", json=self.sample_payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)", data["preview_html"])
        self.assertIn("RFR900NA55RSWLBA574", data["preview_html"])

    def test_03_acceptance_excel_generation(self):
        """Verify 1:1 Excel generation matching 10 golden control points."""
        response = client.post("/api/v1/acceptance/generate-excel", json=self.sample_payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertTrue(data["file_name"].endswith(".xlsx"))

        # Verify generated file structure
        file_path = BASE_DIR / "outputs" / data["file_name"]
        self.assertTrue(file_path.exists())

        wb = openpyxl.load_workbook(file_path, data_only=True)
        self.assertIn("Nghiệm Thu", wb.sheetnames)
        ws = wb["Nghiệm Thu"]
        self.assertEqual(ws.freeze_panes, "A6")
        self.assertEqual(ws["A1"].value, "BIÊN BẢN KIỂM TRA & NGHIỆM THU THIẾT BỊ SAU SỬA CHỮA")
        self.assertEqual(ws["B6"].value, "RFR900NA55RSWLBA574")
        self.assertEqual(ws["B7"].value, "EF550RA4LAWBHI677")

        # The template keeps a fixed empty note area before signatures.
        note_header_row = 6 + len(self.sample_payload["hang_muc"])
        note_body_row = note_header_row + 1
        signature_row = note_header_row + 3
        self.assertEqual(ws[f"A{note_header_row}"].value, "GHI CHÚ / 備註:")
        self.assertIsNone(ws[f"A{note_body_row}"].value)
        self.assertEqual(ws.row_dimensions[note_body_row].height, 81.0)
        self.assertEqual(ws.row_dimensions[signature_row].height, 107.25)
        self.assertEqual(ws.sheet_format.defaultRowHeight, 14.4)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(ws.print_area, "'Nghiệm Thu'!$A$1:$G$30")
        for row in range(note_body_row + 1, 31):
            for col in range(1, 8):
                self.assertIn(ws.cell(row, col).fill.fill_type, ("solid", None))

        # Every photo-sheet note box stays empty; item prices and repair text
        # must never leak into this column.
        rf_sheet = wb["RFR900NA55RSWLBA574"]
        ef_sheet = wb["EF550RA4LAWBHI677"]
        self.assertIsNone(rf_sheet["D6"].value)
        self.assertIsNone(ef_sheet["F6"].value)
        self.assertEqual(rf_sheet.page_margins.left, 0.4)
        self.assertEqual(ef_sheet.page_margins.header, 0.2)
        self.assertTrue(rf_sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertTrue(ef_sheet.sheet_properties.pageSetUpPr.fitToPage)

    def test_04_acceptance_word_generation(self):
        """Verify administrative Word (.docx) document generation."""
        response = client.post("/api/v1/acceptance/generate-word", json=self.sample_payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertTrue(data["file_name"].endswith(".docx"))

        file_path = BASE_DIR / "outputs" / data["file_name"]
        self.assertTrue(file_path.exists())

    def test_05_ai_parse_acceptance(self):
        """Verify AI parsing of delivery sheet text."""
        raw_text = """
        Hardware Delivery Receipt
        合約編號 Số Hợp Đồng ：HM-FINEPRO-20261009.001
        Bên Mua: CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)
        1. RFR900NA55RSWLBA574 - CNC - Cò bắn không nhả - No error founds - 550.000 VND
        2. EF550RA4LAWBHI677 - CNC - Màn hình lỗi - Barcode - 8.000.000 VND
        TỔNG GIÁ TRỊ: 8.550.000 VND
        """
        response = client.post("/api/v1/ai/parse-acceptance", json={"raw_text": raw_text})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertIsNotNone(data["data"])
        self.assertTrue(len(data["data"]["hang_muc"]) >= 2)

    @patch("app.controllers.acceptance_controller.process_uploaded_repair_file")
    def test_06_upload_repair_workbook(self, process_mock):
        """Verify multipart upload returns a downloadable acceptance workbook."""
        process_mock.return_value = {
            "file_name": "Bien_Ban_Nghiem_Thu_Test.xlsx",
            "download_url": "/api/v1/acceptance/download/Bien_Ban_Nghiem_Thu_Test.xlsx",
            "source_file_name": "Bien_ban_sua_chua.xlsx",
            "gemini_model": "gemini-test",
            "assigned_count": 2,
            "blank_count": 10,
            "ignored_note_count": 1,
            "assigned_products": ["SERIAL-01", "SERIAL-02"],
            "blank_products": [],
        }

        response = client.post(
            "/api/v1/acceptance/process-repair-file",
            files={
                "file": (
                    "Bien_ban_sua_chua.xlsx",
                    b"PK\x03\x04mock-workbook",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["assigned_count"], 2)
        self.assertEqual(data["ignored_note_count"], 1)
        self.assertTrue(data["download_url"].endswith("Bien_Ban_Nghiem_Thu_Test.xlsx"))
        process_mock.assert_called_once()

    def test_07_repair_upload_service_pipeline(self):
        """Verify upload validation, template copy and response metadata."""
        from app.services import acceptance_repair_upload_service as service

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            outputs = root / "outputs"
            template = root / "acceptance_template.xlsx"
            template.write_bytes(b"PK\x03\x04template")

            def fake_generate(source_path, target_path, workspace):
                self.assertTrue(source_path.is_file())
                self.assertTrue(target_path.is_file())
                self.assertEqual(workspace, root)
                return {
                    "saved_path": str(target_path),
                    "gemini_model": "gemini-test",
                    "note_text": "1. SERIAL-01 - Ghi chú\n2. Ghi chú chung",
                    "assignments": {"SERIAL-01": "Ghi chú"},
                    "blank_detail_sheets": ["SERIAL-02"],
                }

            fake_module = SimpleNamespace(generate=fake_generate)
            with patch.object(service, "OUTPUTS_DIR", outputs), patch.object(
                service, "REPAIR_ACCEPTANCE_TEMPLATE", template
            ), patch.object(service, "WORKING_DIR", root), patch.object(
                service, "_load_generator_module", return_value=fake_module
            ):
                result = service.process_uploaded_repair_file(
                    "bien_ban_sua_chua.xlsx", b"PK\x03\x04uploaded"
                )

            self.assertEqual(result["assigned_count"], 1)
            self.assertEqual(result["ignored_note_count"], 1)
            self.assertEqual(result["blank_count"], 1)
            target_path = root / "Ho So Du An" / "Nghiem Thu" / result["file_name"]
            self.assertTrue(target_path.is_file())
            self.assertFalse(result["is_linked_contract"])

    def test_07b_repair_upload_with_linked_contract(self):
        """Verify acceptance workbook can optionally link to an existing contract."""
        from app.services import acceptance_repair_upload_service as service

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "acceptance_template.xlsx"
            template.write_bytes(b"PK\x03\x04template")

            def fake_generate(source_path, target_path, workspace, **kwargs):
                return {
                    "saved_path": str(target_path),
                    "gemini_model": "gemini-test",
                    "note_text": "1. SERIAL-01 - Ghi chú",
                    "assignments": {"SERIAL-01": "Ghi chú"},
                    "blank_detail_sheets": [],
                }

            fake_module = SimpleNamespace(generate=fake_generate)
            with patch.object(service, "REPAIR_ACCEPTANCE_TEMPLATE", template), patch.object(
                service, "WORKING_DIR", root
            ), patch.object(
                service, "_load_generator_module", return_value=fake_module
            ):
                # 1. Test with linked contract HD-20261013-001
                result_linked = service.process_uploaded_repair_file(
                    "bien_ban_sua_chua.xlsx", b"PK\x03\x04uploaded", contract_code="HD-20261013-001"
                )
                self.assertTrue(result_linked["is_linked_contract"])
                self.assertEqual(result_linked["contract_code"], "HD-20261013-001")
                self.assertIn("HD-20261013-001", result_linked["file_name"])

                # 2. Test without contract ("không thì thôi")
                result_unlinked = service.process_uploaded_repair_file(
                    "bien_ban_sua_chua.xlsx", b"PK\x03\x04uploaded", contract_code=""
                )
                self.assertFalse(result_unlinked["is_linked_contract"])
                self.assertEqual(result_unlinked["contract_code"], "")

                # 3. Test with non-existent contract ("không thì thôi")
                result_nonexistent = service.process_uploaded_repair_file(
                    "bien_ban_sua_chua.xlsx", b"PK\x03\x04uploaded", contract_code="HD-NON-EXISTENT-999"
                )
                self.assertFalse(result_nonexistent["is_linked_contract"])
                self.assertEqual(result_nonexistent["contract_code"], "")

    def test_08_repair_upload_rejects_non_xlsx(self):
        """Reject non-Excel uploads before invoking Gemini or writing an output."""
        response = client.post(
            "/api/v1/acceptance/process-repair-file",
            files={"file": ("ghi_chu.txt", b"not-an-excel-file", "text/plain")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "error")
        self.assertIn(".xlsx", data["message"])

    def test_09_repair_upload_preserves_all_14_source_records(self):
        """Regression: an uploaded 14-row repair report must yield 14 acceptance items."""
        from app.services import acceptance_repair_upload_service as service

        source = BASE_DIR.parent / "Bieu Mau" / "Biên bản sữa chữa_TEST_14_BAN_GHI.xlsx"
        if not source.is_file():
            self.skipTest("Bieu Mau test fixture not found")

        source_wb = openpyxl.load_workbook(source, data_only=True)
        source_ws = source_wb["CX Technology (VN)"]
        expected = []
        for row in range(12, source_ws.max_row + 1):
            stt = source_ws.cell(row, 1).value
            if not isinstance(stt, (int, float)):
                continue
            expected.append(
                {
                    "serial": str(source_ws.cell(row, 2).value).strip(),
                    "loi_ban_dau": source_ws.cell(row, 3).value,
                    "bo_phan": source_ws.cell(row, 4).value,
                    "hang_muc_kiem_tra": source_ws.cell(row, 5).value,
                }
            )
        source_wb.close()
        self.assertEqual(len(expected), 14)

        with tempfile.TemporaryDirectory() as temp_dir:
            outputs = Path(temp_dir) / "outputs"
            generator = service._load_generator_module()

            def fake_call_gemini(note_text, valid_sheets, workspace):
                assignments = []
                for sheet_name in valid_sheets:
                    source_line = generator._line_for_serial(note_text, sheet_name)
                    if source_line:
                        assignments.append(
                            {
                                "sheet_name": sheet_name,
                                "serial": sheet_name,
                                "note": "test",
                                "source_line": source_line,
                            }
                        )
                return {"assignments": assignments, "ignored_notes": []}, "gemini-test"

            generator._call_gemini = fake_call_gemini
            with patch.object(service, "OUTPUTS_DIR", outputs), patch.object(
                service, "_load_generator_module", return_value=generator
            ):
                result = service.process_uploaded_repair_file(source.name, source.read_bytes())

            generated = outputs / result["file_name"]
            wb = openpyxl.load_workbook(generated, data_only=False)
            ws = wb["Nghiệm Thu"]
            self.assertEqual(result["record_count"], 14)
            self.assertEqual(ws.freeze_panes, "A15")
            self.assertEqual(ws.sheet_view.view, "pageBreakPreview")
            self.assertFalse(bool(ws.sheet_view.showGridLines))
            self.assertEqual(ws.print_area, "'Nghiệm Thu'!$A$1:$G$31")
            self.assertIn("HM-FINEPRO-20261009.001", ws["A3"].value)
            self.assertIn("09/10/2026", ws["E3"].value)
            self.assertIn("A20:G20", {str(item) for item in ws.merged_cells.ranges})
            self.assertIn("A21:G21", {str(item) for item in ws.merged_cells.ranges})
            self.assertIn("A23:C23", {str(item) for item in ws.merged_cells.ranges})
            actual = []
            for row in range(6, ws.max_row + 1):
                serial = ws.cell(row, 2).value
                if not serial:
                    break
                actual.append(
                    {
                        "serial": str(serial).strip(),
                        "loi_ban_dau": ws.cell(row, 3).value,
                        "bo_phan": ws.cell(row, 4).value,
                        "hang_muc_kiem_tra": ws.cell(row, 5).value,
                    }
                )

            detail_sheets = [name for name in wb.sheetnames if name != "Nghiệm Thu"]
            populated_notes = []
            blank_notes = []
            for name in detail_sheets:
                detail = wb[name]
                note_cell = detail.cell(6, detail.max_column)
                self.assertEqual(note_cell.font.name, "Times New Roman")
                note = note_cell.value
                (populated_notes if note else blank_notes).append(name)
            wb.close()

        self.assertEqual(actual, expected)
        self.assertEqual(detail_sheets, [item["serial"] for item in expected])
        self.assertEqual(len(populated_notes), 7)
        self.assertEqual(len(blank_notes), 7)
        self.assertIn("RFR900NA55RSVTEST901", populated_notes)
        self.assertIn("EF550RA4LAWTEST902", blank_notes)


if __name__ == "__main__":
    unittest.main()
