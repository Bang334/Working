# -*- coding: utf-8 -*-
"""Test suite for Multi-file Document AI Extraction and Payment Request Generation."""

import io
import pytest
import httpx
from app.main import app
import openpyxl
import docx


@pytest.fixture
def anyio_backend():
    return "asyncio"


def create_dummy_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("BIÊN BẢN NGHIỆM THU VÀ BÀN GIAO", level=1)
    doc.add_paragraph("Thuộc Hợp đồng mua bán số: HMV151020250001 ký ngày 09/12/2025")
    doc.add_paragraph("Bên Mua: CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)")
    doc.add_paragraph("Tổng giá trị hợp đồng: 234.000.000 VNĐ")
    doc.add_paragraph("Khi nghiệm thu đạt tiêu chuẩn: thanh toán 20% còn lại = 46.800.000 VNĐ")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def create_dummy_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Đơn_Xin_TT"
    ws["A1"] = "Đơn Xin Thanh Toán Nội Bộ"
    ws["B2"] = "Mã báo đơn: HVM007-20261021-2"
    ws["B3"] = "Số tiền: 46.800.000 VND"
    ws["B4"] = "Lý do: Thanh toán đợt Giai đoạn 3 (20%) hợp đồng HMV151020250001"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


@pytest.mark.anyio
async def test_generate_payment_request_from_extracted_endpoint():
    """Test generating official Excel payment request from extracted payload."""
    payload = {
        "customer_name_vi": "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)",
        "customer_name_zh": "CX TECHNOLOGY (VN) 股份有限公司",
        "contract_code": "HMV151020250001",
        "project_name_vi": "RFID Hardware & Custom Software ERP",
        "project_name_zh": "RFID硬體與系統客製化軟體",
        "milestone_name": "驗收合格尾款：20%",
        "milestone_percentage": 20,
        "payment_amount": 46800000,
        "total_contract_amount": 234000000,
        "request_date": "21/10/2026",
        "due_date": "31/10/2026",
        "payment_terms_summary": "1. 40% = 93.600.000 VND; 2. 40% = 93.600.000 VND; 3. 20% = 46.800.000 VND."
    }

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.post("/api/v1/reminders/generate-payment-request-extracted", json=payload)
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "success"
        assert "Phiếu_Yêu_Cầu_Thanh_Toán" in data["file_name"]
        assert data["amount"] == 46800000
        assert data["total_amount"] == 234000000
