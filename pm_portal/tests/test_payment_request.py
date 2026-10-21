"""Kiểm tra bộ xuất Phiếu yêu cầu thanh toán Excel 712TR và Word."""

import sqlite3
import sys
import tempfile
import unittest
import hashlib
import asyncio
from zipfile import ZipFile
from pathlib import Path
from unittest.mock import patch

import openpyxl
from docx import Document
from docx.enum.text import WD_COLOR_INDEX


BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from app.services import payment_request_excel_service as service
from app.services.payment_request_word_service import generate_payment_request_word
from app.routes import api_routes


SCHEMA = """
CREATE TABLE contracts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    contract_code TEXT UNIQUE NOT NULL,
    customer_name_vi TEXT NOT NULL,
    customer_name_zh TEXT,
    project_name_vi TEXT,
    project_name_zh TEXT,
    total_amount REAL DEFAULT 0
);
CREATE TABLE payment_milestones (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    contract_id INTEGER,
    contract_code TEXT NOT NULL,
    milestone_name TEXT NOT NULL,
    percentage REAL,
    amount REAL NOT NULL,
    trigger_condition TEXT,
    due_date TEXT,
    invoice_no TEXT,
    status TEXT
);
CREATE TABLE documents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    contract_id INTEGER,
    contract_code TEXT,
    doc_type TEXT NOT NULL,
    doc_name TEXT NOT NULL,
    file_name TEXT NOT NULL,
    file_path TEXT NOT NULL,
    file_size INTEGER DEFAULT 0,
    file_format TEXT,
    description TEXT
);
"""


class TestPaymentRequestExports(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db_path = self.root / "test.db"
        conn = sqlite3.connect(str(self.db_path))
        conn.executescript(SCHEMA)
        conn.execute(
            """
            INSERT INTO contracts (
                contract_code, customer_name_vi, customer_name_zh,
                project_name_vi, project_name_zh, total_amount
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                "HMV22092026SY0001",
                "CÔNG TY TNHH SEIYO VIỆT NAM",
                "西洋責任有限公司",
                "Quy hoạch mạng kết nối máy công nghiệp IIoT tại nhà máy SEIYO",
                "SEIYO 工廠 IIoT 工業機聯網網路規劃專案",
                712020000,
            ),
        )
        contract_id = conn.execute("SELECT id FROM contracts").fetchone()[0]
        milestones = [
            ("Đợt 1", 30, 213606000, "Ký hợp đồng", "25/10/2026", "HD-00190", "URGENT"),
            ("Đợt 2", 60, 427212000, "Lắp đặt xong sản phẩm (T/T 15d)", "25/12/2026", "Chưa xuất", "UPCOMING"),
            ("Đợt 3", 10, 71202000, "Nghiệm thu đạt tiêu chuẩn (T/T 30d)", "25/01/2027", "Chưa xuất", "UPCOMING"),
        ]
        conn.executemany(
            """
            INSERT INTO payment_milestones (
                contract_id, contract_code, milestone_name, percentage, amount,
                trigger_condition, due_date, invoice_no, status
            ) VALUES (?, 'HMV22092026SY0001', ?, ?, ?, ?, ?, ?, ?)
            """,
            [(contract_id, *row) for row in milestones],
        )
        conn.commit()
        conn.close()

    def tearDown(self):
        self.temp_dir.cleanup()

    def _connect(self):
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        return conn

    def test_generates_712tr_excel_and_matching_word(self):
        output_dir = self.root / "outputs"
        with patch.object(service, "get_db_connection", side_effect=self._connect), patch.object(
            service, "OUTPUT_DIR", output_dir
        ), patch.object(service, "WORKING_DIR", self.root):
            result = service.generate_payment_request_excel(
                contract_code="HMV22092026SY0001",
                milestone_name="Đợt 1",
                application_date="15/10/2026",
            )

        excel_path = Path(result["file_path"])
        word_path = Path(result["word_file_path"])
        self.assertTrue(excel_path.exists())
        self.assertTrue(word_path.exists())
        self.assertIn(result["word_download_url"], result["sheet_html"])

        source_wb = openpyxl.load_workbook(service.TEMPLATE_PATH, data_only=False)
        source_ws = source_wb[service.TEMPLATE_SHEET_NAME]
        out_wb = openpyxl.load_workbook(excel_path, data_only=False)
        self.assertEqual(out_wb.sheetnames, [service.TEMPLATE_SHEET_NAME])
        out_ws = out_wb[service.TEMPLATE_SHEET_NAME]

        self.assertEqual(out_ws.print_area, source_ws.print_area)
        self.assertEqual(out_ws.freeze_panes, source_ws.freeze_panes)
        self.assertEqual(out_ws.sheet_view.view, source_ws.sheet_view.view)
        self.assertEqual(out_ws.sheet_view.zoomScale, source_ws.sheet_view.zoomScale)
        self.assertEqual(out_ws.page_setup.scale, source_ws.page_setup.scale)
        self.assertEqual(
            sorted(str(item) for item in out_ws.merged_cells.ranges),
            sorted(str(item) for item in source_ws.merged_cells.ranges),
        )
        for column in "ABCDEFGHI":
            self.assertEqual(out_ws.column_dimensions[column].width, source_ws.column_dimensions[column].width)
        for row in range(1, 30):
            self.assertEqual(out_ws.row_dimensions[row].height, source_ws.row_dimensions[row].height)
            self.assertEqual(out_ws.row_dimensions[row].hidden, source_ws.row_dimensions[row].hidden)

        self.assertEqual(out_ws["C10"].value, "CÔNG TY TNHH SEIYO VIỆT NAM")
        self.assertEqual(out_ws["D15"].value, 712020000)
        self.assertEqual(out_ws["I15"].value, "=(F15+H15)*30%")
        self.assertIn("Thanh toán 30% giá trị hợp đồng", out_ws["B23"].value)

        doc = Document(word_path)
        full_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn("Phiếu yêu cầu thanh toán", full_text)
        self.assertIn("CÔNG TY TNHH SEIYO VIỆT NAM", table_text)
        self.assertIn("213.606.000", table_text)
        self.assertGreaterEqual(len(doc.tables), 5)

    def test_word_template_source_preserves_runs_and_becomes_primary(self):
        output_dir = self.root / "outputs-word-template"
        template_path = self.root / "đơn yêu cầu NCC thanh toán-Gốc.docx"
        template = Document()
        paragraph = template.add_paragraph("Khách hàng: ")
        first_marker_run = paragraph.add_run("{{CUSTOMER_")
        first_marker_run.bold = True
        paragraph.add_run("NAME_VI}}")
        table = template.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Số hợp đồng"
        table.cell(0, 1).text = "[SO_HOP_DONG]"
        table.cell(1, 0).text = "Số tiền"
        amount_run = table.cell(1, 1).paragraphs[0].add_run("SO_TIEN_THANH_TOAN_DINH_DANG")
        amount_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        template.save(template_path)
        original_hash = hashlib.sha256(template_path.read_bytes()).hexdigest()

        with patch.object(service, "get_db_connection", side_effect=self._connect), patch.object(
            service, "OUTPUT_DIR", output_dir
        ), patch.object(service, "WORKING_DIR", self.root), patch.object(
            service, "WORD_TEMPLATE_PATH", template_path
        ):
            result = service.generate_payment_request_excel(
                contract_code="HMV22092026SY0001",
                milestone_name="Đợt 1",
                application_date="15/10/2026",
                template_source="word",
            )

        self.assertEqual(result["template_source"], "word")
        self.assertEqual(result["template_file_name"], template_path.name)
        self.assertEqual(result["primary_file_name"], result["word_file_name"])
        self.assertEqual(original_hash, hashlib.sha256(template_path.read_bytes()).hexdigest())

        generated = Document(result["word_file_path"])
        self.assertIn("CÔNG TY TNHH SEIYO VIỆT NAM", generated.paragraphs[0].text)
        self.assertTrue(generated.paragraphs[0].runs[1].bold)
        self.assertEqual(generated.tables[0].cell(0, 1).text, "HMV22092026SY0001")
        self.assertEqual(generated.tables[0].cell(1, 1).text, "213.606.000")

    def test_word_template_source_requires_configured_file(self):
        missing_template = self.root / "missing.docx"
        with patch.object(service, "WORD_TEMPLATE_PATH", missing_template):
            with self.assertRaisesRegex(FileNotFoundError, "mẫu Word"):
                service.generate_payment_request_excel(
                    contract_code="HMV22092026SY0001",
                    template_source="word",
                )

    def test_word_template_source_accepts_standard_form_without_placeholders(self):
        output_dir = self.root / "outputs-literal-template"
        template_path = self.root / "literal-payment-form.docx"
        generate_payment_request_word(
            {
                "form_no": "FORM-MAU",
                "application_date": "01/01/2026",
                "contract_code": "HD-MAU",
                "customer_name_vi": "KHÁCH HÀNG MẪU",
                "project_name_vi": "Dự án mẫu",
                "project_name_zh": "範例專案",
                "milestone_name": "Đợt mẫu",
                "percentage": 10,
                "percentage_display": "10",
                "total_contract_amount_formatted": "1.000.000",
                "milestone_amount_formatted": "100.000",
                "amount_words_vi": "Một trăm nghìn đồng chẵn.",
                "amount_words_zh": "越南盾拾萬整",
                "due_date": "10/01/2026",
                "invoice_no": "Chưa xuất",
                "notes_text": "Ghi chú mẫu",
            },
            template_path,
        )

        with patch.object(service, "get_db_connection", side_effect=self._connect), patch.object(
            service, "OUTPUT_DIR", output_dir
        ), patch.object(service, "WORKING_DIR", self.root), patch.object(
            service, "WORD_TEMPLATE_PATH", template_path
        ):
            result = service.generate_payment_request_excel(
                contract_code="HMV22092026SY0001",
                milestone_name="Đợt 1",
                application_date="15/10/2026",
                template_source="word",
            )

        generated = Document(result["word_file_path"])
        self.assertIn("-SYTT", generated.paragraphs[5].text)
        self.assertEqual(generated.tables[0].cell(0, 1).text, "CÔNG TY TNHH SEIYO VIỆT NAM")
        self.assertIn("HMV22092026SY0001", generated.tables[0].cell(1, 1).text)
        self.assertEqual(generated.tables[1].cell(1, 8).text, "213.606.000")

    def test_current_custom_word_template_keeps_layout_and_updates_only_data(self):
        template_path = service.WORD_TEMPLATE_PATH
        self.assertTrue(template_path.exists())
        original_hash = hashlib.sha256(template_path.read_bytes()).hexdigest()
        output_path = self.root / "current-template-output.docx"
        data = {
            "form_no": "HMV15102026TEST-SYTT",
            "application_date": "15/10/2026",
            "contract_code": "HD-20261013-001",
            "customer_name_vi": "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM",
            "project_name_vi": "Dự án kiểm thử mẫu Word hiện tại",
            "project_name_zh": "目前 Word 範本測試專案",
            "milestone_name": "Đợt 1",
            "percentage": 35,
            "percentage_display": "35",
            "total_contract_amount_formatted": "807.583.183",
            "milestone_amount_formatted": "282.654.114",
            "due_date": "16/10/2026",
            "invoice_no": "Chưa xuất",
            "notes_text": "Ghi chú kiểm thử",
            "payment_terms": [
                "01  第一期付款條件達成時：支付合約金額的 35%\nĐợt 1: Thanh toán 35% giá trị hợp đồng",
                "02  第二期付款條件達成時：支付合約金額的 55%\nĐợt 2: Thanh toán 55% giá trị hợp đồng",
                "03  第三期付款條件達成時：支付合約金額的 10%\nĐợt 3: Thanh toán 10% giá trị hợp đồng",
            ],
        }

        generate_payment_request_word(data, output_path, template_path=template_path)

        source = Document(template_path)
        generated = Document(output_path)
        self.assertEqual(original_hash, hashlib.sha256(template_path.read_bytes()).hexdigest())
        self.assertEqual(len(generated.tables), len(source.tables))
        # Bảng điều khoản thanh toán có đúng số hàng bằng số lượng điều khoản thực tế
        self.assertEqual(len(generated.tables[3].rows), len(data["payment_terms"]))
        for idx in (0, 1, 2, 4):
            self.assertEqual(len(generated.tables[idx].rows), len(source.tables[idx].rows))
            self.assertEqual(len(generated.tables[idx].columns), len(source.tables[idx].columns))
        self.assertEqual(
            [table._tbl.tblGrid.xml for table in generated.tables],
            [table._tbl.tblGrid.xml for table in source.tables],
        )
        with ZipFile(template_path) as source_zip, ZipFile(output_path) as generated_zip:
            for part_name in source_zip.namelist():
                if part_name != "word/document.xml":
                    self.assertEqual(
                        generated_zip.read(part_name),
                        source_zip.read(part_name),
                        part_name,
                    )
        source_section = source.sections[0]
        generated_section = generated.sections[0]
        for attribute in (
            "page_width",
            "page_height",
            "top_margin",
            "bottom_margin",
            "left_margin",
            "right_margin",
            "header_distance",
            "footer_distance",
        ):
            self.assertEqual(getattr(generated_section, attribute), getattr(source_section, attribute))

        self.assertIn("HMV15102026TEST-SYTT", generated.tables[0].cell(0, 1).text)
        self.assertEqual(
            generated.tables[1].rows[1].cells[1].text,
            "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM",
        )
        self.assertEqual(generated.tables[1].rows[1].cells[-1].text, "15/10/2026")
        self.assertIn("HD-20261013-001", generated.tables[1].rows[2].cells[1].text)
        self.assertIn("Đợt 1", generated.tables[2].rows[1].cells[1].text)
        self.assertEqual(generated.tables[2].rows[1].cells[8].text, "282.654.114")
        self.assertIn("35%", generated.tables[3].rows[0].cells[0].text)
        self.assertIn("10%", generated.tables[3].rows[2].cells[0].text)
        self.assertEqual(len(generated.tables[3].rows), 3)
        first_term_runs = generated.tables[3].rows[0].cells[0].paragraphs[0].runs
        self.assertTrue(first_term_runs[0].bold)
        self.assertEqual(str(first_term_runs[0].font.color.rgb), "2E75B6")
        self.assertFalse(bool(first_term_runs[1].bold))
        self.assertEqual(str(first_term_runs[1].font.color.rgb), "000000")
        self.assertEqual(first_term_runs[1].font.size.pt, 8.0)
        self.assertEqual(generated.tables[4].cell(0, 3).text, source.tables[4].cell(0, 3).text)

    def test_generate_sheet_endpoint_preserves_original_error_message(self):
        with patch.object(
            api_routes,
            "generate_payment_request_excel",
            side_effect=RuntimeError("sentinel-original-error"),
        ):
            result = asyncio.run(
                api_routes.generate_payment_request_sheet_endpoint(
                    {"contract_code": "HMV-TEST"}
                )
            )

        self.assertEqual(result["status"], "error")
    def test_extract_docx_payment_voucher(self):
        from app.services.payment_voucher_ai_service import _extract_docx_text_from_base64
        import base64
        
        # Test document
        doc = Document()
        doc.add_paragraph("BIÊN BẢN TIẾN ĐỘ LẮP ĐẶT VÀ NGHIỆM THU")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Hợp đồng"
        table.rows[0].cells[1].text = "HMV151020250001"
        table.rows[1].cells[0].text = "Số tiền"
        table.rows[1].cells[1].text = "46.800.000 VND"
        
        tmp_doc = self.root / "test_acceptance.docx"
        doc.save(str(tmp_doc))
        
        with open(tmp_doc, "rb") as f:
            b64_val = base64.b64encode(f.read()).decode("utf-8")
            
        extracted_text = _extract_docx_text_from_base64(b64_val)
        self.assertIn("BIÊN BẢN TIẾN ĐỘ LẮP ĐẶT VÀ NGHIỆM THU", extracted_text)
        self.assertIn("HMV151020250001", extracted_text)
        self.assertIn("46.800.000 VND", extracted_text)


if __name__ == "__main__":
    unittest.main()

