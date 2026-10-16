# -*- coding: utf-8 -*-
"""
Unit tests for Dashboard & Live Project Tracker Service.
Validates live metrics aggregation, KPI calculation, real data binding,
and Excel export capabilities.
"""

import unittest
from pathlib import Path

from app.services.dashboard_service import (
    format_currency_vnd,
    format_millions_vnd,
    get_dashboard_overview_data,
    generate_dashboard_excel_report,
)


class TestDashboardService(unittest.TestCase):
    """Test suite for live dashboard analytics and reporting."""

    def test_01_currency_formatting(self):
        """Verify Vietnamese currency and millions formatting."""
        self.assertEqual(format_currency_vnd(0), "0 ₫")
        self.assertEqual(format_currency_vnd(807583184), "807.583.184 ₫")
        self.assertEqual(format_millions_vnd(807583184), "807,6")
        self.assertEqual(format_millions_vnd(1000000000), "1,00")

    def test_02_get_dashboard_overview_real_data(self):
        """Verify get_dashboard_overview_data returns structured live metrics."""
        data = get_dashboard_overview_data()
        self.assertEqual(data["status"], "success")
        self.assertIn("summary", data)
        self.assertIn("portfolio_bars", data)
        self.assertIn("donut", data)
        self.assertIn("action_items", data)
        self.assertIn("contracts", data)

        summary = data["summary"]
        self.assertGreaterEqual(summary["total_projects"], 0)
        self.assertGreaterEqual(summary["total_contract_value"], 0)
        self.assertIn("report_period", summary)
        self.assertIn("report_date", summary)

        # Donut structure
        donut = data["donut"]
        self.assertIn("gradient", donut)
        self.assertIn("active_rate", donut)
        self.assertIn("pct_running", donut)

    def test_03_generate_dashboard_excel_report(self):
        """Verify export_dashboard_excel produces valid xlsx workbook stream."""
        import openpyxl

        stream = generate_dashboard_excel_report()
        self.assertGreater(len(stream.getvalue()), 1000)

        # Load generated workbook to verify sheets and headers
        wb = openpyxl.load_workbook(stream, data_only=True)
        self.assertIn("Tổng Quan Dự Án", wb.sheetnames)
        self.assertIn("Đợt Thanh Toán", wb.sheetnames)

        ws_ov = wb["Tổng Quan Dự Án"]
        self.assertIn("BÁO CÁO TỔNG QUAN DANH MỤC DỰ ÁN", ws_ov["A1"].value)
        self.assertEqual(ws_ov["A4"].value, "STT")
        self.assertIn("Mã Hợp Đồng", ws_ov["B4"].value)

    def test_04_mark_and_unmark_milestone_paid(self):
        """Verify marking milestone as paid updates amounts and reverting works properly."""
        from app.database import get_contracts_dossier, mark_milestone_paid, unmark_milestone_paid

        contracts = get_contracts_dossier()
        if not contracts:
            return

        c = contracts[0]
        milestones = c.get("milestones", [])
        if not milestones:
            return

        ms = milestones[0]
        ms_id = ms["id"]
        original_paid_amount = c["paid_amount"]

        # Mark as paid
        res = mark_milestone_paid(ms_id, paid_date="15/10/2026", invoice_no="HD-TEST-001")
        self.assertIsNotNone(res)
        self.assertEqual(res["milestone"]["status"], "PAID")
        self.assertEqual(res["milestone"]["status_label"], "Đã thu đủ")
        self.assertGreaterEqual(res["new_paid_amount"], float(ms["amount"]))

        # Unmark (revert)
        revert_res = unmark_milestone_paid(ms_id)
        self.assertIsNotNone(revert_res)
        self.assertNotEqual(revert_res["milestone"]["status"], "PAID")
        self.assertIsNone(revert_res["milestone"]["paid_date"])

    def test_05_update_contract_status(self):
        """Verify updating contract status and notes."""
        from app.database import get_contracts_dossier, update_contract_status

        contracts = get_contracts_dossier()
        if not contracts:
            return

        c = contracts[0]
        c_code = c["contract_code"]
        orig_status = c["status"]

        updated = update_contract_status(c_code, "Đang chạy thử nghiệm thu (Testing 7 ngày)", "Ghi chú test")
        self.assertIsNotNone(updated)
        self.assertEqual(updated["status"], "Đang chạy thử nghiệm thu (Testing 7 ngày)")

        # Restore original status
        update_contract_status(c_code, orig_status, c.get("notes"))


if __name__ == "__main__":
    unittest.main()

