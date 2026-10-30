import io,json
from decimal import Decimal
from pathlib import Path
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.models.shipping_documents import ShippingDocument,review_document
from app.services import shipping_documents_ai_service as ai,shipping_documents_service as exporter


def order(n=1):
    return ShippingDocument(buyer_name='KHÁCH KIỂM THỬ / 測試客戶',buyer_address='Bắc Ninh',issue_date='2026-10-30',items=[dict(
        name_vi=f'Thiết bị {i+1}',name_zh='設備',description='Thông số kiểm thử',group='Phần cứng / 硬體' if i<10 else 'Vật tư / 材料',
        quantity=2,unit='CÁI',unit_price=1000,kind='goods',source='test.xlsx!A1:F10') for i in range(n)])


def test_parser_reads_beyond_60_and_labels_formula_without_cache():
    wb=openpyxl.Workbook();ws=wb.active
    ws['A80']='Hạng mục cuối';ws['B80']='=3*4'
    stream=io.BytesIO();wb.save(stream)
    section,parts=ai.extract_file('order.xlsx',stream.getvalue(),'order')
    assert 'A80: Hạng mục cuối' in section['content']
    assert 'B80: CÔNG THỨC CHƯA CÓ GIÁ TRỊ LƯU' in section['content']
    assert parts==[] and section['role']=='order'


def test_missing_money_and_weight_are_not_zero_and_services_excluded():
    d=order(2);d.items[0].unit_price=None
    d.items[1]=d.items[1].model_copy(update={'kind':'service','include_in_packing':False})
    r=review_document(d)
    assert r['total'] is None and r['packing_count']==1
    assert any('Cân nặng' in s for s in r['warnings'])
    d.items[0].unit_price=Decimal(0)
    assert review_document(d)['total']==2000


def test_conflicting_total_reported_instead_of_changing_prices():
    d=order();d.reported_total=Decimal(9999)
    r=review_document(d)
    assert r['total']==2000 and any('khác tổng' in s for s in r['warnings'])
    assert d.items[0].unit_price==1000


def test_reference_only_and_bad_upload_are_rejected():
    with pytest.raises(ValueError,match='Chỉ có file tham khảo'):
        ai.analyze([dict(role='reference')],[],[])
    with pytest.raises(ValueError,match='chưa hỗ trợ'):
        ai.extract_file('data.xls',b'old excel','order')


def test_ai_notes_list_is_normalized_without_changing_order(monkeypatch):
    raw=order().model_dump(mode='json');raw['notes']=['Ghi chú một','Ghi chú hai']
    monkeypatch.setattr(ai,'_call_gemini_multimodal',lambda *args:raw)
    result=ai.analyze([],[],['Đơn hàng'])
    assert result['data']['notes']=='Ghi chú một\nGhi chú hai'
    assert result['review']['total']==2000


def test_multisource_endpoint_passes_roles_and_notes(monkeypatch):
    captured={}
    def fake(sections,parts,texts):
        captured.update(sections=sections,parts=parts,texts=texts)
        data=order();return dict(data=data.model_dump(mode='json'),review=review_document(data))
    monkeypatch.setattr(ai,'analyze',fake)
    client=TestClient(app)
    response=client.post('/api/shipping-documents/analyze',data={'roles':json.dumps(['order','reference']),'texts':json.dumps(['Ghi chú 1','Ghi chú 2'])},
                         files=[('files',('new.txt','đơn mới'.encode(),'text/plain')),('files',('old.txt',b'old template','text/plain'))])
    assert response.status_code==200
    assert [s['role'] for s in captured['sections']]==['order','reference']
    assert captured['texts']==['Ghi chú 1','Ghi chú 2']
    assert client.get('/shipping-documents').status_code==200


def test_validation_and_no_path_traversal():
    client=TestClient(app)
    assert client.post('/api/shipping-documents/analyze',data={'roles':'[]','texts':'[]'}).status_code==422
    assert client.post('/api/shipping-documents/generate',json=order().model_copy(update={'buyer_name':''}).model_dump(mode='json')).status_code==422
    assert client.get('/api/shipping-documents/download/not-a-token/invoice').status_code==404
    assert client.get('/api/shipping-documents/download/'+'a'*32+'/secret').status_code==404


@pytest.mark.parametrize('count',[1,10,21,35])
def test_real_excel_export_exact_rows_formulas_and_center(tmp_path,monkeypatch,count):
    monkeypatch.setattr(exporter,'EXPORT_ROOT',tmp_path)
    d=order(count)
    if count>1:
        d.items[-1].kind='service';d.items[-1].include_in_packing=False
    result=exporter.generate(d)
    for kind,expected,column in [('invoice',count,'L'),('packing',count-(count>1),'G')]:
        path=exporter.download_path(result['job_id'],kind)
        f=openpyxl.load_workbook(path,data_only=False).active
        v=openpyxl.load_workbook(path,data_only=True).active
        info=json.loads((path.parent/'manifest.json').read_text(encoding='utf-8'))[kind]
        assert f.print_area and str(info['lastRow']) in f.print_area
        assert len(info['detailRows'])==expected
        for no,row in enumerate(info['detailRows'],1):
            assert f[f'A{row}'].value==no
            assert f[f'B{row}'].alignment.horizontal=='center'
            assert f[f'B{row}'].alignment.vertical=='center'
            assert f[f'{column}{row}'].data_type=='f'
        end=int(''.join(filter(str.isdigit,info['total'])))
        assert end==info['detailRows'][-1]+1
        if kind=='invoice':
            assert v[info['total']].value==count*2000
            assert sum(c.value=='INVOICE' for row in f for c in row)==1
        else:
            assert v[info['total']].value in [None,'']
            for row in info['detailRows']:assert v[f'E{row}'].value is None and v[f'G{row}'].value in [None,'']
        assert not any(c.data_type=='e' for row in v for c in row)
        assert all(c.value is None for row in f.iter_rows(min_row=info['lastRow']+1) for c in row)


def test_partial_weight_total_and_complete_weight_recalculate(tmp_path,monkeypatch):
    monkeypatch.setattr(exporter,'EXPORT_ROOT',tmp_path)
    d=order(2);d.items[0].unit_weight_kg=Decimal('1.5')
    result=exporter.generate(d)
    path=exporter.download_path(result['job_id'],'packing')
    v=openpyxl.load_workbook(path,data_only=True).active
    assert v['G13'].value==3 and v['G15'].value in [None,'']
    d.items[1].total_weight_kg=Decimal('4.25')
    result=exporter.generate(d);path=exporter.download_path(result['job_id'],'packing')
    v=openpyxl.load_workbook(path,data_only=True).active
    assert v['G15'].value==7.25
