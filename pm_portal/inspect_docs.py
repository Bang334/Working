import sys
import docx
from pathlib import Path

def dump_doc(path, label, out_f):
    p_path = Path(path)
    if not p_path.exists():
        out_f.write(f"File not found: {path}\n")
        return
    doc = docx.Document(path)
    out_f.write(f"\n==================== {label}: {path} ====================\n")
    out_f.write(f"Total Paragraphs: {len(doc.paragraphs)}, Total Tables: {len(doc.tables)}\n")
    
    out_f.write("\n--- PARAGRAPHS OVERVIEW ---\n")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            runs_info = []
            for r in p.runs:
                font_name = r.font.name
                font_size = r.font.size.pt if r.font.size else None
                bold = r.bold
                color = r.font.color.rgb if r.font.color else None
                runs_info.append(f"[{r.text[:15]} | fn={font_name} sz={font_size} b={bold} c={color}]")
            out_f.write(f"P{i:02d}: {txt}\n")
            if runs_info:
                out_f.write(f"     Runs: {' '.join(runs_info[:5])}\n")
                
    for t_idx, t in enumerate(doc.tables):
        out_f.write(f"\n--- TABLE {t_idx} (rows={len(t.rows)}, cols={len(t.columns)}) ---\n")
        for r_idx, row in enumerate(t.rows):
            cell_texts = []
            for c_idx, cell in enumerate(row.cells):
                cell_txt = cell.text.replace("\n", " // ").strip()
                # get formatting of first run
                r_sample = cell.paragraphs[0].runs[0] if cell.paragraphs and cell.paragraphs[0].runs else None
                fmt = f"fn={r_sample.font.name if r_sample else None} sz={r_sample.font.size.pt if r_sample and r_sample.font.size else None} b={r_sample.bold if r_sample else None} c={r_sample.font.color.rgb if r_sample and r_sample.font.color else None}" if r_sample else ""
                cell_texts.append(f"C{c_idx}[{fmt}]: {cell_txt}")
            # deduplicate adjacent merged cell text
            dedup = []
            for ct in cell_texts:
                val = ct.split("]: ", 1)[-1] if "]: " in ct else ct
                if not dedup or (dedup[-1].split("]: ", 1)[-1] if "]: " in dedup[-1] else dedup[-1]) != val:
                    dedup.append(ct)
            out_f.write(f"  Row {r_idx:02d}: {' | '.join(dedup[:4])}\n")

with open("inspect_results_utf8.txt", "w", encoding="utf-8") as f:
    dump_doc("E:/Bien_ban_nghiem_thu_HMV100120260001.docx", "MAU_USER_CUNG_CAP (E:/Bien_ban_nghiem_thu_HMV100120260001.docx)", f)
    dump_doc("C:/Users/BANG/Downloads/02_BIEN_BAN_NGHIEM_THU_BAN_GIAO_CX_.docx", "DU_LIEU_CX (02_BIEN_BAN_NGHIEM_THU_BAN_GIAO_CX_.docx)", f)
    dump_doc("C:/Users/BANG/Downloads/Bien_Ban_Nghiem_Thu_Lap_Dat_HMV151020250001.docx", "BAN_DA_TAO (Bien_Ban_Nghiem_Thu_Lap_Dat_HMV151020250001.docx)", f)

print("COMPLETED DUMPING TO inspect_results_utf8.txt")
