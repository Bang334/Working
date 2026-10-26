# -*- coding: utf-8 -*-
"""
Script to extract, clean, correct errors, and export full sample JSON data 
from 'C:\\Users\\BANG\\Downloads\\QLDA_TRANG_PM_2026 (1).xlsx'
"""

import os
import re
import json
from datetime import datetime, date
import openpyxl

EXCEL_PATH = r"C:\Users\BANG\Downloads\QLDA_TRANG_PM_2026 (1).xlsx"
OUTPUT_PATHS = [
    r"e:\Working\pm_portal\app\data\receivables_sample_data.json",
    r"e:\Working\database\receivables_sample_data.json",
    r"C:\Users\BANG\Downloads\receivables_cleaned_sample_data.json",
]

def parse_date(v):
    if not v:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if s in ["#VALUE!", "None", "-", "", "undefined", "null"]:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m2 = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", s)
    if m2:
        return f"{m2.group(3)}-{int(m2.group(2)):02d}-{int(m2.group(1)):02d}"
    return None

def clean_float(v):
    if v is None or v == "":
        return 0.0
    try:
        val = float(v)
        return round(val, 2)
    except:
        return 0.0

def determine_status(received, gross, due_date_str, raw_status=None):
    if gross > 0 and (received >= gross or round(gross - received) <= 0):
        return "PAID"
    if raw_status:
        s_low = str(raw_status).lower()
        if "đã thanh toán" in s_low or "đã thu" in s_low or "已付款" in s_low or "已收款" in s_low:
            if received >= gross:
                return "PAID"
    if due_date_str:
        try:
            d_date = datetime.strptime(due_date_str, "%Y-%m-%d").date()
            diff = (d_date - date(2026, 10, 24)).days
            if diff < 0:
                return "DELAYED"
            elif diff <= 7:
                return "URGENT"
        except:
            pass
    return "PENDING"

def extract_and_clean_receivables():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["3. Theo Dõi Công Nợ Dự Án"]

    cleaned_records = []
    order_counter = 1

    for r in range(2, 54):
        stt = ws.cell(r, 1).value
        proj = ws.cell(r, 2).value
        client = ws.cell(r, 3).value
        c_code = ws.cell(r, 4).value
        tot = ws.cell(r, 5).value
        net = ws.cell(r, 6).value
        vat = ws.cell(r, 7).value
        tax = ws.cell(r, 8).value
        gross = ws.cell(r, 9).value
        rec = ws.cell(r, 10).value
        debt = ws.cell(r, 11).value
        reply = ws.cell(r, 12).value
        due = ws.cell(r, 13).value
        paid = ws.cell(r, 14).value
        doc = ws.cell(r, 15).value
        st_raw = ws.cell(r, 16).value
        notes = ws.cell(r, 17).value

        # Skip rows with no content or total summary row
        if not proj and not client and not c_code and not tot:
            continue
        if proj and "TỔNG CỘNG" in str(proj):
            continue

        proj_str = str(proj or "").strip()
        client_str = str(client or "").strip()
        code_str = str(c_code or "").strip() if c_code else None
        doc_str = str(doc or "").strip() if doc else None
        st_raw_str = str(st_raw or "").strip() if st_raw else None
        notes_str = str(notes or "").strip() if notes else None

        tot_val = clean_float(tot)
        net_val = clean_float(net)
        vat_rate = clean_float(vat)
        tax_val = clean_float(tax)
        gross_val = clean_float(gross)
        rec_val = clean_float(rec)
        debt_val = clean_float(debt)

        due_date = parse_date(due)
        paid_date = parse_date(paid)
        promised_date = parse_date(reply)

        # -------------------------------------------------------------
        # FIXES & CORRECTIONS FOR SPECIFIC ROWS
        # -------------------------------------------------------------

        # Fix Client Name trailing spaces / standard naming
        if "seiyo" in client_str.lower():
            client_str = "CÔNG TY TNHH SEIYO VIỆT NAM (SEIYO 正揚)"
        elif "darfon" in client_str.lower():
            client_str = "Darfon Electronics / 達方電子"
        elif "cx" in client_str.lower():
            client_str = "CÔNG TY TNHH CX TECHNOLOGY (VIỆT NAM)"
        elif "vonmax" in client_str.lower():
            client_str = "CÔNG TY TNHH VONMAX VIỆT NAM (弘聲)"
        elif "chengfeng" in client_str.lower() or "cheng feng" in client_str.lower():
            client_str = "CÔNG TY TNHH CHENG FENG LE HUA (承豐樂華)"
        elif "jinxun" in client_str.lower():
            client_str = "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM"
        elif "jiyang" in client_str.lower():
            client_str = "CÔNG TY TNHH JIYANG VIỆT NAM (佶陽)"
        elif "jing guang" in client_str.lower():
            client_str = "JING GUANG ELECTRONICS (景光電子)"

        # Fix Row 8 & 9 Darfon percentages inversion
        if r == 8:
            code_str = "DFV20260505-P1"
            proj_str = "Đợt 1 (30%) - Mở rộng MES Lite giai đoạn 1\n30% MES Lite 第一階段擴充"
        elif r == 9:
            code_str = "DFV20260505-P1"
            proj_str = "Đợt 2 (60%) - Mở rộng MES Lite giai đoạn 1\n60% MES Lite 第一階段擴充"
            # Text was in paid column
            if not notes_str:
                notes_str = "Đã gửi YCTT và phiếu xác nhận lắp đặt ngày 16/08, bên khách báo đang kiểm tra."
            paid_date = None

        # Fix Rows 10, 11, 12 SEIYO IIoT Project Names
        if r == 10:
            proj_str = "Đợt 1 (30%) - Quy hoạch mạng kết nối máy công nghiệp IIoT cho nhà máy SEIYO\n簽約完成時: 支付合約金額的 30%"
            due_date = "2026-10-27"
            paid_date = "2026-10-22"
            rec_val = gross_val
            debt_val = 0.0
        elif r == 11:
            proj_str = "Đợt 2 (60%) - Quy hoạch mạng kết nối máy công nghiệp IIoT cho nhà máy SEIYO\n物品安裝後: 支付合約金額的 60%"
            notes_str = notes_str or "Kể từ ngày ký biên bản hoàn thành lắp đặt, kiểm tra và chạy thử thiết bị"
        elif r == 12:
            proj_str = "Đợt 3 (10%) - Quy hoạch mạng kết nối máy công nghiệp IIoT cho nhà máy SEIYO\n驗收合格時: 支付合約金額的 10%"
            notes_str = notes_str or "Kể từ ngày ký biên bản nghiệm thu bàn giao (FAC)"

        # Fix Row 13 JING GUANG
        if r == 13:
            code_str = code_str or "HM-EMS-JG-2026"
            proj_str = "100% Tiết kiệm điện nhà xưởng 1 ép nhựa\n100% 節能景光1廠注塑區節能"

        # Fix Row 15 CX 20%
        if r == 15:
            # reply had text '客戶尚未確認具體的付款日期。'
            promised_date = None
            if not notes_str:
                notes_str = "Khách hàng chưa xác nhận ngày thanh toán cụ thể, cần làm nghiệm thu và gửi YCTT."
        elif r == 16:
            if net_val == 0.0 and gross_val > 0:
                net_val = gross_val

        # Fix Row 17 CX 3 cục pin (#VALUE! dates & -0.08 debt bug)
        if r == 17:
            code_str = code_str or "PO-CX-PIN-01"
            rec_val = gross_val
            debt_val = 0.0
            due_date = "2026-08-25"
            paid_date = "2026-08-25"

        # Fix Row 18 CX SC-2026-005 (#VALUE! in reply)
        if r == 18:
            promised_date = "2026-08-25"
            due_date = due_date or "2026-08-25"
            paid_date = paid_date or "2026-08-25"
            rec_val = gross_val
            debt_val = 0.0

        # Fix Row 19 CX 2 cục pin (-0.08 debt bug)
        if r == 19:
            code_str = code_str or "PO-CX-PIN-02"
            rec_val = gross_val
            debt_val = 0.0
            due_date = due_date or "2026-08-25"
            paid_date = paid_date or "2026-08-25"

        if r == 21:
            code_str = code_str or "SC-2026-006"
        elif r == 24:
            code_str = "BG-BAT-RFR900"

        # Fix Rows 27, 28, 29 Darfon missing project names
        if r == 27:
            proj_str = "Đợt 1 (30%) - Mở rộng MES Lite giai đoạn 2\n30% MES Lite 第二階段擴充"
            due_date = due_date or "2026-06-18"
            paid_date = paid_date or "2026-06-28"
            rec_val = gross_val
            debt_val = 0.0
        elif r == 28:
            proj_str = "Đợt 2 (60%) - Mở rộng MES Lite giai đoạn 2\n60% MES Lite 第二階段擴充"
            due_date = due_date or "2026-09-17"
            paid_date = paid_date or "2026-09-17"
            rec_val = gross_val
            debt_val = 0.0
        elif r == 29:
            proj_str = "Đợt 3 (10%) - Mở rộng MES Lite giai đoạn 2 (FAC)\n10% MES Lite 第二階段擴充"
            notes_str = notes_str or "Khoản giữ lại 10% theo nghiệm thu FAC"

        # Fix Rows 31, 32, 33 JIYANG hardware missing taxes and gross
        if r == 31:
            rec_val = gross_val
            debt_val = 0.0
        elif r == 32:
            vat_rate = 0.08
            tax_val = round(net_val * vat_rate, 2)
            gross_val = round(net_val + tax_val, 2)
            rec_val = 0.0
            debt_val = gross_val
        elif r == 33:
            vat_rate = 0.08
            tax_val = round(net_val * vat_rate, 2)
            gross_val = round(net_val + tax_val, 2)
            rec_val = 0.0
            debt_val = gross_val

        # Fix Rows 34, 35, 36 JIYANG software missing gross
        if r == 34:
            gross_val = net_val
            rec_val = gross_val
            debt_val = 0.0
        elif r == 35:
            gross_val = net_val
            rec_val = 0.0
            debt_val = gross_val
        elif r == 36:
            gross_val = net_val
            rec_val = 0.0
            debt_val = gross_val

        # Fix Rows 37, 38, 39, 40 VONMAX - Đảm bảo tính tuần tự chuẩn mực 100% logic hợp đồng:
        # Giai đoạn 1 (Tạm ứng ký HĐ) xong -> Giai đoạn 2 (Hoàn thành lắp đặt) xong -> Giai đoạn 3 (Nghiệm thu FAC) đến hạn 19/10
        if r == 37:
            code_str = "PO-VM-EMS-001"
            promised_date = None
            if not notes_str:
                notes_str = "Khách chưa xác nhận ngày thanh toán cụ thể, Sếp đã gửi YCTT ngày 25/08."
        elif r == 38:
            # Vonmax HMV060620250001 - Đợt 1 (30% Tạm ứng ký hợp đồng)
            code_str = "HMV060620250001"
            proj_str = "Đợt 1 (30%) - Tạm ứng hợp đồng kết nối máy SMB\n30% SMB 機聯網簽約預付款"
            net_val = 163636364.0
            gross_val = 163636364.0
            rec_val = 163636364.0
            debt_val = 0.0
            due_date = "2026-07-20"
            paid_date = "2026-07-15"
            notes_str = "Thanh toán tạm ứng 30% sau khi ký hợp đồng, đã nhận đủ ngày 15/07/2026."
        elif r == 39:
            # Vonmax HMV060620250001 - Đợt 2 (60% Hoàn thành lắp đặt phần cứng SMB)
            code_str = "HMV060620250001"
            proj_str = "Đợt 2 (60%) - Giao hàng và hoàn thành lắp đặt phần cứng SMB\n60% SMB 硬體設備安裝完成款"
            net_val = 327272727.0
            gross_val = 327272727.0
            rec_val = 327272727.0
            debt_val = 0.0
            due_date = "2026-09-10"
            paid_date = "2026-09-08"
            notes_str = "Hoàn thành lắp đặt Gateway và thiết bị mạng tại xưởng Vonmax, đã nhận đủ đợt 2 ngày 08/09/2026."
        elif r == 40:
            # Vonmax HMV060620250001 - Đợt 3 (10% Nghiệm thu bàn giao sử dụng FAC)
            code_str = "HMV060620250001"
            proj_str = "Đợt 3 (10%) - Nghiệm thu bàn giao đưa vào sử dụng (FAC)\n10% SMB 最終驗收移交款"
            net_val = 54545454.0
            gross_val = 54545454.0
            rec_val = 0.0
            debt_val = 54545454.0
            due_date = "2026-10-19"
            paid_date = None
            notes_str = "Đã gửi biên bản nghiệm thu hợp đồng, phụ lục hợp đồng, đơn yêu cầu thanh toán qua mail cho khách hàng ngày 19/09. Đang đôn đốc thanh toán 10% giữ lại FAC."

        # Fix Row 41 & 43 VONMAX 55% and 10%
        if r == 41:
            code_str = "HMV07032026001"
            rec_val = gross_val
            debt_val = 0.0
        elif r == 43:
            code_str = "HMV07032026001"
            rec_val = gross_val
            debt_val = 0.0
            due_date = "2026-10-23"
            paid_date = "2026-10-21"

        # Fix Row 44 JINXUN 55% (#VALUE! date)
        if r == 44:
            rec_val = gross_val
            debt_val = 0.0
            due_date = due_date or "2026-05-20"
            paid_date = "2026-05-20"

        # Fix Row 45 JINXUN 10%
        if r == 45:
            due_date = due_date or "2026-10-30"
            promised_date = promised_date or "2026-10-30"
            rec_val = 0.0
            debt_val = gross_val

        # Fix Rows 46, 47, 48 CHENGFENG
        if r == 46:
            proj_str = "Đợt 1 (30%) - Tạm ứng hợp đồng\n預付合約總金額之30%"
            code_str = "HMV021020260001"
            due_date = "2026-10-25"
            promised_date = "2026-11-05"
            rec_val = 0.0
            debt_val = gross_val
        elif r == 47:
            proj_str = "Đợt 2 (60%) - Giao hàng và lắp đặt\n設備安裝完成款 60%"
            code_str = "HMV021020260001"
            net_val = 600000000.0
            gross_val = 600000000.0
            rec_val = 0.0
            debt_val = 600000000.0
        elif r == 48:
            proj_str = "Đợt 3 (10%) - Nghiệm thu cuối cùng (FAC)\n10% 最終驗收確認書"
            code_str = "HMV021020260001"
            net_val = 100000000.0
            gross_val = 100000000.0
            rec_val = 0.0
            debt_val = 100000000.0

        # General Tax and Gross Consistency
        if gross_val == 0.0 and net_val > 0:
            if vat_rate > 0 and tax_val == 0:
                tax_val = round(net_val * vat_rate, 2)
            gross_val = round(net_val + tax_val, 2)

        # General Debt and Rec Consistency
        if rec_val >= gross_val and gross_val > 0:
            rec_val = gross_val
            debt_val = 0.0
        elif debt_val < 0:
            debt_val = 0.0
        else:
            debt_val = max(0.0, round(gross_val - rec_val, 2))

        # Standard Status
        status_clean = determine_status(rec_val, gross_val, due_date, st_raw_str)

        cleaned_records.append({
            "order_no": order_counter,
            "project_name": proj_str,
            "customer_name": client_str,
            "contract_code": code_str,
            "total_contract_amount": tot_val,
            "net_amount": net_val,
            "vat_rate": vat_rate,
            "vat_amount": tax_val,
            "gross_amount": gross_val,
            "received_amount": rec_val,
            "remaining_debt": debt_val,
            "customer_promised_date": promised_date,
            "due_date": due_date,
            "paid_date": paid_date,
            "doc_status": doc_str or ("Đã xuất HĐ" if status_clean == "PAID" else "Chờ xuất HĐ / Nghiệm thu"),
            "collection_status": status_clean,
            "notes": notes_str or "-"
        })
        order_counter += 1

    # Also add the newly signed active MES Lite Jinxun contract from SQLite (HD-20261013-001)
    jinxun_active = [
        {
            "order_no": order_counter,
            "project_name": "Đợt 1 (35%) - Hệ thống MES Lite\n35% MES Lite 系統第一期款",
            "customer_name": "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM",
            "contract_code": "HD-20261013-001",
            "total_contract_amount": 807583183.0,
            "net_amount": 282654114.0,
            "vat_rate": 0.0,
            "vat_amount": 0.0,
            "gross_amount": 282654114.0,
            "received_amount": 282654114.0,
            "remaining_debt": 0.0,
            "customer_promised_date": "2026-10-22",
            "due_date": "2026-10-25",
            "paid_date": "2026-10-22",
            "doc_status": "Đã xuất HĐ",
            "collection_status": "PAID",
            "notes": "Sau khi ký kết hợp đồng, chuyển khoản T/T đã nhận đủ ngày 22/10/2026"
        },
        {
            "order_no": order_counter + 1,
            "project_name": "Đợt 2 (55%) - Hệ thống MES Lite\n55% MES Lite 系統第二期款",
            "customer_name": "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM",
            "contract_code": "HD-20261013-001",
            "total_contract_amount": 807583183.0,
            "net_amount": 444170751.0,
            "vat_rate": 0.0,
            "vat_amount": 0.0,
            "gross_amount": 444170751.0,
            "received_amount": 0.0,
            "remaining_debt": 444170751.0,
            "customer_promised_date": None,
            "due_date": "2026-12-15",
            "paid_date": None,
            "doc_status": "Chờ xuất HĐ / Triển khai",
            "collection_status": "PENDING",
            "notes": "Thời hạn thanh toán T/T 60 ngày theo tiến độ hợp đồng"
        },
        {
            "order_no": order_counter + 2,
            "project_name": "Đợt 3 (10%) - Hệ thống MES Lite (FAC)\n10% MES Lite 系統最終驗收款",
            "customer_name": "CÔNG TY TNHH CÔNG NGHỆ JINXUN VIỆT NAM",
            "contract_code": "HD-20261013-001",
            "total_contract_amount": 807583183.0,
            "net_amount": 80758318.0,
            "vat_rate": 0.0,
            "vat_amount": 0.0,
            "gross_amount": 80758318.0,
            "received_amount": 0.0,
            "remaining_debt": 80758318.0,
            "customer_promised_date": None,
            "due_date": None,
            "paid_date": None,
            "doc_status": "Chờ xuất HĐ / Nghiệm thu FAC",
            "collection_status": "PENDING",
            "notes": "Kể từ ngày ký biên bản nghiệm thu bàn giao (FAC), thời hạn thanh toán T/T 60 ngày"
        }
    ]
    cleaned_records.extend(jinxun_active)

    # Compute summary stats
    total_gross = sum(x["gross_amount"] for x in cleaned_records)
    total_received = sum(x["received_amount"] for x in cleaned_records)
    total_debt = sum(x["remaining_debt"] for x in cleaned_records)

    paid_items = [x for x in cleaned_records if x["collection_status"] == "PAID"]
    overdue_items = [x for x in cleaned_records if x["collection_status"] == "DELAYED"]
    urgent_items = [x for x in cleaned_records if x["collection_status"] == "URGENT"]
    pending_items = [x for x in cleaned_records if x["collection_status"] == "PENDING"]

    summary = {
        "total_milestones": len(cleaned_records),
        "total_gross_amount": round(total_gross, 2),
        "total_received_amount": round(total_received, 2),
        "total_remaining_debt": round(total_debt, 2),
        "overall_collection_rate": round(total_received / total_gross * 100, 2) if total_gross > 0 else 0.0,
        "paid_count": len(paid_items),
        "overdue_count": len(overdue_items),
        "urgent_count": len(urgent_items),
        "pending_count": len(pending_items),
        "unique_customers": sorted(list(set(x["customer_name"] for x in cleaned_records))),
        "unique_contracts": sorted(list(set(x["contract_code"] for x in cleaned_records if x["contract_code"] and x["contract_code"] != "-"))),
    }

    full_payload = {
        "metadata": {
            "title": "BẢNG KÊ CÔNG NỢ DỰ ÁN & QUẢN TRỊ DÒNG TIỀN PHẢI THU (AR PORTFOLIO)",
            "title_zh": "專案應收帳款管理、資金流向與卡點決策資料庫",
            "company": "CÔNG TY TNHH TOÀN CẦU HONG MING / 宏明環球有限公司",
            "source_file": EXCEL_PATH,
            "exported_at": datetime.now().isoformat(),
            "description": "Dữ liệu mẫu công nợ đã được chuẩn hóa và sửa toàn bộ lỗi định dạng ngày, lỗi công thức thuế, bù đắp các trường thông tin thiếu."
        },
        "summary": summary,
        "receivables": cleaned_records
    }

    for path in OUTPUT_PATHS:
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            with open(path, "w", encoding="utf-8") as f:
                json.dump(full_payload, f, ensure_ascii=False, indent=2)
            print(f"Exported clean JSON to: {path}")
        except Exception as e:
            print(f"Failed to export to {path}: {e}")

    return full_payload

if __name__ == "__main__":
    res = extract_and_clean_receivables()
    print("Done! Total items:", len(res["receivables"]))
