# -*- coding: utf-8 -*-
"""
Multi-file document extractor service for payment requests and AR reminders.
Supports parsing DOCX, XLSX, PDF, and image files via python libraries and Gemini AI.
"""

import io
import json
import base64
import logging
import re
import urllib.request
import urllib.error
from typing import Any, Dict, List, Tuple
from fastapi import UploadFile
import docx
import openpyxl

from app.config import CANDIDATE_MODELS, DEFAULT_GEMINI_KEYS, DEFAULT_GEMINI_KEY

logger = logging.getLogger(__name__)

SYSTEM_PROMPT = """Bạn là chuyên gia phân tích dữ liệu hợp đồng, nghiệm thu và đề nghị thanh toán của CÔNG TY TNHH TOÀN CẦU HONG MING (宏明環球有限公司).

NHIỆM VỤ:
Bạn được cung cấp nội dung của một hoặc nhiều tài liệu (ví dụ: Biên bản nghiệm thu bàn giao, Đơn xin thanh toán nội bộ, Hợp đồng, Báo giá, hoặc Ghi chú dự án).
Nhiệm vụ của bạn là đọc toàn diện, đối chiếu chéo số liệu giữa tất cả các tài liệu và bóc tách thành một cấu trúc JSON hoàn chỉnh để lập "Phiếu yêu cầu thanh toán (請款單)" gửi cho khách hàng và soạn thư nhắc nợ.

QUY TẮC PHÂN TÍCH & BÓC TÁCH:
1. Thông tin khách hàng (customer_name_vi, customer_name_zh):
   - Nhận diện đúng bên mua/bên thanh toán (Ví dụ: CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN) / CX TECHNOLOGY (VN) 股份有限公司).
   - Nếu có email người nhận (To/CC), trích xuất vào mảng email_to và email_cc.
2. Hợp đồng & Dự án:
   - contract_code: Mã số hợp đồng (Ví dụ: HMV151020250001).
   - contract_date: Ngày ký hợp đồng nếu có (Ví dụ: 09/12/2025).
   - project_name_vi: Tên dự án tiếng Việt (Ví dụ: RFID 硬體 (RFID Hardware) – Phần mềm tùy chỉnh chức năng hệ thống RFID).
   - project_name_zh: Tên dự án tiếng Trung (Ví dụ: RFID硬體（RFID Hardware）－RFID系統客製化軟體).
   - total_contract_amount: Tổng giá trị hợp đồng bằng số nguyên (Ví dụ: 234000000).
3. Đợt thanh toán mục tiêu (đợt cần thu tiền):
   - milestone_name: Tên đợt thanh toán song ngữ (Ví dụ: 驗收合格尾款：20% / Thanh toán đợt Giai đoạn 3 (20% Nghiệm thu đạt tiêu chuẩn)).
   - milestone_percentage: Số phần trăm của đợt này (Ví dụ: 20 hoặc 30 hoặc 40 hoặc 100).
   - payment_amount: Số tiền đề nghị thanh toán đợt này bằng số nguyên (Ví dụ: 46800000). Phải đối soát để payment_amount = total_contract_amount * percentage / 100 nếu hợp đồng quy định theo tỷ lệ %.
   - currency: Loại tiền tệ (VND hoặc USD).
   - request_date: Ngày nộp đơn (dd/mm/yyyy). Nếu không nói rõ thì lấy ngày hiện tại (21/10/2026).
   - due_date: Ngày đến hạn thanh toán (dd/mm/yyyy). Nếu biên bản nghiệm thu ghi 10-15 ngày thì tính từ request_date.
   - payment_method: "匯款\nChuyển khoản".
4. Tóm tắt tiến độ các đợt thanh toán (payment_terms_summary):
   - Trích dẫn đầy đủ toàn bộ các giai đoạn thanh toán (Đợt 1, Đợt 2, Đợt 3...) từ hợp đồng hoặc mục kết luận thanh toán của biên bản nghiệm thu (kèm số tiền và chữ tiếng Việt - tiếng Trung) để điền vào phần Ghi chú (備註) của Phiếu yêu cầu thanh toán.
5. Soạn thư nhắc nợ / email gợi ý:
   - email_subject: Tiêu đề email song ngữ chuẩn mực doanh nghiệp Đài Loan: 【宏明環球】合約編號 {contract_code} {milestone_name} 到期付款提醒 / Yêu cầu thanh toán {milestone_name} ({customer_name_vi}).
   - email_body: Nội dung thư ngắn gọn, lịch thiệp, song ngữ Việt - Trung.

ĐỊNH DẠNG JSON ĐẦU RA BẮT BUỘC (Chỉ trả về JSON thuần túy, không có markdown block):
{
  "customer_name_vi": "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)",
  "customer_name_zh": "CX TECHNOLOGY (VN) 股份有限公司",
  "contract_code": "HMV151020250001",
  "contract_date": "09/12/2025",
  "project_name_vi": "RFID 硬體 (RFID Hardware) – Phần mềm tùy chỉnh chức năng hệ thống RFID",
  "project_name_zh": "RFID硬體（RFID Hardware）－RFID系統客製化軟體",
  "total_contract_amount": 234000000,
  "milestone_name": "驗收合格尾款：20% (Thanh toán đợt 3 Nghiệm thu đạt tiêu chuẩn)",
  "milestone_percentage": 20,
  "payment_amount": 46800000,
  "currency": "VND",
  "request_date": "21/10/2026",
  "due_date": "31/10/2026",
  "payment_method": "匯款\nChuyển khoản",
  "payment_terms_summary": "1. Sau khi ký kết hợp đồng: thanh toán 40% giá trị hợp đồng = 93.600.000 VNĐ.\n1. 簽約完成時：支付合約金額40%＝93,600,000越盾。\n2. Sau khi triển khai sản phẩm đưa vào sử dụng: thanh toán 40% giá trị hợp đồng = 93.600.000 VNĐ.\n2. 產品安裝後：支付合約金額40%＝93,600,000越盾。\n3. Khi nghiệm thu đạt tiêu chuẩn: thanh toán 20% còn lại = 46.800.000 VNĐ.\n3. 驗收合格時：支付尾款20%＝46,800,000越盾。",
  "email_to": ["pur04_vn@mail.cxtechnology.com"],
  "email_cc": ["jenson@youngming-mes.com", "sales@hongming-global.com"],
  "email_subject": "【宏明環球】合約編號 HMV151020250001 驗收合格尾款（20%）付款提醒 / Đề nghị thanh toán đợt 3 HĐ HMV151020250001 (CX TECHNOLOGY)",
  "email_body": "Kính gửi Quý Ban Giám đốc và Phòng Mua hàng CX Technology...\n\n致 CX Technology 貴公司管理部與採購部：...",
  "analyzed_files": [
    "Tên file 1 (mô tả ngắn)",
    "Tên file 2 (mô tả ngắn)"
  ]
}
"""


def _read_docx_content(content_bytes: bytes) -> str:
    """Extract paragraphs and tables from a docx file."""
    doc = docx.Document(io.BytesIO(content_bytes))
    lines = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            lines.append(txt)
    for idx, table in enumerate(doc.tables, 1):
        lines.append(f"\n--- BẢNG {idx} TRONG FILE DOCX ---")
        for row in table.rows:
            row_vals = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append(" | ".join(row_vals))
    return "\n".join(lines)


def _read_xlsx_content(content_bytes: bytes) -> str:
    """Extract sheets and populated cells from an xlsx file."""
    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
    lines = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        lines.append(f"\n--- SHEET: {sname} ---")
        for r in range(1, min(ws.max_row + 1, 60)):
            row_cells = []
            for c in range(1, min(ws.max_column + 1, 20)):
                val = ws.cell(r, c).value
                if val is not None:
                    row_cells.append(f"{ws.cell(r, c).coordinate}: {str(val).strip()}")
            if row_cells:
                lines.append(f"Row {r:2d}: " + " | ".join(row_cells))
    return "\n".join(lines)


async def extract_content_from_files(files: List[UploadFile]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Extract readable text parts and multimodal inline images/PDFs from uploaded files.
    Returns (text_sections, inline_gemini_parts).
    """
    text_sections = []
    inline_gemini_parts = []

    for f in files:
        fname = f.filename or "unknown"
        ext = fname.lower().split(".")[-1] if "." in fname else ""
        content = await f.read()

        if ext == "docx":
            try:
                extracted = _read_docx_content(content)
                text_sections.append({
                    "filename": fname,
                    "type": "docx",
                    "content": extracted
                })
            except Exception as exc:
                logger.warning("Error reading docx %s: %s", fname, exc)
                text_sections.append({"filename": fname, "type": "docx", "content": f"Lỗi đọc file: {exc}"})

        elif ext in ("xlsx", "xls"):
            try:
                extracted = _read_xlsx_content(content)
                text_sections.append({
                    "filename": fname,
                    "type": "xlsx",
                    "content": extracted
                })
            except Exception as exc:
                logger.warning("Error reading xlsx %s: %s", fname, exc)
                text_sections.append({"filename": fname, "type": "xlsx", "content": f"Lỗi đọc file: {exc}"})

        elif ext in ("txt", "csv"):
            try:
                extracted = content.decode("utf-8", errors="ignore")
                text_sections.append({
                    "filename": fname,
                    "type": "text",
                    "content": extracted
                })
            except Exception as exc:
                text_sections.append({"filename": fname, "type": "text", "content": f"Lỗi đọc: {exc}"})

        elif ext in ("png", "jpg", "jpeg", "webp"):
            mime = f"image/{ext if ext != 'jpg' else 'jpeg'}"
            b64_data = base64.b64encode(content).decode("utf-8")
            inline_gemini_parts.append({
                "inlineData": {
                    "mimeType": mime,
                    "data": b64_data
                }
            })
            text_sections.append({
                "filename": fname,
                "type": "image",
                "content": f"[Ảnh đính kèm: {fname}]"
            })

        elif ext == "pdf":
            # Gửi trực tiếp PDF qua inlineData của Gemini
            b64_data = base64.b64encode(content).decode("utf-8")
            inline_gemini_parts.append({
                "inlineData": {
                    "mimeType": "application/pdf",
                    "data": b64_data
                }
            })
            text_sections.append({
                "filename": fname,
                "type": "pdf",
                "content": f"[Tài liệu PDF: {fname}]"
            })

    return text_sections, inline_gemini_parts


def _clean_json_response(raw_text: str) -> str:
    cleaned = re.sub(r"^```(?:json)?\s*", "", raw_text.strip(), flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    match = re.search(r"(\{.*\})", cleaned, re.DOTALL)
    return match.group(1).strip() if match else cleaned.strip()


def extract_payment_info_with_gemini(
    text_sections: List[Dict[str, Any]],
    inline_gemini_parts: List[Dict[str, Any]],
    preferred_model: str = None
) -> Dict[str, Any]:
    """Call Gemini to extract structured payment request information from all uploaded documents."""
    models_to_try = [preferred_model] if preferred_model else []
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    keys_to_try = list(DEFAULT_GEMINI_KEYS)
    if DEFAULT_GEMINI_KEY and DEFAULT_GEMINI_KEY not in keys_to_try:
        keys_to_try.insert(0, DEFAULT_GEMINI_KEY)

    # Ghép văn bản từ tất cả các file
    compiled_texts = []
    compiled_texts.append("DANH SÁCH CÁC TỆP ĐƯỢC TẢI LÊN:")
    for idx, item in enumerate(text_sections, 1):
        compiled_texts.append(f"\n=======================================================")
        compiled_texts.append(f"TỆP {idx}: {item['filename']} (Loại: {item['type']})")
        compiled_texts.append(f"=======================================================")
        compiled_texts.append(item["content"])

    full_text_prompt = "\n".join(compiled_texts)

    content_parts = []
    # Đưa ảnh / PDF trước nếu có
    for part in inline_gemini_parts:
        content_parts.append(part)
    content_parts.append({"text": full_text_prompt})

    last_error = None
    for model_name in models_to_try:
        for api_key in keys_to_try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
            payload = {
                "systemInstruction": {"parts": [{"text": SYSTEM_PROMPT}]},
                "contents": [
                    {
                        "role": "user",
                        "parts": content_parts
                    }
                ],
                "generationConfig": {
                    "temperature": 0.1,
                    "responseMimeType": "application/json"
                }
            }

            try:
                data = json.dumps(payload).encode("utf-8")
                req = urllib.request.Request(
                    url,
                    data=data,
                    headers={"Content-Type": "application/json"},
                    method="POST"
                )

                with urllib.request.urlopen(req, timeout=40) as resp:
                    resp_data = json.loads(resp.read().decode("utf-8"))
                    candidates = resp_data.get("candidates", [])
                    if candidates:
                        raw_content = candidates[0]["content"]["parts"][0]["text"]
                        cleaned_json = _clean_json_response(raw_content)
                        parsed = json.loads(cleaned_json)
                        parsed["_executed_model"] = model_name
                        return parsed

            except urllib.error.HTTPError as http_err:
                status_code = http_err.code
                err_body = http_err.read().decode("utf-8", errors="ignore")
                logger.warning(
                    "Model %s failed with key ...%s (HTTP %s): %s",
                    model_name, api_key[-6:], status_code, err_body[:120]
                )
                last_error = http_err
                if status_code in (429, 403, 400):
                    continue
            except Exception as exc:
                logger.warning("Error with model %s: %s", model_name, exc)
                last_error = exc
                continue

    raise RuntimeError(f"Tất cả mô hình Gemini đều không thể xử lý yêu cầu. Lỗi cuối: {last_error}")
