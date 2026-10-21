# -*- coding: utf-8 -*-
"""Tạo Phiếu yêu cầu thanh toán Excel + Word theo sheet chuẩn 712TR."""

import html
import os
import re
import openpyxl
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from app.config import WORKING_DIR
from app.database import get_db_connection
from app.services.reminder_service import (
    number_to_vietnamese_words,
    number_to_chinese_words,
    format_amount_display,
)
from app.services.payment_request_word_service import generate_payment_request_word

TEMPLATE_PATH = WORKING_DIR / "Bieu Mau" / "đơn yêu cầu NCC thanh toán-Gốc.xlsx"
WORD_TEMPLATE_PATH = WORKING_DIR / "Bieu Mau" / "đơn yêu cầu NCC thanh toán-Gốc.docx"
TEMPLATE_SHEET_NAME = "712TR"
OUTPUT_DIR = WORKING_DIR / "Ho So Du An" / "Yeu Cau Thanh Toan"
os.makedirs(OUTPUT_DIR, exist_ok=True)


def _coalesce(explicit: Any, fallback: Any) -> Any:
    return explicit if explicit is not None else fallback


def _normalize_template_source(value: Optional[str]) -> str:
    source = str(value or "excel").strip().lower()
    if source not in {"excel", "word"}:
        raise ValueError("Nguồn mẫu chỉ nhận 'excel' hoặc 'word'.")
    return source


def _format_percentage(value: Any) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "0"
    return str(int(number)) if number.is_integer() else f"{number:.2f}".rstrip("0").rstrip(".")


def _parse_application_date(value: str) -> Any:
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(value, fmt)
        except (TypeError, ValueError):
            continue
    return value


def _safe_file_component(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", str(value or "")).strip().rstrip(".")
    return cleaned or "Khach_hang"


def _milestone_trigger_zh(milestone: Dict[str, Any], index: int) -> str:
    # Nếu milestone đã có điều kiện tiếng Trung sẵn từ HĐ
    if milestone.get("dieu_kien_zh"):
        return str(milestone["dieu_kien_zh"]).strip()

    source = f"{milestone.get('milestone_name', '')} {milestone.get('trigger_condition', '')} {milestone.get('dieu_kien_vi', '')}".lower()

    # 1. Kiểm tra nghiệm thu / FAC trước
    if "nghiệm thu" in source or "fac" in source or "hoàn công" in source:
        return "最終驗收合格並簽署移交文件 (FAC) 時"
    # 2. Kiểm tra lắp đặt / cài đặt / chạy thử thiết bị
    if "lắp đặt" in source or "cài đặt" in source or "chạy thử" in source:
        return "設備安裝及調測完成時"
    # 3. Kiểm tra giao hàng / bàn giao vật tư
    if "giao hàng" in source or "bàn giao" in source:
        return "設備交付完成時"
    # 4. Kiểm tra thi công / kéo cáp
    if "kéo cáp" in source or "thi công" in source:
        return "線路施工完成時"
    # 5. Ký kết hợp đồng / tạm ứng / đặt cọc / đợt 1
    if "ký kết" in source or "ký hợp đồng" in source or "tạm ứng" in source or "đặt cọc" in source or index == 1:
        return "雙方簽約完成時"

    return f"第{index:02d}期付款條件達成時"


def _build_bilingual_terms_data(
    milestones: Sequence[Dict[str, Any]],
    active_milestone_name: str = ""
) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    active_name_clean = str(active_milestone_name or "").strip().lower()

    for index, milestone in enumerate(milestones, start=1):
        percentage = _format_percentage(milestone.get("percentage"))
        amount = int(round(float(milestone.get("amount") or 0)))
        amount_display = format_amount_display(amount)
        words_vi = number_to_vietnamese_words(amount)
        words_zh = number_to_chinese_words(amount)

        trigger_vi = str(
            milestone.get("trigger_condition")
            or milestone.get("dieu_kien_vi")
            or milestone.get("milestone_name")
            or f"Đợt {index}"
        ).strip()
        trigger_zh = _milestone_trigger_zh(milestone, index)

        ms_title = str(milestone.get("milestone_name") or f"Đợt {index}").strip()
        is_active = bool(
            active_name_clean
            and (ms_title.lower() == active_name_clean or active_name_clean in ms_title.lower())
        )

        is_paid = str(milestone.get("status") or "").upper() == "PAID"

        # 1. Trạng thái thanh toán: CHỈ hiển thị khi đã thanh toán (is_paid), TUYỆT ĐỐI không hiện 'Chưa thanh toán'
        status_vi = " Đã thanh toán." if is_paid else ""
        status_zh = " 已支付。" if is_paid else ""

        # 2. Hạn thanh toán: CHỈ hiển thị cho đợt đang xin thanh toán (is_active).
        # Các đợt chưa đến hạn / chưa xác định được ngày hoàn thành thì TUYỆT ĐỐI không hiển thị hạn thanh toán
        due_date = str(milestone.get("due_date") or "").strip()
        extra_vi = ""
        extra_zh = ""
        if is_active and due_date:
            extra_vi = f" Hạn thanh toán: {due_date}."
            extra_zh = f" 付款到期日：{due_date}。"

        invoice_no = str(milestone.get("invoice_no") or "").strip()
        if invoice_no and invoice_no != "Chưa xuất":
            extra_vi += f" Hóa đơn: {invoice_no}."
            extra_zh += f" 發票編號：{invoice_no}。"

        # Soạn nội dung từng ngôn ngữ sạch (không gắn số mục)
        words_vi_clean = words_vi.rstrip(".")
        line_vi = f"{trigger_vi}: Thanh toán {percentage}% giá trị hợp đồng ({amount_display} VND - {words_vi_clean}.){extra_vi}{status_vi}".strip()
        line_zh = f"{trigger_zh}：支付合約金額的 {percentage}% ({amount_display} VND，{words_zh})。{extra_zh}{status_zh}".strip()

        items.append({
            "index": index,
            "vietnamese": line_vi,
            "chinese": line_zh,
            "is_active": is_active,
            "is_paid": is_paid,
            "due_date": due_date if is_active else "",
        })

    return items


def _build_payment_terms(
    milestones: Sequence[Dict[str, Any]],
    active_milestone_name: str = "",
    vietnamese_first: bool = True
) -> List[str]:
    bilingual_data = _build_bilingual_terms_data(milestones, active_milestone_name)
    terms: List[str] = []
    for item in bilingual_data:
        num_prefix = f"{item['index']:02d}. "
        if vietnamese_first:
            first_line = f"{num_prefix}{item['vietnamese']}"
            second_line = f"    {item['chinese']}"
        else:
            first_line = f"{num_prefix}{item['chinese']}"
            second_line = f"    {item['vietnamese']}"
        terms.append(f"{first_line}\n{second_line}")
    return terms


def _build_notes_text(
    milestones: Sequence[Dict[str, Any]],
    active_milestone_name: str = "",
    vietnamese_first: bool = True
) -> str:
    lines: List[str] = ["備註 :"]
    for term in _build_payment_terms(milestones, active_milestone_name=active_milestone_name, vietnamese_first=vietnamese_first):
        first_line, second_line = term.split("\n", 1)
        lines.extend((first_line, second_line))
    lines.extend(
        [
            "",
            "**以上價格不包含VAT增值稅，稅率依發票開立時的稅率為準。Giá nêu trên chưa bao gồm thuế GTGT (VAT). Thuế suất áp dụng theo mức thuế suất tại thời điểm xuất hóa đơn.",
        ]
    )
    return "\n".join(lines)


def _index_document(
    cursor,
    *,
    contract_id: Optional[int],
    contract_code: str,
    file_path: Path,
    milestone_name: str,
) -> int:
    file_name = file_path.name
    file_format = file_path.suffix.lower().lstrip(".")
    file_size = file_path.stat().st_size
    cursor.execute("SELECT id FROM documents WHERE file_name = ?", (file_name,))
    doc_row = cursor.fetchone()
    description = f"Phiếu yêu cầu thanh toán {milestone_name} ({contract_code}) - {file_format.upper()}"
    if doc_row:
        doc_id = doc_row["id"]
        cursor.execute(
            """
            UPDATE documents
            SET file_path = ?, file_size = ?, contract_id = ?, contract_code = ?,
                file_format = ?, description = ?
            WHERE id = ?
            """,
            (str(file_path), file_size, contract_id, contract_code, file_format, description, doc_id),
        )
        return int(doc_id)
    cursor.execute(
        """
        INSERT INTO documents (
            contract_id, contract_code, doc_type, doc_name, file_name,
            file_path, file_size, file_format, description
        ) VALUES (?, ?, 'YEU_CAU_THANH_TOAN', ?, ?, ?, ?, ?, ?)
        """,
        (
            contract_id,
            contract_code,
            file_path.stem,
            file_name,
            str(file_path),
            file_size,
            file_format,
            description,
        ),
    )
    return int(cursor.lastrowid)


def generate_payment_request_excel(
    contract_code: str,
    milestone_name: str = "Đợt 2",
    customer_name_vi: Optional[str] = None,
    customer_name_zh: Optional[str] = None,
    project_name_vi: Optional[str] = None,
    project_name_zh: Optional[str] = None,
    amount: Optional[float] = None,
    total_contract_amount: Optional[float] = None,
    percentage: Optional[float] = None,
    due_date: Optional[str] = None,
    invoice_no: Optional[str] = None,
    application_date: Optional[str] = None,
    template_source: str = "excel",
) -> Dict[str, Any]:
    """Tạo XLSX + DOCX, cho phép chọn Excel 712TR hoặc DOCX làm nguồn mẫu chính."""
    selected_template_source = _normalize_template_source(template_source)
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Không tìm thấy file mẫu gốc tại: {TEMPLATE_PATH}")
    if selected_template_source == "word" and not WORD_TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            "Không tìm thấy file mẫu Word. Vui lòng đặt mẫu tại đúng đường dẫn: "
            f"{WORD_TEMPLATE_PATH}"
        )

    now = datetime.now()
    app_date = application_date or now.strftime("%d/%m/%Y")

    # 1. Pull fallback info from SQLite contracts and payment_milestones
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM contracts WHERE contract_code = ?", (contract_code,))
    c_row = cursor.fetchone()

    ms_row = None
    all_milestones: List[Dict[str, Any]] = []
    if c_row:
        cursor.execute("""
            SELECT * FROM payment_milestones 
            WHERE contract_code = ? AND (milestone_name = ? OR milestone_name LIKE ?)
            ORDER BY id ASC LIMIT 1
        """, (contract_code, milestone_name, f"%{milestone_name}%"))
        ms_row = cursor.fetchone()

        if not ms_row:
            cursor.execute("""
                SELECT * FROM payment_milestones 
                WHERE contract_code = ? AND status != 'PAID'
                ORDER BY id ASC LIMIT 1
            """, (contract_code,))
            ms_row = cursor.fetchone()

        cursor.execute(
            "SELECT * FROM payment_milestones WHERE contract_code = ? ORDER BY id ASC",
            (contract_code,),
        )
        all_milestones = [dict(row) for row in cursor.fetchall()]

    cust_vi = customer_name_vi or (c_row["customer_name_vi"] if c_row else "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)")
    cust_zh = customer_name_zh or (c_row["customer_name_zh"] if c_row else "越南CX科技股份有限公司")
    proj_vi = project_name_vi or (c_row["project_name_vi"] if c_row else "Dự án RFID Hardware & Custom Software ERP")
    proj_zh = project_name_zh or (c_row["project_name_zh"] if c_row else "RFID硬體（RFID Hardware）－RFID系統客製化軟體功能上線移交使用")
    c_code = contract_code or "HMV151020250001"
    
    u_price = float(_coalesce(total_contract_amount, c_row["total_amount"] if c_row else 234000000.0))
    pct = float(_coalesce(percentage, ms_row["percentage"] if ms_row else 40.0))
    ms_amt = float(_coalesce(amount, ms_row["amount"] if ms_row else (u_price * pct / 100.0)))
    ms_dict = dict(ms_row) if ms_row else {}
    c_dict = dict(c_row) if c_row else {}

    # Ngày nộp đơn (申請日期) luôn là ngày hiện tại (hoặc ngày chỉ định)
    app_dt = now
    if app_date:
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                app_dt = datetime.strptime(str(app_date).strip(), fmt)
                break
            except ValueError:
                continue

    # Xác định thời hạn thanh toán T/T số ngày từ hợp đồng
    m_term = int(ms_dict.get("payment_term_days") or 0)
    cond_text = f"{ms_dict.get('milestone_name', '')} {ms_dict.get('trigger_condition', '')} {milestone_name}".lower()
    is_immediate = (
        "ký" in cond_text
        or "tạm ứng" in cond_text
        or "đặt cọc" in cond_text
        or "ngay" in cond_text
        or ms_dict.get("dot") == 1
        or str(ms_dict.get("milestone_name") or milestone_name).strip() == "Đợt 1"
    )

    if m_term <= 0:
        if is_immediate:
            m_term = 10  # Thanh toán ngay / tạm ứng khi ký hợp đồng: mặc định 10 ngày
        else:
            m_term = 30  # Các đợt tiếp theo: mặc định 30 ngày nếu không có thỏa thuận khác

    # Ngày đến hạn (應付款日期): Nếu caller truyền do người dùng chỉ định thì dùng,
    # ngược lại TÍNH TỪ NGÀY NỘP ĐƠN (app_dt) + số ngày thời hạn thanh toán T/T (m_term)
    if due_date and str(due_date).strip():
        due_dt_str = str(due_date).strip()
    else:
        due_dt_str = (app_dt + timedelta(days=m_term)).strftime("%d/%m/%Y")

    inv_num = str(_coalesce(invoice_no, ms_dict.get("invoice_no") or "Chưa xuất"))
    ms_title = ms_dict.get("milestone_name") or milestone_name

    # Cập nhật lại due_date vào database cho đợt thanh toán này
    if ms_dict.get("id"):
        try:
            cursor.execute(
                """
                UPDATE payment_milestones 
                SET due_date = ?, 
                    payment_term_days = ?,
                    status = CASE WHEN status = 'PAID' THEN 'PAID' ELSE 'URGENT' END,
                    status_label = CASE WHEN status = 'PAID' THEN status_label ELSE 'Đến hạn thanh toán' END
                WHERE id = ?
                """,
                (due_dt_str, m_term, ms_dict["id"])
            )
        except Exception:
            try:
                cursor.execute(
                    """
                    UPDATE payment_milestones 
                    SET due_date = ?, 
                        status = CASE WHEN status = 'PAID' THEN 'PAID' ELSE 'URGENT' END,
                        status_label = CASE WHEN status = 'PAID' THEN status_label ELSE 'Đến hạn thanh toán' END
                    WHERE id = ?
                    """,
                    (due_dt_str, ms_dict["id"])
                )
            except Exception:
                pass

    if not all_milestones:
        all_milestones = [
            {
                "milestone_name": ms_title,
                "percentage": pct,
                "amount": ms_amt,
                "trigger_condition": ms_title,
                "due_date": due_dt_str,
                "invoice_no": inv_num,
                "status": "URGENT",
            }
        ]
    else:
        for item in all_milestones:
            if str(item.get("milestone_name")) == str(ms_title) or (ms_dict.get("id") and item.get("id") == ms_dict.get("id")):
                item.update(
                    {
                        "percentage": pct,
                        "amount": ms_amt,
                        "due_date": due_dt_str,
                        "invoice_no": inv_num,
                    }
                )

    bilingual_terms_data = _build_bilingual_terms_data(all_milestones, active_milestone_name=ms_title)
    payment_terms = _build_payment_terms(all_milestones, active_milestone_name=ms_title, vietnamese_first=True)
    notes_text = _build_notes_text(all_milestones, active_milestone_name=ms_title, vietnamese_first=True)

    wb = openpyxl.load_workbook(str(TEMPLATE_PATH))

    if TEMPLATE_SHEET_NAME not in wb.sheetnames:
        conn.close()
        raise KeyError(f"File mẫu không có sheet bắt buộc '{TEMPLATE_SHEET_NAME}'.")
    ws = wb[TEMPLATE_SHEET_NAME]

    # Form code
    form_no = f"HMV{now.strftime('%d%m%Y')}{c_code[-4:]}-SYTT"
    ws["H9"] = form_no

    # Customer and application date
    ws["C10"] = cust_vi
    ws["I10"] = _parse_application_date(app_date)

    # Project basis
    project_text = (
        f"Căn cứ hợp đồng số : {c_code}\n"
        f"{proj_vi} - {ms_title}\n"
        f"根据合同编号：{c_code}\n"
        f"{proj_zh}"
    )
    ws["C11"] = project_text

    # Due date and method
    ws["C12"] = due_dt_str
    ws["H12"] = "匯款\nChuyển khoản"

    # Line Item (Row 15)
    ws["A15"] = 1
    pct_display = _format_percentage(pct)
    item_text = f"請款 / Thanh toán {ms_title}: {pct_display}%\n合約 / Hợp đồng：{c_code}"
    if inv_num and inv_num != "Chưa xuất":
        item_text += f"\n發票 / Hóa đơn：{inv_num}"
    ws["B15"] = item_text
    ws["C15"] = "VND"
    ws["D15"] = int(u_price)
    ws["E15"] = 1
    ws["F15"] = "=D15*E15"
    ws["G15"] = 0
    ws["H15"] = "=F15*G15"
    ws["I15"] = f"=(F15+H15)*{pct_display}%"

    # Ghi chú lấy từ toàn bộ tiến độ thanh toán của hợp đồng.
    amt_words_vi = number_to_vietnamese_words(int(ms_amt))
    amt_words_zh = number_to_chinese_words(int(ms_amt))
    ws["B23"] = notes_text

    # Chỉ giữ đúng sheet mẫu 712TR; không thay đổi bố cục, style hoặc page setup.
    for s in list(wb.sheetnames):
        if s != TEMPLATE_SHEET_NAME:
            del wb[s]
    wb.active = 0
    if hasattr(wb, "calculation"):
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True

    safe_customer = _safe_file_component(cust_vi)
    base_file_name = f"đơn yêu cầu thanh toán - {safe_customer} - {c_code}"
    file_name = f"{base_file_name}.xlsx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_file = OUTPUT_DIR / file_name

    try:
        wb.save(str(out_file))
    except PermissionError:
        file_name = f"đơn yêu cầu thanh toán - {cust_vi} - {c_code}_v2.xlsx"
        out_file = OUTPUT_DIR / file_name
        wb.save(str(out_file))

    # Giữ bản tham chiếu nhanh tại thư mục làm việc.
    root_copy = WORKING_DIR / file_name
    try:
        wb.save(str(root_copy))
    except Exception:
        pass

    file_size = out_file.stat().st_size
    contract_id = c_row["id"] if c_row else None
    doc_id = _index_document(
        cursor,
        contract_id=contract_id,
        contract_code=c_code,
        file_path=out_file,
        milestone_name=ms_title,
    )

    sheet_data = {
        "form_no": form_no,
        "application_date": app_date,
        "contract_code": c_code,
        "customer_name_vi": cust_vi,
        "customer_name_zh": cust_zh,
        "project_name_vi": proj_vi,
        "project_name_zh": proj_zh,
        "milestone_name": ms_title,
        "percentage": pct,
        "percentage_display": pct_display,
        "total_contract_amount": u_price,
        "total_contract_amount_formatted": format_amount_display(int(u_price)),
        "milestone_amount": ms_amt,
        "milestone_amount_formatted": format_amount_display(int(ms_amt)),
        "amount_words_vi": amt_words_vi,
        "amount_words_zh": amt_words_zh,
        "due_date": due_dt_str,
        "invoice_no": inv_num,
        "notes_text": notes_text,
        "payment_terms": payment_terms,
        "payment_terms_bilingual": bilingual_terms_data,
        "milestones": all_milestones,
        "file_name": file_name,
        "download_url": f"/api/v1/dossier/documents/{doc_id}/download",
        "template_source": selected_template_source,
        "template_file_name": (
            WORD_TEMPLATE_PATH.name if selected_template_source == "word" else TEMPLATE_PATH.name
        ),
    }

    selected_word_template = WORD_TEMPLATE_PATH if selected_template_source == "word" else None
    word_file = generate_payment_request_word(
        sheet_data,
        OUTPUT_DIR / f"{base_file_name}.docx",
        template_path=selected_word_template,
    )
    try:
        generate_payment_request_word(
            sheet_data,
            WORKING_DIR / word_file.name,
            template_path=selected_word_template,
        )
    except Exception:
        pass
    word_doc_id = _index_document(
        cursor,
        contract_id=contract_id,
        contract_code=c_code,
        file_path=word_file,
        milestone_name=ms_title,
    )
    word_download_url = f"/api/v1/dossier/documents/{word_doc_id}/download"
    sheet_data.update(
        {
            "word_file_name": word_file.name,
            "word_file_path": str(word_file),
            "word_download_url": word_download_url,
        }
    )

    conn.commit()
    conn.close()

    sheet_html = render_payment_request_sheet_html(sheet_data)

    primary_is_word = selected_template_source == "word"
    primary_file = word_file if primary_is_word else out_file
    primary_doc_id = word_doc_id if primary_is_word else doc_id
    primary_download_url = word_download_url if primary_is_word else f"/api/v1/dossier/documents/{doc_id}/download"

    return {
        "status": "success",
        "file_name": file_name,
        "file_path": str(out_file),
        "file_size": file_size,
        "doc_id": doc_id,
        "download_url": f"/api/v1/dossier/documents/{doc_id}/download",
        "word_file_name": word_file.name,
        "word_file_path": str(word_file),
        "word_file_size": word_file.stat().st_size,
        "word_doc_id": word_doc_id,
        "word_download_url": word_download_url,
        "template_source": selected_template_source,
        "template_file_name": sheet_data["template_file_name"],
        "primary_file_name": primary_file.name,
        "primary_file_path": str(primary_file),
        "primary_doc_id": primary_doc_id,
        "primary_download_url": primary_download_url,
        "contract_code": c_code,
        "milestone_name": ms_title,
        "sheet_data": sheet_data,
        "sheet_html": sheet_html,
    }


def render_payment_request_sheet_html(data: Dict[str, Any]) -> str:
    """Render high-fidelity HTML table mimicking Excel spreadsheet 1:1 exactly from 'đơn yêu cầu NCC thanh toán-Gốc.xlsx'."""
    cust_vi = html.escape(str(data.get("customer_name_vi") or ""))
    c_code = html.escape(str(data.get("contract_code") or ""))
    proj_vi = html.escape(str(data.get("project_name_vi") or ""))
    proj_zh = html.escape(str(data.get("project_name_zh") or ""))
    ms_name = html.escape(str(data.get("milestone_name") or "Đợt 2"))
    form_no = html.escape(str(data.get("form_no") or ""))
    app_date = html.escape(str(data.get("application_date") or ""))
    due_date = html.escape(str(data.get("due_date") or ""))
    inv_no = html.escape(str(data.get("invoice_no") or ""))
    pct = data.get("percentage", 40)
    amt_formatted = html.escape(str(data.get("milestone_amount_formatted") or "0"))
    total_contract = format_amount_display(int(data.get("total_contract_amount", 0)))
    words_vi = html.escape(str(data.get("amount_words_vi") or ""))
    words_zh = html.escape(str(data.get("amount_words_zh") or ""))
    file_name = html.escape(str(data.get("file_name") or ""))
    download_url = html.escape(str(data.get("download_url") or "#"))
    word_file_name = html.escape(str(data.get("word_file_name") or ""))
    word_download_url = html.escape(str(data.get("word_download_url") or "#"))
    template_source = str(data.get("template_source") or "excel")
    template_file_name = html.escape(str(data.get("template_file_name") or TEMPLATE_PATH.name))
    source_icon = "W" if template_source == "word" else "X"
    source_color = "#1d4ed8" if template_source == "word" else "#166534"
    source_message = (
        f"✓ Word được điền trực tiếp trên mẫu {template_file_name}; kèm bản Excel 712TR"
        if template_source == "word"
        else "✓ Excel dùng đúng sheet 712TR; kèm bản Word song ngữ cùng nội dung"
    )
    notes_html = html.escape(str(data.get("notes_text") or "")).replace("\n", "<br>")
    pct_display = html.escape(str(data.get("percentage_display") or pct))

    inv_display = f"<br>發票：{inv_no}" if inv_no and inv_no != "Chưa xuất" else ""

    return f"""
    <div class="excel-sheet-container" style="background: #ffffff; border: 1.5px solid #047857; border-radius: 8px; box-shadow: 0 8px 30px rgba(0,0,0,0.08); padding: 1.5rem; font-family: 微軟正黑體, 'Microsoft JhengHei', 'Microsoft YaHei', Arial, sans-serif; color: #000000; overflow-x: auto;">
      
      <!-- Excel Action Header Bar -->
      <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 1.5rem; padding-bottom: 0.85rem; border-bottom: 1px solid #e2e8f0; flex-wrap: wrap; gap: 0.75rem;">
        <div style="display: flex; align-items: center; gap: 0.6rem;">
          <span style="display: inline-flex; align-items: center; justify-content: center; width: 28px; height: 28px; background: {source_color}; color: #ffffff; border-radius: 6px; font-weight: 800; font-size: 0.85rem;">{source_icon}</span>
          <div>
            <strong style="color: #0f172a; font-size: 0.98rem;">{word_file_name if template_source == "word" else file_name}</strong>
            <div style="font-size: 0.75rem; color: {source_color}; font-weight: 600;">{source_message}</div>
          </div>
        </div>

        <div style="display: flex; gap: 0.5rem;">
          <a href="{download_url}" download class="btn btn-primary btn-sm" style="background: #166534; border-color: #15803d; font-size: 0.82rem; font-weight: 700; display: inline-flex; align-items: center; gap: 0.4rem; padding: 0.4rem 0.85rem; color: #ffffff; border-radius: 6px; text-decoration: none;">
            <svg class="icon" viewBox="0 0 24 24" style="width: 15px; height: 15px;"><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4M7 10l5 5 5-5M12 15V3"/></svg>
            <span>Tải File Excel (.xlsx) Này</span>
          </a>
          <a href="{word_download_url}" download class="btn btn-primary btn-sm" style="background: #1d4ed8; border-color: #1e40af; font-size: 0.82rem; font-weight: 700; display: inline-flex; align-items: center; gap: 0.4rem; padding: 0.4rem 0.85rem; color: #ffffff; border-radius: 6px; text-decoration: none;" title="{word_file_name}">
            <span style="font-weight: 900;">W</span>
            <span>Tải File Word (.docx)</span>
          </a>
          <button type="button" class="btn btn-outline btn-sm" onclick="ReminderEditor.scrollToEmailSection()" style="font-size: 0.82rem; font-weight: 600; display: inline-flex; align-items: center; gap: 0.35rem; padding: 0.4rem 0.85rem; border: 1px solid #2563eb; color: #1d4ed8; border-radius: 6px; background: #eff6ff;">
            <span>Chuyển Sang Bước Gửi Mail 👉</span>
          </button>
        </div>
      </div>

      <!-- Master Spreadsheet Page - 100% Exact 1:1 Replica -->
      <div style="max-width: 900px; margin: 0 auto; background: #ffffff; border: 1px solid #cbd5e1; padding: 35px 40px; box-shadow: 0 4px 16px rgba(0,0,0,0.06);">
        
        <!-- Row 2: Header Company Name Chinese (24pt bold, centered) -->
        <div style="font-family: 微軟正黑體, 'Microsoft JhengHei', sans-serif; font-size: 24px; font-weight: bold; text-align: center; color: #000000; letter-spacing: 2px;">
          浤 詺 資 訊 責 任 有 限 公 司
        </div>

        <!-- Row 3: Header Company Name Vietnamese (16pt bold, centered) -->
        <div style="font-family: Arial, sans-serif; font-size: 16px; font-weight: bold; text-align: center; color: #000000; margin-top: 4px;">
          CÔNG TY TNHH TOÀN CẦU HONG MING
        </div>

        <!-- Row 5: Document Title Chinese (24pt bold, centered) -->
        <div style="font-family: 'Microsoft YaHei', 微軟正黑體, sans-serif; font-size: 24px; font-weight: bold; text-align: center; color: #000000; letter-spacing: 6px; margin-top: 18px;">
          請 款 單
        </div>

        <!-- Row 6: Document Title Vietnamese (18pt bold, centered) -->
        <div style="font-family: Arial, sans-serif; font-size: 18px; font-weight: bold; text-align: center; color: #000000; margin-top: 4px;">
          Phiếu yêu cầu thanh toán
        </div>

        <!-- Row 8: Salute (14pt, left-aligned) -->
        <div style="font-family: 'Microsoft YaHei', 微軟正黑體, sans-serif; font-size: 14px; text-align: left; color: #000000; margin-top: 16px; padding-left: 2px;">
          CÔNG TY TNHH TOÀN CẦU HONG MING 致敬：
        </div>

        <!-- Row 9: Form NO. (Right-aligned under G9:I9) -->
        <div style="text-align: right; margin-top: 4px; font-family: Arial, sans-serif; font-size: 12px; color: #000000;">
          <strong>Form NO.</strong>&nbsp;&nbsp;<span>{form_no}</span>
        </div>

        <!-- Rows 10-12: Metadata (NO BORDERS, exactly matching template rows 10 to 12) -->
        <table style="width: 100%; border-collapse: collapse; margin-top: 8px; margin-bottom: 12px; font-family: 微軟正黑體, 'Microsoft JhengHei', 'Microsoft YaHei', sans-serif;">
          <!-- Row 10: Customer Name & Application Date -->
          <tr>
            <td style="width: 20%; min-width: 145px; text-align: right; font-weight: bold; font-size: 13px; line-height: 1.4; padding: 5px 12px 5px 0; vertical-align: middle; border: none; white-space: nowrap;">
              客 戶 名 稱：<br>Tên khách hàng：
            </td>
            <td style="width: 50%; font-family: 'Times New Roman', serif; font-size: 14px; text-align: left; vertical-align: middle; padding: 5px 4px; border: none; font-weight: normal;">
              {cust_vi}
            </td>
            <td style="width: 15%; min-width: 120px; text-align: right; font-weight: bold; font-size: 13px; line-height: 1.4; padding: 5px 8px 5px 0; vertical-align: middle; border: none; white-space: nowrap;">
              申 請 日 期：<br>Ngày nộp đơn：
            </td>
            <td style="width: 15%; font-size: 13px; text-align: center; vertical-align: middle; padding: 5px 4px; border: none; white-space: nowrap;">
              {app_date}
            </td>
          </tr>

          <!-- Row 11: Project Basis / Contract Basis -->
          <tr>
            <td style="text-align: right; font-weight: bold; font-size: 14px; line-height: 1.4; padding: 6px 12px 6px 0; vertical-align: top; border: none; white-space: nowrap;">
              專 案 名 稱：<br>Tên dự án：
            </td>
            <td colspan="3" style="font-family: 'Times New Roman', serif; font-size: 13px; line-height: 1.5; text-align: left; padding: 6px 4px; border: none;">
              Căn cứ hợp đồng số : {c_code}<br>
              {proj_vi} - {ms_name}<br>
              根据合同编号：{c_code}<br>
              {proj_zh}
            </td>
          </tr>

          <!-- Row 12: Due Date & Payment Method -->
          <tr>
            <td style="text-align: right; font-weight: bold; font-size: 14px; line-height: 1.4; padding: 5px 12px 5px 0; vertical-align: middle; border: none; white-space: nowrap;">
              應 付 款 日 期：<br>Ngày đến hạn thanh toán:
            </td>
            <td style="font-family: Calibri, Arial, sans-serif; font-size: 15px; text-align: left; vertical-align: middle; padding: 5px 4px; border: none; white-space: nowrap;">
              {due_date}
            </td>
            <td style="text-align: right; font-weight: bold; font-size: 14px; line-height: 1.4; padding: 5px 8px 5px 0; vertical-align: middle; border: none; white-space: nowrap;">
              付 款 方 式：<br>Phương thức thanh toán：
            </td>
            <td style="font-size: 13px; text-align: left; vertical-align: middle; padding: 5px 4px; border: none; line-height: 1.3; white-space: nowrap;">
              匯款<br>Chuyển khoản
            </td>
          </tr>
        </table>

        <!-- Row 14: Main Table Header (Solid Yellow #FFFF00, exact labels from template) -->
        <table style="width: 100%; border-collapse: collapse; border: 1.5px solid #000000; font-size: 11.5px; margin-top: 6px; font-family: 微軟正黑體, 'Microsoft JhengHei', sans-serif;">
          <thead>
            <tr style="background-color: #ffff00; color: #000000; font-weight: bold; text-align: center; height: 38px;">
              <th style="border: 1px solid #000000; width: 4%; padding: 4px 1px; white-space: nowrap;">項<br>STT</th>
              <th style="border: 1px solid #000000; width: 23%; padding: 4px 6px;">請款項目<br>Hạng mục thanh toán</th>
              <th style="border: 1px solid #000000; width: 6%; padding: 4px 2px; white-space: nowrap;">幣別<br>Tiền tệ</th>
              <th style="border: 1px solid #000000; width: 13%; padding: 4px 4px; white-space: nowrap;">單價<br>Đơn giá</th>
              <th style="border: 1px solid #000000; width: 7.5%; padding: 4px 2px; white-space: nowrap;">數量<br>Số lượng</th>
              <th style="border: 1px solid #000000; width: 14%; padding: 4px 4px; white-space: nowrap;">合同金額<br>Số tiền hợp đồng</th>
              <th style="border: 1px solid #000000; width: 8%; padding: 4px 2px; white-space: nowrap;">稅率<br>Thuế suất</th>
              <th style="border: 1px solid #000000; width: 9.5%; padding: 4px 3px; white-space: nowrap;">稅額 VAT</th>
              <th style="border: 1px solid #000000; width: 15%; padding: 4px 4px; white-space: nowrap;">請款金額<br>Số tiền xin TT</th>
            </tr>
          </thead>
          <tbody>
            <!-- Row 15: Primary Active Milestone Line Item -->
            <tr style="height: 64px;">
              <td style="border: 1px solid #000000; text-align: center; font-size: 12px;">1</td>
              <td style="border: 1px solid #000000; padding: 6px 8px; text-align: left; line-height: 1.45; font-size: 12px;">
                請款 / Thanh toán {ms_name}: {pct_display}%<br>
                合約：{c_code}{inv_display}
              </td>
              <td style="border: 1px solid #000000; text-align: center; font-size: 12px;">VND</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px; font-size: 12px;">{total_contract}</td>
              <td style="border: 1px solid #000000; text-align: center; font-size: 12px;">1</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px; font-size: 12px;">{total_contract}</td>
              <td style="border: 1px solid #000000; text-align: center; font-size: 12px;">0%</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px; font-size: 12px;">0</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px; font-weight: bold; font-size: 13px;">{amt_formatted}</td>
            </tr>

            <!-- Rows 16 to 20: Template Empty Placeholder Rows -->
            <tr style="height: 28px; color: #475569;">
              <td style="border: 1px solid #000000; text-align: center;">2</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
            </tr>
            <tr style="height: 28px; color: #475569;">
              <td style="border: 1px solid #000000; text-align: center;">3</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
            </tr>
            <tr style="height: 28px; color: #475569;">
              <td style="border: 1px solid #000000; text-align: center;">4</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
            </tr>
            <tr style="height: 28px; color: #475569;">
              <td style="border: 1px solid #000000; text-align: center;">5</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
            </tr>
            <tr style="height: 28px; color: #475569;">
              <td style="border: 1px solid #000000; text-align: center;">6</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000;"></td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
              <td style="border: 1px solid #000000; text-align: right; padding-right: 6px;">0</td>
            </tr>

            <!-- Row 21: Total (合計 Tổng cộng) matching template -->
            <tr style="height: 34px;">
              <td colspan="5" style="border: 1.5px solid #000000; text-align: center; font-weight: bold; font-size: 16px;">
                合計 Tổng cộng
              </td>
              <td colspan="4" style="border: 1.5px solid #000000; text-align: center; font-weight: bold; font-size: 18px; color: #000000;">
                {amt_formatted}
              </td>
            </tr>
          </tbody>
        </table>

        <!-- Row 23: Remarks (備註) Merged Box matching template 1:1 -->
        <div style="border: 1.5px solid #000000; padding: 12px 14px; margin-top: 14px; margin-bottom: 14px; font-size: 11px; line-height: 1.6; text-align: left; background: #ffffff;">
          <div>{notes_html}</div>
        </div>

        <!-- Rows 25-26: Bank Information Table matching template 1:1 -->
        <table style="width: 100%; border-collapse: collapse; border: 1.5px solid #000000; font-size: 12px; margin-bottom: 22px; font-family: 微軟正黑體, 'Microsoft JhengHei', sans-serif;">
          <!-- Row 25 -->
          <tr style="height: 30px;">
            <td style="border: 1px solid #000000; width: 18%; text-align: center; font-weight: bold;">銀行帳戶名稱</td>
            <td style="border: 1px solid #000000; width: 44%; padding: 4px 10px; font-weight: normal;">CÔNG TY TNHH TOÀN CẦU HONG MING</td>
            <td style="border: 1px solid #000000; width: 15%; text-align: center; font-weight: bold;">銀行帳號</td>
            <td style="border: 1px solid #000000; width: 23%; text-align: center; padding: 4px 6px;">604080006546(VND)</td>
          </tr>
          <!-- Row 26 -->
          <tr style="height: 30px;">
            <td style="border: 1px solid #000000; text-align: center; font-weight: bold;">銀行名稱</td>
            <td colspan="3" style="border: 1px solid #000000; padding: 4px 10px; font-weight: normal;">CATHAY UNITED BANK – HO CHI MINH CITY BRANCH</td>
          </tr>
        </table>

        <!-- Row 28: Signatures (NO BORDER, matching template 1:1) -->
        <table style="width: 100%; border-collapse: collapse; margin-top: 25px; margin-bottom: 50px; font-family: 微軟正黑體, 'Microsoft JhengHei', sans-serif;">
          <tr>
            <td style="width: 50%; text-align: center; font-size: 14px; font-weight: bold; padding-bottom: 70px; border: none;">
              買方主管
            </td>
            <td style="width: 50%; text-align: center; font-size: 14px; font-weight: bold; padding-bottom: 70px; border: none;">
              申請單位
            </td>
          </tr>
        </table>

      </div>
    </div>
    """
