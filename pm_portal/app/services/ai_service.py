"""AI Service for parsing unstructured business text into contract data via Google Gemini API."""

import json
import logging
import re
import urllib.error
import urllib.request
from typing import Any, Dict, List, Optional

from app.config import CANDIDATE_MODELS, DEFAULT_GEMINI_KEY, DEFAULT_MODEL, DEFAULT_GEMINI_KEYS
from app.services.payment_schedule_service import resolve_payment_schedule
from app.utils.currency_helper import (
    format_currency_vnd,
    number_to_chinese_words,
    number_to_vietnamese_words,
    parse_positive_integer,
)

logger = logging.getLogger(__name__)

SYSTEM_PROMPT = """Bạn là trợ lý AI chuyên nghiệp về bóc tách và chuẩn hóa dữ liệu hợp đồng mua bán thiết bị / phần mềm song ngữ Việt - Trung cho CÔNG TY TNHH TOÀN CẦU HONG MING (宏明環球有限公司).

NHIỆM VỤ CỦA BẠN:
Phân tích văn bản, danh sách thiết bị hoặc bảng dữ liệu do người dùng cung cấp và trích xuất thành định dạng JSON chuẩn xác 100%.

QUY TẮC CỐT LÕI:
1. Bóc tách toàn bộ danh mục sản phẩm (bao gồm STT, Tên sản phẩm song ngữ nếu có, Đơn giá, Số lượng, Thành tiền). Giữ nguyên đầy đủ tiếng Việt và chữ Hán Phồn thể nếu có trong từng dòng sản phẩm.
2. Trích xuất thông tin Bên A (Khách hàng / 買方) bao gồm: Tên công ty (Việt & Trung), Địa chỉ (Việt & Trung), Mã số thuế, Người đại diện, Chức vụ (Việt & Trung).
3. Nếu người dùng cung cấp cả tiếng Việt và tiếng Trung thì điền chính xác vào các trường tương ứng. Nếu không có tiếng Trung thì để chuỗi rỗng. Tuyệt đối không tự suy diễn hoặc bịa đặt thông tin.
4. Tự động tính toán hoặc đọc chính xác Tổng tiền hợp đồng, Tổng số lượng.
5. Dự án: Tự động trích xuất hoặc đặt tên dự án phù hợp với danh mục sản phẩm (VD: "Hệ thống quản lý kho WMS & Phần mềm MES Pro").
6. Tiến độ thanh toán: HỖ TRỢ ĐẶC BIỆT CÚ PHÁP VIẾT TẮT SỐ NGÀY:
   - Số `0`: quy ước là "ngay lập tức" (thanh toán ngay sau khi hai bên ký kết hợp đồng / 雙方簽約後即期以電匯（T/T）方式支付, "thoi_han_tt_ngay": 0).
   - Số nguyên dương N (ví dụ `30`, `60`, `90`): quy ước là thời hạn thanh toán N ngày ("thoi_han_tt_ngay": N, điều kiện tiếng Việt có "thời hạn thanh toán T/T N ngày", tiếng Trung có "付款期限為 T/T N 天").
   - VÍ DỤ CÚ PHÁP VIẾT TẮT: Khi người dùng nhập "35% 0, 55% 60, 10% 60" hoặc "35% 0 55% 60 10% 60", AI phải hiểu ngay và bóc tách thành 3 đợt:
     + Đợt 1 (35%): thoi_han_tt_ngay: 0, tieu_de: "Tạm ứng ký hợp đồng", tieu_de_trung: "簽約定金", dieu_kien_viet: "Thanh toán tạm ứng 35% ngay sau khi hai bên ký kết hợp đồng ngay lập tức", dieu_kien_trung: "雙方簽約後即期以電匯（T/T）方式支付"
     + Đợt 2 (55%): thoi_han_tt_ngay: 60, tieu_de: "Lắp đặt thiết bị", tieu_de_trung: "設備安裝款", dieu_kien_viet: "Sau khi bàn giao thiết bị và hoàn thành cài đặt hệ thống, thời hạn thanh toán T/T 60 ngày", dieu_kien_trung: "於設備安裝及系統建置完成並經雙方簽署確認文件後，付款期限為 T/T 60 天"
     + Đợt 3 (10%): thoi_han_tt_ngay: 60, tieu_de: "Nghiệm thu cuối cùng (FAC)", tieu_de_trung: "驗收尾款", dieu_kien_viet: "Sau khi khách hàng nghiệm thu hoàn tất và ký biên bản bàn giao (FAC), thời hạn thanh toán T/T 60 ngày", dieu_kien_trung: "於最終驗收（FAC）完成並簽署驗收交付文件後，付款期限為 T/T 60 天"
   - Giữ nguyên thứ tự và số đợt (1 đến 4 đợt). Với mỗi đợt, luôn trả "thoi_han_tt_ngay" (số nguyên hoặc null).
   - Nếu đầu vào không nêu tiến độ thì trả về danh sách rỗng; hệ thống sẽ đối chiếu mẫu 備註 trong workbook đã phê duyệt. Tuyệt đối không tự đặt tỷ lệ mặc định.

ĐỊNH DẠNG JSON ĐẦU RA BẮT BUỘC:
{
  "so_hop_dong": "HD-20261013-001",
  "ngay_ky": "13",
  "thang_ky": "10",
  "nam_ky": "2026",
  "ben_a": {
    "ten_cong_ty": "Tên công ty tiếng Việt",
    "ten_tieng_trung": "Tên công ty tiếng Trung nếu có",
    "dia_chi": "Địa chỉ tiếng Việt",
    "dia_chi_trung": "Địa chỉ tiếng Trung nếu có",
    "mst": "Mã số thuế",
    "nguoi_dai_dien": "Họ và tên người đại diện",
    "chuc_vu": "Chức vụ tiếng Việt",
    "chuc_vu_trung": "Chức vụ tiếng Trung"
  },
  "ben_b": {
    "ten_cong_ty": "CÔNG TY TNHH TOÀN CẦU HONG MING",
    "ten_tieng_trung": "宏明環球有限公司",
    "dia_chi": "Lầu 2, số 68 Nguyễn Huệ, Phường Sài Gòn, Thành phố Hồ Chí Minh, Việt Nam",
    "dia_chi_trung": "越南胡志明市西貢坊阮惠路68號2樓",
    "mst": "0318280158",
    "nguoi_dai_dien": "TANG WEI LI",
    "chuc_vu": "TỔNG GIÁM ĐỐC",
    "chuc_vu_trung": "總經理"
  },
  "du_an": {
    "ten_du_an_viet": "Tên dự án tiếng Việt",
    "ten_du_an_trung": "Tên dự án/báo giá bằng tiếng Trung (Phồn thể, BẮT BUỘC có - nếu bản gốc chỉ có tiếng Việt thì dịch chính xác sang tiếng Trung, ví dụ: 'Hệ thống MES Lite' -> 'MES Lite 系統' hoặc 'MES Lite 系統導入專案報價單')"
  },
  "hang_muc": [
    {
      "stt": 1,
      "ten_san_pham": "Tên sản phẩm song ngữ",
      "don_gia": "185.000.000 VND",
      "so_luong": "1 SET",
      "tong_gia": "185.000.000 VND"
    }
  ],
  "tong_so_luong": "36 SET",
  "tong_tien_truoc_giam": "1.382.100.000 VND",
  "tong_tien_sau_giam": "1.382.100.000 VND",
  "tong_tien_hop_dong_so": "1.382.100.000",
  "phan_bo_thanh_toan": [],
  "thoi_gian_giao_hang_ngay": 80,
  "thoi_gian_bao_hanh_thang": 12
}

CHỈ TRẢ VỀ DUY NHẤT KHỐI JSON HỢP LỆ (BẮT ĐẦU BẰNG { VÀ KẾT THÚC BẰNG }), KHÔNG VIẾT BẤT KỲ LỜI MỞ ĐẦU HAY GIẢI THÍCH NÀO NGOÀI JSON.
"""


import base64
import io

def extract_text_from_excel(file_bytes: bytes) -> str:
    """Extract rows and tables from Excel file bytes (.xlsx, .xlsm, .xltx)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        chunks = []
        for name in wb.sheetnames:
            ws = wb[name]
            chunks.append(f"=== BẢNG TÍNH / SHEET: {name} ===")
            for row in ws.iter_rows(values_only=True):
                if any(row):
                    line = " | ".join(str(c).strip() if c is not None else "" for c in row)
                    chunks.append(line)
        return "\n".join(chunks)
    except Exception as exc:
        logger.warning("Error reading Excel bytes with openpyxl: %s", exc)
        return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text and tables from Word DOCX file bytes."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        chunks = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                chunks.append(txt)
        for table in doc.tables:
            for row in table.rows:
                row_vals = [cell.text.strip() for cell in row.cells]
                if any(row_vals):
                    chunks.append(" | ".join(row_vals))
        return "\n".join(chunks)
    except Exception as exc:
        logger.warning("Error reading DOCX bytes: %s", exc)
        return ""


def parse_contract_text(
    raw_input_text: str = "",
    file_bytes: bytes = None,
    file_mime_type: str = None,
    file_name: str = None,
    api_key: str = None,
    preferred_model: str = None
) -> Dict[str, Any]:
    """Parse unstructured contract text or multimodal files (Image, Excel, Word) using Google Gemini API."""
    keys_to_try = [api_key] if api_key else ([DEFAULT_GEMINI_KEY] + [k for k in DEFAULT_GEMINI_KEYS if k != DEFAULT_GEMINI_KEY])

    models_to_try = []
    if preferred_model and preferred_model in CANDIDATE_MODELS:
        models_to_try.append(preferred_model)
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    schedule_source_text = raw_input_text or ""
    fname_lower = (file_name or "").lower()
    fmime_lower = (file_mime_type or "").lower()

    is_image = (
        ("image" in fmime_lower)
        or any(fname_lower.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp"])
    )
    is_excel = (
        ("spreadsheet" in fmime_lower)
        or any(fname_lower.endswith(ext) for ext in [".xlsx", ".xls", ".xlsm", ".xltx"])
    )
    is_docx = fname_lower.endswith(".docx")

    if file_bytes and is_image:
        prompt_text = (
            f"{SYSTEM_PROMPT}\n\n"
            f"NHIỆM VỤ: Hãy phân tích kỹ bức ảnh đính kèm (ảnh chụp bảng báo giá, danh mục thiết bị, đơn hàng PO hoặc chứng từ).\n"
            f"Trích xuất chính xác 100% thông tin pháp nhân Bên A, danh mục thiết bị, số lượng, đơn giá, tổng tiền hợp đồng và quy đổi theo đúng cấu trúc JSON yêu cầu."
        )
        if raw_input_text and raw_input_text.strip():
            prompt_text += f"\n\nCHÚ THÍCH THÊM CỦA NGƯỜI DÙNG KÈM THEO:\n\"\"\"\n{raw_input_text.strip()}\n\"\"\""

        mime_to_send = "image/jpeg"
        if "png" in fname_lower or "png" in fmime_lower:
            mime_to_send = "image/png"
        elif "webp" in fname_lower or "webp" in fmime_lower:
            mime_to_send = "image/webp"

        parts = [
            {"text": prompt_text},
            {
                "inlineData": {
                    "mimeType": mime_to_send,
                    "data": base64.b64encode(file_bytes).decode("ascii")
                }
            }
        ]
    elif file_bytes and is_excel:
        excel_text = extract_text_from_excel(file_bytes)
        schedule_source_text = f"{schedule_source_text}\n{excel_text}"
        prompt_text = (
            f"{SYSTEM_PROMPT}\n\n"
            f"DỮ LIỆU ĐƯỢC ĐỌC TỪ TỆP BẢNG TÍNH EXCEL ĐÍNH KÈM ({file_name or 'Bao_Gia.xlsx'}):\n\"\"\"\n{excel_text}\n\"\"\""
        )
        if raw_input_text and raw_input_text.strip():
            prompt_text += f"\n\nCHÚ THÍCH THÊM CỦA NGƯỜI DÙNG KÈM THEO:\n\"\"\"\n{raw_input_text.strip()}\n\"\"\""
        parts = [{"text": prompt_text}]
    elif file_bytes and is_docx:
        docx_text = extract_text_from_docx(file_bytes)
        schedule_source_text = f"{schedule_source_text}\n{docx_text}"
        prompt_text = (
            f"{SYSTEM_PROMPT}\n\n"
            f"DỮ LIỆU ĐƯỢC ĐỌC TỪ TỆP WORD DOCX ĐÍNH KÈM ({file_name or 'Hop_Dong.docx'}):\n\"\"\"\n{docx_text}\n\"\"\""
        )
        if raw_input_text and raw_input_text.strip():
            prompt_text += f"\n\nCHÚ THÍCH THÊM CỦA NGƯỜI DÙNG KÈM THEO:\n\"\"\"\n{raw_input_text.strip()}\n\"\"\""
        parts = [{"text": prompt_text}]
    else:
        prompt = f"{SYSTEM_PROMPT}\n\nDỮ LIỆU ĐẦU VÀO CẦN BÓC TÁCH:\n\"\"\"\n{raw_input_text}\n\"\"\""
        parts = [{"text": prompt}]

    payload = {
        "contents": [{"parts": parts}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    last_error = None
    executed_model = None
    parsed = None

    for key in keys_to_try:
        for model in models_to_try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={key}"
            logger.info("Attempting Gemini model: %s (key %s...)", model, key[:8])
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"}
            )
            # Set 25s timeout for AI generation
            timeout_sec = 25
            try:
                with urllib.request.urlopen(req, timeout=timeout_sec) as response:
                    result = json.loads(response.read().decode("utf-8"))
                    candidate = result.get("candidates", [])[0]
                    text_out = candidate["content"]["parts"][0]["text"].strip()

                    # Clean markdown fences if any
                    if text_out.startswith("```json"):
                        text_out = text_out[7:]
                    elif text_out.startswith("```"):
                        text_out = text_out[3:]
                    if text_out.endswith("```"):
                        text_out = text_out[:-3]
                    text_out = text_out.strip()

                    parsed = json.loads(text_out)
                    executed_model = model
                    break
            except urllib.error.HTTPError as exc:
                if exc.code == 403:
                    logger.warning("Key %s... returned 403. Rotating to next key.", key[:8])
                    break
                elif exc.code in (404, 429, 503):
                    logger.warning("Model %s returned HTTP %d on key %s... Trying next model.", model, exc.code, key[:8])
                    last_error = exc
                    continue
                else:
                    last_error = exc
            except Exception as exc:
                logger.warning("Gemini network error with model %s (%s). Transitioning...", model, exc)
                last_error = exc
        if parsed is not None:
            break

    if parsed is None:
        raise RuntimeError(f"Không thể kết nối với Gemini AI ({last_error})")

    # Xử lý hậu kỳ AN TOÀN bên ngoài vòng lặp gọi AI:
    try:
        parsed["phan_bo_thanh_toan"] = resolve_payment_schedule(
            schedule_source_text, parsed
        )
    except Exception as exc:
        logger.warning("Lỗi resolve_payment_schedule (%s). Giữ nguyên tiến độ AI/mặc định.", exc)
        if "phan_bo_thanh_toan" not in parsed:
            parsed["phan_bo_thanh_toan"] = []

    try:
        _enrich_contract_calculations(parsed)
    except Exception as exc:
        logger.warning("Lỗi _enrich_contract_calculations: %s", exc)

    parsed["_executed_model"] = executed_model
    logger.info("Parsed successfully using model: %s", executed_model)
    return parsed


def _enrich_contract_calculations(data: Dict[str, Any]) -> None:
    """Enrich contract dictionary with calculated words and installment amounts."""
    # Compute total contract amount in integer
    raw_total = data.get("tong_tien_hop_dong_so", "")
    total_val = parse_positive_integer(raw_total)
    
    # If raw_total was empty, calculate from line items
    if total_val == 0 and "hang_muc" in data:
        for item in data["hang_muc"]:
            item_total = parse_positive_integer(item.get("tong_gia", 0))
            total_val += item_total
        data["tong_tien_hop_dong_so"] = f"{total_val:,}".replace(",", ".")
        data["tong_tien_sau_giam"] = format_currency_vnd(total_val)
        data["tong_tien_truoc_giam"] = format_currency_vnd(total_val)

    # Text in Vietnamese and Chinese
    data["tong_tien_chu_viet"] = number_to_vietnamese_words(total_val)
    data["tong_tien_chu_trung"] = number_to_chinese_words(total_val)

    # Compute installment amounts
    installments = data.get("phan_bo_thanh_toan", [])
    for inst in installments:
        pct = float(inst.get("phan_tram", 0))
        inst_amount = int(total_val * pct / 100.0)
        inst["so_tien"] = format_currency_vnd(inst_amount)


def summarize_project_title_with_ai(
    product_names: List[str],
    api_key: str = None,
    preferred_model: str = None
) -> Dict[str, str]:
    """
    Sử dụng Gemini AI để phân tích toàn bộ danh sách sản phẩm và khái quát hóa
    thành đúng 1 tên dự án / báo giá ngắn gọn song ngữ (Việt - Trung):
    - Tiếng Việt: Bắt đầu bằng 'Hệ thống...', 'Hạ tầng...', hoặc 'Giải pháp...'
    - Tiếng Trung: Ngắn gọn (Phồn thể) và BẮT BUỘC kết thúc bằng '報價單'.
    """
    if not product_names:
        return {"ten_du_an_viet": "", "ten_du_an_trung": ""}

    cleaned_names = [str(n).strip() for n in product_names if str(n).strip()]
    if not cleaned_names:
        return {"ten_du_an_viet": "", "ten_du_an_trung": ""}

    items_text = "\n".join(f"- {name}" for name in cleaned_names)

    prompt = f"""
Bạn là chuyên gia phân tích kỹ thuật và soạn thảo hợp đồng song ngữ Việt - Trung cho doanh nghiệp.
Dưới đây là danh sách toàn bộ các sản phẩm, thiết bị, dịch vụ hoặc phần mềm trong báo giá/hợp đồng:
\"\"\"
{items_text}
\"\"\"

Nhiệm vụ của bạn:
Phân tích toàn bộ danh sách sản phẩm trên và TỰ ĐỘNG KHÁI QUÁT HÓA thành đúng 1 tên dự án/báo giá ngắn gọn, súc tích và trang trọng theo chuẩn doanh nghiệp:
1. "ten_du_an_viet": Tên tiếng Việt ngắn gọn, BẮT BUỘC bắt đầu bằng "Hệ thống...", "Hạ tầng...", "Gói thiết bị..." hoặc "Giải pháp..." (Ví dụ: "Hệ thống MES Lite", "Hạ tầng mạng kết nối máy móc toàn xưởng", "Hệ thống camera giám sát và an ninh", "Hệ thống quản lý kho thông minh WMS"...)
2. "ten_du_an_trung": Tên tiếng Trung (Phồn thể) tương ứng ngắn gọn và BẮT BUỘC kết thúc bằng "報價單" (Ví dụ: "MES Lite 系統報價單", "機聯網網路佈建報價單", "智慧廠區監控系統報價單", "WMS 智慧倉庫管理系統報價單"...)

Chỉ trả về DUY NHẤT một khối JSON hợp lệ theo định dạng sau, không kèm bất kỳ giải thích nào:
{{
  "ten_du_an_viet": "Hệ thống...",
  "ten_du_an_trung": "...報價單"
}}
"""

    keys_to_try = [api_key] if api_key else ([DEFAULT_GEMINI_KEY] + [k for k in DEFAULT_GEMINI_KEYS if k != DEFAULT_GEMINI_KEY])
    models_to_try = []
    if preferred_model and preferred_model in CANDIDATE_MODELS:
        models_to_try.append(preferred_model)
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    for key in keys_to_try:
        for model in models_to_try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={key}"
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"}
            )
            try:
                with urllib.request.urlopen(req, timeout=15) as resp:
                    resp_data = json.loads(resp.read().decode("utf-8"))
                    text = resp_data["candidates"][0]["content"]["parts"][0]["text"]
                    clean_text = re.sub(r"^```(?:json)?\s*", "", text.strip())
                    clean_text = re.sub(r"\s*```$", "", clean_text).strip()
                    res_json = json.loads(clean_text)
                    if isinstance(res_json, dict) and res_json.get("ten_du_an_viet"):
                        return {
                            "ten_du_an_viet": str(res_json.get("ten_du_an_viet", "")).strip(),
                            "ten_du_an_trung": str(res_json.get("ten_du_an_trung", "")).strip()
                        }
            except Exception as e:
                logger.warning("Gemini model %s failed for project title summarization: %s", model, e)
                continue

    # Fallback tự nhiên nếu không kết nối được AI
    first_item = cleaned_names[0]
    base_name = first_item.split("\n")[0].split("-")[0].strip()
    return {
        "ten_du_an_viet": f"Hệ thống {base_name}",
        "ten_du_an_trung": f"{base_name} 系統報價單"
    }
