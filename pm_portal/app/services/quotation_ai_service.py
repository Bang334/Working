"""AI Service for parsing unstructured text or files into quotation data via Google Gemini API."""

import base64
import json
import logging
import re
import urllib.error
import urllib.request
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

import openpyxl
from docx import Document

from app.config import CANDIDATE_MODELS, DEFAULT_GEMINI_KEY, DEFAULT_GEMINI_KEYS, DEFAULT_MODEL
from app.utils.currency_helper import parse_positive_integer

logger = logging.getLogger(__name__)

SYSTEM_PROMPT = """Bạn là trợ lý AI chuyên gia bóc tách dữ liệu báo giá thiết bị IoT nhà máy, hệ thống RFID và quản lý năng lượng EMS cho CÔNG TY TNHH TOÀN CẦU HONG MING (宏明環球有限公司).

NHIỆM VỤ CỦA BẠN:
Phân tích văn bản yêu cầu báo giá, danh sách thiết bị từ email, chat WeChat/Zalo hoặc bảng tính khảo sát, hợp đồng và trích xuất thành định dạng JSON chuẩn xác 100% theo Biểu Mẫu Báo Giá Gốc (BaoGiaGoc.xlsx).

QUY TẮC CỐT LÕI:
1. Bóc tách từng hạng mục sản phẩm theo đúng 6 cột chi tiết của mẫu gốc:
   - product_name: Tên Sản Phẩm (產品名稱) - song ngữ Việt - Trung (vd: THẺ RFID làm bằng silicon / 硅膠型RFID標籤)
   - category: Hạng Mục / Quy cách kỹ thuật (項目類別) - chi tiết model, kích thước, tần số
   - content: Nội Dung Sản Phẩm (產品內容) - mã part number, dải mã số serial, phạm vi cung cấp
   - qty: Số lượng (數量) - số nguyên > 0
   - lead_time: Thời gian hoàn thành (完成時間) - ví dụ: "60天內", "15-20天", "7-10天"
   - price: Đơn giá (單價) - số tiền nguyên (nếu không có thì để 0)
2. Trích xuất thông tin bên mua (買方資訊):
   - customer_name_vi: Tên công ty tiếng Việt
   - customer_name_zh: Tên công ty tiếng Trung
   - contact_person: Người liên lạc (聯絡人) kèm phòng ban
   - phone: Điện thoại liên lạc (聯絡電話)
   - address: Địa chỉ liên lạc (聯絡地址)
   - email: Email liên lạc (電子郵件)
   - tax_code: Mã số thuế (MST)
3. Tiêu đề và ngày tháng:
   - quote_title: Tiêu đề báo giá song ngữ (vd: BÁO GIÁ THẺ RFID SILICON / 硅膠型RFID標籤報價單)
   - quotation_no: Số báo giá định dạng HM-QT-YYYYMM-001
   - quotation_date: Ngày báo giá định dạng DD/MM/YYYY
4. Trích xuất 5 Điều khoản tuyên bố đặc biệt của báo giá (◆ 報價特別聲明):
   - delivery_term: Điều 1 - Thời gian giao hàng
   - validity_term: Điều 2 - Thời hạn hiệu lực
   - warranty_term: Điều 3 - Thời hạn bảo hành
   - vat_term: Điều 4 - Thuế VAT
   - payment_term: Điều 5 - Phương thức thanh toán
5. Trả về đúng định dạng JSON thuần túy, không bọc giải thích rườm rà.

ĐỊNH DẠNG JSON ĐẦU RA BẮT BUỘC:
{
  "customer_name_vi": "Tên công ty tiếng Việt",
  "customer_name_zh": "Tên công ty tiếng Trung",
  "contact_person": "Người liên lạc",
  "phone": "Số điện thoại",
  "address": "Địa chỉ",
  "email": "Email",
  "tax_code": "Mã số thuế",
  "quote_title": "BÁO GIÁ THIẾT BỊ / 設備報價單",
  "quotation_no": "HM-QT-202610-001",
  "quotation_date": "27/10/2026",
  "delivery_term": "1. 以上報價交期：自收到付款之日起60天內交貨。\\nNgày giao hàng theo báo giá trên: Trong vòng 60 ngày kể từ ngày nhận được thanh toán.",
  "validity_term": "2. 報價有效期限：報價單有效期限至30天後，報價簽回視同專案確認，不可任意終止、取消本專案。\\nThời hạn hiệu lực của báo giá: Báo giá có hiệu lực đến hết ngày quy định.",
  "warranty_term": "3. 上述報價中的硬體設備自交貨日起保固六個月/十二個月；因人為因素造成之碰撞、破裂、損壞等情況，不在保固範圍內\\nThiết bị phần cứng trong báo giá trên được bảo hành theo quy định kể từ ngày giao hàng.",
  "vat_term": "4.以上報價不含增值稅 VAT，稅率則按當時發票開出時之規定。\\nGiá trên chưa bao gồm VAT. Thuế suất sẽ áp dụng theo quy định tại thời điểm xuất hóa đơn.",
  "payment_term": "5. 付款方式 : 訂金 50% , 交貨 50% , 現金/匯款 付款 TT30Day。\\nPhương thức thanh toán: Đặt cọc 50%, thanh toán 50% khi giao hàng, TT30 ngày.",
  "items": [
    {
      "stt": 1,
      "product_name": "Tên sản phẩm tiếng Việt & Trung",
      "category": "Hạng mục / Quy cách kỹ thuật",
      "content": "Nội dung sản phẩm / Model chi tiết",
      "qty": 500,
      "lead_time": "60天內",
      "price": 39000
    }
  ]
}
"""


def extract_text_from_excel(file_bytes: bytes) -> str:
    """Read cell values from uploaded Excel workbook for quotation parsing."""
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        lines = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            lines.append(f"--- SHEET: {sheetname} ---")
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
                if non_empty:
                    lines.append(" | ".join(non_empty))
        return "\n".join(lines[:250])
    except Exception as exc:
        logger.warning("Failed to extract Excel content: %s", exc)
        return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Read paragraphs and tables from uploaded DOCX file."""
    try:
        doc = Document(BytesIO(file_bytes))
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_vals = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_vals:
                    lines.append(" | ".join(row_vals))
        return "\n".join(lines[:300])
    except Exception as exc:
        logger.warning("Failed to extract DOCX content: %s", exc)
        return ""


def fallback_heuristic_parse(raw_text: str) -> Dict[str, Any]:
    """Fallback parser using regex when Gemini API is offline or unavailable."""
    today_str = datetime.now().strftime("%d/%m/%Y")
    now_ym = datetime.now().strftime("%Y%m")

    cust_vi = "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)"
    cust_zh = "台灣精星科技(越南)有限公司"
    proj = "BÁO GIÁ THIẾT BỊ & LINH KIỆN / 設備與零件報價單"
    quote_no = f"HM-QT-{now_ym}-001"
    contact = "Phòng Mua Hàng / 採購處"
    phone = "+84966687109"
    addr = "KCX Tân Thuận, Quận 7, TP. Hồ Chí Minh"
    email = "pur02_vn@mail.cxtechnology.com"

    # Search for customer
    cust_match = re.search(r"(?:công ty|khách hàng|kính gửi|to|customer)\s*[:：\-]?\s*([^\n\r,]+)", raw_text, re.IGNORECASE)
    if cust_match:
        cust_vi = cust_match.group(1).strip()

    tax_match = re.search(r"(?:mst|tax|mã số thuế)\s*[:：\-]?\s*([0-9\-]+)", raw_text, re.IGNORECASE)
    tax_code = tax_match.group(1).strip() if tax_match else "0301445789"

    items = []
    lines = raw_text.splitlines()
    for line in lines:
        cleaned = line.strip("-* \t•")
        if not cleaned or len(cleaned) < 3:
            continue

        qty_price_match = re.search(r"(?:số lượng|sl|qty)?\s*[:：]?\s*(\d+)\s*(bộ|cái|hệ|cuộn|gói|set|ea)?(?:\s*[x*]\s*([0-9.,]+))?", cleaned, re.IGNORECASE)
        if qty_price_match:
            qty = int(qty_price_match.group(1)) if qty_price_match.group(1) else 1
            unit = qty_price_match.group(2) or "Bộ"
            price = parse_positive_integer(qty_price_match.group(3)) if qty_price_match.group(3) else 0
            pname = cleaned.split(":")[0].strip() if ":" in cleaned else cleaned
            items.append({
                "stt": len(items) + 1,
                "product_name": pname,
                "category": f"Quy cách kỹ thuật {pname}",
                "content": f"Model & Phụ kiện {pname}",
                "qty": qty,
                "lead_time": "60天內",
                "price": price
            })
        elif len(items) < 8 and ("thiết bị" in cleaned.lower() or "cáp" in cleaned.lower() or "rfid" in cleaned.lower() or "gateway" in cleaned.lower() or "đồng hồ" in cleaned.lower()):
            items.append({
                "stt": len(items) + 1,
                "product_name": cleaned,
                "category": f"Quy cách {cleaned}",
                "content": f"Chi tiết {cleaned}",
                "qty": 1,
                "lead_time": "60天內",
                "price": 0
            })

    if not items:
        items = [
            {
                "stt": 1,
                "product_name": "THẺ RFID làm bằng silicon\n硅膠型RFID標籤",
                "category": "THẺ RFID SILICON (KHÔNG KHÁNG KIM LOẠI) 90x18mm, NXP UCODE 8 BLACK",
                "content": "OPP9018 (D0010121 ~ D0010620)",
                "qty": 500,
                "lead_time": "60天內",
                "price": 39000
            }
        ]

    return {
        "customer_name_vi": cust_vi,
        "customer_name_zh": cust_zh,
        "contact_person": contact,
        "phone": phone,
        "address": addr,
        "email": email,
        "tax_code": tax_code,
        "quote_title": proj,
        "quotation_no": quote_no,
        "quotation_date": today_str,
        "delivery_term": "1. 以上報價交期：自收到付款之日起60天內交貨。\nNgày giao hàng theo báo giá trên: Trong vòng 60 ngày kể từ ngày nhận được thanh toán.",
        "validity_term": "2. 報價有效期限：報價單有效期限至30天後，報價簽回視同專案確認，不可任意終止、取消本專案。\n    Thời hạn hiệu lực của báo giá: Báo giá có hiệu lực 30 ngày kể từ ngày phát hành.",
        "warranty_term": "3. 上述報價中的硬體設備自交貨日起保固六個月；因人為因素造成之碰撞、破裂、損壞等情況，不在保固範圍內\n  Thiết bị phần cứng trong báo giá trên được bảo hành sáu tháng kể từ ngày giao hàng; các hư hỏng do va chạm, vỡ, gãy do con người gây ra không nằm trong phạm vi bảo hành",
        "vat_term": "4.以上報價不含增值稅 VAT，稅率則按當時發票開出時之規定。\nGiá trên chưa bao gồm VAT. Thuế suất sẽ áp dụng theo quy định tại thời điểm xuất hóa đơn.",
        "payment_term": "5. 付款方式 : 訂金 50% , 交貨 50%  , 現金/匯款 付款  TT30Day。\n     Phương thức thanh toán: Đặt cọc 50%, thanh toán 50% khi giao hàng, thanh toán bằng tiền mặt/chuyển khoản, TT30 ngày.",
        "items": items,
        "_executed_model": "Heuristic Parser"
    }


def parse_quotation_text(
    raw_input_text: str,
    file_bytes: Optional[bytes] = None,
    file_mime_type: Optional[str] = None,
    file_name: Optional[str] = None,
    preferred_model: Optional[str] = None
) -> Dict[str, Any]:
    """Parse quotation raw input text or attachments using Gemini API with model rotation."""
    keys_to_try = [k for k in DEFAULT_GEMINI_KEYS if k]
    if DEFAULT_GEMINI_KEY and DEFAULT_GEMINI_KEY not in keys_to_try:
        keys_to_try.insert(0, DEFAULT_GEMINI_KEY)

    models_to_try = []
    if preferred_model and preferred_model in CANDIDATE_MODELS:
        models_to_try.append(preferred_model)
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    fname_lower = (file_name or "").lower()
    fmime_lower = (file_mime_type or "").lower()
    is_image = any(ext in fname_lower or ext in fmime_lower for ext in ("jpg", "jpeg", "png", "webp"))
    is_excel = any(ext in fname_lower or ext in fmime_lower for ext in ("xls", "xlsx", "spreadsheet"))
    is_docx = any(ext in fname_lower or ext in fmime_lower for ext in ("doc", "docx", "word"))

    parts = []
    if file_bytes and is_image:
        mime_to_send = "image/jpeg"
        if "png" in fname_lower or "png" in fmime_lower:
            mime_to_send = "image/png"
        elif "webp" in fname_lower or "webp" in fmime_lower:
            mime_to_send = "image/webp"

        prompt_text = f"{SYSTEM_PROMPT}\n\nHÃY BÓC TÁCH BẢNG BÁO GIÁ / DANH MỤC THIẾT BỊ TỪ ẢNH ĐÍNH KÈM:"
        if raw_input_text and raw_input_text.strip():
            prompt_text += f"\n\nGhi chú bổ sung:\n{raw_input_text.strip()}"

        parts = [
            {"text": prompt_text},
            {
                "inlineData": {
                    "mimeType": mime_to_send,
                    "data": base64.b64encode(file_bytes).decode("ascii")
                }
            }
        ]
    elif file_bytes and is_excel:
        excel_text = extract_text_from_excel(file_bytes)
        prompt_text = (
            f"{SYSTEM_PROMPT}\n\nDỮ LIỆU ĐỌC TỪ BẢNG TÍNH EXCEL ĐÍNH KÈM ({file_name or 'Bao_Gia.xlsx'}):\n\"\"\"\n{excel_text}\n\"\"\""
        )
        if raw_input_text and raw_input_text.strip():
            prompt_text += f"\n\nCHÚ THÍCH CỦA NGƯỜI DÙNG:\n\"\"\"\n{raw_input_text.strip()}\n\"\"\""
        parts = [{"text": prompt_text}]
    elif file_bytes and is_docx:
        docx_text = extract_text_from_docx(file_bytes)
        prompt_text = (
            f"{SYSTEM_PROMPT}\n\nDỮ LIỆU ĐỌC TỪ TỆP WORD ĐÍNH KÈM ({file_name or 'Bao_Gia.docx'}):\n\"\"\"\n{docx_text}\n\"\"\""
        )
        if raw_input_text and raw_input_text.strip():
            prompt_text += f"\n\nCHÚ THÍCH CỦA NGƯỜI DÙNG:\n\"\"\"\n{raw_input_text.strip()}\n\"\"\""
        parts = [{"text": prompt_text}]
    else:
        prompt = f"{SYSTEM_PROMPT}\n\nDỮ LIỆU ĐẦU VÀO CẦN BÓC TÁCH BÁO GIÁ:\n\"\"\"\n{raw_input_text}\n\"\"\""
        parts = [{"text": prompt}]

    payload = {
        "contents": [{"parts": parts}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    parsed = None
    executed_model = None

    if keys_to_try:
        for key in keys_to_try:
            for model in models_to_try:
                url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={key}"
                logger.info("Quotation AI: attempting model %s", model)
                req = urllib.request.Request(
                    url,
                    data=json.dumps(payload).encode("utf-8"),
                    headers={"Content-Type": "application/json"}
                )
                try:
                    with urllib.request.urlopen(req, timeout=25) as response:
                        result = json.loads(response.read().decode("utf-8"))
                        candidate = result.get("candidates", [])[0]
                        text_out = candidate["content"]["parts"][0]["text"].strip()
                        if text_out.startswith("```json"):
                            text_out = text_out[7:]
                        elif text_out.startswith("```"):
                            text_out = text_out[3:]
                        if text_out.endswith("```"):
                            text_out = text_out[:-3]
                        parsed = json.loads(text_out.strip())
                        executed_model = model
                        break
                except Exception as exc:
                    logger.warning("Quotation AI failed with model %s: %s", model, exc)
                    continue
            if parsed:
                break

    if not parsed:
        logger.info("Falling back to heuristic quotation parser.")
        parsed = fallback_heuristic_parse(raw_input_text)
        executed_model = parsed.get("_executed_model", "Heuristic Parser")

    parsed["_executed_model"] = executed_model
    return parsed
