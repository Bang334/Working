# -*- coding: utf-8 -*-
"""AI Service for parsing payment voucher / application request data via Gemini API (Text + PDF + Images + Excel)."""

import base64
import io
import json
import logging
import re
import urllib.error
import urllib.request
from datetime import datetime
from typing import Any, Dict, List, Optional, Union

import openpyxl

from app.config import CANDIDATE_MODELS, DEFAULT_GEMINI_KEY, DEFAULT_GEMINI_KEYS
from app.services.beneficiary_service import find_matching_beneficiary, list_beneficiaries
from app.services.payment_voucher_excel_service import get_next_report_code
from app.services.reminder_service import (
    format_amount_display,
    number_to_chinese_words,
    number_to_vietnamese_words,
)
from app.utils.currency_helper import parse_positive_integer

logger = logging.getLogger(__name__)

PAYMENT_VOUCHER_SYSTEM_PROMPT = """Bạn là chuyên viên tài chính kế toán cao cấp của CÔNG TY TNHH TOÀN CẦU HONG MING (宏明環球有限公司).
NHIỆM VỤ:
Phân tích tài liệu đầu vào (Biên bản nghiệm thu / tiến độ lắp đặt Word, Hóa đơn điện tử PDF, Ảnh chụp hóa đơn/biên lai, Bảng tính Excel, Hợp đồng, Đơn đặt hàng PO, hoặc Tin nhắn văn bản) và trích xuất thành JSON chuẩn xác 100% để lập "ĐƠN XIN THANH TOÁN" (浤詺付款申请单).

KHẢ NĂNG THÍCH ỨNG ĐA DẠNG LOẠI TÀI LIỆU / BIỂU MẪU (MULTI-FORM ADAPTIVE):
1. BIÊN BẢN TIẾN ĐỘ LẮP ĐẶT / BIÊN BẢN NGHIỆM THU (Acceptance Confirmation / 驗收確認書 / 驗收單):
   - Nhận diện đợt/giai đoạn nghiệm thu thanh toán (ví dụ: Giai đoạn 3 / 第三期, tỷ lệ 20%...).
   - Bóc tách chính xác số tiền thanh toán của đợt đó (amount). Lưu ý: Nếu có bảng phân bổ tiến độ thanh toán, ưu tiên lấy số tiền theo đợt nghiệm thu tương ứng.
   - Đơn vị thụ hưởng (beneficiary_name): 
     * Bên yêu cầu thanh toán / Bên cung cấp dịch vụ (thường là Bên B trong biên bản nghiệm thu).
     * Nếu Bên B là CÔNG TY TNHH TOÀN CẦU HONG MING (hoặc 宏明環球有限公司), ghi tên công ty chính xác.
   - Lý do thanh toán (reason_vi & reason_zh):
     * reason_vi: Diễn giải rõ ràng theo chuẩn kế toán: "Thanh toán đợt [kỳ/giai đoạn] ([tỷ lệ]%) theo Biên bản nghiệm thu và Hợp đồng số [số HĐ] ngày [ngày ký HĐ] cho [tên hạng mục/phần mềm/dự án]."
     * reason_zh: Diễn giải tiếng Trung Phồn thể chuẩn mực doanh nghiệp Hong Ming: "依據驗收確認書及合約編號 [số HĐ]（日期 [ngày ký]），支付第[kỳ]期（[tỷ lệ]%）驗收款項，用於[tên dự án tiếng Trung]之相關費用。"
     * full_reason: Ghép nối 2 dòng: reason_vi trước, reason_zh sau (ngăn cách bằng dấu xuống dòng \\n).
2. HÓA ĐƠN GTGT (VAT Invoice / 發票) hoặc BIÊN LAI:
   - Trích xuất Bên bán làm beneficiary_name, STK, ngân hàng, tổng tiền thanh toán sau thuế.
   - Lý do: Ghi rõ số hóa đơn, ký hiệu mẫu số, ngày hóa đơn, nội dung mua sắm hàng hóa hoặc dịch vụ.
3. HỢP ĐỒNG KINH TẾ (Contract / 合約) & ĐƠN ĐẶT HÀNG (PO / 採購單):
   - Trích xuất đợt thanh toán đến hạn hoặc tổng tiền PO, tên đối tác/nhà cung cấp, số HĐ/PO.
4. BẢNG KÊ CHI PHÍ (Hải quan, vận tải, nhiều hóa đơn con):
   - Bóc tách các khoản con vào breakdown_items (desc, amount) tối đa 4 dòng, tổng tiền amount là tổng cộng.
5. TIN NHẮN / CHỈ THỊ VĂN BẢN (WeChat, Zalo, Email, Tờ trình nội bộ):
   - Trích xuất thông tin người nhận, STK, ngân hàng, số tiền, mục đích chi.

QUY TẮC BẮT BUỘC VỀ DỮ LIỆU ĐẦU RA:
- beneficiary_name: Tên công ty/cá nhân thụ hưởng (Bên nhận tiền).
- account_number: Số tài khoản ngân hàng (nếu tài liệu không có STK thì để chuỗi rỗng "").
- bank_name: Tên ngân hàng nhận tiền (nếu tài liệu không có thì để chuỗi rỗng "").
- amount: Số tiền thanh toán (nguyên dương, ví dụ 46800000, 2592000...).
- currency: 'VND' hoặc 'USD' (mặc định VND).
- reason_vi: Diễn giải tiếng Việt đầy đủ, rõ ràng số hợp đồng/hóa đơn, đợt thanh toán, tên dự án.
- reason_zh: Diễn giải tiếng Trung Phồn thể chuẩn mực Đài Loan.
- full_reason: reason_vi + "\\n" + reason_zh.
- report_code: Định dạng `报单号：{mã_nhân_viên}-{YYYYMMDD}-{số_thứ_tự_trong_ngày}` (mặc định: `报单号：HVM007-{YYYYMMDD}-1`).
- request_date: Ngày lập đơn định dạng YYYY-MM-DD.
- prepared_by: Người lập đơn (mặc định: "NGUYỄN SỸ KIM BẰNG").
- breakdown_items: Mảng các dòng chi tiết con nếu có (tối đa 4 dòng), hoặc [] nếu chỉ có 1 khoản duy nhất.

ĐỊNH DẠNG JSON ĐẦU RA BẮT BUỘC:
{
  "beneficiary_name": "CÔNG TY TNHH TOÀN CẦU HONG MING",
  "account_number": "",
  "bank_name": "",
  "amount": 46800000,
  "currency": "VND",
  "breakdown_items": [],
  "reason_vi": "Thanh toán đợt Giai đoạn 3 (20%) theo Biên bản nghiệm thu và Hợp đồng số HMV151020250001 ngày 09/12/2025 cho Phần mềm tùy chỉnh chức năng hệ thống RFID.",
  "reason_zh": "依據驗收確認書及合約編號 HMV151020250001（日期 2025/12/09），支付第三期（20%）驗收款項，用於RFID 系統功能客製化軟體之相關費用。",
  "full_reason": "Thanh toán đợt Giai đoạn 3 (20%) theo Biên bản nghiệm thu và Hợp đồng số HMV151020250001 ngày 09/12/2025 cho Phần mềm tùy chỉnh chức năng hệ thống RFID.\\n依據驗收確認書及合約編號 HMV151020250001（日期 2025/12/09），支付第三期（20%）驗收款項，用於RFID 系統功能客製化軟體之相關費用。",
  "report_code": "报单号：HVM007-20261021-1",
  "request_date": "2026-10-21",
  "prepared_by": "NGUYỄN SỸ KIM BẰNG"
}

LƯU Ý: Chỉ trả về mã JSON thuần túy, không có markdown block hay lời chào bên ngoài.
"""

TRANSLATE_REASON_SYSTEM_PROMPT = """Bạn là trợ lý dịch thuật tài chính - hành chính song ngữ Việt - Trung cho CÔNG TY TNHH TOÀN CẦU HONG MING.
Nhiệm vụ: Nhận một đoạn lý do thanh toán bằng tiếng Việt và dịch sang Tiếng Trung Phồn thể chuẩn phong cách chứng từ công ty Hong Ming.
Ví dụ:
Tiếng Việt: "Thanh toán chi phí Thu Thập dữ liệu máy dập Hóa Đơn: 055"
Tiếng Trung: "支付沖壓機資料收集費用，發票號碼：055"

Trả về JSON có dạng:
{
  "reason_vi": "nội dung tiếng Việt đã chuẩn hóa",
  "reason_zh": "dịch sang tiếng Trung phồn thể chuẩn mực",
  "full_reason": "nội dung tiếng Việt\\nnội dung tiếng Trung"
}
Chỉ trả về JSON thuần túy.
"""


def _clean_json_response(raw_text: str) -> str:
    """Loại bỏ markdown backticks và bóc tách chuỗi JSON."""
    cleaned = re.sub(r"^```(?:json)?\s*", "", raw_text.strip(), flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    match = re.search(r"(\{.*\})", cleaned, re.DOTALL)
    return match.group(1).strip() if match else cleaned.strip()


def _extract_excel_text_from_base64(b64_str: str) -> str:
    """Đọc nhanh các ô dữ liệu từ file Excel base64 để làm ngữ cảnh cho AI."""
    try:
        clean_b64 = re.sub(r"^data:.*?;base64,", "", b64_str.strip())
        raw_bytes = base64.b64decode(clean_b64)
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        lines = []
        for sname in wb.sheetnames[:3]:
            ws = wb[sname]
            lines.append(f"--- Bảng tính: {sname} ---")
            for r in range(1, min(ws.max_row + 1, 40)):
                row_vals = [
                    f"{openpyxl.utils.get_column_letter(c)}{r}: {ws.cell(r, c).value}"
                    for c in range(1, min(ws.max_column + 1, 15))
                    if ws.cell(r, c).value is not None
                ]
                if row_vals:
                    lines.append(" | ".join(row_vals))
        return "\n".join(lines)
    except Exception as exc:
        logger.warning("Could not read excel from base64: %s", exc)
        return ""


def _extract_docx_text_from_base64(b64_str: str) -> str:
    """Trích xuất toàn bộ văn bản và bảng biểu từ file Word (.docx) base64."""
    try:
        import docx
        clean_b64 = re.sub(r"^data:.*?;base64,", "", b64_str.strip())
        raw_bytes = base64.b64decode(clean_b64)
        doc = docx.Document(io.BytesIO(raw_bytes))
        chunks = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                chunks.append(txt)
        for t_idx, table in enumerate(doc.tables):
            chunks.append(f"\n[BẢNG THÔNG TIN {t_idx + 1}]")
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    chunks.append(" | ".join(cells))
        return "\n".join(chunks)
    except Exception as exc:
        logger.warning("Could not read docx from base64: %s", exc)
        return ""


def _call_gemini_multimodal(
    system_prompt: str,
    parts: List[Dict[str, Any]],
    preferred_model: Optional[str] = None
) -> Dict[str, Any]:
    """Hàm gọi API Gemini hỗ trợ đa phương thức (Text + File PDF + Image) xoay vòng model & api key an toàn."""
    models_to_try = [preferred_model] if preferred_model else []
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    keys_to_try = list(DEFAULT_GEMINI_KEYS)
    if DEFAULT_GEMINI_KEY and DEFAULT_GEMINI_KEY not in keys_to_try:
        keys_to_try.insert(0, DEFAULT_GEMINI_KEY)

    last_error = None
    for model_name in models_to_try:
        for api_key in keys_to_try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
            payload = {
                "systemInstruction": {"parts": [{"text": system_prompt}]},
                "contents": [
                    {
                        "role": "user",
                        "parts": parts
                    }
                ],
                "generationConfig": {
                    "temperature": 0.1,
                    "responseMimeType": "application/json"
                }
            }

            try:
                data = json.dumps(payload).encode("utf-8")
                req = urllib.request.Request(
                    url,
                    data=data,
                    headers={"Content-Type": "application/json"},
                    method="POST"
                )

                with urllib.request.urlopen(req, timeout=40) as resp:
                    resp_data = json.loads(resp.read().decode("utf-8"))
                    candidates = resp_data.get("candidates", [])
                    if candidates:
                        raw_content = candidates[0]["content"]["parts"][0]["text"]
                        cleaned_json = _clean_json_response(raw_content)
                        parsed = json.loads(cleaned_json)
                        parsed["_executed_model"] = model_name
                        return parsed

            except urllib.error.HTTPError as http_err:
                status_code = http_err.code
                err_body = http_err.read().decode("utf-8", errors="ignore")
                logger.warning(
                    "Model %s failed with key ...%s (HTTP %s): %s",
                    model_name, api_key[-6:], status_code, err_body[:100]
                )
                last_error = http_err
            except Exception as exc:
                logger.warning("Error with model %s: %s", model_name, exc)
                last_error = exc

    if last_error:
        raise last_error
    raise RuntimeError("Không thể kết nối đến AI Gemini.")


def extract_payment_voucher_data(
    raw_text: Optional[str] = None,
    preferred_model: Optional[str] = None,
    file_base64: Optional[str] = None,
    file_mime: Optional[str] = None,
    file_name: Optional[str] = None
) -> Dict[str, Any]:
    """
    Trích xuất toàn diện thông tin đơn xin thanh toán từ văn bản thô hoặc file đính kèm (PDF, Ảnh, Excel).
    Tự động đối chiếu Sổ tay tài khoản thụ hưởng để điền STK & ngân hàng chuẩn nhất.
    """
    has_text = bool(raw_text and raw_text.strip())
    has_file = bool(file_base64 and file_base64.strip())

    if not has_text and not has_file:
        raise ValueError("Vui lòng cung cấp văn bản hoặc đính kèm tệp (PDF, Ảnh, Excel) cần trích xuất.")

    # 1. Chuẩn bị ngữ cảnh các đối tác thường dùng để đưa vào AI nhận diện tốt hơn
    known_beneficiaries = list_beneficiaries(limit=25)
    known_list_str = "\n".join([
        f"- {b['unit_name']} (STK: {b['account_number']}, NH: {b['bank_name']}, Tên ngắn: {b.get('short_name') or ''})"
        for b in known_beneficiaries
    ])

    today_str = datetime.now().strftime("%Y-%m-%d")
    today_code = datetime.now().strftime("%Y%m%d")

    # Tính mã báo đơn tiếp theo cho nhân viên HVM007
    auto_code_info = get_next_report_code(employee_code="HVM007", target_date=today_str)
    suggested_code = auto_code_info["full_code"]

    prompt_context = f"""DANH SÁCH ĐƠN VỊ THỤ HƯỞNG ĐÃ LƯU TRONG HỆ THỐNG:
{known_list_str}

Hôm nay là ngày: {today_str}
Gợi ý mã báo đơn tiếp theo: {suggested_code}
"""
    if has_text:
        prompt_context += f"\nCHỈ THỊ / GHI CHÚ BỔ SUNG:\n{raw_text.strip()}\n"

    # 2. Xây dựng parts cho Gemini API
    parts: List[Dict[str, Any]] = []

    if has_file:
        clean_b64 = re.sub(r"^data:.*?;base64,", "", file_base64.strip())
        fname_lower = (file_name or "").lower()
        mime = (file_mime or "").lower()

        # 1. Kiểm tra file Word (.docx, .doc)
        if fname_lower.endswith((".docx", ".doc")) or "wordprocessingml" in mime or "msword" in mime:
            docx_text = _extract_docx_text_from_base64(file_base64)
            prompt_context += f"\nNỘI DUNG TRÍCH XUẤT TỪ VĂN BẢN WORD ĐÍNH KÈM ({file_name}):\n{docx_text}\n"
            parts.append({
                "text": f"Hãy đọc kỹ toàn bộ nội dung văn bản Word đính kèm ({file_name}), kết hợp với ngữ cảnh sau để trích xuất đầy đủ thông tin lập Đơn Xin Thanh Toán:\n\n{prompt_context}"
            })

        # 2. Kiểm tra file Excel (.xlsx, .xls, .xlsm)
        elif fname_lower.endswith((".xlsx", ".xls", ".xlsm")) or "spreadsheet" in mime or "excel" in mime:
            excel_text = _extract_excel_text_from_base64(file_base64)
            prompt_context += f"\nNỘI DUNG TRÍCH XUẤT TỪ BẢNG TÍNH EXCEL ĐÍNH KÈM ({file_name}):\n{excel_text}\n"
            parts.append({
                "text": f"Hãy đọc kỹ bảng tính Excel đính kèm ({file_name}), kết hợp với ngữ cảnh sau để trích xuất đầy đủ thông tin lập Đơn Xin Thanh Toán:\n\n{prompt_context}"
            })

        # 3. Kiểm tra file Text / CSV / JSON
        elif fname_lower.endswith((".txt", ".csv", ".json")) or "text/" in mime:
            try:
                raw_bytes = base64.b64decode(clean_b64)
                txt_content = raw_bytes.decode("utf-8", errors="ignore")
                prompt_context += f"\nNỘI DUNG TỆP VĂN BẢN ĐÍNH KÈM ({file_name}):\n{txt_content}\n"
                parts.append({"text": prompt_context})
            except Exception:
                parts.append({"text": prompt_context})

        # 4. File PDF hoặc Ảnh: gửi inlineData qua Gemini Multimodal Vision API
        else:
            if fname_lower.endswith(".pdf") or "pdf" in mime:
                final_mime = "application/pdf"
            elif fname_lower.endswith(".png") or "png" in mime:
                final_mime = "image/png"
            elif fname_lower.endswith((".jpg", ".jpeg")) or "jpeg" in mime or "jpg" in mime:
                final_mime = "image/jpeg"
            elif fname_lower.endswith(".webp") or "webp" in mime:
                final_mime = "image/webp"
            else:
                final_mime = "application/pdf" if fname_lower.endswith(".pdf") else "image/jpeg"

            parts.append({
                "inlineData": {
                    "mimeType": final_mime,
                    "data": clean_b64
                }
            })
            parts.append({
                "text": f"Hãy đọc kỹ tệp đính kèm ({file_name or 'Tài liệu'}), kết hợp với ngữ cảnh sau để trích xuất đầy đủ thông tin lập Đơn Xin Thanh Toán:\n\n{prompt_context}"
            })
    else:
        parts.append({"text": prompt_context})

    # 3. Gọi Gemini API
    parsed = _call_gemini_multimodal(PAYMENT_VOUCHER_SYSTEM_PROMPT, parts, preferred_model)

    # 4. Hậu kiểm và đối chiếu danh bạ thụ hưởng
    ben_name = parsed.get("beneficiary_name", "")
    acc_num = parsed.get("account_number", "")
    bank = parsed.get("bank_name", "")

    matched_ben = find_matching_beneficiary(ben_name) or (find_matching_beneficiary(raw_text) if has_text else None)
    if matched_ben:
        parsed["beneficiary_name"] = matched_ben["unit_name"]
        if not acc_num or len(acc_num.strip()) < 5:
            parsed["account_number"] = matched_ben["account_number"]
        if not bank or len(bank.strip()) < 3:
            parsed["bank_name"] = matched_ben["bank_name"]
        parsed["beneficiary_id"] = matched_ben["id"]
        parsed["beneficiary_short_name"] = matched_ben.get("short_name") or ""
        parsed["_matched_from_directory"] = True
    else:
        parsed["_matched_from_directory"] = False

    # 5. Chuẩn hóa số tiền & tính chữ
    raw_amount = parsed.get("amount")
    amount_int = parse_positive_integer(raw_amount) if raw_amount else 0
    parsed["amount"] = amount_int
    parsed["amount_formatted"] = format_amount_display(amount_int)
    parsed["currency"] = str(parsed.get("currency") or "VND").upper()
    parsed["amount_in_words_vi"] = number_to_vietnamese_words(amount_int)
    parsed["amount_in_words_zh"] = number_to_chinese_words(amount_int)

    # Tỷ giá USD mặc định 26.000
    exchange_rate = 26000
    parsed["exchange_rate"] = exchange_rate
    parsed["amount_usd"] = round(amount_int / exchange_rate, 2) if amount_int > 0 else 0.0

    # Đảm bảo full_reason song ngữ
    r_vi = parsed.get("reason_vi", "").strip()
    r_zh = parsed.get("reason_zh", "").strip()
    if r_vi and r_zh:
        parsed["full_reason"] = f"{r_vi}\n{r_zh}"
    elif r_vi:
        parsed["full_reason"] = r_vi
    elif r_zh:
        parsed["full_reason"] = r_zh

    raw_code = str(parsed.get("report_code") or "").strip()
    if not raw_code or "HMV002-" in raw_code:
        req_d = parsed.get("request_date") or today_str
        parsed["report_code"] = get_next_report_code(employee_code="HVM007", target_date=req_d)["full_code"]
    elif not raw_code.startswith("报单号：") and not raw_code.startswith("报单"):
        parsed["report_code"] = f"报单号：{raw_code}"

    if not parsed.get("request_date"):
        parsed["request_date"] = today_str

    if not parsed.get("prepared_by"):
        parsed["prepared_by"] = "NGUYỄN SỸ KIM BẰNG"

    return parsed


def translate_payment_reason(reason_vi: str, preferred_model: Optional[str] = None) -> Dict[str, str]:
    """Hỗ trợ dịch nhanh nội dung lý do thanh toán sang tiếng Trung song ngữ."""
    if not reason_vi or not reason_vi.strip():
        return {"reason_vi": "", "reason_zh": "", "full_reason": ""}

    user_text = f"Nội dung cần dịch sang tiếng Trung song ngữ:\n{reason_vi.strip()}"
    res = _call_gemini_multimodal(TRANSLATE_REASON_SYSTEM_PROMPT, [{"text": user_text}], preferred_model)
    return {
        "reason_vi": res.get("reason_vi", reason_vi.strip()),
        "reason_zh": res.get("reason_zh", ""),
        "full_reason": res.get("full_reason", f"{reason_vi.strip()}\n{res.get('reason_zh', '')}")
    }
