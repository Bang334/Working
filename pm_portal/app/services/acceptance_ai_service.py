"""AI Service for parsing unstructured acceptance notes or delivery sheets via Gemini API."""

import json
import logging
import re
import urllib.error
import urllib.request
from typing import Any, Dict, List, Optional

from app.config import CANDIDATE_MODELS, DEFAULT_GEMINI_KEY, DEFAULT_MODEL, DEFAULT_GEMINI_KEYS
from app.utils.currency_helper import (
    format_currency_vnd,
    number_to_chinese_words,
    number_to_vietnamese_words,
    parse_positive_integer,
)

logger = logging.getLogger(__name__)

ACCEPTANCE_SYSTEM_PROMPT = """Bạn là chuyên gia phân tích và bóc tách dữ liệu nghiệm thu thiết bị / biên bản giao nhận kỹ thuật song ngữ Việt - Trung cho CÔNG TY TNHH TOÀN CẦU HONG MING (宏明環球有限公司).

NHIỆM VỤ:
Phân tích văn bản đầu vào (danh sách thiết bị sửa chữa, đơn ký nhận bàn giao phần cứng, báo giá giao nhận hoặc ghi chú nghiệm thu) và trích xuất thành định dạng JSON chuẩn 100%.

QUY TẮC CỐT LÕI:
1. Trích xuất danh sách thiết bị đầy đủ:
   - stt: Số thứ tự (1, 2, 3...)
   - serial_number: Mã sản phẩm hoặc số Serial (VD: RFR900NA55RSWLBA574, EF550RA4LAWBHI677)
   - loi_ban_dau: Lỗi ban đầu khách yêu cầu kiểm tra (song ngữ nếu có, VD: Cò bắn không nhả / 扳機扣住不鬆)
   - bo_phan: Bộ phận gửi sửa (VD: CNC, Xi treo / 吊鍍)
   - hang_muc_kiem_tra: Hạng mục đã kiểm tra & xử lý (VD: Đã kiểm tra nhưng không có lỗi / 已檢查，未發現異常, hoặc Kiểm tra hiển thị màn hình & đầu quét barcode)
   - ket_qua: Kết quả kiểm tra (mặc định "Đạt / 合格")
   - don_gia: Đơn giá nếu có
   - tong_gia: Thành tiền nếu có
2. Trích xuất thông tin Hợp đồng / Biên bản:
   - so_hop_dong: Mã đơn / Số HĐ (VD: HM-FINEPRO-20261009.001 hoặc HM-20261009.001)
   - so_bien_ban: Số biên bản (VD: 001/2026/BBNT-SC)
   - ngay_nghiem_thu: Ngày (DD/MM/YYYY)
   - ben_a: Tên công ty khách hàng (VD: CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN))
3. Tổng tiền: Đọc hoặc tính tổng tiền bằng số (VD: 13700000).
4. Ghi chú (ghi_chu): BỎ TRỐNG HOÀN TOÀN (""), TUYỆT ĐỐI KHÔNG ghi giá tiền và KHÔNG tự sinh cam kết hay lặp lại linh kiện vào ghi chú.

ĐỊNH DẠNG JSON ĐẦU RA BẮT BUỘC:
{
  "tieu_de_viet": "BIÊN BẢN KIỂM TRA & NGHIỆM THU THIẾT BỊ SAU SỬA CHỮA",
  "tieu_de_trung": "維修設備檢查及驗收記錄",
  "so_hop_dong": "HM-FINEPRO-20261009.001",
  "so_bien_ban": "001/2026/BBNT-SC",
  "ngay_nghiem_thu": "09/10/2026",
  "ben_a": {
    "ten_cong_ty": "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)",
    "ten_tieng_trung": "CX Technology (VN) 股份有限公司",
    "dia_chi": "KCX Tân Thuận, Quận 7, TP.HCM"
  },
  "hang_muc": [
    {
      "stt": 1,
      "serial_number": "RFR900NA55RSWLBA574",
      "loi_ban_dau": "Cò bắn không nhả\\n扳機扣住不鬆",
      "bo_phan": "CNC",
      "hang_muc_kiem_tra": "Đã kiểm tra nhưng không có lỗi\\n已檢查，未發現異常",
      "ket_qua": "Đạt / 合格",
      "don_gia": "550.000 VND",
      "tong_gia": "550.000 VND",
      "can_anh": true
    }
  ],
  "tong_tien_so": "13.700.000",
  "ghi_chu": ""
}
CHỈ TRẢ VỀ JSON HỢP LỆ, KHÔNG KÈM THEO BẤT KỲ VĂN BẢN GIẢI THÍCH NÀO NGOÀI JSON.
"""


def parse_acceptance_text(raw_input_text: str, preferred_model: str = None) -> Dict[str, Any]:
    """Parse raw acceptance or delivery note text with Gemini API, fallback to local heuristics."""
    models_to_try = []
    if preferred_model and preferred_model in CANDIDATE_MODELS:
        models_to_try.append(preferred_model)
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    keys_to_try = [DEFAULT_GEMINI_KEY] + [k for k in DEFAULT_GEMINI_KEYS if k != DEFAULT_GEMINI_KEY]

    last_error = None
    for key in keys_to_try:
        for model_name in models_to_try:
            try:
                logger.info("Attempting acceptance parsing with model: %s (key %s...)", model_name, key[:8])
                parsed = _call_gemini_acceptance(raw_input_text, model_name, key=key)
                if parsed and isinstance(parsed, dict) and parsed.get("hang_muc"):
                    parsed["_executed_model"] = model_name
                    return parsed
            except urllib.error.HTTPError as exc:
                if exc.code == 403:
                    logger.warning("Key %s... returned 403. Rotating to next key.", key[:8])
                    break
                elif exc.code in (404, 429, 503):
                    logger.warning("Model %s returned HTTP %d on key %s... Trying next model.", model_name, exc.code, key[:8])
                    last_error = exc
                    continue
                else:
                    last_error = exc
            except Exception as exc:
                logger.warning("Model %s failed (%s). Trying next model.", model_name, exc)
                last_error = exc

    logger.warning("All Gemini models/keys failed (%s). Using fallback local parser.", last_error)
    return _fallback_local_acceptance_parse(raw_input_text)


def _call_gemini_acceptance(text: str, model_name: str, key: str = None) -> Dict[str, Any]:
    """Execute Gemini generateContent HTTP request."""
    api_key = key or DEFAULT_GEMINI_KEY
    endpoint = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
    payload = {
        "contents": [
            {
                "role": "user",
                "parts": [
                    {"text": ACCEPTANCE_SYSTEM_PROMPT},
                    {"text": f"Dữ liệu cần bóc tách:\n---\n{text}\n---"}
                ]
            }
        ],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    req_data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        endpoint,
        data=req_data,
        headers={"Content-Type": "application/json"},
        method="POST"
    )

    # Dynamic timeout: fast 6s timeout for 3.8/3.7, 25s for standard models
    timeout_sec = 6 if ("3.8" in model_name or "3.7" in model_name) else 25
    with urllib.request.urlopen(req, timeout=timeout_sec) as resp:
        res_bytes = resp.read()
        res_json = json.loads(res_bytes.decode("utf-8"))

    candidates = res_json.get("candidates") or []
    if not candidates:
        raise ValueError("No candidates returned from Gemini API")

    content_text = candidates[0]["content"]["parts"][0]["text"]
    cleaned_json_str = _clean_json_codeblock(content_text)
    return json.loads(cleaned_json_str)


def _clean_json_codeblock(raw: str) -> str:
    """Strip markdown code fence."""
    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text)
    return text.strip()


def _fallback_local_acceptance_parse(raw_text: str) -> Dict[str, Any]:
    """Deterministic local parser for delivery receipt text or table (TSV / Excel / Text)."""
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]

    # Extract contract number
    so_hd = "HM-FINEPRO-20261013.001"
    match_hd = re.search(r"(?:Hợp Đồng|合約編號|PO|HĐ)[：:\s]*([A-Za-z0-9\.\-_]+)", raw_text, re.IGNORECASE)
    if match_hd:
        so_hd = match_hd.group(1).strip()

    # Extract Party A
    party_a = "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)"
    party_a_trung = "CX Technology (VN) 股份有限公司"
    party_a_addr = "KCX Tân Thuận, Quận 7, TP.HCM"
    match_a = re.search(r"(?:Bên Mua|甲方|Khách hàng)[：:\s]*([^\n\r]+)", raw_text, re.IGNORECASE)
    if match_a:
        cand = match_a.group(1).strip()
        cand = re.sub(r"^[\(（][^）\)]*[\)）]", "", cand).strip()
        cand = re.sub(r"^(?:Bên Mua|买方|買方|Khách hàng)[:：\s]*", "", cand).strip()
        if len(cand) > 3:
            party_a = cand
            if "minh phát" in party_a.lower():
                party_a_trung = "MINH PHÁT 科技責任有限公司"
                party_a_addr = "KCN Hòa Phú, Bắc Ninh"
            elif "cx" not in party_a.lower():
                party_a_trung = party_a
                party_a_addr = ""

    # Extract Party B if specified
    ben_b_custom = None
    match_b = re.search(r"(?:Bên Bán|乙方|Đơn vị sửa chữa)[：:\s]*([^\n\r]+)", raw_text, re.IGNORECASE)
    if match_b:
        cand_b = match_b.group(1).strip()
        cand_b = re.sub(r"^[\(（][^）\)]*[\)）]", "", cand_b).strip()
        cand_b = re.sub(r"^(?:Bên Bán|卖方|賣方)[:：\s]*", "", cand_b).strip()
        if len(cand_b) > 3:
            ben_b_custom = {
                "ten_cong_ty": cand_b,
                "ten_tieng_trung": "宏明工業設備有限公司" if "thiết bị công nghiệp" in cand_b.lower() else "宏明環球有限公司",
                "dia_chi": "Lầu 2, số 68 Nguyễn Huệ, Phường Sài Gòn, TP. Hồ Chí Minh",
                "mst": "0318280158",
                "nguoi_dai_dien": "TANG WEI LI",
                "chuc_vu": "Tổng Giám Đốc"
            }

    # Extract total amount if explicitly stated in text
    total_val = 0
    match_tot = re.search(r"(?:TỔNG THÀNH TIỀN|TỔNG GIÁ TRỊ|TỔNG TIỀN|總計|總價)[^0-9]*([\d\.,]+)", raw_text, re.IGNORECASE)
    if match_tot:
        total_val = parse_positive_integer(match_tot.group(1))

    # 1. Attempt TSV / Tabular data parsing (Direct Excel copy-paste)
    tsv_items = []
    for line in lines:
        if "\t" in line:
            parts = [p.strip() for p in line.split("\t") if p.strip()]
            if not parts:
                continue
            # Header check
            if any(h in parts[0].upper() for h in ["STT", "NO.", "MÃ"]) or any("TÊN SẢN PHẨM" in p.upper() for p in parts):
                continue
            # Total check
            if any(t in parts[0].upper() for t in ["TỔNG", "TOTAL", "總計", "總價"]):
                if total_val == 0:
                    for p in parts:
                        nums = re.findall(r"(\d{1,3}(?:[.,]\d{3})+|\d{6,10})", p)
                        if nums:
                            total_val = parse_positive_integer(nums[-1])
                continue

            if parts[0].isdigit() and len(parts) >= 3:
                stt = int(parts[0])
                sn = parts[1]
                loi = parts[2] if len(parts) > 2 else ""
                bp = parts[3] if len(parts) > 3 else "CNC"
                kq_or_hm = parts[4] if len(parts) > 4 else "Đạt / 合格"

                # Standardize inspection / repair detail
                hm_lower = kq_or_hm.lower()
                if "no error" in hm_lower:
                    hm = "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常 (No error found)"
                elif "barcode" in hm_lower:
                    hm = "Kiểm tra hiển thị màn hình & đầu quét Barcode\n檢查螢幕顯示及條碼掃描模組"
                elif "adapter" in hm_lower:
                    hm = "Kiểm tra và xử lý gá chuyển đổi Mounting Adapter\n檢查及更換轉接支架"
                elif "power" in hm_lower or "button" in hm_lower:
                    hm = "Kiểm tra và xử lý cụm nút nguồn Power Button\n檢修電源按鍵模組"
                elif "signal" in hm_lower:
                    hm = "Kiểm tra và xử lý tín hiệu kết nối Signal Error\n檢修訊號接收模組"
                elif "scanner" in hm_lower:
                    hm = "Kiểm tra và thay thế linh kiện đầu đọc Scanner Error\n更換掃描頭模組"
                elif "touch" in hm_lower:
                    hm = "Kiểm tra và thay thế cụm màn hình cảm ứng Touchscreen\n更換觸控螢幕模組"
                elif "connection" in hm_lower:
                    hm = "Kiểm tra bo mạch và kết nối truyền thông Connection Error\n檢修通訊連線模組"
                elif "battery" in hm_lower:
                    hm = "Kiểm tra và thay thế cụm pin sạc Battery Error\n檢修及更換電池模組"
                else:
                    hm = kq_or_hm

                # Extract prices from remaining parts
                price_candidates = []
                for p in parts[4:]:
                    nums = re.findall(r"(\d{1,3}(?:[.,]\d{3})+|\d{5,10})", p)
                    for n in nums:
                        val = parse_positive_integer(n)
                        if val > 10000:
                            price_candidates.append(val)

                don_gia_str = ""
                tong_gia_str = ""
                if len(price_candidates) >= 2:
                    don_gia_str = f"{price_candidates[-2]:,}".replace(",", ".") + " VND"
                    tong_gia_str = f"{price_candidates[-1]:,}".replace(",", ".") + " VND"
                elif len(price_candidates) == 1:
                    tong_gia_str = f"{price_candidates[0]:,}".replace(",", ".") + " VND"

                tsv_items.append({
                    "stt": stt,
                    "serial_number": sn,
                    "loi_ban_dau": loi,
                    "bo_phan": bp,
                    "hang_muc_kiem_tra": hm,
                    "ket_qua": "Đạt / 合格",
                    "don_gia": don_gia_str,
                    "tong_gia": tong_gia_str,
                    "can_anh": False
                })

    if tsv_items:
        items = tsv_items
    else:
        # 2. Line-by-line regex parser
        items = []
        item_counter = 1
        for line in lines:
            sn_match = re.search(r"([A-Z0-9]{10,25})", line)
            if sn_match:
                sn = sn_match.group(1)
                numbers = re.findall(r"(\d{1,3}(?:[.,]\d{3})+|\d{6,9})", line)
                price_str = ""
                if numbers:
                    val = parse_positive_integer(numbers[-1])
                    if val > 10000:
                        price_str = f"{val:,}".replace(",", ".") + " VND"
                        total_val += val

                bp = "CNC"
                if "xi treo" in line.lower() or "吊鍍" in line:
                    bp = "Xi treo\n吊鍍"

                loi = "Cần kiểm tra kỹ thuật"
                if "cò" in line.lower() or "扳機" in line or "扳机" in line or "bản" in line.lower():
                    loi = "Có bản không nhả\n扳扣拉不鬆"
                    if "sạc" in line.lower() or "充電" in line or "充电" in line:
                        loi = "Cò bắn & Sạc không vô điện\n扳機故障及無法充電"
                elif "màn hình" in line.lower() or "螢幕" in line:
                    loi = "Màn hình lỗi\n螢幕故障"
                elif "nút nguồn" in line.lower() or "按鍵" in line:
                    loi = "Lỗi nút nguồn\n電源按鍵故障"
                elif "tín hiệu" in line.lower() or "訊號" in line:
                    loi = "Không nhận tín hiệu\n無法接收訊號"
                elif "đầu đọc" in line.lower() or "掃描" in line:
                    loi = "Lỗi đầu đọc\n掃描頭故障"
                elif "kết nối" in line.lower() or "連線" in line:
                    loi = "Không kết nối được\n無法連線"
                elif "pin" in line.lower() or "電池" in line:
                    loi = "Pin không sạc được\n電池無法充電"

                hm = "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常"
                if "barcode" in line.lower():
                    hm = "Kiểm tra hiển thị màn hình & đầu quét barcode\n檢查螢幕顯示及條碼掃描功能"
                elif "adapter" in line.lower():
                    hm = "Kiểm tra và xử lý gá chuyển đổi Mounting Adapter\n檢查及更換轉接支架"

                items.append({
                    "stt": item_counter,
                    "serial_number": sn,
                    "loi_ban_dau": loi,
                    "bo_phan": bp,
                    "hang_muc_kiem_tra": hm,
                    "ket_qua": "Đạt / 合格",
                    "don_gia": price_str,
                    "tong_gia": price_str,
                    "can_anh": False
                })
                item_counter += 1

    if not items:
        # Default sample fallback
        items = [
            {
                "stt": 1,
                "serial_number": "RFR900NA55RSWLBA574",
                "loi_ban_dau": "Cò bắn không nhả\n扳機扣住不鬆",
                "bo_phan": "CNC",
                "hang_muc_kiem_tra": "Đã kiểm tra nhưng không có lỗi\n已檢查，未發現異常",
                "ket_qua": "Đạt / 合格",
                "don_gia": "550.000 VND",
                "tong_gia": "550.000 VND",
                "can_anh": False
            }
        ]
        total_val = 550000

    # Auto sum total if not parsed
    if total_val == 0 and items:
        sum_p = 0
        for it in items:
            p = parse_positive_integer(it.get("tong_gia") or it.get("don_gia") or 0)
            sum_p += p
        if sum_p > 0:
            total_val = sum_p

    # Build smart note 3 (Chi tiết kết quả kiểm tra & linh kiện thay thế)
    no_error_stts = []
    repair_notes = []
    for it in items:
        hm = it.get("hang_muc_kiem_tra", "")
        sn = it.get("serial_number", "")
        stt = it.get("stt", "")
        if "không có lỗi" in hm.lower() or "no error" in hm.lower():
            no_error_stts.append(str(stt))
        else:
            first_hm = hm.split("\n")[0].split("/")[0].strip()
            repair_notes.append(f"   - STT {stt} ({sn}): {first_hm}")

    ghi_chu_3_lines = ["Chi tiết kết quả kiểm tra & linh kiện thay thế / 檢驗結果與更換零件明細:"]
    if no_error_stts:
        stt_joined = ", ".join(no_error_stts)
        ghi_chu_3_lines.append(f"   - STT {stt_joined}: Không phát hiện lỗi phát sinh thêm (No error founds) / 檢查未發現異常。")
    for rn in repair_notes:
        ghi_chu_3_lines.append(rn)
    smart_ghi_chu_3 = "\n".join(ghi_chu_3_lines)

    result_dict = {
        "_executed_model": "LocalHeuristicsParser",
        "tieu_de_viet": "BIÊN BẢN KIỂM TRA & NGHIỆM THU THIẾT BỊ SAU SỬA CHỮA",
        "tieu_de_trung": "維修設備檢查及驗收記錄",
        "so_hop_dong": so_hd,
        "so_bien_ban": "001/2026/BBNT-SC",
        "ngay_nghiem_thu": datetime.now().strftime("%d/%m/%Y"),
        "ben_a": {
            "ten_cong_ty": party_a,
            "ten_tieng_trung": party_a_trung,
            "dia_chi": party_a_addr
        },
        "hang_muc": items,
        "tong_so_luong": f"{len(items)} Thiết bị",
        "tong_tien_so": f"{total_val:,}".replace(",", "."),
        "ghi_chu_1": "Thiết bị đã được kiểm tra chức năng, bảo trì/sửa chữa và chạy thử đạt tiêu chuẩn kỹ thuật trước khi bàn giao / 設備已完成功能檢查、維修更換零件並測試正常運作後交付。",
        "ghi_chu_3": smart_ghi_chu_3
    }
    if ben_b_custom:
        result_dict["ben_b"] = ben_b_custom

    return result_dict


INSTALLATION_SYSTEM_PROMPT = """Bạn là chuyên gia phân tích và bóc tách dữ liệu Biên bản nghiệm thu lắp đặt & bàn giao đưa vào sử dụng song ngữ Việt - Trung cho CÔNG TY TNHH TOÀN CẦU HONG MING (宏明環球有限公司).

NHIỆM VỤ:
Phân tích văn bản đầu vào (nội dung hợp đồng, bảng hạng mục bàn giao lắp đặt, ghi chú nghiệm thu tiến độ) và trích xuất thành định dạng JSON chuẩn 100%.

QUY TẮC CỐT LÕI:
1. Trích xuất thông tin hợp đồng & các bên:
   - so_hop_dong: Mã hợp đồng (VD: HMV151020250001, HMV100120260001)
   - so_bien_ban: Số biên bản (VD: 002_BIEN_BAN_NGHIEM_THU_BAN_GIAO_CX_ hoặc BB-NT-2026-___)
   - ngay_ky_hop_dong: Ngày ký HĐ (DD/MM/YYYY)
   - ngay_nghiem_thu: Ngày nghiệm thu (DD/MM/YYYY)
   - gia_tri_hop_dong: Tổng giá trị HĐ (VD: 234.000.000 VNĐ)
   - ten_he_thong: Tên hệ thống / dự án tiếng Việt
   - ten_he_thong_trung: Tên hệ thống / dự án tiếng Trung
   - ben_a: { ten_cong_ty, ten_tieng_trung, dia_chi, mst }
   - ben_b: { ten_cong_ty, ten_tieng_trung, dia_chi, mst, nguoi_dai_dien, chuc_vu }
2. Trích xuất căn cứ pháp lý:
   - can_cu_phu_luc: Căn cứ phụ lục báo giá
   - can_cu_phu_luc_trung: Tiếng Trung
   - can_cu_ky_thuat: Căn cứ biên bản kỹ thuật / chạy thử
   - can_cu_ky_thuat_trung: Tiếng Trung
3. Trích xuất danh mục hạng mục nghiệm thu lắp đặt (hang_muc):
   - stt: Số thứ tự (1, 2, 3...)
   - hang_muc: Tên hạng mục ngắn gọn song ngữ (CHỈ lấy cột Hạng mục / 項目類別, TUYỆT ĐỐI KHÔNG ghép thêm tên sản phẩm hay nội dung dài dòng vào). Ví dụ:
     "1. Xác nhận mở vụ án / 立案確認, Yêu cầu xác nhận / 需求確認"
     "2. Mỗi máy có thể quét các nhân viên vận hành máy khác nhau khi vào và ra khỏi trạm.\\n進出站時每個機台都可以掃不同的機台操作者Máy quét mã vạch được tối ưu hóa để lựa chọn máy bằng cách nhấp chuột.掃描機台條碼優化為點選機台方式。"
   - xac_nhan_lap_dat: "Xác nhận thiết bị đã được đưa vào sử dụng.\\n確認已投入使用"
   - ket_qua_xac_nhan: "Đã nghiệm thu, xác nhận đạt yêu cầu.\\n驗收合格，確認符合需求"
4. Trích xuất đề nghị thanh toán & các giai đoạn:
   - giai_doan_thanh_toan: Tên giai đoạn đề nghị thanh toán đợt này (VD: Giai đoạn 2 (Sau khi triển khai lắp đặt đưa vào sử dụng - 40%))
   - so_tien_thanh_toan: Số tiền đề nghị thanh toán đợt này (VD: 93.600.000 VND)
   - so_tien_chu_vn: Bằng chữ tiếng Việt
   - so_tien_chu_tw: Bằng chữ tiếng Trung
   - payment_stages: [ { stage, condition, rate, amount } ]

ĐỊNH DẠNG JSON ĐẦU RA BẮT BUỘC:
{
  "so_hop_dong": "HMV151020250001",
  "so_bien_ban": "002_BIEN_BAN_NGHIEM_THU_BAN_GIAO_CX_",
  "ngay_ky_hop_dong": "09/12/2025",
  "ngay_nghiem_thu": "09/10/2026",
  "gia_tri_hop_dong": "234.000.000 VNĐ",
  "ten_he_thong": "RFID 硬體 (RFID Hardware) – Phần mềm tùy chỉnh chức năng hệ thống RFID",
  "ten_he_thong_trung": "RFID 硬體（RFID Hardware）－ RFID 系統客製化軟體",
  "ben_a": {
    "ten_cong_ty": "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)",
    "ten_tieng_trung": "CX TECHNOLOGY (VN) 股份有限公司",
    "dia_chi": "Số Lô CT 10b-12-14-16-18-20-22-24a, Khu C, KCX Tân Thuận, Phường Tân Thuận, Tp. Hồ Chí Minh",
    "mst": "0300737556"
  },
  "ben_b": {
    "ten_cong_ty": "CÔNG TY TNHH TOÀN CẦU HONG MING",
    "ten_tieng_trung": "宏明環球有限公司",
    "dia_chi": "Lầu 2, số 68 Nguyễn Huệ, Phường Sài Gòn, Thành phố Hồ Chí Minh, Việt Nam",
    "mst": "0318280158",
    "nguoi_dai_dien": "Ông TANG WEI LI",
    "chuc_vu": "Tổng giám đốc"
  },
  "can_cu_phu_luc": "Phụ lục 1 – Báo giá khai phá chức năng ưu hóa tùy chỉnh (V1.2.3_VND) kèm theo hợp đồng",
  "can_cu_phu_luc_trung": "附件一「RFID系統客製化功能開發報價單（V1.2.3_VND）」",
  "can_cu_ky_thuat": "Biên bản kiểm tra kỹ thuật ngày 04/03/2026 xác nhận hoàn thành lắp đặt và kết quả chạy thử, điều chỉnh, theo dõi vận hành sau đó.",
  "can_cu_ky_thuat_trung": "依據 2026/03/04《技術檢驗單》所確認之安裝完成情況，以及其後之試運行、調整及運行觀察結果。",
  "giai_doan_thanh_toan": "Giai đoạn 2 (Sau khi triển khai lắp đặt đưa vào sử dụng - 40%)",
  "so_tien_thanh_toan": "93.600.000 VND",
  "so_tien_chu_vn": "Chín mươi ba triệu sáu trăm nghìn đồng chẵn.",
  "so_tien_chu_tw": "玖仟參佰陸拾萬越盾整",
  "hang_muc": [
    {
      "stt": 1,
      "ten_san_pham": "Concept Development",
      "hang_muc": "1. Xác nhận mở vụ án / 立案確認\\n2. Yêu cầu xác nhận / 需求確認",
      "noi_dung": "Nộp hồ sơ dự án, họp thảo luận: lập kế hoạch kiến trúc dự án, lập kế hoạch quy trình và xác nhận nhu cầu.",
      "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\\n確認已投入使用",
      "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu và chính thức đưa vào sử dụng.\\n驗收合格，確認正式投入使用。"
    }
  ],
  "payment_stages": [
    {
      "stage": "GIAI ĐOẠN 1\\n第一期",
      "condition": "Sau khi ký kết hợp đồng\\n簽約完成時",
      "rate": "40%",
      "amount": "93.600.000 VND"
    },
    {
      "stage": "GIAI ĐOẠN 2 (ĐỢT NÀY)\\n第二期（本期款）",
      "condition": "Sau khi triển khai sản phẩm đưa vào sử dụng\\n產品安裝後",
      "rate": "40%",
      "amount": "93.600.000 VND"
    },
    {
      "stage": "GIAI ĐOẠN 3 (QUYẾT TOÁN)\\n第三期（尾款）",
      "condition": "Khi nghiệm thu đạt tiêu chuẩn\\n驗收合格時",
      "rate": "20%",
      "amount": "46.800.000 VND"
    },
    {
      "stage": "TỔNG CỘNG HỢP ĐỒNG\\n合約總計",
      "condition": "100% Giá trị Hợp đồng",
      "rate": "100%",
      "amount": "234.000.000 VNĐ"
    }
  ]
}
CHỈ TRẢ VỀ JSON HỢP LỆ, KHÔNG KÈM THEO BẤT KỲ VĂN BẢN GIẢI THÍCH NÀO NGOÀI JSON.
"""


def parse_installation_acceptance_text(raw_input_text: str, preferred_model: str = None) -> Dict[str, Any]:
    """Parse raw installation acceptance text via Gemini API, fallback to heuristics."""
    models_to_try = []
    if preferred_model and preferred_model in CANDIDATE_MODELS:
        models_to_try.append(preferred_model)
    for m in CANDIDATE_MODELS:
        if m not in models_to_try:
            models_to_try.append(m)

    api_keys_to_try = [k for k in DEFAULT_GEMINI_KEYS if k]
    if not api_keys_to_try and DEFAULT_GEMINI_KEY:
        api_keys_to_try.append(DEFAULT_GEMINI_KEY)

    last_error = None
    for api_key in api_keys_to_try:
        for model in models_to_try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
            payload = {
                "contents": [
                    {
                        "role": "user",
                        "parts": [
                            {"text": INSTALLATION_SYSTEM_PROMPT},
                            {"text": f"VĂN BẢN ĐẦU VÀO CẦN PHÂN TÍCH:\n\n{raw_input_text}"}
                        ]
                    }
                ],
                "generationConfig": {
                    "temperature": 0.1,
                    "responseMimeType": "application/json"
                }
            }

            try:
                req = urllib.request.Request(
                    url,
                    data=json.dumps(payload).encode("utf-8"),
                    headers={"Content-Type": "application/json"},
                    method="POST"
                )
                with urllib.request.urlopen(req, timeout=30) as response:
                    res_body = response.read().decode("utf-8")
                    res_json = json.loads(res_body)
                    text_content = res_json["candidates"][0]["content"]["parts"][0]["text"]
                    clean_text = text_content.strip()
                    if clean_text.startswith("```json"):
                        clean_text = clean_text[7:]
                    if clean_text.endswith("```"):
                        clean_text = clean_text[:-3]

                    parsed = json.loads(clean_text.strip())
                    parsed["_executed_model"] = model
                    
                    # Sanitize & normalize payment_stages
                    raw_stages = parsed.get("payment_stages") or []
                    if not raw_stages:
                        raw_stages = _extract_payment_stages_from_text(raw_input_text)
                    
                    cleaned_stages = []
                    for idx, st in enumerate(raw_stages, start=1):
                        s_name = str(st.get("stage") or "").strip()
                        # Filter out summary/total rows
                        if any(kw in s_name.lower() for kw in ["tổng cộng", "tổng giá trị", "合計", "總計"]) or str(st.get("rate") or "").strip() == "100%":
                            continue
                        
                        m = re.search(r'(?:GIAI\s*ĐOẠN|ĐỢT)\s*([0-9]+)', s_name, re.IGNORECASE)
                        s_num = int(m.group(1)) if m else idx
                        
                        is_sel = st.get("selected")
                        if is_sel is None:
                            is_sel = (s_num in (2, 3))
                        
                        cleaned_stages.append({
                            "stage_id": str(st.get("stage_id") or s_num),
                            "stage_num": s_num,
                            "stage": s_name or f"GIAI ĐOẠN {s_num}",
                            "condition": str(st.get("condition") or "").strip(),
                            "rate": str(st.get("rate") or "").strip(),
                            "amount": str(st.get("amount") or "").strip(),
                            "selected": bool(is_sel)
                        })
                    
                    if not cleaned_stages:
                        cleaned_stages = _extract_payment_stages_from_text(raw_input_text)
                        
                    parsed["payment_stages"] = cleaned_stages
                    return parsed
            except Exception as e:
                last_error = e
                logger.warning("Gemini model %s with key %s... failed: %s", model, api_key[:8] if api_key else "", e)

    # Heuristic fallback if Gemini is offline
    logger.info("Using local heuristic fallback for installation acceptance: %s", last_error)
    return _fallback_parse_installation_text(raw_input_text)


def extract_text_from_upload_file(file_bytes: bytes, filename: str) -> str:
    """Extract clean text content from uploaded files (.docx, .xlsx, .txt, .md, .json, .csv)."""
    import io
    fname = (filename or "").lower()

    if fname.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        table_lines = []
        for t_idx, tbl in enumerate(doc.tables):
            for r in tbl.rows:
                cells = [c.text.replace("\n", " ").strip() for c in r.cells]
                # dedup merged cells
                dedup = []
                for cell in cells:
                    if not dedup or cell != dedup[-1]:
                        dedup.append(cell)
                if any(dedup):
                    table_lines.append(" | ".join(dedup))
        return "\n".join(paragraphs + ["\n--- BẢNG DỮ LIỆU ---"] + table_lines)

    elif fname.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"--- SHEET: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if non_empty:
                    lines.append(" | ".join(non_empty))
        return "\n".join(lines)

    elif fname.endswith((".txt", ".md", ".json", ".csv", ".log")):
        for enc in ["utf-8", "utf-8-sig", "cp1258", "latin-1"]:
            try:
                return file_bytes.decode(enc)
            except Exception:
                continue
        return file_bytes.decode("utf-8", errors="replace")

    elif fname.endswith(".pdf"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = [p.extract_text() for p in reader.pages if p.extract_text()]
            return "\n".join(pages_text)
        except Exception:
            return ""

    return ""


def _extract_payment_stages_from_text(raw_text: str) -> List[Dict[str, Any]]:
    """Extract all contract payment stages from raw text with robust regex."""
    stages = []
    # Match pattern: 1. Sau khi ký kết... 40% = 93.600.000 VNĐ ... 1. 簽約完成時...
    pattern = r'(?P<idx>[1-5])\.\s*(?P<vn>[^\n]+?thanh toán\s*(?P<rate>\d+[\.,]?\d*%).*?=\s*(?P<amount>[\d\.,]+)\s*(?:VNĐ|VND|越盾)[^\n]*)(?:\n(?P=idx)\.\s*(?P<tw>[^\n]+))?'
    matches = list(re.finditer(pattern, raw_text, re.IGNORECASE))
    
    zh_stage_map = {1: "第一期", 2: "第二期", 3: "第三期", 4: "第四期", 5: "第五期"}
    
    if matches:
        for m in matches:
            idx = int(m.group("idx"))
            rate = m.group("rate").replace(",", ".")
            amount_str = m.group("amount").strip()
            vn_text = m.group("vn").strip()
            tw_text = (m.group("tw") or "").strip()
            
            cond_vn = re.sub(r'[:：]?\s*thanh toán.*$', '', vn_text, flags=re.IGNORECASE).strip()
            cond_tw = re.sub(r'[:：]?\s*支付.*$', '', tw_text, flags=re.IGNORECASE).strip()
            condition_combined = f"{cond_vn}\n{cond_tw}" if cond_tw else cond_vn
            
            stages.append({
                "stage_id": str(idx),
                "stage_num": idx,
                "stage": f"GIAI ĐOẠN {idx}\n{zh_stage_map.get(idx, f'第{idx}期')}",
                "condition": condition_combined,
                "rate": rate,
                "amount": f"{amount_str} VND",
                "selected": (idx == 2 or idx == 3) # default selected for installation/acceptance
            })

    if not stages:
        # Default 3 stages for CX Technology HMV151020250001
        stages = [
            {
                "stage_id": "1",
                "stage": "GIAI ĐOẠN 1\n第一期",
                "condition": "Sau khi ký kết hợp đồng\n簽約完成時",
                "rate": "40%",
                "amount": "93.600.000 VND",
                "selected": False
            },
            {
                "stage_id": "2",
                "stage": "GIAI ĐOẠN 2\n第二期",
                "condition": "Sau khi triển khai sản phẩm đưa vào sử dụng\n產品安裝後",
                "rate": "40%",
                "amount": "93.600.000 VND",
                "selected": True
            },
            {
                "stage_id": "3",
                "stage": "GIAI ĐOẠN 3\n第三期",
                "condition": "Khi nghiệm thu đạt tiêu chuẩn\n驗收合格時",
                "rate": "20%",
                "amount": "46.800.000 VND",
                "selected": True
            }
        ]
    return stages


def _fallback_parse_installation_text(raw_input_text: str) -> Dict[str, Any]:
    """Parse text using local rules and regex when Gemini API is unavailable."""
    m_hd = re.search(r'(HMV\d+)', raw_input_text)
    so_hop_dong = m_hd.group(1) if m_hd else "HMV151020250001"

    m_val = re.search(r'([\d\.,]+)\s*(?:VNĐ|VND|越盾)', raw_input_text)
    gia_tri_hop_dong = f"{m_val.group(1)} VNĐ" if m_val else "234.000.000 VNĐ"

    payment_stages = _extract_payment_stages_from_text(raw_input_text)

    # Extract items if table lines exist
    items = []
    lines = raw_input_text.splitlines()
    for line in lines:
        if "|" in line:
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 3 and parts[0].isdigit():
                # Prefer column 2 (Hạng mục) if available
                hm_val = parts[2] if len(parts) >= 3 else parts[1]
                items.append({
                    "stt": int(parts[0]),
                    "hang_muc": hm_val,
                    "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用",
                    "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"
                })

    if not items:
        items = [
            {"stt": 1, "hang_muc": "1. Xác nhận mở vụ án / 立案確認, Yêu cầu xác nhận / 需求確認", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"},
            {"stt": 2, "hang_muc": "2. Mỗi máy có thể quét các nhân viên vận hành máy khác nhau khi vào và ra khỏi trạm.\n進出站時每個機台都可以掃不同的機台操作者Máy quét mã vạch được tối ưu hóa để lựa chọn máy bằng cách nhấp chuột.掃描機台條碼優化為點選機台方式。", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"},
            {"stt": 3, "hang_muc": "3. Quét nhiều ID nhân viên để vào và ra khỏi trạm.\n進出站可刷複數人員工號。", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入 sử dụng", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"},
            {"stt": 4, "hang_muc": "4. QR條碼掃描\nQuét mã vạch QR", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"},
            {"stt": 5, "hang_muc": "5. 特定人員解除綁定權限控管\nHủy ràng buộc quyền cho nhân viên cụ thể", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"},
            {"stt": 6, "hang_muc": "6. RFID離線模式可以自動切換\nChế độ ngoại tuyến RFID có thể chuyển đổi tự động", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"},
            {"stt": 7, "hang_muc": "7. 增加Run Card查詢條件\nThêm điều kiện truy vấn Run Card", "xac_nhan_lap_dat": "Xác nhận thiết bị đã được đưa vào sử dụng.\n確認已投入使用", "ket_qua_xac_nhan": "Đã nghiệm thu, xác nhận đạt yêu cầu.\n驗收合格，確認符合需求"}
        ]

    return {
        "_executed_model": "LocalHeuristicsFallback",
        "so_hop_dong": so_hop_dong,
        "so_bien_ban": f"002_BIEN_BAN_NGHIEM_THU_BAN_GIAO_{so_hop_dong}_",
        "ngay_ky_hop_dong": "09/12/2025",
        "ngay_nghiem_thu": datetime.now().strftime("%d/%m/%Y"),
        "gia_tri_hop_dong": gia_tri_hop_dong,
        "ten_he_thong": "RFID 硬體 (RFID Hardware) – Phần mềm tùy chỉnh chức năng hệ thống RFID",
        "ten_he_thong_trung": "RFID 硬體（RFID Hardware）－ RFID 系統客製化軟體",
        "ben_a": {
            "ten_cong_ty": "CÔNG TY CỔ PHẦN CX TECHNOLOGY (VN)",
            "ten_tieng_trung": "CX TECHNOLOGY (VN) 股份有限公司",
            "dia_chi": "Số Lô CT 10b-12-14-16-18-20-22-24a, Khu C, KCX Tân Thuận, Phường Tân Thuận, Tp. Hồ Chí Minh",
            "mst": "0300737556"
        },
        "ben_b": {
            "ten_cong_ty": "CÔNG TY TNHH TOÀN CẦU HONG MING",
            "ten_tieng_trung": "宏明環球有限公司",
            "dia_chi": "Lầu 2, số 68 Nguyễn Huệ, Phường Sài Gòn, Thành phố Hồ Chí Minh, Việt Nam",
            "mst": "0318280158",
            "nguoi_dai_dien": "Ông TANG WEI LI",
            "chuc_vu": "Tổng giám đốc"
        },
        "can_cu_phu_luc": "Phụ lục 1 – Báo giá khai phá chức năng ưu hóa tùy chỉnh (V1.2.3_VND) kèm theo hợp đồng",
        "can_cu_phu_luc_trung": "附件一「RFID系統客製化功能開發報價單（V1.2.3_VND）」",
        "can_cu_ky_thuat": "Biên bản kiểm tra kỹ thuật ngày 04/03/2026 xác nhận hoàn thành lắp đặt và kết quả chạy thử, điều chỉnh, theo dõi vận hành sau đó.",
        "can_cu_ky_thuat_trung": "依據 2026/03/04《技術檢驗單》所確認之安裝完成情況，以及其後之試運行、調整及運行觀察結果。",
        "giai_doan_thanh_toan": "Giai đoạn 2 & 3 (Triển khai lắp đặt và nghiệm thu hoàn thành - 60%)",
        "so_tien_thanh_toan": "140.400.000 VND",
        "so_tien_chu_vn": "Một trăm bốn mươi triệu bốn trăm nghìn đồng chẵn.",
        "so_tien_chu_tw": "越南盾壹億肆仟零肆拾萬整",
        "hang_muc": items,
        "payment_stages": payment_stages
    }


def parse_installation_acceptance_file(file_bytes: bytes, filename: str, preferred_model: str = None) -> Dict[str, Any]:
    """Extract text from uploaded document and parse into structured installation acceptance data."""
    extracted_text = extract_text_from_upload_file(file_bytes, filename)
    if not extracted_text.strip():
        logger.warning("Empty or unsupported file content: %s", filename)
        return _fallback_parse_installation_text("")
    return parse_installation_acceptance_text(extracted_text, preferred_model=preferred_model)

