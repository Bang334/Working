"""Quotation Excel workbook (.xlsx) generator for Hong Ming Global."""

import os
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from app.config import OUTPUTS_DIR, WORKING_DIR
from app.utils.currency_helper import (
    format_currency_vnd,
    number_to_chinese_words,
    number_to_vietnamese_words,
    parse_positive_integer,
)

# Colors & Fonts
FONT_FAMILY = "Times New Roman"
FONT_CORP_NAME = Font(name=FONT_FAMILY, size=13, bold=True, color="FF1F4E79")
FONT_CORP_SUB = Font(name=FONT_FAMILY, size=10, bold=True, color="FF2E75B6")
FONT_CORP_INFO = Font(name=FONT_FAMILY, size=9, bold=False, color="FF64748B")

FONT_DOC_TITLE = Font(name=FONT_FAMILY, size=15, bold=True, color="FF1F4E79")
FONT_DOC_SUB = Font(name=FONT_FAMILY, size=11, bold=True, color="FF2E75B6")

FONT_META_LABEL = Font(name=FONT_FAMILY, size=10, bold=True, color="FF1E293B")
FONT_META_VALUE = Font(name=FONT_FAMILY, size=10, bold=False, color="FF0F172A")

FONT_HDR = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFFFF")
FONT_DATA_STT = Font(name=FONT_FAMILY, size=10, bold=False, color="FF64748B")
FONT_DATA = Font(name=FONT_FAMILY, size=10, bold=False, color="FF0F172A")
FONT_DATA_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="FF0F172A")

FONT_SUM_LABEL = Font(name=FONT_FAMILY, size=10, bold=True, color="FF1E293B")
FONT_SUM_VALUE = Font(name=FONT_FAMILY, size=10, bold=True, color="FF1E293B")
FONT_GRAND_LABEL = Font(name=FONT_FAMILY, size=11, bold=True, color="FF1F4E79")
FONT_GRAND_VALUE = Font(name=FONT_FAMILY, size=12, bold=True, color="FF15803D")

FONT_WORDS_LABEL = Font(name=FONT_FAMILY, size=9.5, bold=True, color="FF334155")
FONT_WORDS_VALUE = Font(name=FONT_FAMILY, size=9.5, italic=True, color="FF1E293B")

FONT_TERMS_HDR = Font(name=FONT_FAMILY, size=11, bold=True, color="FF1F4E79")
FONT_TERMS_ITEM = Font(name=FONT_FAMILY, size=9.5, bold=False, color="FF1E293B")

FONT_SIG_HDR = Font(name=FONT_FAMILY, size=10.5, bold=True, color="FF1F4E79")
FONT_SIG_SUB = Font(name=FONT_FAMILY, size=9, italic=True, color="FF64748B")

# Fills
FILL_DARK_BLUE = PatternFill("solid", fgColor="FF1F4E79")
FILL_HEADER_BLUE = PatternFill("solid", fgColor="FF2E75B6")
FILL_LIGHT_BLUE = PatternFill("solid", fgColor="FFE0F2FE")
FILL_LIGHT_GRAY = PatternFill("solid", fgColor="FFF8FAFC")
FILL_SUMMARY_BG = PatternFill("solid", fgColor="FFF1F5F9")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFFFF")

# Borders
BORDER_THIN = Border(
    left=Side(style='thin', color='FFCBD5E1'),
    right=Side(style='thin', color='FFCBD5E1'),
    top=Side(style='thin', color='FFCBD5E1'),
    bottom=Side(style='thin', color='FFCBD5E1')
)
BORDER_HEADER = Border(
    left=Side(style='thin', color='FF1F4E79'),
    right=Side(style='thin', color='FF1F4E79'),
    top=Side(style='thin', color='FF1F4E79'),
    bottom=Side(style='thin', color='FF1F4E79')
)


def apply_cell_style(cell, font=None, fill=None, border=None, alignment=None, number_format=None):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format


def generate_quotation_excel(data: Dict[str, Any]) -> Dict[str, Any]:
    """Generate professional bilingual Quotation Excel workbook (.xlsx) based on BaoGiaGoc.xlsx."""
    target_dir = WORKING_DIR / "Ho So Du An" / "Bao Gia"
    target_dir.mkdir(parents=True, exist_ok=True)

    # Prepare data
    cust_vi = str(data.get("customer_name_vi") or data.get("customer_name") or "KHÁCH HÀNG").strip()
    cust_zh = str(data.get("customer_name_zh") or data.get("customer_name_cn") or "").strip()
    tax_code = str(data.get("tax_code") or "").strip()
    quote_no = str(data.get("quotation_no") or data.get("quote_no") or "HM-QT-001").strip()
    quote_date = str(data.get("quotation_date") or data.get("quote_date") or datetime.now().strftime("%d/%m/%Y")).strip()
    project_name = str(data.get("project_name") or "Dự án Cung Cấp Thiết Bị & Dịch Vụ IoT").strip()
    contact_person = str(data.get("contact_person") or "").strip()
    phone = str(data.get("phone") or "").strip()
    email = str(data.get("email") or "").strip()
    address = str(data.get("address") or "").strip()
    currency = str(data.get("currency") or "VND").upper().strip()

    validity = str(data.get("validity") or "30 ngày kể từ ngày phát hành").strip()
    lead_time = str(data.get("lead_time") or "60 - 80 ngày sau khi xác nhận đơn hàng").strip()
    warranty = str(data.get("warranty") or "12 tháng kể từ ngày giao hàng").strip()
    payment_terms = str(data.get("payment_terms") or "Tạm ứng 50%, Giao hàng 50% trong vòng 10 ngày").strip()
    notes = str(data.get("notes") or "").strip()

    raw_items = data.get("items") or []
    if not raw_items:
        raw_items = [{"name": "Thiết bị / Dịch vụ kỹ thuật", "qty": 1, "price": 0}]

    items = []
    subtotal = 0
    for idx, it in enumerate(raw_items, start=1):
        name = str(it.get("name") or it.get("product_name") or f"Thiết bị {idx}").strip()
        cat = str(it.get("category") or it.get("item_category") or it.get("unit") or "Thiết bị chính hãng").strip()
        content = str(it.get("content") or it.get("product_content") or it.get("model") or "").strip()
        item_lead = str(it.get("lead_time") or it.get("delivery_time") or lead_time).strip()

        qty_val = it.get("qty") or it.get("quantity") or 1
        try:
            qty = int(qty_val)
        except (ValueError, TypeError):
            qty = 1
        if qty <= 0:
            qty = 1

        price_val = it.get("price") or it.get("unit_price") or 0
        try:
            price = float(price_val) if "." in str(price_val) else int(price_val)
        except (ValueError, TypeError):
            price = 0

        tot = qty * price
        subtotal += tot
        items.append({
            "stt": idx,
            "name": name,
            "category": cat,
            "content": content,
            "qty": qty,
            "lead_time": item_lead,
            "price": price,
            "total": tot
        })

    # Template path
    template_path = Path(r"D:\Working\Bieu Mau\BaoGiaGoc.xlsx")
    if not template_path.is_file():
        alt_path = WORKING_DIR / "Bieu Mau" / "BaoGiaGoc.xlsx"
        if alt_path.is_file():
            template_path = alt_path

    if template_path.is_file():
        wb = openpyxl.load_workbook(str(template_path))
        ws = wb.active

        # Update Buyer Info in F2:H7
        display_cust = f"{cust_vi} ({cust_zh})" if cust_zh and cust_zh != cust_vi else cust_vi
        ws["H2"] = display_cust
        if contact_person:
            ws["H3"] = contact_person
        if phone:
            ws["H4"] = phone
        if address:
            ws["H6"] = address
        if email:
            ws["H7"] = email

        # Parse quotation base date
        base_date = None
        for fmt in ("%d/%m/%Y", "%Y/%m/%d", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                base_date = datetime.strptime(quote_date, fmt)
                break
            except ValueError:
                continue
        if not base_date:
            base_date = datetime.now()

        formatted_quote_date = base_date.strftime("%Y/%m/%d")

        # Update Quotation Title in A9 (Tự động đổi theo tên dự án)
        if project_name:
            p_upper = project_name.upper().strip()
            if p_upper.startswith("BÁO GIÁ"):
                ws["A9"] = f"{p_upper} / 報價單" if "/ 報價單" not in p_upper else p_upper
            else:
                ws["A9"] = f"BÁO GIÁ {p_upper} / 報價單"

        # Update Quotation Date in A10 (Tự động đổi theo ngày làm đơn)
        ws["A10"] = f"Ngày báo giá報價日期：{formatted_quote_date}"

        # Clean merged cells from row 12 downwards to prevent overlaps
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row >= 12:
                ws.unmerge_cells(str(rng))

        # Template row 12 is item 1, row 13 is empty row
        # Delete row 13 first
        ws.delete_rows(13)
        num_items = len(items)
        if num_items > 1:
            ws.insert_rows(13, amount=(num_items - 1))

        start_r = 12
        f_calibri = Font(name="Calibri", size=14)
        f_arial = Font(name="Arial", size=13)
        al_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        al_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        al_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='A0A0A0'),
            right=Side(style='thin', color='A0A0A0'),
            top=Side(style='thin', color='A0A0A0'),
            bottom=Side(style='thin', color='A0A0A0')
        )
        num_fmt = "$#,##0.00" if currency == "USD" else "#,##0"

        for idx, it in enumerate(items):
            r = start_r + idx
            ws.row_dimensions[r].height = 48.0

            ws[f"A{r}"] = it["stt"]
            ws[f"B{r}"] = it["name"]
            ws[f"C{r}"] = it["category"]
            ws[f"E{r}"] = it["content"]
            ws[f"F{r}"] = it["qty"]
            ws[f"G{r}"] = it["lead_time"]
            ws[f"H{r}"] = it["price"]
            ws[f"I{r}"] = f"=H{r}*F{r}"

            ws.merge_cells(f"C{r}:D{r}")

            for col_l in ["A", "B", "C", "D", "E", "F", "G", "H", "I"]:
                ws[f"{col_l}{r}"].border = thin_border

            ws[f"A{r}"].font = f_calibri
            ws[f"A{r}"].alignment = al_center

            ws[f"B{r}"].font = f_arial
            ws[f"B{r}"].alignment = al_left

            ws[f"C{r}"].font = f_arial
            ws[f"C{r}"].alignment = al_left

            ws[f"E{r}"].font = f_arial
            ws[f"E{r}"].alignment = al_left

            ws[f"F{r}"].font = f_calibri
            ws[f"F{r}"].alignment = al_center
            ws[f"F{r}"].number_format = "#,##0"

            ws[f"G{r}"].font = f_calibri
            ws[f"G{r}"].alignment = al_center

            ws[f"H{r}"].font = f_calibri
            ws[f"H{r}"].alignment = al_right
            ws[f"H{r}"].number_format = num_fmt

            ws[f"I{r}"].font = f_calibri
            ws[f"I{r}"].alignment = al_right
            ws[f"I{r}"].number_format = num_fmt

        # Blank Row "以下空白"
        blank_r = start_r + num_items
        ws.row_dimensions[blank_r].height = 24.6
        ws[f"A{blank_r}"] = "以下空白 Chỗ trống sau đây"
        ws.merge_cells(f"A{blank_r}:I{blank_r}")
        ws[f"A{blank_r}"].font = Font(name="微軟正黑體", size=14)
        ws[f"A{blank_r}"].alignment = Alignment(horizontal="center", vertical="center")
        for col_l in ["A", "B", "C", "D", "E", "F", "G", "H", "I"]:
            ws[f"{col_l}{blank_r}"].border = thin_border

        # Summary Row: 此致 & TỔNG THÀNH TIỀN
        sum_r = blank_r + 1
        ws.row_dimensions[sum_r].height = 32.0

        ws[f"A{sum_r}"] = f"此致\n         {display_cust}"
        ws.merge_cells(f"A{sum_r}:B{sum_r}")
        ws[f"A{sum_r}"].font = Font(name="微軟正黑體", size=12, bold=True, color="0000FF")
        ws[f"A{sum_r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws[f"H{sum_r}"] = "TỔNG THÀNH TIỀN \n總價"
        ws[f"H{sum_r}"].font = Font(name="Arial", size=14, bold=True, color="FF0000")
        ws[f"H{sum_r}"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        ws[f"H{sum_r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[f"H{sum_r}"].border = thin_border

        ws[f"I{sum_r}"] = f"=SUM(I{start_r}:I{start_r + num_items - 1})"
        ws[f"I{sum_r}"].font = Font(name="Arial", size=14, bold=True)
        # Format universal VND to avoid WPS bug "VND 13.70 (VND 0.000)"
        ws[f"I{sum_r}"].number_format = "$#,##0.00" if currency == "USD" else '#,##0 "VND"'
        ws[f"I{sum_r}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"I{sum_r}"].border = thin_border

        # -------------------------------------------------------------
        # EXACT TERMS & CONFIRMATION BOX FROM BaoGiaGoc.xlsx (Row 16 to 21)
        # -------------------------------------------------------------
        r_hdr = sum_r + 1
        r_t1 = sum_r + 2
        r_t2 = sum_r + 3
        r_t3 = sum_r + 4
        r_t4 = sum_r + 5
        r_t5 = sum_r + 6

        # Set exact row heights from template
        ws.row_dimensions[r_hdr].height = 34.8
        ws.row_dimensions[r_t1].height = 58.2
        ws.row_dimensions[r_t2].height = 58.2
        ws.row_dimensions[r_t3].height = 58.2
        ws.row_dimensions[r_t4].height = 58.2
        ws.row_dimensions[r_t5].height = 316.0

        # Row Header: A{r_hdr} & E{r_hdr}
        ws[f"A{r_hdr}"] = "※客戶確認 Xác nhận khách hàng"
        ws[f"A{r_hdr}"].font = Font(name="Arial", size=14, bold=False, color="FF9C6500")
        ws[f"A{r_hdr}"].alignment = Alignment(horizontal="left", vertical="center")

        ws[f"E{r_hdr}"] = "◆ 報價特別聲明 Tuyên bố đặc biệt của báo giá :"
        ws[f"E{r_hdr}"].font = Font(name="微軟正黑體", size=12, bold=True)
        ws[f"E{r_hdr}"].alignment = Alignment(horizontal="left", vertical="center")

        # Row 1: Sign note in A{r_t1}
        ws[f"A{r_t1}"] = "(請簽名用印 Vui lòng ký tên và đóng dấu)"
        ws[f"A{r_t1}"].font = Font(name="Arial", size=12, bold=False, color="FF9C6500")
        ws[f"A{r_t1}"].alignment = Alignment(horizontal="left", vertical="top")

        # 1. Term 1: Delivery / Lead time
        if "以上報價交期" in lead_time:
            t1 = lead_time
        elif "07-10" in lead_time or "7-10" in lead_time:
            t1 = (
                "1. 以上報價交期：自收到付款之日起07-10個工作日內交貨。\n"
                "Ngày giao hàng theo báo giá trên: Trong vòng 07-10 ngày làm việc kể từ ngày nhận được thanh toán."
            )
        elif "30-45" in lead_time:
            t1 = (
                "1. 以上報價交期：自收到付款之日起30-45天內交貨。\n"
                "Ngày giao hàng theo báo giá trên: Trong vòng 30-45 ngày kể từ ngày nhận được thanh toán."
            )
        elif lead_time and lead_time != "60 - 80 ngày sau khi xác nhận đơn hàng" and "60" not in lead_time:
            t1 = (
                f"1. 以上報價交期：自收到付款之日起{lead_time}交貨。\n"
                f"Ngày giao hàng theo báo giá trên: Trong vòng {lead_time} kể từ ngày nhận được thanh toán."
            )
        else:
            t1 = (
                "1. 以上報價交期：自收到付款之日起60天內交貨。\n"
                "Ngày giao hàng theo báo giá trên: Trong vòng 60 ngày kể từ ngày nhận được thanh toán."
            )

        ws[f"E{r_t1}"] = t1
        ws.merge_cells(f"E{r_t1}:I{r_t1}")
        ws[f"E{r_t1}"].font = Font(name="微軟正黑體", size=12, bold=True)
        ws[f"E{r_t1}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # 2. Term 2: Quotation Validity - TỰ ĐỘNG TÍNH NGÀY HẾT HẠN THEO NGÀY BÁO GIÁ
        expiry_dt = None
        # Match explicit dates: YYYY/MM/DD or YYYY-MM-DD
        m_ymd = re.search(r'(\d{4}[/-]\d{1,2}[/-]\d{1,2})', validity)
        if m_ymd:
            try:
                expiry_dt = datetime.strptime(m_ymd.group(1).replace("-", "/"), "%Y/%m/%d")
            except ValueError:
                pass

        if not expiry_dt:
            # Match explicit dates: DD/MM/YYYY or DD-MM-YYYY
            m_dmy = re.search(r'(\d{1,2}[/-]\d{1,2}[/-]\d{4})', validity)
            if m_dmy:
                try:
                    expiry_dt = datetime.strptime(m_dmy.group(1).replace("-", "/"), "%d/%m/%Y")
                except ValueError:
                    pass

        if not expiry_dt:
            # Match number of days in validity (e.g. "15 ngày", "30 ngày", "60 ngày", "15天")
            m_days = re.search(r'(\d+)\s*(?:ngày|day|天|days)', validity, re.IGNORECASE)
            if m_days:
                days_add = int(m_days.group(1))
            elif "15" in validity:
                days_add = 15
            elif "30" in validity:
                days_add = 30
            elif "60" in validity:
                days_add = 60
            else:
                # Default Hong Ming validity period from BaoGiaGoc.xlsx (21/09 -> 05/10 is exactly 14 days)
                days_add = 14
            expiry_dt = base_date + timedelta(days=days_add)

        expiry_zh = expiry_dt.strftime("%Y/%m/%d")
        expiry_vi = expiry_dt.strftime("%d/%m/%Y")

        t2 = (
            f"2. 報價有效期限：報價單有效期限至{expiry_zh}，報價簽回視同專案確認，不可任意終止、取消本專案。\n"
            f"    Thời hạn hiệu lực của báo giá: Báo giá có hiệu lực đến hết ngày {expiry_vi}. Báo giá đã ký được coi là xác nhận của dự án. Dự án không thể bị chấm dứt hoặc hủy bỏ theo ý muốn."
        )

        ws[f"E{r_t2}"] = t2
        ws.merge_cells(f"E{r_t2}:I{r_t2}")
        ws[f"E{r_t2}"].font = Font(name="微軟正黑體", size=12, bold=True)
        ws[f"E{r_t2}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # 3. Term 3: Hardware Warranty
        if "上述報價中的" in warranty:
            t3 = warranty
        elif "3" in warranty or "03" in warranty:
            t3 = (
                "3. 上述報價中的硬體設備自交貨日起保固三個月；因人為因素造成之碰撞、破裂、損壞等情況，不在保固範圍內\n"
                "  Thiết bị phần cứng trong báo giá trên được bảo hành ba tháng kể từ ngày giao hàng; các hư hỏng do va chạm, vỡ, gãy do con người gây ra không nằm trong phạm vi bảo hành"
            )
        elif "12" in warranty:
            t3 = (
                "3. 上述報價中的硬體設備自交貨日起保固十二個月；因人為因素造成之碰撞、破裂、損壞等情況，不在保固範圍內\n"
                "  Thiết bị phần cứng trong báo giá trên được bảo hành 12 tháng kể từ ngày giao hàng; các hư hỏng do va chạm, vỡ, gãy do con người gây ra không nằm trong phạm vi bảo hành"
            )
        else:
            t3 = (
                "3. 上述報價中的硬體設備自交貨日起保固六個月；因人為因素造成之碰撞、破裂、損壞等情況，不在保固範圍內\n"
                "  Thiết bị phần cứng trong báo giá trên được bảo hành sáu tháng kể từ ngày giao hàng; các hư hỏng do va chạm, vỡ, gãy do con người gây ra không nằm trong phạm vi bảo hành"
            )

        ws[f"E{r_t3}"] = t3
        ws.merge_cells(f"E{r_t3}:I{r_t3}")
        ws[f"E{r_t3}"].font = Font(name="微軟正黑體", size=12, bold=True)
        ws[f"E{r_t3}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # 4. Term 4: VAT (CHỮ ĐỎ FFFF0000 100% MẪU GỐC)
        t4 = (
            "4.以上報價不含增值稅 VAT，稅率則按當時發票開出時之規定。\n"
            "Giá trên chưa bao gồm VAT. Thuế suất sẽ áp dụng theo quy định tại thời điểm xuất hóa đơn."
        )
        ws[f"E{r_t4}"] = t4
        ws.merge_cells(f"E{r_t4}:I{r_t4}")
        ws[f"E{r_t4}"].font = Font(name="微軟正黑體", size=12, bold=False, color="FFFF0000")
        ws[f"E{r_t4}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # 5. Term 5: Payment & Acceptance (CHỮ XANH DƯƠNG FF0000FF, 18 DÒNG ĐẦY ĐỦ 100% PHÁP LÝ GỐC)
        official_term_5 = (
            "5. 一、付款方式：\n"
            "訂單總金額超過 30,000,000 越南盾（VND 30,000,000）時，甲方須於訂購前支付訂單總金額之 50% 作為預付款；剩餘 50% 應於收到貨物之日起 10 日內，以銀行轉帳方式支付予乙方。\n"
            "Đối với đơn hàng có tổng giá trị ttừ 30.000.000 VNĐ trở lên. : Bên A thanh toán trước 50% tổng giá trị đơn hàng trước khi đặt hàng. 50% còn lại được thanh toán bằng hình thức chuyển khoản ngân hàng trong vòng 10 ngày kể từ ngày Bên A nhận hàng.\n"
            "訂單總金額低於 30,000,000 越南盾（VND 30,000,000）時，甲方應於收到貨物之日起 10 日內，以銀行轉帳方式支付訂單總金額之 100%。\n"
            "Đối với đơn hàng có tổng giá trị dưới 30.000.000 VNĐ: Bên A thanh toán 100% tổng giá trị đơn hàng bằng hình thức chuyển khoản ngân hàng trong vòng 10 ngày kể từ ngày nhận hàng.\n"
            "二、驗收條款：\n"
            "甲方應於收到貨物之日起 7 日內完成檢查及驗收。如甲方於上述期限內未就產品品質、數量或規格提出書面異議，則視為貨物已完成交付並驗收合格。\n"
            "Bên A có trách nhiệm kiểm tra và nghiệm thu hàng hóa trong vòng 07 ngày kể từ ngày nhận hàng. Nếu trong thời hạn nêu trên, Bên A không có bất kỳ khiếu nại hoặc ý kiến phản hồi bằng văn bản nào liên quan đến chất lượng, số lượng hoặc quy cách sản phẩm, hàng hóa được xem là đã bàn giao và nghiệm thu thành công\n"
            "甲方應依本合約約定之付款條件履行付款義務，不得以未另行簽署驗收文件或未提出驗收確認為由延遲付款。\n"
            "Bên A có trách nhiệm thực hiện thanh toán đúng theo điều kiện thanh toán đã thỏa thuận trong hợp đồng, không được trì hoãn thanh toán với lý do chưa ký biên bản nghiệm thu riêng hoặc chưa xác nhận nghiệm thu."
        )

        if payment_terms and "一、付款方式" in payment_terms and "二、驗收條款" in payment_terms:
            t5 = payment_terms.strip()
            if not t5.startswith("5."):
                t5 = "5. " + t5
        else:
            t5 = official_term_5

        if notes:
            t5 += f"\n\n備註 Ghi chú bổ sung: {notes}"

        ws[f"E{r_t5}"] = t5
        ws.merge_cells(f"E{r_t5}:I{r_t5}")
        ws[f"E{r_t5}"].font = Font(name="微軟正黑體", size=12, bold=False, color="FF0000FF")
        ws[f"E{r_t5}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Outer box medium borders (matching BaoGiaGoc.xlsx)
        m_side = Side(style='medium', color='FF000000')
        for row_idx in range(r_hdr, r_t5 + 1):
            ws.cell(row_idx, 1).border = Border(
                left=m_side,
                right=ws.cell(row_idx, 1).border.right if ws.cell(row_idx, 1).border else None,
                top=ws.cell(row_idx, 1).border.top if ws.cell(row_idx, 1).border else None,
                bottom=ws.cell(row_idx, 1).border.bottom if ws.cell(row_idx, 1).border else None
            )
            ws.cell(row_idx, 9).border = Border(
                right=m_side,
                left=ws.cell(row_idx, 9).border.left if ws.cell(row_idx, 9).border else None,
                top=ws.cell(row_idx, 9).border.top if ws.cell(row_idx, 9).border else None,
                bottom=ws.cell(row_idx, 9).border.bottom if ws.cell(row_idx, 9).border else None
            )

        for col_idx in range(1, 10):
            ws.cell(r_t5, col_idx).border = Border(
                bottom=m_side,
                left=ws.cell(r_t5, col_idx).border.left if ws.cell(r_t5, col_idx).border else None,
                right=ws.cell(r_t5, col_idx).border.right if ws.cell(r_t5, col_idx).border else None,
                top=ws.cell(r_t5, col_idx).border.top if ws.cell(r_t5, col_idx).border else None
            )

        # Print area
        ws.print_area = f"$A$1:$I${r_t5}"

    else:
        # Fallback programmatic generator if template missing
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Báo Giá"
        ws["A1"] = "CÔNG TY TNHH TOÀN CẦU HONG MING"
        ws["A2"] = f"BÁO GIÁ {project_name}"
        ws["A3"] = f"Khách hàng: {cust_vi}"

    # Save File
    from app.utils.file_helper import document_filename
    timestamp = str(data.get("_file_timestamp") or "")
    if not re.fullmatch(r"\d{8}_\d{6}_\d{6}", timestamp):
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    file_name = document_filename("Bao_Gia", cust_vi, quote_no, "xlsx", timestamp)
    output_path = target_dir / file_name

    try:
        wb.save(str(output_path))
    except PermissionError:
        file_name = f"{output_path.stem}_v2.xlsx"
        output_path = target_dir / file_name
        wb.save(str(output_path))

    # Index into SQLite documents table
    file_size = output_path.stat().st_size if output_path.is_file() else 0
    try:
        from app.database import index_single_document
        index_single_document(
            contract_code="",
            doc_type="BAO_GIA",
            doc_name=f"Báo Giá Excel ({quote_no}) - {cust_vi}",
            file_name=file_name,
            file_path=str(output_path),
            file_size=file_size,
            file_format="xlsx",
            description=f"Bảng báo giá thiết bị Excel: {cust_vi} [Số: {quote_no}]"
        )
    except Exception:
        pass

    total_display = f"{subtotal:,.2f} USD" if currency == "USD" else f"{int(subtotal):,}".replace(",", ".") + " VND"

    return {
        "status": "success",
        "message": f"Đã xuất thành công Báo giá Excel (.xlsx) chuẩn mẫu BaoGiaGoc cho {cust_vi}!",
        "file_name": file_name,
        "file_path": str(output_path),
        "download_url": f"/api/v1/quotation/download/{file_name}",
        "total_amount": subtotal,
        "total_amount_formatted": total_display,
        "customer_name": cust_vi,
        "quotation_no": quote_no,
        "currency": currency,
    }
