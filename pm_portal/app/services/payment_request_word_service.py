# -*- coding: utf-8 -*-
"""Xuất Phiếu yêu cầu thanh toán song ngữ sang Word theo bố cục sheet 712TR."""

import copy
import os
import re
import tempfile
from zipfile import ZipFile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.table import _Cell


FONT_LATIN = "Times New Roman"
FONT_CJK = "Microsoft JhengHei"
YELLOW = "FFFF00"
BLACK = "000000"


def _template_values(data: Dict[str, Any]) -> Dict[str, str]:
    """Các trường được Word mẫu hỗ trợ, gồm cả bí danh Việt/Anh thông dụng."""
    percentage = str(data.get("percentage_display") or data.get("percentage") or "")
    amount = str(data.get("milestone_amount_formatted") or "")
    invoice_no = str(data.get("invoice_no") or "")
    payment_item = (
        f"請款 / Thanh toán {data.get('milestone_name', '')}: {percentage}%\n"
        f"合約 / Hợp đồng：{data.get('contract_code', '')}"
    )
    if invoice_no and invoice_no != "Chưa xuất":
        payment_item += f"\n發票 / Hóa đơn：{invoice_no}"

    canonical = {
        "FORM_NO": data.get("form_no", ""),
        "APPLICATION_DATE": data.get("application_date", ""),
        "CUSTOMER_NAME_VI": data.get("customer_name_vi", ""),
        "CUSTOMER_NAME_ZH": data.get("customer_name_zh", ""),
        "CONTRACT_CODE": data.get("contract_code", ""),
        "PROJECT_NAME_VI": data.get("project_name_vi", ""),
        "PROJECT_NAME_ZH": data.get("project_name_zh", ""),
        "MILESTONE_NAME": data.get("milestone_name", ""),
        "PERCENTAGE": percentage,
        "TOTAL_CONTRACT_AMOUNT": data.get("total_contract_amount", ""),
        "TOTAL_CONTRACT_AMOUNT_FORMATTED": data.get("total_contract_amount_formatted", ""),
        "MILESTONE_AMOUNT": data.get("milestone_amount", ""),
        "MILESTONE_AMOUNT_FORMATTED": amount,
        "AMOUNT_WORDS_VI": data.get("amount_words_vi", ""),
        "AMOUNT_WORDS_ZH": data.get("amount_words_zh", ""),
        "DUE_DATE": data.get("due_date", ""),
        "INVOICE_NO": invoice_no,
        "PAYMENT_ITEM": payment_item,
        "NOTES": data.get("notes_text", ""),
    }
    aliases = {
        "SO_PHIEU": "FORM_NO",
        "MAU_SO": "FORM_NO",
        "NGAY_NOP_DON": "APPLICATION_DATE",
        "NGAY_LAP": "APPLICATION_DATE",
        "TEN_KHACH_HANG": "CUSTOMER_NAME_VI",
        "TEN_KHACH_HANG_VI": "CUSTOMER_NAME_VI",
        "TEN_KHACH_HANG_ZH": "CUSTOMER_NAME_ZH",
        "SO_HOP_DONG": "CONTRACT_CODE",
        "TEN_DU_AN": "PROJECT_NAME_VI",
        "TEN_DU_AN_VI": "PROJECT_NAME_VI",
        "TEN_DU_AN_ZH": "PROJECT_NAME_ZH",
        "DOT_THANH_TOAN": "MILESTONE_NAME",
        "TY_LE_THANH_TOAN": "PERCENTAGE",
        "TY_LE": "PERCENTAGE",
        "TONG_GIA_TRI_HOP_DONG": "TOTAL_CONTRACT_AMOUNT",
        "TONG_GIA_TRI_DINH_DANG": "TOTAL_CONTRACT_AMOUNT_FORMATTED",
        "SO_TIEN_THANH_TOAN": "MILESTONE_AMOUNT",
        "SO_TIEN_THANH_TOAN_DINH_DANG": "MILESTONE_AMOUNT_FORMATTED",
        "SO_TIEN": "MILESTONE_AMOUNT_FORMATTED",
        "BANG_CHU_VI": "AMOUNT_WORDS_VI",
        "BANG_CHU_ZH": "AMOUNT_WORDS_ZH",
        "HAN_THANH_TOAN": "DUE_DATE",
        "SO_HOA_DON": "INVOICE_NO",
        "NOI_DUNG_THANH_TOAN": "PAYMENT_ITEM",
        "GHI_CHU": "NOTES",
    }
    values = {key: str(value if value is not None else "") for key, value in canonical.items()}
    values.update({alias: values[target] for alias, target in aliases.items()})
    return values


def _replace_token_across_runs(paragraph, token: str, value: str) -> int:
    """Thay token dù Word đã chia token thành nhiều run; style run đầu được giữ nguyên."""
    count = 0
    while paragraph.runs:
        full_text = "".join(run.text for run in paragraph.runs)
        start = full_text.find(token)
        if start < 0:
            break
        end = start + len(token)
        cursor = 0
        start_index = end_index = None
        start_offset = end_offset = 0
        for index, run in enumerate(paragraph.runs):
            run_end = cursor + len(run.text)
            if start_index is None and cursor <= start < run_end:
                start_index = index
                start_offset = start - cursor
            if cursor < end <= run_end:
                end_index = index
                end_offset = end - cursor
                break
            cursor = run_end
        if start_index is None or end_index is None:
            break
        first_run = paragraph.runs[start_index]
        last_run = paragraph.runs[end_index]
        prefix = first_run.text[:start_offset]
        suffix = last_run.text[end_offset:]
        if start_index == end_index:
            first_run.text = prefix + value + suffix
        else:
            first_run.text = prefix + value
            for index in range(start_index + 1, end_index):
                paragraph.runs[index].text = ""
            last_run.text = suffix
        count += 1
    return count


def _normalized_marker(value: str) -> str:
    return re.sub(r"[^0-9A-ZÀ-Ỹ_]+", "_", str(value or "").strip().upper()).strip("_")


def _replace_in_template_paragraph(paragraph, values: Dict[str, str]) -> int:
    count = 0
    for key, value in values.items():
        for token in (f"{{{{{key}}}}}", f"[{key}]", f"<<{key}>>"):
            count += _replace_token_across_runs(paragraph, token, value)

    # Hỗ trợ ô tô vàng/chữ đỏ dùng trực tiếp tên trường, không có dấu ngoặc.
    for run in paragraph.runs:
        color = run.font.color.rgb if (run.font and run.font.color) else None
        is_marker_style = (run.font and run.font.highlight_color is not None) or (color is not None and str(color).upper() == "FF0000")
        marker = _normalized_marker(run.text)
        if is_marker_style and marker in values:
            run.text = values[marker]
            count += 1
    return count


def _replace_in_template_table(table, values: Dict[str, str]) -> int:
    count = 0
    seen_cells = set()
    for row in table.rows:
        for cell in row.cells:
            cell_key = id(cell._tc)
            if cell_key in seen_cells:
                continue
            seen_cells.add(cell_key)
            for paragraph in cell.paragraphs:
                count += _replace_in_template_paragraph(paragraph, values)
            for nested_table in cell.tables:
                count += _replace_in_template_table(nested_table, values)
    return count


def _fill_word_template(doc: Document, data: Dict[str, Any]) -> int:
    values = _template_values(data)
    count = sum(_replace_in_template_paragraph(paragraph, values) for paragraph in doc.paragraphs)
    count += sum(_replace_in_template_table(table, values) for table in doc.tables)
    seen_story_parts = set()
    for section in doc.sections:
        references = []
        for tag_name, kind in (("w:headerReference", "header"), ("w:footerReference", "footer")):
            for reference in section._sectPr.findall(qn(tag_name)):
                references.append((kind, reference.get(qn("w:type"), "default")))
        for kind, reference_type in references:
            attribute = kind if reference_type == "default" else f"{reference_type}_page_{kind}"
            story = getattr(section, attribute)
            story_key = id(story.part)
            if story_key in seen_story_parts:
                continue
            seen_story_parts.add(story_key)
            count += sum(_replace_in_template_paragraph(paragraph, values) for paragraph in story.paragraphs)
            count += sum(_replace_in_template_table(table, values) for table in story.tables)
    return count


def _replace_paragraph_value(paragraph, value: Any) -> None:
    """Đổi nội dung nhưng giữ nguyên style của run đầu tại ô mẫu."""
    text = str(value if value is not None else "")
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _replace_cell_value(cell, value: Any) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    _replace_paragraph_value(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        _replace_paragraph_value(paragraph, "")


def _unique_row_cells(row) -> List[Any]:
    """Trả về các ô hiển thị thực, bỏ các bản lặp do merge ngang/dọc."""
    cells: List[Any] = []
    seen_elements = set()
    for cell in row.cells:
        element = cell._tc
        if element in seen_elements:
            continue
        seen_elements.add(element)
        cells.append(cell)
    return cells


def _physical_table_rows(table) -> List[List[Any]]:
    """Đọc các w:tc thật theo từng w:tr, ổn định giữa các bản python-docx.

    Một số phiên bản python-docx diễn giải bảng có gridBefore/gridAfter khác nhau
    và có thể đưa ô đầu của hàng sau vào ``row.cells`` của hàng trước. Duyệt trực
    tiếp XML giữ đúng thứ tự ô vật lý của mẫu Word.
    """
    return [
        [_Cell(tc, table) for tc in table_row.tc_lst]
        for table_row in table._tbl.tr_lst
    ]


def _find_table_with_labels(doc: Document, required_labels: Sequence[Sequence[str]]):
    for table in doc.tables:
        table_text = "\n".join(
            cell.text
            for row in _physical_table_rows(table)
            for cell in row
        )
        if all(any(label in table_text for label in alternatives) for alternatives in required_labels):
            return table
    return None


def _replace_value_after_label(table, labels: Sequence[str], value: Any) -> bool:
    for cells in _physical_table_rows(table):
        for index, cell in enumerate(cells[:-1]):
            if any(label in cell.text for label in labels):
                _replace_cell_value(cells[index + 1], value)
                return True
    return False


def _replace_form_number(doc: Document, value: Any) -> bool:
    replacement = f"Form NO. {value}"
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().lower().startswith("form no."):
            _replace_paragraph_value(paragraph, replacement)
            return True
    for table in doc.tables:
        for row in table.rows:
            for cell in _unique_row_cells(row):
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip().lower().startswith("form no."):
                        _replace_paragraph_value(paragraph, replacement)
                        return True
    return False


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _fill_payment_terms_table(table, data: Dict[str, Any]) -> int:
    bilingual_items = data.get("payment_terms_bilingual") or []
    terms = [str(value) for value in (data.get("payment_terms") or [])]
    if len(table.rows) == 1:
        _replace_cell_value(table.cell(0, 0), data.get("notes_text", ""))
        return 1
    if not terms and not bilingual_items:
        return 0

    count = len(bilingual_items) if bilingual_items else len(terms)
    # Đồng bộ số hàng khớp 1:1 với số lượng điều khoản: "Bao nhiêu điều khoản bấy nhiêu hàng"
    while len(table.rows) > count and len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < count:
        new_tr = copy.deepcopy(table.rows[-1]._tr)
        table._tbl.append(new_tr)

    for index in range(count):
        row = table.rows[index]
        cells = _unique_row_cells(row)
        if not cells:
            continue
        cell = cells[0]

        if index < len(bilingual_items):
            item = bilingual_items[index]
            chinese_body = item.get("chinese", "").strip()
            vietnamese_body = item.get("vietnamese", "").strip()
        else:
            term = terms[index]
            part1, part2 = (term.split("\n", 1) + [""])[:2]
            p1_clean = re.sub(r"^\d{2}[.\s]+", "", part1.strip()).strip()
            p2_clean = re.sub(r"^\d{2}[.\s]+", "", part2.strip()).strip()
            has_cjk_p1 = bool(re.search(r"[\u4e00-\u9fff]", p1_clean))
            if has_cjk_p1:
                chinese_body, vietnamese_body = p1_clean, p2_clean
            else:
                chinese_body, vietnamese_body = p2_clean, p1_clean

        prefix = f"{index + 1:02d}  "

        # Nền xen kẽ: lẻ (1, 3, 5...) màu trắng FFFFFF, chẵn (2, 4...) màu xanh nhạt F3F8FC
        fill_color = "FFFFFF" if (index % 2 == 0) else "F3F8FC"
        _set_cell_shading(cell, fill_color)

        # Mẫu hiện tại dùng 2 run cho dòng Trung: số thứ tự xanh đậm và nội dung
        # 8 pt màu đen. Chỉ thay text trong từng run để không làm toàn dòng nhận
        # nhầm style của số thứ tự và khiến bảng phình sang trang 2.
        first_paragraph = cell.paragraphs[0]
        if not first_paragraph.runs:
            first_paragraph.add_run()
        first_paragraph.runs[0].text = prefix
        if len(first_paragraph.runs) >= 2:
            first_paragraph.runs[1].text = chinese_body
            for run in first_paragraph.runs[2:]:
                run.text = ""
        else:
            first_paragraph.add_run(chinese_body)

        if len(cell.paragraphs) >= 2:
            _replace_paragraph_value(cell.paragraphs[1], vietnamese_body)
            for paragraph in cell.paragraphs[2:]:
                _replace_paragraph_value(paragraph, "")
        elif vietnamese_body:
            cell.add_paragraph(vietnamese_body)
    return 1



def _fill_known_payment_request_layout(doc: Document, data: Dict[str, Any]) -> int:
    """Điền mẫu 請款單 chuẩn dù mẫu dùng dữ liệu minh họa thay vì placeholder."""
    metadata_table = _find_table_with_labels(
        doc,
        (
            ("Tên khách hàng", "客 戶 名 稱"),
            ("Tên dự án", "專 案 名 稱"),
            ("Ngày đến hạn thanh toán", "應 付 款 日 期"),
        ),
    )
    payment_table = _find_table_with_labels(
        doc,
        (
            ("Hạng mục thanh toán", "請款項目"),
            ("Số tiền hợp đồng", "合同金額"),
            ("Số tiền xin TT", "請款金額"),
        ),
    )
    if metadata_table is None or payment_table is None:
        return 0
    if len(payment_table.rows) < 3 or len(payment_table.rows[1].cells) < 9:
        return 0

    percentage = str(data.get("percentage_display") or data.get("percentage") or "")
    invoice_no = str(data.get("invoice_no") or "")
    payment_item = (
        f"請款 / Thanh toán {data.get('milestone_name', '')}：{percentage}%\n"
        f"合約 / Hợp đồng：{data.get('contract_code', '')}"
    )
    if invoice_no and invoice_no != "Chưa xuất":
        payment_item += f"\n發票 / Hóa đơn：{invoice_no}"
    project_text = (
        f"Căn cứ hợp đồng số : {data.get('contract_code', '')}\n"
        f"Thực hiện các điều khoản của Dự án “{data.get('project_name_vi', '')}”.\n"
        f"根據合同編號：{data.get('contract_code', '')}\n"
        f"於 {data.get('project_name_zh', '')}相關條款。"
    )

    replacement_count = int(_replace_form_number(doc, data.get("form_no", "")))
    replacement_count += int(
        _replace_value_after_label(
            metadata_table,
            ("Tên khách hàng", "客 戶 名 稱"),
            data.get("customer_name_vi", ""),
        )
    )
    replacement_count += int(
        _replace_value_after_label(
            metadata_table,
            ("Ngày nộp đơn", "申 請 日 期"),
            data.get("application_date", ""),
        )
    )
    replacement_count += int(
        _replace_value_after_label(
            metadata_table,
            ("Tên dự án", "專 案 名 稱"),
            project_text,
        )
    )
    replacement_count += int(
        _replace_value_after_label(
            metadata_table,
            ("Ngày đến hạn thanh toán", "應 付 款 日 期"),
            data.get("due_date", ""),
        )
    )

    row = payment_table.rows[1]
    row_values = (
        "1",
        payment_item,
        "VND",
        data.get("total_contract_amount_formatted", ""),
        "1",
        data.get("total_contract_amount_formatted", ""),
        "0%",
        "-",
        data.get("milestone_amount_formatted", ""),
    )
    for index, value in enumerate(row_values):
        _replace_cell_value(row.cells[index], value)
    total_cells = _unique_row_cells(payment_table.rows[2])
    if len(total_cells) >= 2:
        _replace_cell_value(total_cells[-1], data.get("milestone_amount_formatted", ""))

    payment_table_index = next(
        index
        for index, table in enumerate(doc.tables)
        if table._tbl is payment_table._tbl
    )
    terms_table = next(
        (
            table
            for table in doc.tables[payment_table_index + 1 :]
            if table.rows and len(_unique_row_cells(table.rows[0])) == 1
        ),
        None,
    )
    if terms_table is not None:
        replacement_count += _fill_payment_terms_table(terms_table, data)
    return replacement_count + 1


def _set_run_font(run, name: str = FONT_LATIN, size: float = 9, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_fonts.set(qn("w:cs"), name)


def _set_paragraph(
    paragraph,
    text: str,
    *,
    size: float = 9,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    font: str = FONT_LATIN,
    space_before: float = 0,
    space_after: float = 0,
    line_spacing: float = 1.0,
) -> None:
    paragraph.clear()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for index, part in enumerate(str(text or "").split("\n")):
        if index:
            paragraph.add_run().add_break()
        _set_run_font(paragraph.add_run(part), name=font, size=size, bold=bold)


def _set_cell_text(
    cell,
    text: str,
    *,
    size: float = 8,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    font: str = FONT_LATIN,
    vertical=WD_CELL_VERTICAL_ALIGNMENT.CENTER,
) -> None:
    cell.vertical_alignment = vertical
    _set_paragraph(cell.paragraphs[0], text, size=size, bold=bold, align=align, font=font)
    _set_cell_margins(cell, top=55, start=55, bottom=55, end=55)


def _set_cell_margins(cell, top: int = 45, start: int = 45, bottom: int = 45, end: int = 45) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_width(cell, width_inches: float) -> None:
    width_twips = int(width_inches * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_inches)


def _set_table_borders(table, *, color: str = BLACK, size: int = 6, val: str = "single") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), val)
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def _remove_table_borders(table) -> None:
    _set_table_borders(table, color="FFFFFF", size=0, val="nil")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_row_minimum_height(row, points: float) -> None:
    row.height = Pt(points)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _set_table_layout_fixed(table) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_table_column_widths(table, widths: Sequence[float]) -> None:
    total_twips = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")

    grid_columns = table._tbl.tblGrid.gridCol_lst
    for index, width in enumerate(widths):
        width_twips = int(width * 1440)
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(width_twips))
        table.columns[index].width = Inches(width)
        for cell in table.columns[index].cells:
            _set_cell_width(cell, width)


def _format_number(value: Any) -> str:
    try:
        return f"{int(round(float(value))):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "0"


def _add_metadata_table(doc: Document, data: Dict[str, Any]) -> None:
    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_layout_fixed(table)
    _remove_table_borders(table)
    widths = [1.35, 3.50, 1.25, 1.10]
    _set_table_column_widths(table, widths)
    for row in table.rows:
        _set_cant_split(row)

    _set_cell_text(table.cell(0, 0), "客 戶 名 稱：\nTên khách hàng：", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font=FONT_CJK)
    _set_cell_text(table.cell(0, 1), str(data.get("customer_name_vi") or ""), size=9.5, font=FONT_LATIN)
    _set_cell_text(table.cell(0, 2), "申 請 日 期：\nNgày nộp đơn：", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font=FONT_CJK)
    _set_cell_text(table.cell(0, 3), str(data.get("application_date") or ""), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    _set_cell_text(table.cell(1, 0), "專 案 名 稱：\nTên dự án：", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font=FONT_CJK)
    project_cell = table.cell(1, 1).merge(table.cell(1, 3))
    project_text = (
        f"Căn cứ hợp đồng số: {data.get('contract_code', '')}\n"
        f"{data.get('project_name_vi', '')} - {data.get('milestone_name', '')}\n"
        f"根據合同編號：{data.get('contract_code', '')}\n"
        f"{data.get('project_name_zh', '')}"
    )
    _set_cell_text(project_cell, project_text, size=8.5, vertical=WD_CELL_VERTICAL_ALIGNMENT.CENTER)

    _set_cell_text(table.cell(2, 0), "應 付 款 日 期：\nNgày đến hạn thanh toán：", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font=FONT_CJK)
    _set_cell_text(table.cell(2, 1), str(data.get("due_date") or ""), size=9.5)
    _set_cell_text(table.cell(2, 2), "付 款 方 式：\nPhương thức thanh toán：", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font=FONT_CJK)
    _set_cell_text(table.cell(2, 3), "匯款\nChuyển khoản", size=9, font=FONT_CJK)


def _add_payment_table(doc: Document, data: Dict[str, Any]) -> None:
    table = doc.add_table(rows=3, cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_layout_fixed(table)
    _set_table_borders(table, size=8)
    widths = [0.30, 1.65, 0.42, 0.82, 0.48, 1.00, 0.63, 0.72, 1.18]
    _set_table_column_widths(table, widths)
    for row in table.rows:
        _set_cant_split(row)

    headers = [
        "項\nSTT",
        "請款項目\nHạng mục thanh toán",
        "幣別\nTiền tệ",
        "單價\nĐơn giá",
        "數量\nSố lượng",
        "合同金額\nSố tiền hợp đồng",
        "稅率\nThuế suất",
        "稅額 VAT",
        "請款金額\nSố tiền xin TT",
    ]
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        _set_cell_shading(cell, YELLOW)
        _set_cell_text(cell, text, size=7.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    _set_repeat_table_header(table.rows[0])
    _set_row_minimum_height(table.rows[0], 30)

    total_contract = _format_number(data.get("total_contract_amount"))
    amount = _format_number(data.get("milestone_amount"))
    pct = data.get("percentage_display") or str(data.get("percentage") or "0")
    invoice = str(data.get("invoice_no") or "")
    item_text = (
        f"請款 / Thanh toán {data.get('milestone_name', '')}：{pct}%\n"
        f"合約 / Hợp đồng：{data.get('contract_code', '')}"
    )
    if invoice and invoice != "Chưa xuất":
        item_text += f"\n發票 / Hóa đơn：{invoice}"

    values = ["1", item_text, "VND", total_contract, "1", total_contract, "0%", "-", amount]
    aligns = [
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.RIGHT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.RIGHT,
    ]
    for index, (value, align) in enumerate(zip(values, aligns)):
        _set_cell_text(table.cell(1, index), value, size=7.1, align=align, font=FONT_CJK if index in (1, 2) else FONT_LATIN)
    _set_row_minimum_height(table.rows[1], 44)

    total_label = table.cell(2, 0).merge(table.cell(2, 4))
    total_value = table.cell(2, 5).merge(table.cell(2, 8))
    _set_cell_text(total_label, "合計 Tổng cộng", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    _set_cell_text(total_value, amount, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _set_row_minimum_height(table.rows[2], 20)


def _add_notes_table(doc: Document, data: Dict[str, Any]) -> None:
    terms = [str(value) for value in (data.get("payment_terms") or [])]
    if not terms:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_layout_fixed(table)
        _set_table_borders(table, size=8)
        _set_table_column_widths(table, [7.20])
        _set_cell_text(
            table.cell(0, 0),
            str(data.get("notes_text") or ""),
            size=7.35,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            font=FONT_CJK,
            vertical=WD_CELL_VERTICAL_ALIGNMENT.TOP,
        )
        _set_cell_margins(table.cell(0, 0), top=80, start=90, bottom=80, end=90)
        return

    bilingual_items = data.get("payment_terms_bilingual") or []
    count = len(bilingual_items) if bilingual_items else len(terms)
    table = doc.add_table(rows=count, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_layout_fixed(table)
    _set_table_borders(table, size=4, color="D9D9D9")
    _set_table_column_widths(table, [7.20])
    for index in range(count):
        cell = table.cell(index, 0)
        if index < len(bilingual_items):
            item = bilingual_items[index]
            chinese_body = item.get("chinese", "").strip()
            vietnamese_body = item.get("vietnamese", "").strip()
        else:
            term = terms[index]
            part1, part2 = (term.split("\n", 1) + [""])[:2]
            p1_clean = re.sub(r"^\d{2}[.\s]+", "", part1.strip()).strip()
            p2_clean = re.sub(r"^\d{2}[.\s]+", "", part2.strip()).strip()
            has_cjk_p1 = bool(re.search(r"[\u4e00-\u9fff]", p1_clean))
            if has_cjk_p1:
                chinese_body, vietnamese_body = p1_clean, p2_clean
            else:
                chinese_body, vietnamese_body = p2_clean, p1_clean

        prefix = f"{index + 1:02d}  "

        fill_color = "FFFFFF" if (index % 2 == 0) else "F3F8FC"
        _set_cell_shading(cell, fill_color)
        _set_cell_margins(cell, top=65, bottom=65, start=80, end=80)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        p = cell.paragraphs[0]
        p.clear()
        r_num = p.add_run(prefix)
        _set_run_font(r_num, name="Arial", size=9.0, bold=True)
        r_num.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

        r_zh = p.add_run(chinese_body)
        _set_run_font(r_zh, name=FONT_CJK, size=8.0, bold=False)

        if vietnamese_body:
            p_vi = cell.add_paragraph()
            p_vi.paragraph_format.left_indent = Pt(18)
            p_vi.paragraph_format.space_before = Pt(0)
            p_vi.paragraph_format.space_after = Pt(0)
            r_vi = p_vi.add_run(vietnamese_body)
            _set_run_font(r_vi, name=FONT_LATIN, size=8.0, bold=False)



def _add_bank_table(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_layout_fixed(table)
    _set_table_borders(table, size=8)
    widths = [1.45, 2.85, 1.05, 1.85]
    _set_table_column_widths(table, widths)
    for row in table.rows:
        _set_cant_split(row)
    _set_cell_text(table.cell(0, 0), "銀行帳戶名稱", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    _set_cell_text(table.cell(0, 1), "CÔNG TY TNHH TOÀN CẦU HONG MING", size=8)
    _set_cell_text(table.cell(0, 2), "銀行帳號", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    _set_cell_text(table.cell(0, 3), "604080006546(VND)", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(table.cell(1, 0), "銀行名稱", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    bank_cell = table.cell(1, 1).merge(table.cell(1, 3))
    _set_cell_text(bank_cell, "CATHAY UNITED BANK - HO CHI MINH CITY BRANCH", size=8)


def _add_signatures(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_layout_fixed(table)
    _remove_table_borders(table)
    _set_table_column_widths(table, [3.60, 3.60])
    _set_cell_text(table.cell(0, 0), "買方主管", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    _set_cell_text(table.cell(0, 1), "申請單位", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK)
    _set_row_minimum_height(table.rows[0], 48)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.9)
    section.bottom_margin = Cm(0.9)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(8.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_after = Pt(0)


def _save_from_template_preserving_package(
    doc: Document,
    template_path: Path,
    output_path: Path,
) -> Path:
    """Lưu bản điền dữ liệu và giữ mọi package part ngoài document.xml y nguyên mẫu."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_paths: List[Path] = []
    try:
        with tempfile.NamedTemporaryFile(
            dir=str(output_path.parent),
            prefix=f".{output_path.stem}-edited-",
            suffix=".docx",
            delete=False,
        ) as handle:
            edited_path = Path(handle.name)
        temp_paths.append(edited_path)
        doc.save(str(edited_path))

        with tempfile.NamedTemporaryFile(
            dir=str(output_path.parent),
            prefix=f".{output_path.stem}-rebuilt-",
            suffix=".docx",
            delete=False,
        ) as handle:
            rebuilt_path = Path(handle.name)
        temp_paths.append(rebuilt_path)

        with ZipFile(str(template_path), "r") as source_zip, ZipFile(
            str(edited_path), "r"
        ) as edited_zip, ZipFile(str(rebuilt_path), "w") as output_zip:
            edited_document_xml = edited_zip.read("word/document.xml")
            for item in source_zip.infolist():
                content = (
                    edited_document_xml
                    if item.filename == "word/document.xml"
                    else source_zip.read(item.filename)
                )
                output_zip.writestr(item, content)

        actual_path = output_path
        try:
            os.replace(str(rebuilt_path), str(actual_path))
        except PermissionError:
            actual_path = output_path.with_name(f"{output_path.stem}_v2{output_path.suffix}")
            os.replace(str(rebuilt_path), str(actual_path))
        return actual_path
    finally:
        for temp_path in temp_paths:
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except OSError:
                pass


def generate_payment_request_word(
    data: Dict[str, Any],
    output_path: Path,
    template_path: Optional[Path] = None,
) -> Path:
    """Tạo Word từ mẫu DOCX nếu được chọn; nếu không, dựng theo bố cục sheet 712TR."""
    if template_path is not None:
        if not template_path.exists():
            raise FileNotFoundError(f"Không tìm thấy file mẫu Word tại: {template_path}")
        doc = Document(str(template_path))
        replacement_count = _fill_word_template(doc, data)
        if replacement_count == 0:
            replacement_count = _fill_known_payment_request_layout(doc, data)
        if replacement_count == 0:
            raise ValueError(
                "Mẫu Word không khớp cấu trúc Phiếu yêu cầu thanh toán và không có placeholder được hỗ trợ. "
                "Hãy dùng dạng {{TEN_TRUONG}}, [TEN_TRUONG], <<TEN_TRUONG>> "
                "hoặc tên trường được tô vàng/chữ đỏ."
            )
    else:
        doc = Document()
        _configure_document(doc)

        _set_paragraph(doc.add_paragraph(), "浤 詺 資 訊 責 任 有 限 公 司", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK, space_after=0)
        _set_paragraph(doc.add_paragraph(), "CÔNG TY TNHH TOÀN CẦU HONG MING", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font="Arial", space_after=4)
        _set_paragraph(doc.add_paragraph(), "請 款 單", size=19, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_CJK, space_after=0)
        _set_paragraph(doc.add_paragraph(), "Phiếu yêu cầu thanh toán", size=13.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font="Arial", space_after=5)

        _set_paragraph(doc.add_paragraph(), "CÔNG TY TNHH TOÀN CẦU HONG MING 致敬：", size=9.5, font=FONT_CJK, space_after=0)
        _set_paragraph(doc.add_paragraph(), f"Form NO.     {data.get('form_no', '')}", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font="Arial", space_after=2)

        _add_metadata_table(doc, data)
        _add_payment_table(doc, data)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1.5)
        _add_notes_table(doc, data)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1.5)
        _add_bank_table(doc)
        _add_signatures(doc)

    if template_path is not None:
        return _save_from_template_preserving_package(doc, template_path, output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    actual_path = output_path
    try:
        doc.save(str(actual_path))
    except PermissionError:
        actual_path = output_path.with_name(f"{output_path.stem}_v2{output_path.suffix}")
        doc.save(str(actual_path))
    return actual_path
