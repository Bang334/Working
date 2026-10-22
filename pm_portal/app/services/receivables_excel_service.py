# -*- coding: utf-8 -*-
"""Service module to export Project Receivables (AR) into Excel file matching 1:1 Sheet 3 format."""

import io
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.receivables_service import list_receivables, get_receivables_stats


def export_receivables_excel(
    customer: Optional[str] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    overdue_only: bool = False
) -> bytes:
    """Generate Excel workbook matching Sheet 3 'Theo Dõi Công Nợ Dự Án' with live formulas."""
    items = list_receivables(customer=customer, status=status, search=search, overdue_only=overdue_only)
    stats = get_receivables_stats()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3. Theo Dõi Công Nợ Dự Án"

    # Font definitions
    FONT_NAME = "Segoe UI"
    font_title = Font(name=FONT_NAME, size=15, bold=True, color="1F4E79")
    font_subtitle = Font(name=FONT_NAME, size=10, italic=True, color="555555")
    font_header = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    font_data = Font(name=FONT_NAME, size=9, color="111827")
    font_data_bold = Font(name=FONT_NAME, size=9, bold=True, color="111827")
    font_total = Font(name=FONT_NAME, size=10, bold=True, color="1F4E79")

    # Fills
    fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_total = PatternFill(start_color="E6F0FA", end_color="E6F0FA", fill_type="solid")

    # Status fills
    fill_paid = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    fill_overdue = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
    fill_urgent = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")

    # Borders
    thin_border_side = Side(border_style="thin", color="D0D7DE")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="1F4E79"))
    border_total = Border(left=thin_border_side, right=thin_border_side, top=Side(border_style="thin", color="1F4E79"), bottom=Side(border_style="double", color="1F4E79"))

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1. Table Column Headers (Bắt đầu trực tiếp từ Row 1, giữ thuần túy table)
    headers = [
        ("STT\n序號", 8),
        ("Dự án\n專案名稱", 28),
        ("Khách hàng\n客戶名稱", 22),
        ("Hợp đồng\n合約編號", 18),
        ("Giá trị HĐ (chưa VAT)\n合約總額 (未稅)", 22),
        ("Số tiền Đợt TT (chưa VAT)\n本期請款 (未稅)", 22),
        ("Thuế suất\n稅率", 12),
        ("Thuế VAT\n營業稅", 20),
        ("Tổng tiền TT (gồm VAT)\n應收總額 (含稅)", 24),
        ("Số tiền đã nhận\n已收金額", 22),
        ("Số tiền còn lại\n剩餘尾款", 22),
        ("Ngày KH hẹn TT\n預計付款日", 16),
        ("Ngày đến hạn\n到期日", 16),
        ("Ngày đã thanh toán\n實際付款日", 16),
        ("Trạng thái CT\n單據狀態", 18),
        ("Trạng thái thu\n收款狀態", 18),
        ("Ghi chú tiến độ\n備註", 32),
    ]

    header_row = 1
    ws.row_dimensions[header_row].height = 42
    for col_idx, (header_text, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_header
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # 2. Populate Data Rows (Bắt đầu từ Row 2)
    start_row = 2
    for idx, item in enumerate(items, start=1):
        row_num = start_row + idx - 1

        # Zebra striping
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra if is_even else None

        # Values
        stt = item.get("order_no") or idx
        proj = item.get("project_name") or ""
        cust = item.get("customer_name") or ""
        contract = item.get("contract_code") or ""
        total_contract = item.get("total_contract_amount") or 0.0
        net_val = item.get("net_amount") or 0.0
        vat_rate = item.get("vat_rate") or 0.0
        rec_val = item.get("received_amount") or 0.0
        promised = item.get("customer_promised_date") or ""
        due = item.get("due_date") or ""
        paid = item.get("paid_date") or ""
        doc_status = item.get("doc_status") or ""
        status_label = item.get("status_label") or item.get("collection_status") or ""
        notes = item.get("notes") or ""

        # Dynamic row height calculation (anti-clipping for wrapped text)
        text_lines = max(
            len(str(proj).split("\n")),
            len(str(cust).split("\n")),
            len(str(doc_status).split("\n")),
            len(str(notes).split("\n")),
            1
        )
        if len(str(proj)) > 30 or len(str(notes)) > 35:
            text_lines = max(text_lines, 2)
        if len(str(proj)) > 60 or len(str(notes)) > 70:
            text_lines = max(text_lines, 3)

        ws.row_dimensions[row_num].height = max(26, min(120, text_lines * 19 + 6))

        # Cells with live Excel formulas
        # F: net, G: vat_rate, H: vat_amt (=F*G), I: gross (=F+H), J: rec, K: remaining (=I-J)
        ws.cell(row=row_num, column=1, value=stt).alignment = align_center
        ws.cell(row=row_num, column=2, value=proj).alignment = align_left
        ws.cell(row=row_num, column=3, value=cust).alignment = align_left
        ws.cell(row=row_num, column=4, value=contract).alignment = align_center
        
        c5 = ws.cell(row=row_num, column=5, value=total_contract)
        c5.alignment = align_right
        c5.number_format = '#,##0" ₫"'

        c6 = ws.cell(row=row_num, column=6, value=net_val)
        c6.alignment = align_right
        c6.number_format = '#,##0" ₫"'

        c7 = ws.cell(row=row_num, column=7, value=vat_rate)
        c7.alignment = align_center
        c7.number_format = '0%'

        c8 = ws.cell(row=row_num, column=8, value=f"=F{row_num}*G{row_num}")
        c8.alignment = align_right
        c8.number_format = '#,##0" ₫"'

        c9 = ws.cell(row=row_num, column=9, value=f"=F{row_num}+H{row_num}")
        c9.alignment = align_right
        c9.number_format = '#,##0" ₫"'

        c10 = ws.cell(row=row_num, column=10, value=rec_val)
        c10.alignment = align_right
        c10.number_format = '#,##0" ₫"'

        c11 = ws.cell(row=row_num, column=11, value=f"=I{row_num}-J{row_num}")
        c11.alignment = align_right
        c11.number_format = '#,##0" ₫"'

        ws.cell(row=row_num, column=12, value=promised).alignment = align_center
        ws.cell(row=row_num, column=13, value=due).alignment = align_center
        ws.cell(row=row_num, column=14, value=paid).alignment = align_center
        ws.cell(row=row_num, column=15, value=doc_status).alignment = align_left
        
        c16 = ws.cell(row=row_num, column=16, value=status_label)
        c16.alignment = align_center
        
        # Color status column
        col_status_code = item.get("collection_status")
        if col_status_code == "PAID":
            c16.fill = fill_paid
        elif col_status_code == "DELAYED":
            c16.fill = fill_overdue
        elif col_status_code == "URGENT":
            c16.fill = fill_urgent

        ws.cell(row=row_num, column=17, value=notes).alignment = align_left

        # Apply fonts, borders, fills
        for c in range(1, 18):
            cell = ws.cell(row=row_num, column=c)
            cell.font = font_data
            cell.border = border_cell
            if row_fill and c != 16:  # preserve status fill
                cell.fill = row_fill

    last_data_row = start_row + len(items) - 1
    total_row = last_data_row + 1

    # 4. Summary Row (TỔNG CỘNG / 合計)
    if items:
        ws.row_dimensions[total_row].height = 30
        ws.merge_cells(f"A{total_row}:D{total_row}")
        total_label_cell = ws.cell(row=total_row, column=1, value="TỔNG CỘNG / 合計")
        total_label_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_label_cell.font = font_total

        # Formulas for columns 5, 6, 8, 9, 10, 11
        ws.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{last_data_row})")
        ws.cell(row=total_row, column=6, value=f"=SUM(F{start_row}:F{last_data_row})")
        ws.cell(row=total_row, column=7, value="-").alignment = align_center
        ws.cell(row=total_row, column=8, value=f"=SUM(H{start_row}:H{last_data_row})")
        ws.cell(row=total_row, column=9, value=f"=SUM(I{start_row}:I{last_data_row})")
        ws.cell(row=total_row, column=10, value=f"=SUM(J{start_row}:J{last_data_row})")
        ws.cell(row=total_row, column=11, value=f"=SUM(K{start_row}:K{last_data_row})")

        for c in range(1, 18):
            cell = ws.cell(row=total_row, column=c)
            cell.font = font_total
            cell.fill = fill_total
            cell.border = border_total
            if c in (5, 6, 8, 9, 10, 11):
                cell.number_format = '#,##0" ₫"'
                cell.alignment = align_right

    # 5. KÍCH HOẠT AUTOFILTER TRÊN TOÀN BỘ 17 CỘT TIÊU ĐỀ
    if last_data_row >= 1:
        ws.auto_filter.ref = f"A1:Q{last_data_row}"

    # 6. Freeze Panes & Sheet View Settings (Ghim cố định STT & Tên Dự án, cuộn ngang dọc không mất)
    ws.freeze_panes = "C2"
    ws.views.sheetView[0].showGridLines = True

    # 7. Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
