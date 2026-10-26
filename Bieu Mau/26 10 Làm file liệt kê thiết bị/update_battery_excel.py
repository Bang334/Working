# -*- coding: utf-8 -*-
"""
Script bổ sung 8 viên pin vào file Excel thiết bị DFV
Đường dẫn file: E:\Working\Bieu Mau\26 10 Làm file liệt kê thiết bị\DANH_SACH_THIET_BI_SUA_CHUA_THAY_THE_DFV_DARFON.xlsx
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

EXCEL_PATH = r"E:\Working\Bieu Mau\26 10 Làm file liệt kê thiết bị\DANH_SACH_THIET_BI_SUA_CHUA_THAY_THE_DFV_DARFON.xlsx"
BAT_DIR = r"E:\Working\Bieu Mau\26 10 Làm file liệt kê thiết bị\batteries"

BATTERIES = [
    {
        "stt": 12,
        "img": "bat_1_C629B2309000251.jpg",
        "sn": "C629B2309000251",
    },
    {
        "stt": 13,
        "img": "bat_2_C629B2309000613.jpg",
        "sn": "C629B2309000613",
    },
    {
        "stt": 14,
        "img": "bat_3_C629B2309001047.jpg",
        "sn": "C629B2309001047",
    },
    {
        "stt": 15,
        "img": "bat_4_C629B2309000337.jpg",
        "sn": "C629B2309000337",
    },
    {
        "stt": 16,
        "img": "bat_5_C629B2309001215.jpg",
        "sn": "C629B2309001215",
    },
    {
        "stt": 17,
        "img": "bat_6_C629B2309000561.jpg",
        "sn": "C629B2309000561",
    },
    {
        "stt": 18,
        "img": "bat_7_C629B2309001198.jpg",
        "sn": "C629B2309001198",
    },
    {
        "stt": 19,
        "img": "bat_8_C629B2305001525.jpg",
        "sn": "C629B2305001525",
    }
]

def update_workbook():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    FONT_FAMILY = "Microsoft JhengHei"
    font_data = Font(name=FONT_FAMILY, size=9.5, color="111111")
    font_data_bold = Font(name=FONT_FAMILY, size=9.5, bold=True, color="111111")
    font_sn = Font(name="Consolas", size=10.5, bold=True, color="002060")
    font_hs = Font(name="Consolas", size=10.5, bold=True, color="0B5394")
    font_fault = Font(name=FONT_FAMILY, size=9, bold=False, color="900C3F")
    font_summary = Font(name=FONT_FAMILY, size=10.5, bold=True, color="1F4E79")
    
    fill_zebra_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra_2 = PatternFill(start_color="F9FBFC", end_color="F9FBFC", fill_type="solid")
    fill_hs = PatternFill(start_color="E8F4F8", end_color="E8F4F8", fill_type="solid")
    fill_fault = PatternFill(start_color="FDF2E9", end_color="FDF2E9", fill_type="solid")
    fill_summary = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")
    
    bd_thin = Side(border_style="thin", color="D0D7DE")
    border_cell = Border(left=bd_thin, right=bd_thin, top=bd_thin, bottom=bd_thin)
    border_summary = Border(
        top=Side(style="thin", color="1F4E79"),
        bottom=Side(style="double", color="1F4E79"),
        left=bd_thin,
        right=bd_thin
    )
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # ---------------------------------------------------------
    # 1. CẬP NHẬT SHEET 1: DANH SÁCH THIẾT BỊ (明細表)
    # ---------------------------------------------------------
    ws1 = wb["Danh Sách Thiết Bị (明細表)"]
    
    # Unmerge dòng 16 cũ nếu đang gộp
    for rng in list(ws1.merged_cells.ranges):
        if rng.coord in ["A16:C16", "D16:H16"]:
            ws1.unmerge_cells(rng.coord)
            
    # Bổ sung 8 viên pin từ Hàng 16 đến Hàng 23
    for idx, bat in enumerate(BATTERIES):
        row = 16 + idx
        ws1.row_dimensions[row].height = 120
        fill_row = fill_zebra_2 if (row % 2 == 1) else fill_zebra_1
        
        # Col 1: STT
        c1 = ws1.cell(row=row, column=1, value=bat["stt"])
        c1.alignment = align_center
        c1.font = font_data_bold
        c1.fill = fill_row
        c1.border = border_cell
        
        # Col 2: Ảnh (Cắt sắc nét, nhúng nguyên bản)
        img_path = os.path.join(BAT_DIR, bat["img"])
        c2 = ws1.cell(row=row, column=2, value="")
        c2.alignment = align_center
        c2.fill = fill_row
        c2.border = border_cell
        c2.hyperlink = img_path
        
        xl_img = XLImage(img_path)
        # Tỷ lệ pin dọc: rộng ~86px, cao ~118px
        xl_img.width = 86
        xl_img.height = 118
        ws1.add_image(xl_img, f"B{row}")
        
        # Col 3: Chủng loại
        c3 = ws1.cell(row=row, column=3, value="Pin sạc máy tính bảng\n平板專用鋰離子電池")
        c3.alignment = align_center
        c3.font = font_data_bold
        c3.fill = fill_row
        c3.border = border_cell
        
        # Col 4: Tên & Model
        c4 = ws1.cell(row=row, column=4, value="Rechargeable Li-ion Battery\nModel: CLP629 (10000mAh)")
        c4.alignment = align_left
        c4.font = font_data_bold
        c4.fill = fill_row
        c4.border = border_cell
        
        # Col 5: P/N
        c5 = ws1.cell(row=row, column=5, value="1ICP6/45/116-2")
        c5.alignment = align_center
        c5.font = font_data
        c5.fill = fill_row
        c5.border = border_cell
        
        # Col 6: S/N
        c6 = ws1.cell(row=row, column=6, value=bat["sn"])
        c6.alignment = align_center
        c6.font = font_sn
        c6.fill = fill_row
        c6.border = border_cell
        
        # Col 7: Mã HS Code
        c7 = ws1.cell(row=row, column=7, value="8507.60.90")
        c7.alignment = align_center
        c7.font = font_hs
        c7.fill = fill_hs
        c7.border = border_cell
        
        # Col 8: Hiện trạng lỗi
        c8 = ws1.cell(
            row=row, 
            column=8, 
            value="Pin phồng rộp (电池鼓包), không đảm bảo an toàn gửi kèm máy. Đề nghị cấp mới.\n電池鼓包，無法隨機寄送。申請原廠補發新電池。"
        )
        c8.alignment = align_left
        c8.font = font_fault
        c8.fill = fill_fault
        c8.border = border_cell

    # Dòng tổng cộng mới tại Hàng 24
    sum_row = 24
    ws1.row_dimensions[sum_row].height = 30
    ws1.merge_cells(f"A{sum_row}:C{sum_row}")
    c_sum = ws1.cell(row=sum_row, column=1, value="TỔNG CỘNG / 合計")
    c_sum.alignment = align_center
    c_sum.font = font_summary
    c_sum.fill = fill_summary
    c_sum.border = border_summary
    for col in range(2, 4):
        ws1.cell(row=sum_row, column=col).fill = fill_summary
        ws1.cell(row=sum_row, column=col).border = border_summary
        
    ws1.merge_cells(f"D{sum_row}:H{sum_row}")
    sum_text = (
        "19 Thiết bị & Linh kiện (3 Máy tính bảng HT10 + 4 Đế sạc Cradle S + 4 Nguồn Adapter + 8 Pin sạc HT10)\n"
        "合計 19 件設備及配件（3台HT10平板 + 4個充電底座 + 4個12V 2A電源適應器 + 8顆HT10專用鋰電池）"
    )
    c_sum_detail = ws1.cell(row=sum_row, column=4, value=sum_text)
    c_sum_detail.alignment = align_left
    c_sum_detail.font = font_summary
    c_sum_detail.fill = fill_summary
    c_sum_detail.border = border_summary
    for col in range(5, 9):
        ws1.cell(row=sum_row, column=col).fill = fill_summary
        ws1.cell(row=sum_row, column=col).border = border_summary
        
    ws1.print_area = f"A1:H{sum_row}"

    # ---------------------------------------------------------
    # 2. CẬP NHẬT SHEET 2: TỔNG HỢP HS CODE & NHẬP KHẨU
    # ---------------------------------------------------------
    ws2 = wb["Tổng Hợp HS Code & Nhập Khẩu"]
    
    # Cập nhật số lượng pin tại Hàng 5 từ 3 lên 8
    ws2["D5"] = 8
    ws2["D5"].font = font_data_bold
    ws2["D5"].alignment = align_center
    
    # Cập nhật diễn giải tại I5
    ws2["I5"] = (
        "Đổi mới/Bổ sung: 8 viên pin HT10 bị phồng nghiêm trọng (电池鼓包) bao gồm pin của các máy gửi về "
        "và phần pin bổ sung của A Bằng (阿鵬). Đề nghị nhà máy cấp bù 8 viên pin mới nguyên seal."
    )
    ws2["I5"].alignment = align_left
    ws2["I5"].font = font_data
    
    # Dòng tổng kết Hàng 8
    ws2["D8"] = "=SUM(D4:D7)"
    ws2["E8"] = "19 Hạng mục thiết bị & linh kiện (Bao gồm 11 thiết bị/phụ kiện kiểm tra bảo hành + 8 viên pin cấp mới thay thế)"
    
    # Lưu file an toàn
    save_path = EXCEL_PATH
    try:
        wb.save(save_path)
        print(f"Successfully updated file: {save_path}")
    except PermissionError:
        save_path = EXCEL_PATH.replace(".xlsx", "_updated.xlsx")
        wb.save(save_path)
        print(f"File was locked, saved to: {save_path}")

if __name__ == "__main__":
    update_workbook()
